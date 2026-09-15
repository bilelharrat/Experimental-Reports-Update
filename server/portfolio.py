"""Private portfolio: positions, founder updates, KPI history, marks, alerts, reserves and LP tear sheets.

Everything here is record-backed. KPIs come from a founder update (regex-extracted, with
the excerpt kept) or from a manual entry; marks are entered by a person with their basis.
Nothing is projected or estimated — a missing figure stays missing and the alert says so.
"""
from __future__ import annotations

import io
import json
import math
import re
import threading
import uuid
from datetime import datetime, timedelta, timezone
from typing import Any

from docx import Document
from docx.shared import Pt

from . import company_paths, decisions_store, intake_decks, storage

_LOCK = threading.RLock()

def _reserves_file():
    return storage.DATA_DIR / "settings" / "reserves.yaml"

POSITION_KEYS = {
    "round": str,
    "security": str,
    "entry_date": str,
    "invested_usd": float,
    "ownership_pct": float,
    "entry_post_money_usd": float,
    "planned_follow_on_usd": float,
    "board_seat": bool,
    "lead": bool,
    "notes": str,
}

KPI_KEYS = ("arr_usd", "revenue_usd", "burn_usd_month", "cash_usd", "runway_months", "headcount", "customers")

MARK_BASES = ("last_round", "lp_report", "secondary", "manual", "cost")

RUNWAY_HIGH_MONTHS = 9
RUNWAY_MEDIUM_MONTHS = 12
STALE_DAYS = 60
BURN_JUMP_PCT = 25.0

_CASH_PATTERN = (
    r"\b(?:cash(?!\s+(?:burn|flow|collect\w*|from|receipts|conversion))|cash balance|in the bank|bank balance)\b"
    + intake_decks._gap(30, "burn|arr|revenue|runway")
    + intake_decks.MONEY
)

_KPI_BOUNDS: dict[str, dict[str, Any]] = {
    "arr_usd": {"min": 0},
    "revenue_usd": {"min": 0},
    "burn_usd_month": {"min": 0, "hint": "enter 0 if cash-flow positive"},
    "cash_usd": {"min": 0},
    "runway_months": {"min": 0},
    "headcount": {"min": 0, "integer": True},
    "customers": {"min": 0, "integer": True},
}
_POSITION_BOUNDS: dict[str, dict[str, Any]] = {
    "invested_usd": {"min": 0},
    "entry_post_money_usd": {"min": 0},
    "planned_follow_on_usd": {"min": 0},
    "ownership_pct": {"min": 0, "max": 100},
}


def _now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _company_dir(company_id: str):
    return company_paths.company_dir(company_id)


def _path(company_id: str):
    return _company_dir(company_id) / "portfolio.json"


def has_record(company_id: str) -> bool:
    try:
        return _path(company_id).exists()
    except ValueError:
        return False


def _number(key: str, value: Any, *, min: float | None = None, max: float | None = None, integer: bool = False, hint: str = "") -> float | int:
    if isinstance(value, bool):
        raise ValueError(f"{key} must be a number")
    try:
        number = float(str(value).replace(",", "").replace("$", "").strip())
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{key} must be a number") from exc
    if not math.isfinite(number):
        raise ValueError(f"{key} must be a finite number")
    suffix = f" ({hint})" if hint else ""
    if min is not None and number < min:
        raise ValueError(f"{key} must be at least {min:g}{suffix}")
    if max is not None and number > max:
        raise ValueError(f"{key} must be at most {max:g}{suffix}")
    if integer:
        if not number.is_integer():
            raise ValueError(f"{key} must be a whole number")
        return int(number)
    return number


def _finite(value: Any) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return float(value) if math.isfinite(value) else None


def _empty(company_id: str) -> dict:
    return {"company_id": company_id, "position": {}, "kpis": [], "updates": [], "marks": []}


def _load(company_id: str) -> dict:
    path = _path(company_id)
    if not path.exists():
        return _empty(company_id)
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return _empty(company_id)
    base = _empty(company_id)
    for key in ("position", "kpis", "updates", "marks"):
        if key in data and isinstance(data[key], type(base[key])):
            base[key] = data[key]
    return base


def _save(company_id: str, payload: dict) -> None:
    path = _path(company_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    tmp.replace(path)


def _coerce(key: str, value: Any, kind: type) -> Any:
    if value is None or value == "":
        return None
    if kind is bool:
        return bool(value) if not isinstance(value, str) else value.strip().lower() in {"1", "true", "yes", "y"}
    if kind is float:
        return _number(key, value, **_POSITION_BOUNDS.get(key, {}))
    return str(value).strip()[:2000]


def _parse_date(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return _now()
    for fmt in ("%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%SZ", "%Y-%m-%dT%H:%M:%S%z"):
        try:
            parsed = datetime.strptime(text.replace("+00:00", "Z") if fmt.endswith("Z") else text, fmt)
            if parsed.tzinfo is None:
                parsed = parsed.replace(tzinfo=timezone.utc)
            return parsed.astimezone(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
        except ValueError:
            continue
    raise ValueError("as_of must be an ISO date (YYYY-MM-DD)")


def _dt(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return None


# ---- Positions -------------------------------------------------------------------------


def update_position(company_id: str, patch: dict, *, updated_by: str | None = None) -> dict:
    if not isinstance(patch, dict):
        raise ValueError("patch must be an object")
    unknown = sorted(set(patch) - set(POSITION_KEYS))
    if unknown:
        raise ValueError(f"Unknown position fields: {', '.join(unknown)}")
    with _LOCK:
        record = _load(company_id)
        position = dict(record["position"])
        for key, value in patch.items():
            coerced = _coerce(key, value, POSITION_KEYS[key])
            if coerced is None:
                position.pop(key, None)
            else:
                position[key] = coerced
        position["updated_at"] = _now()
        if updated_by:
            position["updated_by"] = updated_by
        record["position"] = position
        _save(company_id, record)
    return get_portfolio(company_id)


# ---- KPIs ------------------------------------------------------------------------------


def _kpi_row(company_id: str, row: dict, *, source: str, created_by: str | None, excerpts: dict | None = None) -> dict:
    clean: dict[str, Any] = {
        "id": f"kpi-{uuid.uuid4().hex[:10]}",
        "as_of": _parse_date(row.get("as_of")),
        "recorded_at": _now(),
        "source": source,
        "note": str(row.get("note") or "")[:1000],
    }
    if created_by:
        clean["created_by"] = created_by
    has_value = False
    for key in KPI_KEYS:
        value = row.get(key)
        if value is None or value == "":
            continue
        clean[key] = _number(key, value, **_KPI_BOUNDS[key])
        has_value = True
    if not has_value:
        raise ValueError("A KPI row needs at least one figure")
    if excerpts:
        clean["excerpts"] = excerpts
    return clean


def add_kpi(company_id: str, row: dict, *, created_by: str | None = None) -> dict:
    with _LOCK:
        record = _load(company_id)
        record["kpis"].append(_kpi_row(company_id, row, source="manual", created_by=created_by))
        _save(company_id, record)
    return get_portfolio(company_id)


# ---- Founder updates -------------------------------------------------------------------


class _TextSlide:
    def __init__(self, text: str):
        self.text = text
        self.notes = ""
        self.slide_no = None


def extract_update_fields(text: str) -> list[dict]:
    """Deterministic extraction over a founder update: same rules as deck intake plus cash."""
    fields = intake_decks.extract_fields([_TextSlide(text)])
    flat = " ".join(text.split())
    match = re.search(_CASH_PATTERN, flat, flags=re.I)
    if match:
        pairs = intake_decks._money_pairs(match)
        usd = intake_decks._money_usd(*pairs[0]) if pairs else None
        fields.append({
            "name": "cash",
            "value": " ".join(match.group(0).split()),
            "usd": usd,
            "page": None,
            "excerpt": intake_decks._excerpt(flat, match.start(), match.end()),
        })
    return fields


_FIELD_TO_KPI = {
    "arr": "arr_usd",
    "revenue": "revenue_usd",
    "burn": "burn_usd_month",
    "cash": "cash_usd",
    "runway": "runway_months",
    "headcount": "headcount",
    "customers": "customers",
}


def add_update(
    company_id: str,
    *,
    text: str,
    as_of: str | None = None,
    source: str = "email",
    subject: str = "",
    created_by: str | None = None,
) -> dict:
    body = str(text or "").strip()
    if not body:
        raise ValueError("Update text is empty")
    fields = extract_update_fields(body)
    kpi_values: dict[str, Any] = {"as_of": as_of}
    excerpts: dict[str, str] = {}
    for field in fields:
        key = _FIELD_TO_KPI.get(field["name"])
        if not key:
            continue
        value = field["usd"] if field["usd"] is not None else field["value"]
        try:
            number = float(str(value).replace(",", ""))
        except ValueError:
            continue
        if not math.isfinite(number) or number < 0:
            continue
        kpi_values[key] = value
        excerpts[key] = field["excerpt"]
    update = {
        "id": f"upd-{uuid.uuid4().hex[:10]}",
        "as_of": _parse_date(as_of),
        "received_at": _now(),
        "source": str(source or "email")[:40],
        "subject": str(subject or "")[:200],
        "text": body[:20000],
        "extracted": fields,
        "kpi_id": None,
    }
    if created_by:
        update["created_by"] = created_by
    with _LOCK:
        record = _load(company_id)
        if len(kpi_values) > 1:
            row = _kpi_row(company_id, kpi_values, source=f"founder_update:{update['id']}", created_by=created_by, excerpts=excerpts)
            record["kpis"].append(row)
            update["kpi_id"] = row["id"]
        record["updates"].append(update)
        _save(company_id, record)
    return get_portfolio(company_id)


# ---- Marks -----------------------------------------------------------------------------


def add_mark(company_id: str, *, value_usd: Any, basis: str, as_of: str | None = None, note: str = "", created_by: str | None = None) -> dict:
    basis = str(basis or "").strip().lower()
    if basis not in MARK_BASES:
        raise ValueError(f"basis must be one of {', '.join(MARK_BASES)}")
    value = _number("value_usd", value_usd, min=0)
    mark = {
        "id": f"mark-{uuid.uuid4().hex[:10]}",
        "as_of": _parse_date(as_of),
        "recorded_at": _now(),
        "value_usd": value,
        "basis": basis,
        "note": str(note or "")[:1000],
    }
    if created_by:
        mark["created_by"] = created_by
    with _LOCK:
        record = _load(company_id)
        record["marks"].append(mark)
        _save(company_id, record)
    return get_portfolio(company_id)


def remove_item(company_id: str, kind: str, item_id: str) -> bool:
    if kind not in {"kpis", "updates", "marks"}:
        raise ValueError("kind must be kpis, updates or marks")
    with _LOCK:
        record = _load(company_id)
        before = len(record[kind])
        record[kind] = [item for item in record[kind] if item.get("id") != item_id]
        if kind == "updates":
            gone = {item.get("kpi_id") for item in record["updates"]}
            # keep KPI rows that came from remaining updates; drop the one tied to the removed update
            record["kpis"] = [k for k in record["kpis"] if not (str(k.get("source", "")).startswith("founder_update:") and k["source"].split(":", 1)[1] == item_id)]
            _ = gone
        changed = len(record[kind]) != before
        if changed:
            _save(company_id, record)
    return changed


# ---- Derived views ---------------------------------------------------------------------


def _latest(items: list[dict]) -> dict | None:
    if not items:
        return None
    return sorted(items, key=lambda item: (item.get("as_of") or "", item.get("recorded_at") or item.get("received_at") or ""))[-1]


def alerts_for(record: dict, *, now: datetime | None = None) -> list[dict]:
    now = now or datetime.now(timezone.utc)
    alerts: list[dict] = []
    kpis = sorted(record.get("kpis") or [], key=lambda item: (item.get("as_of") or "", item.get("recorded_at") or ""))

    def _cash_and_burn(row: dict) -> tuple[float, float] | None:
        cash, burn = _finite(row.get("cash_usd")), _finite(row.get("burn_usd_month"))
        return (cash, burn) if cash and burn and cash > 0 and burn > 0 else None

    src = next((k for k in reversed(kpis) if _finite(k.get("runway_months")) is not None or _cash_and_burn(k)), None)
    if src:
        runway = _finite(src.get("runway_months"))
        as_of = str(src.get("as_of") or "")[:10]
        if runway is not None:
            if runway < RUNWAY_HIGH_MONTHS:
                alerts.append({"kind": "runway", "severity": "high", "label": f"Runway {runway:g} months", "detail": f"Reported as of {as_of} ({src.get('source')})."})
            elif runway < RUNWAY_MEDIUM_MONTHS:
                alerts.append({"kind": "runway", "severity": "medium", "label": f"Runway {runway:g} months", "detail": f"Reported as of {as_of}."})
        else:
            cash, burn = _cash_and_burn(src)
            implied = cash / burn
            severity = "high" if implied < RUNWAY_HIGH_MONTHS else "medium" if implied < RUNWAY_MEDIUM_MONTHS else None
            if severity:
                shown = math.floor(implied * 10) / 10
                alerts.append({"kind": "runway", "severity": severity, "label": f"Cash ÷ burn ≈ {shown:g} months", "detail": f"Derived from cash and monthly burn reported as of {as_of}; the founder did not state runway."})
    if kpis:
        burns = [k for k in kpis if _finite(k.get("burn_usd_month"))]
        if len(burns) >= 2:
            prev, cur = burns[-2]["burn_usd_month"], burns[-1]["burn_usd_month"]
            if prev > 0:
                jump = (cur - prev) / prev * 100
                if jump >= BURN_JUMP_PCT:
                    alerts.append({"kind": "burn", "severity": "medium", "label": f"Burn up {jump:.0f}%", "detail": f"${prev:,.0f} → ${cur:,.0f} per month between {burns[-2]['as_of'][:10]} and {burns[-1]['as_of'][:10]}."})
    last_touch = max([_dt(u.get("as_of")) for u in record.get("updates") or []] + [_dt(k.get("as_of")) for k in kpis] + [None], key=lambda d: d or datetime.min.replace(tzinfo=timezone.utc))
    if last_touch is None:
        alerts.append({"kind": "stale", "severity": "low", "label": "No founder update on file", "detail": "Paste the latest update or record KPIs by hand."})
    elif now - last_touch > timedelta(days=STALE_DAYS):
        alerts.append({"kind": "stale", "severity": "low", "label": f"No update in {(now - last_touch).days} days", "detail": f"Last data point {last_touch.date().isoformat()}."})
    return alerts


def _summary(company_id: str, record: dict, *, now: datetime | None = None) -> dict:
    company = storage.get_company(company_id) or {"id": company_id, "name": company_id}
    position = record.get("position") or {}
    latest_kpi = _latest(record.get("kpis") or [])
    latest_mark = _latest(record.get("marks") or [])
    invested = position.get("invested_usd")
    moic = None
    if _finite(invested) and latest_mark and _finite(latest_mark.get("value_usd")) is not None:
        moic = round(latest_mark["value_usd"] / invested, 2)
    decisions = decisions_store.list_decisions(company_id).get("items") or []
    latest_decision = decisions[0] if decisions else None
    return {
        "company_id": company_id,
        "company_name": company.get("name") or company_id,
        "ticker": company.get("ticker"),
        "position": position,
        "latest_kpi": latest_kpi,
        "latest_mark": latest_mark,
        "moic": moic,
        "latest_decision": {"verdict": latest_decision.get("verdict"), "decided_at": latest_decision.get("decided_at")} if latest_decision else None,
        "alerts": alerts_for(record, now=now),
        "kpi_count": len(record.get("kpis") or []),
        "update_count": len(record.get("updates") or []),
    }


def get_portfolio(company_id: str) -> dict:
    with _LOCK:
        record = _load(company_id)
    out = _summary(company_id, record)
    out["kpis"] = sorted(record["kpis"], key=lambda item: item.get("as_of") or "")
    out["updates"] = sorted(record["updates"], key=lambda item: item.get("as_of") or "", reverse=True)
    out["marks"] = sorted(record["marks"], key=lambda item: item.get("as_of") or "")
    return out


def _portfolio_company_ids() -> list[str]:
    ids: list[str] = []
    for company in storage.list_companies():
        cid = company.get("id")
        if not cid:
            continue
        try:
            if has_record(cid):
                ids.append(cid)
                continue
            items = decisions_store.list_decisions(cid).get("items") or []
        except ValueError:
            continue
        if items and items[0].get("verdict") == "invest":
            ids.append(cid)
    return sorted(set(ids))


def dashboard(*, now: datetime | None = None) -> dict:
    now = now or datetime.now(timezone.utc)
    rows = []
    for cid in _portfolio_company_ids():
        with _LOCK:
            record = _load(cid)
        rows.append(_summary(cid, record, now=now))
    invested = sum(r["position"].get("invested_usd") or 0 for r in rows)
    marked_rows = [r for r in rows if r["latest_mark"]]
    marked_value = sum(r["latest_mark"]["value_usd"] for r in marked_rows)
    moic_rows = [r for r in marked_rows if (r["position"].get("invested_usd") or 0) > 0]
    moic_value = sum(r["latest_mark"]["value_usd"] for r in moic_rows)
    marked_cost = sum(r["position"]["invested_usd"] for r in moic_rows)
    alerts = [dict(a, company_id=r["company_id"], company_name=r["company_name"]) for r in rows for a in r["alerts"]]
    severity_rank = {"high": 0, "medium": 1, "low": 2}
    alerts.sort(key=lambda a: (severity_rank.get(a["severity"], 9), a["company_name"]))
    reserves = reserves_plan(rows=rows)
    return {
        "generated_at": _now(),
        "companies": rows,
        "totals": {
            "company_count": len(rows),
            "invested_usd": invested,
            "marked_count": len(marked_rows),
            "marked_value_usd": marked_value if marked_rows else None,
            "marked_cost_usd": marked_cost if moic_rows else None,
            "marked_moic": round(moic_value / marked_cost, 2) if marked_cost else None,
            "marked_uncosted_count": len(marked_rows) - len(moic_rows),
            "alert_count": len(alerts),
            "high_alert_count": sum(1 for a in alerts if a["severity"] == "high"),
        },
        "alerts": alerts,
        "reserves": reserves,
    }


# ---- Reserves / follow-on planner --------------------------------------------------------


def get_reserves_settings() -> dict:
    data = storage._read_yaml_lenient(_reserves_file(), {}) or {}
    return {
        "fund_size_musd": data.get("fund_size_musd"),
        "reserve_pct": data.get("reserve_pct"),
        "notes": data.get("notes") or "",
        "updated_at": data.get("updated_at"),
    }


def save_reserves_settings(patch: dict) -> dict:
    with storage._WRITE_LOCK:
        storage._quarantine_unparseable(_reserves_file())
        current = get_reserves_settings()
        for key in ("fund_size_musd", "reserve_pct"):
            if key in patch:
                current[key] = _coerce(key, patch[key], float)
        if current.get("fund_size_musd") is not None and not current["fund_size_musd"] > 0:
            raise ValueError("fund_size_musd must be greater than 0")
        if "notes" in patch:
            current["notes"] = str(patch["notes"] or "")[:2000]
        if current.get("reserve_pct") is not None and not 0 <= current["reserve_pct"] <= 100:
            raise ValueError("reserve_pct must be between 0 and 100")
        current["updated_at"] = _now()
        _reserves_file().parent.mkdir(parents=True, exist_ok=True)
        storage._write_yaml(_reserves_file(), current)
    return reserves_plan()


def reserves_plan(*, rows: list[dict] | None = None) -> dict:
    settings = get_reserves_settings()
    if rows is None:
        rows = [_summary(cid, _load(cid)) for cid in _portfolio_company_ids()]
    planned = [
        {
            "company_id": r["company_id"],
            "company_name": r["company_name"],
            "invested_usd": r["position"].get("invested_usd"),
            "planned_follow_on_usd": r["position"].get("planned_follow_on_usd"),
            "ownership_pct": r["position"].get("ownership_pct"),
            "runway_alert": next((a for a in r["alerts"] if a["kind"] == "runway"), None),
        }
        for r in rows
    ]
    pool = None
    if settings.get("fund_size_musd") is not None and settings.get("reserve_pct") is not None:
        pool = settings["fund_size_musd"] * 1e6 * settings["reserve_pct"] / 100
    committed = sum(p["planned_follow_on_usd"] or 0 for p in planned)
    return {
        "settings": settings,
        "reserve_pool_usd": pool,
        "planned_follow_on_usd": committed,
        "remaining_usd": (pool - committed) if pool is not None else None,
        "companies": planned,
    }


# ---- LP tear sheet (DOCX) ----------------------------------------------------------------


def _fmt_money(value: Any) -> str:
    if value is None:
        return "—"
    try:
        value = float(value)
    except (TypeError, ValueError):
        return "—"
    if not math.isfinite(value):
        return "—"
    if abs(value) >= 1e9:
        return f"${value / 1e9:,.2f}B"
    if abs(value) >= 1e6:
        return f"${value / 1e6:,.1f}M"
    if abs(value) >= 1e3:
        return f"${value / 1e3:,.0f}K"
    return f"${value:,.0f}"


def tear_sheet_docx(company_id: str) -> bytes:
    data = get_portfolio(company_id)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Helvetica"
    style.font.size = Pt(10)
    doc.add_heading(f"{data['company_name']} — LP tear sheet", level=1)
    doc.add_paragraph(f"Generated {_now()[:10]} · every figure below is on record; blanks are not estimated.")

    position = data["position"]
    doc.add_heading("Position", level=2)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in (
        ("Round", position.get("round") or "—"),
        ("Security", position.get("security") or "—"),
        ("Entry date", (position.get("entry_date") or "—")[:10]),
        ("Invested", _fmt_money(position.get("invested_usd"))),
        ("Ownership", f"{position['ownership_pct']:.2f}%" if _finite(position.get("ownership_pct")) is not None else "—"),
        ("Entry post-money", _fmt_money(position.get("entry_post_money_usd"))),
        ("Board seat", "Yes" if position.get("board_seat") else "No"),
        ("Lead", "Yes" if position.get("lead") else "No"),
        ("Latest mark", f"{_fmt_money(data['latest_mark']['value_usd'])} ({data['latest_mark']['basis']}, {data['latest_mark']['as_of'][:10]})" if data["latest_mark"] else "No mark recorded"),
        ("MOIC", f"{data['moic']:.2f}x" if _finite(data["moic"]) is not None else "—"),
    ):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)

    doc.add_heading("KPI history", level=2)
    kpis = data["kpis"]
    if not kpis:
        doc.add_paragraph("No KPIs recorded.")
    else:
        headers = ["As of", "ARR", "Revenue", "Burn / mo", "Cash", "Runway", "Headcount", "Customers", "Source"]
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Grid Accent 1"
        for i, h in enumerate(headers):
            table.rows[0].cells[i].text = h
        for row in kpis:
            cells = table.add_row().cells
            cells[0].text = (row.get("as_of") or "")[:10]
            cells[1].text = _fmt_money(row.get("arr_usd"))
            cells[2].text = _fmt_money(row.get("revenue_usd"))
            cells[3].text = _fmt_money(row.get("burn_usd_month"))
            cells[4].text = _fmt_money(row.get("cash_usd"))
            cells[5].text = f"{row['runway_months']:g} mo" if _finite(row.get("runway_months")) is not None else "—"
            cells[6].text = str(row.get("headcount") or "—")
            cells[7].text = str(row.get("customers") or "—")
            cells[8].text = str(row.get("source") or "")

    doc.add_heading("Marks", level=2)
    if not data["marks"]:
        doc.add_paragraph("No marks recorded.")
    else:
        for mark in data["marks"]:
            doc.add_paragraph(f"{mark['as_of'][:10]}: {_fmt_money(mark['value_usd'])} — {mark['basis']}{(' · ' + mark['note']) if mark.get('note') else ''}", style="List Bullet")

    doc.add_heading("Open alerts", level=2)
    if not data["alerts"]:
        doc.add_paragraph("None.")
    for alert in data["alerts"]:
        doc.add_paragraph(f"[{alert['severity'].upper()}] {alert['label']} — {alert['detail']}", style="List Bullet")

    doc.add_heading("Latest founder update", level=2)
    if data["updates"]:
        latest = data["updates"][0]
        doc.add_paragraph(f"{latest['as_of'][:10]} · {latest.get('subject') or latest.get('source')}")
        doc.add_paragraph(latest["text"][:3000])
    else:
        doc.add_paragraph("None on file.")

    record = decisions_store.render_decision_record_md(company_id, max_chars=2000)
    if record:
        doc.add_heading("Decision record", level=2)
        doc.add_paragraph(record)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
