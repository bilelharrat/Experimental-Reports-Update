"""Turn an attached document into text the model can actually read.

Claude's Read tool opens plain text, images and PDFs, but a ``.docx`` is
a ZIP of XML — Read reports it as binary, which is how an analyst who
attached a Word file got "I wasn't able to open that .docx directly"
instead of an answer. So anything whose bytes aren't already readable is
converted here, once, at staging time, and the text is written beside
the file as ``<sha>.extract.txt``.

The ask then inlines that text into the prompt, which buys three things
beyond the fix: the answer no longer depends on the Read tool at all
(iOS Ask runs with tools off), it no longer depends on Claude (Gemini
can answer from the same text while Claude is resting), and it costs one
fewer round trip.

Formats that carry no extractable text — images, and PDFs that are pure
scans — are left to the model's own eyes and reported by
``needs_native_read``.
"""

from __future__ import annotations

import io
import logging
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

logger = logging.getLogger(__name__)


# ---- Caps ---------------------------------------------------------------

# What we keep on disk per attachment...
MAX_TEXT_CHARS = 400_000
# ...and what we're willing to spend on the prompt itself. A real term
# sheet runs ~55k characters, so the per-file cap has to clear that
# comfortably or the answer is about half a document. Anything past it
# stays in the sidecar, which the prompt points at.
INLINE_CHARS_PER_FILE = 120_000
INLINE_CHARS_TOTAL = 240_000

_XLSX_MAX_ROWS = 400
_XLSX_MAX_COLS = 40
_PDF_MAX_PAGES = 300
_CONVERT_TIMEOUT_S = 25.0

TEXT_SUFFIX = ".extract.txt"


# ---- Public API ---------------------------------------------------------


def extract(data: bytes, mime: str, *, filename: str | None = None) -> str:
    """Best-effort plain text for ``data``.

    Returns "" when the format carries no text we can pull out (an image)
    or when conversion failed (a scanned PDF, a legacy .doc on a box with
    no converter). Never raises: an attachment that resists extraction is
    still worth staging.
    """
    try:
        text = _dispatch(data, mime, filename)
    except Exception:  # noqa: BLE001 - one bad file must not fail the upload
        logger.warning("attachment text extraction failed (%s)", mime, exc_info=True)
        return ""
    return _tidy(text)[:MAX_TEXT_CHARS]


def text_path(attachments_dir: Path, stored_name: str) -> Path:
    """Where ``<sha>.docx``'s extracted text lives."""
    return Path(attachments_dir) / (_stem(stored_name) + TEXT_SUFFIX)


def write_text(attachments_dir: Path, stored_name: str, text: str, *, display_name: str) -> int:
    """Write the sidecar for ``stored_name``. Returns the character count."""
    if not text:
        return 0
    path = text_path(attachments_dir, stored_name)
    path.parent.mkdir(parents=True, exist_ok=True)
    body = f"# {display_name}\n\n{text}\n"
    if len(text) >= MAX_TEXT_CHARS:
        body += f"\n[truncated at {MAX_TEXT_CHARS:,} characters]\n"
    tmp = path.with_suffix(".tmp")
    tmp.write_text(body, encoding="utf-8")
    tmp.replace(path)
    return len(text)


def read_text(attachments_dir: Path, stored_name: str) -> tuple[str, str]:
    """``(display_name, text)`` from a sidecar, or ``("", "")``."""
    path = text_path(attachments_dir, stored_name)
    try:
        raw = path.read_text(encoding="utf-8")
    except OSError:
        return "", ""
    name = ""
    if raw.startswith("# "):
        head, _, rest = raw.partition("\n")
        name = head[2:].strip()
        raw = rest.lstrip("\n")
    return name, raw.strip()


def needs_native_read(attachments_dir: Path, stored_names: list[str]) -> bool:
    """True when some attachment can only be understood by looking at it —
    an image, or a PDF nothing could be read out of. Those questions have
    to go to Claude; everything else any engine can answer from the text.
    """
    return any(
        not text_path(attachments_dir, name).exists()
        for name in stored_names
    )


def prompt_block(attachments_dir: Path, stored_names: list[str]) -> str:
    """The block appended to a question that carries files.

    Text we already have is inlined (that is the whole point — the model
    should not have to open anything); the rest is named as a path under
    ``attachments/``, which is inside the session workdir Claude is given
    with ``--add-dir``.
    """
    if not stored_names:
        return ""
    attachments_dir = Path(attachments_dir)
    sections: list[str] = []
    budget = INLINE_CHARS_TOTAL
    for stored in stored_names:
        display, text = read_text(attachments_dir, stored)
        label = display or stored
        if not text:
            sections.append(
                f"- {label}: `attachments/{stored}` — open it with the Read tool."
            )
            continue
        allowance = min(INLINE_CHARS_PER_FILE, max(budget, 0))
        shown = text[:allowance]
        budget -= len(shown)
        note = ""
        if len(shown) < len(text):
            note = (
                f"\n[{label} truncated here — the full text is in "
                f"`attachments/{text_path(attachments_dir, stored).name}`]"
            )
        sections.append(
            f"--- BEGIN {label} ---\n{shown}{note}\n--- END {label} ---"
        )
    return (
        "\n\nAttached to this question ("
        + str(len(stored_names))
        + " file"
        + ("" if len(stored_names) == 1 else "s")
        + "):\n\n"
        + "\n\n".join(sections)
    )


# ---- Extraction ---------------------------------------------------------


def _dispatch(data: bytes, mime: str, filename: str | None) -> str:
    if mime.startswith("image/"):
        return ""
    if mime == "application/pdf":
        return _from_pdf(data)
    if mime.endswith("wordprocessingml.document"):
        return _from_docx(data)
    if mime.endswith("presentationml.presentation"):
        return _from_pptx(data)
    if mime.endswith("spreadsheetml.sheet"):
        return _from_xlsx(data)
    if mime == "text/html":
        return _from_html(data)
    if mime in ("application/rtf", "text/rtf"):
        return _from_rtf(data, filename)
    if mime == "application/msword":
        return _convert_with_cli(data, filename or "document.doc")
    return _decode(data)


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    out: list[str] = []
    for index, page in enumerate(reader.pages[:_PDF_MAX_PAGES], start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - a bad page shouldn't lose the rest
            text = ""
        if text.strip():
            out.append(f"## Page {index}\n{text.strip()}")
    return "\n\n".join(out)


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_runs(el) -> str:
    parts: list[str] = []
    for node in el.iter():
        if node.tag == _W + "t":
            parts.append(node.text or "")
        elif node.tag == _W + "tab":
            parts.append("\t")
        elif node.tag in (_W + "br", _W + "cr"):
            parts.append("\n")
    return "".join(parts)


def _docx_walk(el, lines: list[str]) -> None:
    for child in el:
        tag = child.tag
        if tag == _W + "p":
            lines.append(_docx_runs(child))
        elif tag == _W + "tbl":
            # Term sheets keep their numbers in tables, so rows stay rows.
            for row in child.findall(_W + "tr"):
                cells = [
                    " ".join(_docx_runs(cell).split())
                    for cell in row.findall(_W + "tc")
                ]
                lines.append(" | ".join(cells))
            lines.append("")
        else:
            # Content controls (w:sdt) and section wrappers hold real
            # paragraphs one level down.
            _docx_walk(child, lines)


def _from_docx(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as bundle:
        if "word/document.xml" not in bundle.namelist():
            return ""
        root = ET.fromstring(bundle.read("word/document.xml"))
    body = root.find(_W + "body")
    lines: list[str] = []
    _docx_walk(body if body is not None else root, lines)
    return "\n".join(lines)


def _from_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        lines = [f"## Slide {index}"]
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    lines.append(
                        " | ".join(" ".join(c.text.split()) for c in row.cells)
                    )
        notes = slide.notes_slide if slide.has_notes_slide else None
        if notes is not None and notes.notes_text_frame.text.strip():
            lines.append(f"Speaker notes: {notes.notes_text_frame.text.strip()}")
        if len(lines) > 1:
            out.append("\n".join(lines))
    return "\n\n".join(out)


_S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_PKG_R = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def _from_xlsx(data: bytes) -> str:
    """Sheet by sheet, rows as tab-separated lines.

    Done with the standard library on purpose: openpyxl isn't a
    dependency, and a cap table only needs its values, not its styles.
    """
    with zipfile.ZipFile(io.BytesIO(data)) as bundle:
        names = set(bundle.namelist())
        if "xl/workbook.xml" not in names:
            return ""
        shared: list[str] = []
        if "xl/sharedStrings.xml" in names:
            root = ET.fromstring(bundle.read("xl/sharedStrings.xml"))
            for item in root:
                shared.append("".join(t.text or "" for t in item.iter(_S + "t")))
        targets: dict[str, str] = {}
        if "xl/_rels/workbook.xml.rels" in names:
            rels = ET.fromstring(bundle.read("xl/_rels/workbook.xml.rels"))
            for rel in rels:
                target = rel.get("Target") or ""
                if target.startswith("/"):
                    target = target[1:]
                elif not target.startswith("xl/"):
                    target = "xl/" + target.lstrip("./")
                targets[rel.get("Id") or ""] = target
        workbook = ET.fromstring(bundle.read("xl/workbook.xml"))
        out: list[str] = []
        sheets = workbook.find(_S + "sheets")
        for sheet in list(sheets if sheets is not None else []):
            path = targets.get(sheet.get(_R + "id") or "")
            if not path or path not in names:
                continue
            rows = _xlsx_rows(ET.fromstring(bundle.read(path)), shared)
            if rows:
                out.append(f"## {sheet.get('name') or path}\n" + "\n".join(rows))
        return "\n\n".join(out)


def _xlsx_rows(sheet_root, shared: list[str]) -> list[str]:
    rows: list[str] = []
    data_el = sheet_root.find(_S + "sheetData")
    rows_el = list(data_el if data_el is not None else [])
    for row in rows_el[:_XLSX_MAX_ROWS]:
        cells: list[str] = []
        for cell in list(row)[:_XLSX_MAX_COLS]:
            cells.append(_xlsx_cell(cell, shared))
        while cells and not cells[-1]:
            cells.pop()
        if cells:
            rows.append("\t".join(cells))
    return rows


def _xlsx_cell(cell, shared: list[str]) -> str:
    kind = cell.get("t") or "n"
    if kind == "s":
        value = cell.findtext(_S + "v") or ""
        try:
            return shared[int(value)]
        except (ValueError, IndexError):
            return ""
    if kind == "inlineStr":
        node = cell.find(_S + "is")
        return "".join(t.text or "" for t in node.iter(_S + "t")) if node is not None else ""
    return (cell.findtext(_S + "v") or "").strip()


def _from_html(data: bytes) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(data), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    return soup.get_text("\n")


_RTF_CONTROL = re.compile(r"\\(?:[a-z]+-?\d*|[^a-z])")


def _from_rtf(data: bytes, filename: str | None) -> str:
    converted = _convert_with_cli(data, filename or "document.rtf")
    if converted:
        return converted
    # No converter on the box: strip the control words and keep the prose.
    text = _decode(data)
    text = re.sub(r"\{\\\*.*?\}", " ", text, flags=re.S)
    text = _RTF_CONTROL.sub(" ", text)
    return text.replace("{", " ").replace("}", " ")


def _convert_with_cli(data: bytes, filename: str) -> str:
    """macOS ``textutil`` first, then LibreOffice — the same ladder
    ``docx_pdf`` climbs for the other direction.
    """
    import shutil

    suffix = Path(filename).suffix or ".doc"
    with tempfile.TemporaryDirectory(prefix="bsh-attach-") as tmp:
        src = Path(tmp) / f"input{suffix}"
        src.write_bytes(data)
        textutil = shutil.which("textutil")
        if textutil:
            try:
                done = subprocess.run(
                    [textutil, "-convert", "txt", "-stdout", "-encoding", "30", str(src)],
                    capture_output=True,
                    timeout=_CONVERT_TIMEOUT_S,
                )
                if done.returncode == 0 and done.stdout.strip():
                    return done.stdout.decode("utf-8", "replace")
            except (OSError, subprocess.SubprocessError):
                logger.warning("textutil conversion failed for %s", filename, exc_info=True)
        from .docx_pdf import soffice_path

        soffice = soffice_path()
        if soffice:
            out_dir = Path(tmp) / "out"
            out_dir.mkdir(exist_ok=True)
            try:
                subprocess.run(
                    [soffice, "--headless", "--convert-to", "txt:Text",
                     "--outdir", str(out_dir), str(src)],
                    capture_output=True,
                    timeout=_CONVERT_TIMEOUT_S,
                )
            except (OSError, subprocess.SubprocessError):
                logger.warning("soffice conversion failed for %s", filename, exc_info=True)
            else:
                for produced in out_dir.glob("*.txt"):
                    return produced.read_text(encoding="utf-8", errors="replace")
    return ""


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", "replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1", "replace")


def _tidy(text: str) -> str:
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    lines = [line.rstrip() for line in text.split("\n")]
    out: list[str] = []
    blanks = 0
    for line in lines:
        if line:
            blanks = 0
            out.append(line)
            continue
        blanks += 1
        if blanks <= 1:
            out.append("")
    return "\n".join(out).strip()


def _stem(stored_name: str) -> str:
    return stored_name.rsplit(".", 1)[0] if "." in stored_name else stored_name


__all__ = [
    "MAX_TEXT_CHARS",
    "TEXT_SUFFIX",
    "extract",
    "needs_native_read",
    "prompt_block",
    "read_text",
    "text_path",
    "write_text",
]
