"""Deterministic fact check: every figure a memo states must be on file.

What counts as "on file" (the corpus): the curated fact ledger, the
research folder's documents and digests, the company's registry entry,
the source cache (what retrieval fetched — ``source_cache``), the sources
frozen with this run, and the firm's own records (transcripts, founder
updates, KPI rows, reference calls, evidence excerpts). Model output is
never corpus: the analysis-pass artifacts, the spine's pins and the
memo's own source titles cannot vouch for a number the model wrote.

Each figure ends up in one bucket:

- ``verified``    the source the sentence cites carries it;
- ``supported``   some source on file carries it — the same spelling, or
                  the same value in another spelling within one rounding
                  step ("$2.6B" is carried by "$2,580 million");
- ``derived``     a calculation note carries it;
- ``unsupported`` none of the above — the finding the gate reports.

The extractor is conservative on purpose, like the pin check: money,
percentages, multiples, scaled amounts ("40 million"), thousands-separated
counts and counted nouns ("30 customers", "95+ patents") are checked;
bare small integers, years, ids and citation tokens are not, because a
false positive costs a repair round and a false negative costs nothing
the memo did not already have.

Those four buckets are what the repair enforces, and they are loose on
purpose (see above). Next to them the check reports, never enforces, an
honest split a reader can trust (``tiers`` in ``to_dict``;
:func:`summarize_fact_check` is the reader-facing summary):

- ``verified``         a cited source that is not the subject's own site
                       carries the figure;
- ``found_elsewhere``  independent evidence on file carries it (and
                       ``in_context`` counts those where a word from the
                       memo's own phrasing sits within ~120 characters of
                       the number in the source);
- ``derived``          a calculation note carries it;
- ``company_reported`` only pages on the subject company's own domain do;
- ``registry_only``    only registry fields that carry ``source_refs`` do;
- ``not_traced``       nothing above — including figures the loose match
                       accepts only through the whole registry dump, the
                       company description, or a bare number that is a
                       date, a page, a section or a footnote ("30 June",
                       "p. 12", "No. 7", "§ 4", "[12]").

The run's source URLs are audited the same way (:func:`audit_source_urls`):
each is marked seen or unseen, and a bare homepage reused by two or more
sources is flagged. Nothing is ever stripped — a dropped URL would fail the
renderer's source-URL rule.
"""
from __future__ import annotations

import bisect
import json
import logging
import math
import os
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

import yaml

from . import company_paths, numbers_lint, research_store, source_cache, storage

logger = logging.getLogger(__name__)

# Below this much evidence text (ledger, documents, retrieved sources,
# firm records) the check still reports, but nothing is enforced: with
# nothing to check against, every figure is "unsupported" and a repair
# would strip a memo of its numbers for no gain.
THIN_CORPUS_CHARS = 2_000
# Enforcement in "auto" mode needs a corpus this rich and this much of
# the memo already traceable before an unsupported figure is worth a
# repair round: past that point the figure, not the corpus, is the problem.
AUTO_ENFORCE_MIN_CHARS = 50_000
AUTO_ENFORCE_MIN_COVERAGE = 60
MAX_FED_FINDINGS = 12
MAX_DOC_CHARS = 300_000
# Report-only tiers: rows kept for the viewer, and how close a word from the
# memo's phrasing must sit to the number in a source to be "in context".
MAX_FIGURE_TIER_ROWS = 80
CONTEXT_WINDOW_CHARS = 120
# Positions of one number checked per figure for the in-context tier; a
# common number ("10") can occur thousands of times in a large corpus.
MAX_CONTEXT_POSITIONS = 400

_CITATION_RE = re.compile(r"\[(?:[SC]\d+)(?:\s*,\s*[SC]\d+)*\]")
_CITATION_ID_RE = re.compile(r"[SC]\d+")
_URL_RE = re.compile(r"https?://\S+", re.I)
_YEAR_RE = re.compile(r"^(?:19|20)\d{2}$")

_CURRENCY = r"(?:NT\$|NTD\s?|US\$|U\.S\.\$|USD\s?|\$|€|EUR\s?|£|GBP\s?|¥|JPY\s?|RMB\s?|CN¥|HK\$|A\$|C\$|S\$)"
_NUM = r"(?:\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?)"
_SCALE = r"(?:trillion|billion|million|thousand|tn|bn|mm|mn|k|m|b|t)"
_PCT = r"(?:%|percent(?:age points)?|pct\.?|pp|bps|basis points)"
_COUNT_NOUNS = (
    "customers|clients|patents|employees|staff|engineers|headcount|countries|"
    "markets|sites|stores|units|vehicles|contracts|partners|users|subscribers|"
    "deployments|locations|facilities|plants|hospitals|banks|logos|enterprises|"
    "installations|devices|sensors|satellites|aircraft|ships|vessels|trucks|"
    "robots|patients|trials|months|enterprises|accounts|merchants|suppliers|"
    "distributors|dealers|branches|offices|schools|universities|cities|states|"
    "projects|pilots|licenses|licences|shareholders|investors|founders|"
    "team members|people|megawatts|gigawatts|mw|gw|kwh|mwh|gwh|tons|tonnes|"
    "barrels|acres|hectares|square feet|sq ft|km|miles|nodes|towers|rigs|wells"
)

# A multiple followed on the same line by another number is a dimension
# ("2x 4"), not a figure; on the next line it is the next cell of a table.
# A year after it is what the multiple is of ("12x 2025E EBITDA").
_NOT_A_DIMENSION = r"(?![ \t]?[$€£¥])(?![ \t]?(?!(?:19|20)\d\d(?!\d))\d)"
# Digit runs this long are identifiers (hashes, tracking codes), never
# figures, and the figure pattern's time grows with their square.
_LONG_DIGITS_RE = re.compile(r"\d{25,}")
_FIGURE_RE = re.compile(
    rf"(?P<pct_range>{_NUM}\s?(?:–|—|-|to)\s?{_NUM}\s?{_PCT})(?![\w])"
    rf"|(?P<mult_range>{_NUM}\s?(?:–|—|-|to)\s?{_NUM}\s?(?:x|×))(?![\w]){_NOT_A_DIMENSION}"
    rf"|(?P<range>{_CURRENCY}?\s?{_NUM}\s?(?:–|—|-|to)\s?{_CURRENCY}?\s?{_NUM}\s?{_SCALE}?)(?![\w%.])"
    rf"|(?P<money>{_CURRENCY}\s?{_NUM}\s?{_SCALE}?)(?![\w%])"
    rf"|(?P<pct>{_NUM}\s?{_PCT})(?![\w])"
    rf"|(?P<mult>{_NUM}\s?(?:x|×))(?![\w]){_NOT_A_DIMENSION}"
    rf"|(?P<scaled>{_NUM}\s?{_SCALE})(?![\w])"
    rf"|(?P<thousands>\d{{1,3}}(?:,\d{{3}})+)(?![\w.]\d)"
    rf"|(?P<plus>\d{{2,}}(?:,\d{{3}})*\+)(?![\w])"
    rf"|(?P<counted>\d{{2,}}(?:,\d{{3}})*)(?=\s+(?:[a-z][\w-]*\s+){{0,2}}(?:{_COUNT_NOUNS})\b)",
    re.I,
)
_BARE_RE = re.compile(r"(?<![\w.,])\d{2,}(?:,\d{3})*(?:\.\d+)?(?![\w.])")
# A count preceded by one of these is an ordinal, a label or a window, not
# a fact about the company ("top 10 customers", "past 12 months", "page 12").
_NON_FACT_LEADERS = {
    "top", "first", "last", "next", "past", "prior", "previous", "another",
    "every", "each", "per", "page", "figure", "table", "section", "chapter",
    "step", "phase", "round", "series", "tier", "level", "grade", "class",
    "q", "fy", "h", "week", "day", "year", "item", "note", "rule", "form",
}
_SCALES = {
    "k": 1e3, "thousand": 1e3,
    "m": 1e6, "mm": 1e6, "mn": 1e6, "million": 1e6,
    "b": 1e9, "bn": 1e9, "billion": 1e9,
    "t": 1e12, "tn": 1e12, "trillion": 1e12,
}


@dataclass
class Figure:
    raw: str
    value: float
    klass: str  # "amount" (money or count), "pct", "mult"
    start: int
    end: int
    mantissa: str
    plus: bool = False

    @property
    def forms(self) -> set[str]:
        return numbers_lint._digits_forms(self.raw.rstrip("+"))


def _strip_citations(text: str) -> str:
    """Blank citation tokens and URLs so their digits are never figures."""
    out = _CITATION_RE.sub(lambda m: " " * len(m.group(0)), text)
    return _URL_RE.sub(lambda m: " " * len(m.group(0)), out)


def _parse_number(raw: str) -> tuple[float, str] | None:
    """``(value, mantissa)`` of a number with optional currency/scale."""
    text = raw.strip().rstrip("+").lower()
    text = re.sub(rf"^{_CURRENCY.lower()}\s?", "", text, flags=re.I)
    match = re.match(rf"({_NUM})\s?([a-z%.]*(?:\s+[a-z]+)*)$", text)
    if not match:
        return None
    mantissa = match.group(1)
    try:
        value = float(mantissa.replace(",", ""))
    except ValueError:
        return None
    unit = (match.group(2).split() or [""])[0].rstrip(".")
    return value * _SCALES.get(unit, 1.0), mantissa


def _significant(mantissa: str) -> tuple[int, int]:
    """Bounds on how many significant digits ``mantissa`` states.

    "500" may be "about five hundred" (2 significant, once the leading
    digit is joined by its neighbour) or exactly 500 (3); "2.58" is 3. The
    lower bound is never below 2: one significant digit would let $450M
    stand in for $500M, and those were different facts in the run that
    motivated this module.
    """
    digits = mantissa.replace(",", "").replace(".", "").lstrip("0")
    total = max(1, len(digits))
    nonzero = max(1, len(digits.rstrip("0")))
    lo = max(2, nonzero)
    return lo, max(lo, total)


def _rounds_to(value: float, target: float, sig: int) -> bool:
    if value == 0 or target == 0:
        return value == target
    if (value < 0) != (target < 0):
        return False
    value, target = abs(value), abs(target)
    magnitude = math.floor(math.log10(value))
    step = 10 ** (magnitude - sig + 1)
    rounded = math.floor(value / step + 0.5) * step
    return math.isclose(rounded, target, rel_tol=1e-9, abs_tol=1e-12)


def _value_supported(figure: Figure, values: dict[str, list[float]]) -> bool:
    """Whether some corpus value of a compatible class rounds to the figure."""
    candidates = values.get(figure.klass) or []
    if not candidates:
        return False
    lo, hi = _significant(figure.mantissa)
    target = figure.value
    for value in candidates:
        if math.isclose(value, target, rel_tol=1e-9, abs_tol=1e-12):
            return True
        for sig in range(lo, hi + 1):
            if _rounds_to(value, target, sig):
                return True
    return False


def extract_figures(text: str) -> list[Figure]:
    """Figures the memo asserts in ``text``, left to right."""
    cleaned = _strip_citations(text or "")
    cleaned = _LONG_DIGITS_RE.sub(lambda m: " " * len(m.group(0)), cleaned)
    out: list[Figure] = []
    for match in _FIGURE_RE.finditer(cleaned):
        kind = match.lastgroup or ""
        raw = match.group(0).strip()
        # A pattern may take the space before the number with it ("of
        # 20-25M"); the figure, and its glue test, start at the number.
        lead = len(match.group(0)) - len(match.group(0).lstrip())
        start, end = match.start() + lead, match.end()
        before = cleaned[max(0, start - 3):start]
        # "10-K", "FY24", "Q3", "S-1", "H1" style identifiers are not figures:
        # a letter or hyphen glued to the number. Any whitespace separates —
        # scraped comps tables put one cell per line ("$489M\n1.2x"), and
        # treating the newline as glue dropped every multiple in them. (\Z,
        # not $: "$" also matches before a trailing newline.)
        # A hyphen before a currency sign is a range's dash ("$315M-$1.4B"),
        # not glue.
        if re.search(r"[A-Za-z]\Z", before) or (
            before.endswith("-") and not re.match(_CURRENCY, raw, re.I)
        ):
            continue
        if kind in {"counted", "plus", "thousands"}:
            leader = re.search(r"([A-Za-z]+)\s+$", cleaned[max(0, start - 16):start])
            if leader and leader.group(1).lower() in _NON_FACT_LEADERS:
                continue
        if kind == "range":
            out.extend(_range_figures(raw, start, cleaned))
            continue
        if kind in {"pct_range", "mult_range"}:
            out.extend(_unit_range_figures(raw, start, kind))
            continue
        parsed = _parse_number(raw)
        if parsed is None:
            continue
        value, mantissa = parsed
        if kind == "pct":
            unit = raw.lower()
            if "bps" in unit or "basis" in unit:
                value = value / 100.0
            klass = "pct"
        elif kind == "mult":
            if value > 1000:
                continue
            klass = "mult"
        else:
            if kind in {"thousands", "counted", "plus"} and _YEAR_RE.match(mantissa.replace(",", "")):
                continue
            if kind == "counted" and value < 10:
                continue
            klass = "amount"
        out.append(
            Figure(
                raw=raw,
                value=value,
                klass=klass,
                start=start,
                end=end,
                mantissa=mantissa,
                plus=raw.endswith("+"),
            )
        )
    return out


def _range_figures(raw: str, start: int, text: str) -> list[Figure]:
    """"$20–25M" / "$1.2 to $1.5B" / "20-25%" as two figures sharing the
    scale and currency the range states once. A range with neither a
    currency nor a scale is a bare pair ("3-5") and is skipped."""
    match = re.match(
        rf"(?P<c1>{_CURRENCY})?\s?(?P<n1>{_NUM})\s?(?:–|—|-|to)\s?(?P<c2>{_CURRENCY})?\s?(?P<n2>{_NUM})\s?(?P<scale>{_SCALE})?$",
        raw,
        re.I,
    )
    if not match:
        return []
    currency = match.group("c1") or match.group("c2")
    scale = match.group("scale") or ""
    if not currency and not scale:
        return []
    out: list[Figure] = []
    for key in ("n1", "n2"):
        mantissa = match.group(key)
        if _YEAR_RE.match(mantissa) and not currency and not scale:
            continue
        parsed = _parse_number(f"{mantissa}{scale}")
        if parsed is None:
            continue
        piece = f"{currency or ''}{mantissa}{scale}"
        offset = raw.find(mantissa)
        out.append(
            Figure(
                raw=piece,
                value=parsed[0],
                klass="amount",
                start=start + max(0, offset),
                end=start + max(0, offset) + len(mantissa),
                mantissa=mantissa,
            )
        )
    return out


def _unit_range_figures(raw: str, start: int, kind: str) -> list[Figure]:
    """"20-25%" / "1.5–2.0x" as two figures sharing the trailing unit."""
    match = re.match(rf"(?P<n1>{_NUM})\s?(?:–|—|-|to)\s?(?P<n2>{_NUM})\s?(?P<unit>.+)$", raw, re.I)
    if not match:
        return []
    unit = match.group("unit").strip()
    out: list[Figure] = []
    for key in ("n1", "n2"):
        mantissa = match.group(key)
        parsed = _parse_number(mantissa)
        if parsed is None:
            continue
        value = parsed[0]
        if kind == "pct_range":
            klass = "pct"
            if "bps" in unit.lower() or "basis" in unit.lower():
                value = value / 100.0
        else:
            klass = "mult"
            if value > 1000:
                continue
        offset = raw.find(mantissa)
        out.append(
            Figure(
                raw=f"{mantissa}{unit if klass == 'pct' and unit == '%' else (' ' + unit if klass == 'pct' else unit)}",
                value=value,
                klass=klass,
                start=start + max(0, offset),
                end=start + max(0, offset) + len(mantissa),
                mantissa=mantissa,
            )
        )
    return out


@dataclass
class TextIndex:
    """What one text can vouch for: the exact spellings of its figures
    (in every canonical form) and their values by class. Membership is
    O(1), so a 200-figure memo checks against a multi-megabyte corpus in
    milliseconds instead of re-scanning it per figure.

    ``strict=True`` builds the report-only index behind the honest tiers:
    a bare number that is a date, a page, a section or a footnote marker
    ("30 June", "June 30", "2026-06-13", "p. 12", "No. 7", "§ 4", "[12]")
    vouches for nothing there. The default (loose) index is what the
    repair enforces and is unchanged."""

    forms: set[str] = field(default_factory=set)
    values: dict[str, list[float]] = field(default_factory=lambda: {"amount": [], "pct": [], "mult": []})

    @classmethod
    def of(cls, text: str, *, strict: bool = False) -> "TextIndex":
        loose, tight = _index_pair(text)
        return tight if strict else loose

    def merge(self, other: "TextIndex") -> None:
        self.forms |= other.forms
        for klass, found in other.values.items():
            self.values.setdefault(klass, []).extend(found)

    def supports(self, figure: Figure) -> bool:
        if figure.forms & self.forms:
            return True
        return _value_supported(figure, self.values)


def corpus_values(text: str, *, strict: bool = False) -> dict[str, list[float]]:
    """Every number a source text carries, by class — the permissive side
    of the match, so a memo's "30 customers" is carried by a KPI row that
    says only "30". ``strict`` drops bare numbers that are dates, pages,
    sections or footnote markers (report-only; see :class:`TextIndex`)."""
    loose, tight = _index_pair(text)
    return (tight if strict else loose).values


_MONTH_WORDS = (
    "january|february|march|april|may|june|july|august|september|october|"
    "november|december|jan|feb|mar|apr|jun|jul|aug|sept|sep|oct|nov|dec"
)
# A bare number right after one of these is a date, a page, a section or a
# footnote ("June 30", "p. 12", "No. 7", "§ 4", "[12]", "^3"), not a fact
# about the company. Anchored at the end of the text before the number.
_NOISE_BEFORE_RE = re.compile(
    rf"(?:\b(?:{_MONTH_WORDS})\.?|\bpages?|\bpp?\.|\bnos?\.|§+|¶|\bfn\.?|"
    r"\bfootnotes?|\bnotes?|\bref\.?|\[\^?|\^)\s*$",
    re.I,
)
# ... or right before one of these ("30 June", the "]" closing "[12]").
_NOISE_AFTER_RE = re.compile(rf"^(?:\s*(?:{_MONTH_WORDS})\b|\])", re.I)
_DATE_SPAN_RE = re.compile(r"\b\d{4}-\d{1,2}-\d{1,2}\b|\b\d{1,2}/\d{1,2}/\d{2,4}\b")


def _bare_is_noise(text: str, start: int, end: int, date_starts: list[int], date_spans: list[tuple[int, int]]) -> bool:
    """Whether the bare number at ``text[start:end]`` is a date, page,
    section or footnote marker rather than a quantity."""
    if _NOISE_BEFORE_RE.search(text[max(0, start - 16):start]):
        return True
    if _NOISE_AFTER_RE.match(text[end:end + 16]):
        return True
    index = bisect.bisect_right(date_starts, start) - 1
    return index >= 0 and date_spans[index][0] <= start < date_spans[index][1]


def _covered(position: int, starts: list[int], spans: list[tuple[int, int]]) -> bool:
    """Whether ``position`` falls inside one of the figure spans (sorted by
    start). A figure span is a few dozen characters at most, so only the
    spans starting within 64 characters before ``position`` can hold it."""
    index = bisect.bisect_right(starts, position) - 1
    while index >= 0 and starts[index] > position - 64:
        start, end = spans[index]
        if start <= position < end:
            return True
        index -= 1
    return False


def _index_pair(text: str) -> tuple[TextIndex, TextIndex]:
    """The loose (enforced) and the strict (report-only) index of one
    text, from a single scan. The loose index is exactly what
    ``TextIndex.of`` and ``corpus_values`` always built."""
    cleaned = _strip_citations(text or "")
    loose, tight = TextIndex(), TextIndex()
    spans: list[tuple[int, int]] = []
    for figure in extract_figures(cleaned):
        loose.values[figure.klass].append(figure.value)
        tight.values[figure.klass].append(figure.value)
        forms = figure.forms
        loose.forms |= forms
        tight.forms |= forms
        spans.append((figure.start, figure.end))
    spans.sort()
    starts = [start for start, _ in spans]
    date_spans: list[tuple[int, int]] | None = None
    date_starts: list[int] = []
    for match in _BARE_RE.finditer(cleaned):
        raw = match.group(0)
        form = raw.replace(",", "").lower()
        loose.forms.add(form)
        if date_spans is None:
            date_spans = [m.span() for m in _DATE_SPAN_RE.finditer(cleaned)]
            date_starts = [span[0] for span in date_spans]
        noise = _bare_is_noise(cleaned, match.start(), match.end(), date_starts, date_spans)
        if not noise:
            tight.forms.add(form)
        if _covered(match.start(), starts, spans):
            continue
        if _YEAR_RE.match(raw):
            continue
        try:
            value = float(raw.replace(",", ""))
        except ValueError:
            continue
        loose.values["amount"].append(value)
        if not noise:
            tight.values["amount"].append(value)
    return loose, tight


# ---- corpus ----------------------------------------------------------------


# Report-only tiers a corpus text can belong to. "independent" evidence and
# the subject's own pages ("company") come from the texts the loose index
# is built from; "registry" is the registry fields that carry source_refs,
# kept apart from the loose corpus (which reads the whole registry dump).
# "context" texts (the registry dump, the company description) feed only
# the loose index.
HONEST_TEXT_TIERS = ("independent", "company", "registry")


@dataclass
class Corpus:
    texts: list[tuple[str, str]] = field(default_factory=list)  # (label, text)
    evidence_labels: list[str] = field(default_factory=list)
    evidence_chars: int = 0
    index: TextIndex = field(default_factory=TextIndex)
    # --- report-only honesty tiers (never enforced) ---
    text_tiers: list[str] = field(default_factory=list)  # parallel to ``texts``
    tier_texts: list[tuple[str, str, str]] = field(default_factory=list)  # (tier, label, text), not in ``texts``
    tier_index: dict[str, TextIndex] = field(default_factory=dict)
    company_hosts: set[str] = field(default_factory=set)
    company_tokens: set[str] = field(default_factory=set)

    def add(self, label: str, text: str, *, evidence: bool = True, tier: str | None = None) -> None:
        text = str(text or "").strip()
        if not text:
            return
        if len(text) > MAX_DOC_CHARS:
            text = text[:MAX_DOC_CHARS]
        self.texts.append((label, text))
        self.text_tiers.append(tier or ("independent" if evidence else "context"))
        if evidence:
            self.evidence_labels.append(label)
            self.evidence_chars += len(text)

    def add_tier_text(self, tier: str, label: str, text: str) -> None:
        """A text that counts only for a report-only tier (the registry
        fields that carry source_refs), never for the enforced match."""
        text = str(text or "").strip()
        if text and tier in HONEST_TEXT_TIERS:
            self.tier_texts.append((tier, label, text[:MAX_DOC_CHARS]))

    def _tiers(self) -> list[str]:
        tiers = list(self.text_tiers[: len(self.texts)])
        evidence = set(self.evidence_labels)
        for label, _ in self.texts[len(tiers):]:
            tiers.append("independent" if label in evidence else "context")
        return tiers

    def finish(self) -> "Corpus":
        index = TextIndex()
        strict: dict[str, TextIndex] = {}
        for (_, text), tier in zip(self.texts, self._tiers()):
            loose, tight = _index_pair(text)
            index.merge(loose)
            if tier in HONEST_TEXT_TIERS:
                strict.setdefault(tier, TextIndex()).merge(tight)
        for tier, _, text in self.tier_texts:
            strict.setdefault(tier, TextIndex()).merge(TextIndex.of(text, strict=True))
        index.values = {klass: sorted(set(found)) for klass, found in index.values.items()}
        for tier_index in strict.values():
            tier_index.values = {klass: sorted(set(found)) for klass, found in tier_index.values.items()}
        self.index = index
        self.tier_index = strict
        return self

    def tier_texts_of(self, tier: str) -> list[str]:
        """The texts of one report-only tier (for the in-context check)."""
        out = [text for (_, text), t in zip(self.texts, self._tiers()) if t == tier]
        out.extend(text for t, _, text in self.tier_texts if t == tier)
        return out

    @property
    def thin(self) -> bool:
        return self.evidence_chars < THIN_CORPUS_CHARS

    def summary(self) -> str:
        parts = [f"{len(self.texts)} texts", f"{self.evidence_chars:,} evidence chars"]
        return ", ".join(parts)

    def describe(self) -> list[dict]:
        return [
            {"label": label, "chars": len(text), "tier": tier}
            for (label, text), tier in zip(self.texts, self._tiers())
        ]


_DIGEST_FILES = {"known_sources.md"}


def _research_folder_texts(research_dir: Path | None) -> list[tuple[str, str]]:
    """Text of every document and digest in the research folder, extracted
    the way the Gemini path reads them (PDF, DOCX, PPTX, text)."""
    if research_dir is None:
        return []
    root = Path(research_dir)
    if not root.is_dir():
        return []
    from . import memo_engine

    out: list[tuple[str, str]] = []
    try:
        entries = sorted(p for p in root.iterdir() if p.is_file())
    except OSError:
        return []
    for path in entries:
        name = path.name
        if name.startswith("index.yaml") or name.endswith(".progress.jsonl") or name in _DIGEST_FILES:
            continue
        try:
            text = memo_engine.file_text(path)
        except Exception:  # noqa: BLE001
            logger.warning("fact check: could not read %s", path, exc_info=True)
            continue
        if text.strip():
            out.append((f"research file: {name}", text))
    return out


def _host_of(url: Any) -> str:
    return storage._normalize_host(url) if url else ""


def _on_hosts(host: str, hosts: set[str]) -> bool:
    """Whether ``host`` is one of ``hosts`` or a subdomain of one."""
    return bool(host) and any(host == h or host.endswith("." + h) for h in hosts)


def _record_hosts(record: dict) -> set[str]:
    """The subject company's own web domains, from its registry record."""
    return {
        host
        for host in (_host_of(record.get("website")), _host_of(record.get("logo_domain")))
        if host and "." in host
    }


def _record_name_tokens(record: dict, company_id: str) -> set[str]:
    """Words that name the subject company. They sit next to nearly every
    number in its coverage, so they never count as metric context."""
    names = [company_id, record.get("name"), record.get("legal_name"), *(record.get("aliases") or [])]
    tokens: set[str] = set()
    for name in names:
        if isinstance(name, str):
            tokens |= {_stem(token) for token in re.findall(r"[a-z][a-z0-9&'-]+", name.lower())}
    return tokens


def _url_from_label(label: str) -> str:
    """The URL a ``cached source …: <url>`` / ``run source …: <url>`` label names."""
    tail = label.split(": ", 1)[-1].strip() if ": " in label else ""
    return tail if tail.lower().startswith(("http://", "https://")) else ""


def _source_text_tier(label: str, company_hosts: set[str]) -> str:
    """A retrieved page on the subject's own domain is company-reported;
    anything else retrieval found is independent evidence."""
    host = _host_of(_url_from_label(label))
    return "company" if _on_hosts(host, company_hosts) else "independent"


def _registry_sourced_texts(record: Any) -> list[tuple[str, str]]:
    """``(path, text)`` for every registry field that carries
    ``source_refs`` — the only registry content the report-only
    ``registry`` tier reads (the loose corpus reads the whole dump)."""
    out: list[tuple[str, str]] = []

    def flatten(node: Any) -> list[str]:
        if isinstance(node, dict):
            parts: list[str] = []
            for key, value in node.items():
                if key != "source_refs":
                    parts.extend(flatten(value))
            return parts
        if isinstance(node, list):
            parts = []
            for item in node:
                parts.extend(flatten(item))
            return parts
        if node is None or isinstance(node, bool):
            return []
        return [str(node)]

    def walk(node: Any, path: str) -> None:
        if isinstance(node, dict):
            refs = node.get("source_refs")
            if isinstance(refs, list) and any(refs):
                text = " ; ".join(part for part in flatten(node) if part.strip())
                if text:
                    out.append((path or "entry", text))
            for key, value in node.items():
                if key != "source_refs":
                    walk(value, f"{path}.{key}" if path else str(key))
        elif isinstance(node, list):
            for index, item in enumerate(node):
                walk(item, f"{path}[{index}]")

    walk(record, "")
    return out


def _firm_record_rows(texts: list[str], labels: list[str]) -> list[tuple[str, str]]:
    """``(label, text)`` pairs from ``numbers_lint._corpus``. It returns more
    texts than labels (the company record is two texts — description and
    round — under one label, and several kinds share one label), so the
    old index-clamped pairing labelled the round as evidence. The company
    record's two texts are paired explicitly; the rest keep the clamped
    pairing (all of them are evidence either way)."""
    if len(texts) == len(labels):
        return list(zip(labels, texts))
    rows: list[tuple[str, str]] = []
    rest_texts, rest_labels = list(texts), list(labels)
    if rest_labels and rest_labels[0] == "company record" and len(rest_texts) >= 2:
        rows = [("company record", rest_texts[0]), ("company record", rest_texts[1])]
        rest_texts, rest_labels = rest_texts[2:], rest_labels[1:]
    for index, text in enumerate(rest_texts):
        label = rest_labels[min(index, len(rest_labels) - 1)] if rest_labels else "firm record"
        rows.append((label, text))
    return rows


def build_corpus(
    company_id: str,
    *,
    run_dir: Path | None = None,
    research_dir: Path | None = None,
) -> Corpus:
    """Everything on file for ``company_id`` that can vouch for a figure."""
    corpus = Corpus()
    if research_dir is None:
        try:
            research_dir = research_store.RESEARCH_ROOT / company_paths.storage_key(company_id)
        except ValueError:
            research_dir = None
    for label, text in _research_folder_texts(research_dir):
        corpus.add(label, text)
    try:
        from . import claude_runner

        entry = claude_runner._extract_company_registry_entry_yaml(
            storage.COMPANIES_FILE, company_id
        )
        if entry:
            corpus.add("registry entry", entry, evidence=False)
            record = yaml.safe_load(entry)
            if isinstance(record, dict):
                corpus.company_hosts = _record_hosts(record)
                corpus.company_tokens = _record_name_tokens(record, company_id)
                for path, text in _registry_sourced_texts(record):
                    corpus.add_tier_text("registry", f"registry field: {path}", text)
    except Exception:  # noqa: BLE001
        logger.debug("fact check: no registry entry for %s", company_id, exc_info=True)
    for label, text in source_cache.corpus_texts(company_id):
        corpus.add(label, text, tier=_source_text_tier(label, corpus.company_hosts))
    if run_dir is not None:
        cached_files = {
            str(row.get("file"))
            for row in source_cache.list_sources(company_id)
        }
        for label, text in source_cache.run_source_texts(run_dir):
            file_hint = label.split("run source ", 1)[-1].split(":", 1)[0]
            if f"{file_hint}.txt" in cached_files:
                continue  # the company cache already carries this text
            corpus.add(label, text, tier=_source_text_tier(label, corpus.company_hosts))
    try:
        texts, labels = numbers_lint._corpus(company_id)
    except Exception:  # noqa: BLE001
        logger.debug("fact check: firm corpus unavailable for %s", company_id, exc_info=True)
        texts, labels = [], []
    for label, text in _firm_record_rows(list(texts), list(labels)):
        if not str(text or "").strip():
            continue
        corpus.add(f"firm record: {label}", text, evidence=label != "company record")
    return corpus.finish()


# ---- the check --------------------------------------------------------------


@dataclass
class FactFinding:
    code: str
    section_id: str
    location: str
    figure: str
    snippet: str
    detail: str
    severity: str = "P0"
    citations: list[str] = field(default_factory=list)
    # The figure sits where a reader takes it as the memo's headline claim
    # (executive summary, the key-metrics table, the recommendation /
    # decision section). Report field; see ``should_enforce``.
    headline: bool = False

    def to_dict(self) -> dict:
        return {
            "code": self.code,
            "section_id": self.section_id,
            "location": self.location,
            "figure": self.figure,
            "snippet": self.snippet,
            "suggestion": self.detail,
            "severity": self.severity,
            "citations": list(self.citations),
            "headline": self.headline,
        }


@dataclass
class FactCheckResult:
    checked: int = 0
    verified: int = 0
    supported: int = 0
    derived: int = 0
    unsupported: int = 0
    unverifiable_citations: int = 0
    findings: list[FactFinding] = field(default_factory=list)
    corpus_summary: str = ""
    corpus_texts: list[dict] = field(default_factory=list)
    evidence_chars: int = 0
    thin_corpus: bool = False
    repair_feed: bool = False
    repair_feed_reason: str = ""
    error: str | None = None
    # Report-only (never enforced): the honest split of ``checked`` — see
    # the module docstring — and the figures a reader should look at.
    tiers: dict[str, Any] = field(default_factory=dict)
    figure_tiers: list[dict] = field(default_factory=list)
    # Report-only: seen/unseen status of the package's source URLs.
    source_urls: dict | None = None

    @property
    def coverage_pct(self) -> int | None:
        if not self.checked:
            return None
        return round((self.checked - self.unsupported) / self.checked * 100)

    @property
    def unsupported_findings(self) -> list[FactFinding]:
        return [f for f in self.findings if f.code == "unsupported_figure"]

    @property
    def comparison_findings(self) -> list[FactFinding]:
        """Comparisons whose own figures contradict them (``comparison_errors``)."""
        return [f for f in self.findings if f.code == "comparison_error"]

    @property
    def unsupported_headline_findings(self) -> list[FactFinding]:
        """Unsupported figures in the headline sections (see
        ``headline_section``) — the enforcement floor's trigger."""
        return [f for f in self.unsupported_findings if f.headline]

    @property
    def status(self) -> str:
        if self.error:
            return "error"
        if not self.checked:
            return "skipped"
        if not self.unsupported:
            return "pass"
        return "warn" if self.thin_corpus or not self.repair_feed else "fail"

    def to_dict(self) -> dict:
        unsupported = self.unsupported_findings
        return {
            "status": self.status,
            "checked": self.checked,
            "verified": self.verified,
            "supported": self.supported,
            "derived": self.derived,
            "unsupported": self.unsupported,
            "unverifiable_citations": self.unverifiable_citations,
            "coverage_pct": self.coverage_pct,
            "evidence_chars": self.evidence_chars,
            "thin_corpus": self.thin_corpus,
            "repair_feed": self.repair_feed,
            "repair_feed_reason": self.repair_feed_reason,
            "corpus": self.corpus_texts,
            "p0_count": len(unsupported) if not self.thin_corpus else 0,
            "unsupported_headline": len(self.unsupported_headline_findings),
            # Every headline one, whatever the cut below (the run's warning
            # lists them).
            "headline_findings": [f.to_dict() for f in self.unsupported_headline_findings[:40]],
            "comparison_findings": [f.to_dict() for f in self.comparison_findings[:20]],
            "comparison_count": len(self.comparison_findings),
            "finding_count": len(self.findings),
            "findings": [f.to_dict() for f in self.findings[:60]],
            "error": self.error,
            # Report-only additions (2026-09-22). Readers go through
            # ``summarize_fact_check``; nothing here feeds the repair.
            "tiers": dict(self.tiers),
            "figure_tiers": [dict(row) for row in self.figure_tiers[:MAX_FIGURE_TIER_ROWS]],
            "source_urls": self.source_urls,
        }

    def summary_lines(self) -> list[str]:
        """Findings as repair feedback, shaped like the pin-check lines so
        the error→section mapper routes each to its owning section. Only
        unsupported figures are fed — headline ones first; a citation
        mismatch is a note for the analyst, not a rewrite instruction.
        Unsupported figures are fed only when enforcement is on; a
        comparison its own figures contradict ("$597M, still below
        $579.37M") is arithmetic, not evidence, and is always fed."""
        lines = [
            f"fact check comparison_error in section {finding.section_id}: "
            f"\"{finding.snippet}\" — {finding.detail}"
            for finding in self.comparison_findings[:MAX_FED_FINDINGS]
        ]
        if not self.repair_feed:
            return lines
        ordered = self.unsupported_headline_findings + [
            f for f in self.unsupported_findings if not f.headline
        ]
        for finding in ordered[: max(0, MAX_FED_FINDINGS - len(lines))]:
            lines.append(
                f"fact check {finding.code} in section {finding.section_id}: "
                f"\"{finding.snippet}\" — {finding.detail}"
            )
        return lines


def _loc(value: Any) -> str:
    if isinstance(value, dict):
        en = value.get("en")
        return en if isinstance(en, str) else ""
    return value if isinstance(value, str) else ""


_SKIP_KEYS = {"id", "type", "component", "level", "source_ids", "zh", "kind", "slug", "style"}


def _iter_block_strings(value: Any, path: str = ""):
    """``(path, text)`` for every English string in a block, skipping ids
    and Chinese twins. Dicts with an ``en`` key yield only that string."""
    if isinstance(value, dict):
        if isinstance(value.get("en"), str):
            yield path, value["en"]
            return
        for key, item in value.items():
            if key in _SKIP_KEYS:
                continue
            yield from _iter_block_strings(item, f"{path}.{key}" if path else str(key))
    elif isinstance(value, list):
        for index, item in enumerate(value):
            yield from _iter_block_strings(item, f"{path}[{index}]")
    elif isinstance(value, str):
        yield path, value


def _citations_in(text: str) -> list[str]:
    ids: list[str] = []
    for token in _CITATION_RE.findall(text or ""):
        ids.extend(_CITATION_ID_RE.findall(token))
    return ids


def _calc_texts(package: dict) -> dict[str, str]:
    """What each calculation note DERIVES: its result, its meaning, and the
    right-hand side of every "=" in its formula. Inputs and the left-hand
    side are what the note consumes; those figures must be sourced on
    their own, or a note could launder any number by using it once."""
    out: dict[str, str] = {}
    for calc in package.get("calculations") or []:
        if not isinstance(calc, dict):
            continue
        calc_id = str(calc.get("id") or "").strip()
        if not calc_id:
            continue
        parts = []
        for key in ("result", "meaning"):
            parts.append(_loc(calc.get(key)) or str(calc.get(key) or ""))
        formula = _loc(calc.get("formula")) or str(calc.get("formula") or "")
        for segment in formula.split("=")[1:]:
            # "13 × $200B = $2.6T; $2.6T ÷ $1.75T = 1.5x": keep "$2.6T" and
            # "1.5x", not the "$2.6T ÷ $1.75T" that feeds the second step.
            parts.append(re.split(r"[;,]|\band\b", segment, maxsplit=1)[0] if "÷" in segment or "×" in segment or "*" in segment or "/" in segment else segment)
        out[calc_id] = " \n ".join(parts)
    return out


def _source_texts(package: dict, company_id: str, run_dir: Path | None) -> dict[str, str]:
    """Text on file for each package source, by id, when its URL (or a
    strong title match) points at a cached page."""
    out: dict[str, str] = {}
    run_texts: dict[str, str] = {}
    if run_dir is not None:
        base = Path(run_dir) / source_cache.SOURCES_DIRNAME
        for row in source_cache.run_manifest(run_dir):
            canon = source_cache.canonical_url(row.get("url"))
            if canon and row.get("file"):
                try:
                    run_texts[canon] = (base / str(row["file"])).read_text(encoding="utf-8")
                except OSError:
                    continue
    for source in package.get("sources") or []:
        if not isinstance(source, dict):
            continue
        source_id = str(source.get("id") or "").strip()
        if not source_id:
            continue
        url = source.get("url")
        record = source_cache.find_by_url(company_id, url) if url else None
        text = ""
        if record is not None:
            text = source_cache.source_text(company_id, record)
        canon = source_cache.canonical_url(url) if url else None
        if not text and canon and canon in run_texts:
            text = run_texts[canon]
        if not text and not url:
            match = source_cache.match_title(company_id, _loc(source.get("title")) or "", min_score=0.75)
            if match is not None:
                text = source_cache.source_text(company_id, match)
        if text:
            out[source_id] = text
    return out


# Where a figure is a headline claim: the section ids and the component
# slugs of the executive summary, the key-metrics table and the
# recommendation / decision. ``_HEADLINE_SECTION_RE`` (below) is the
# looser name match the report-only tiers use.
_HEADLINE_SECTION_IDS = frozenset(
    {"executive_summary", "investment_decision", "recommendation", "decision"}
)
_HEADLINE_COMPONENTS = frozenset(
    {"key_metrics_snapshot", "key_metrics", "recommendation", "investment_decision", "deal_terms"}
)


def headline_section(section_id: str, component: Any = None) -> bool:
    """True when a figure here is one a reader takes as the memo's
    conclusion: the executive summary, the key-metrics table, the
    recommendation / investment-decision section."""
    sid = str(section_id or "").strip().lower()
    if sid in _HEADLINE_SECTION_IDS or re.search(r"exec|summary|recommend|decision", sid):
        return True
    components = component if isinstance(component, list) else [component]
    return any(str(c or "").strip().lower() in _HEADLINE_COMPONENTS for c in components)


def should_enforce(
    corpus: Corpus,
    coverage: int | None,
    *,
    headline_unsupported: int = 0,
) -> tuple[bool, str]:
    """Whether unsupported figures are fed to the repair, and why.

    The auto policy enforces when the corpus is rich (evidence chars and
    coverage above their floors). The enforcement FLOOR (2026-09-23): an
    unsupported figure in a headline section is fed whatever the coverage
    — a memo whose executive summary states "8 to 10 months of runway"
    from no source is not "the corpus is short", it is an invented
    headline. A thin corpus (nothing to check against) and an explicit
    ``BSH_MEMO_FACT_CHECK_REPAIR=0`` still switch enforcement off."""
    raw = os.environ.get("BSH_MEMO_FACT_CHECK_REPAIR", "auto").strip().lower()
    if raw in {"0", "false", "no", "off"}:
        return False, "BSH_MEMO_FACT_CHECK_REPAIR=0: report only"
    if corpus.thin:
        return False, f"corpus too thin to enforce ({corpus.evidence_chars:,} evidence chars < {THIN_CORPUS_CHARS:,})"
    if raw in {"1", "true", "yes", "on"}:
        return True, "BSH_MEMO_FACT_CHECK_REPAIR=1"
    if headline_unsupported > 0:
        return True, (
            f"headline floor: {headline_unsupported} unsupported figure"
            f"{'s' if headline_unsupported != 1 else ''} in the executive summary, "
            "key metrics or recommendation — enforced regardless of coverage"
            + (f" ({coverage}% traceable)" if coverage is not None else "")
        )
    return _auto_policy(corpus, coverage)


def _repair_policy(corpus: Corpus, coverage: int | None) -> tuple[bool, str]:
    """The pre-floor policy (kept for callers and tests that pass no
    headline count): ``should_enforce`` with no headline figures."""
    return should_enforce(corpus, coverage, headline_unsupported=0)


def _auto_policy(corpus: Corpus, coverage: int | None) -> tuple[bool, str]:
    if corpus.evidence_chars < AUTO_ENFORCE_MIN_CHARS:
        return False, (
            f"auto: corpus has {corpus.evidence_chars:,} evidence chars, enforcement starts at "
            f"{AUTO_ENFORCE_MIN_CHARS:,}"
        )
    if coverage is not None and coverage < AUTO_ENFORCE_MIN_COVERAGE:
        return False, (
            f"auto: only {coverage}% of figures traceable, enforcement starts at "
            f"{AUTO_ENFORCE_MIN_COVERAGE}% (the corpus, not the memo, is short)"
        )
    return True, "auto: corpus rich enough to enforce"


def check_package(
    package: dict,
    corpus: Corpus,
    *,
    source_texts: dict[str, str] | None = None,
) -> FactCheckResult:
    """Check every figure in the package's sections against the corpus."""
    result = FactCheckResult(
        corpus_summary=corpus.summary(),
        corpus_texts=corpus.describe(),
        evidence_chars=corpus.evidence_chars,
        thin_corpus=corpus.thin,
    )
    if not isinstance(package, dict):
        result.error = "memo package must be a JSON object"
        return result
    source_texts = source_texts or {}
    # One scan per cited-source text yields both the enforced (loose) index
    # and the strict one the report-only tiers read.
    source_pairs = {sid: _index_pair(text) for sid, text in source_texts.items()}
    source_index = {sid: pair[0] for sid, pair in source_pairs.items()}
    calc_texts = _calc_texts(package)
    calc_index = {cid: TextIndex.of(text) for cid, text in calc_texts.items()}
    corpus_label = corpus.summary()
    # Report-only: any failure here switches the tally off, never the check.
    honest: _HonestTally | None
    try:
        honest = _HonestTally(package, corpus, {sid: pair[1] for sid, pair in source_pairs.items()})
    except Exception:  # noqa: BLE001
        logger.warning("fact check: honest tiers unavailable", exc_info=True)
        honest = None

    for s_index, section in enumerate(package.get("sections") or []):
        if not isinstance(section, dict):
            continue
        section_id = str(section.get("id") or f"sections[{s_index}]")
        for b_index, block in enumerate(section.get("blocks") or []):
            if not isinstance(block, dict) or block.get("type") == "heading":
                continue
            location_base = f"sections[{s_index}].blocks[{b_index}]"
            for path, text in _iter_block_strings(block):
                figures = extract_figures(text)
                if not figures:
                    continue
                location = f"{location_base}.{path}" if path else location_base
                for error in comparison_errors(text):
                    result.findings.append(
                        FactFinding(
                            code="comparison_error",
                            section_id=section_id,
                            location=location,
                            figure="",
                            snippet=error["snippet"],
                            detail=f"{error['detail']}; correct the comparison or the figure",
                            severity="P0",
                            headline=headline_section(section_id, block.get("component")),
                        )
                    )
                cited = _citations_in(text)
                cited_sources = [c for c in cited if c.startswith("S")]
                cited_calcs = [c for c in cited if c.startswith("C")]
                for figure in figures:
                    result.checked += 1
                    bucket = "unsupported"
                    with_text = [sid for sid in cited_sources if sid in source_index]
                    # 1. Derived by a note the sentence cites.
                    if any(cid in calc_index and calc_index[cid].supports(figure) for cid in cited_calcs):
                        bucket = "derived"
                    # 2. Verified: the cited source's own text carries it.
                    elif with_text and any(source_index[sid].supports(figure) for sid in with_text):
                        bucket = "verified"
                    else:
                        if cited_sources and not with_text:
                            result.unverifiable_citations += 1
                        # 3. Derived by any note (the sentence forgot the [C#]).
                        if any(index.supports(figure) for index in calc_index.values()):
                            bucket = "derived"
                        # 4. Supported: anything on file carries it.
                        elif corpus.index.supports(figure):
                            bucket = "supported"
                    if honest is not None:
                        try:
                            honest.classify(
                                figure,
                                bucket=bucket,
                                text=text,
                                section_id=section_id,
                                location=location,
                                cited_sources=cited_sources,
                            )
                        except Exception:  # noqa: BLE001
                            logger.warning("fact check: honest tiers failed", exc_info=True)
                            honest = None
                    if bucket == "derived":
                        result.derived += 1
                        continue
                    if bucket == "verified":
                        result.verified += 1
                        continue
                    if bucket == "supported":
                        result.supported += 1
                        if with_text:
                            result.findings.append(
                                FactFinding(
                                    code="citation_mismatch",
                                    section_id=section_id,
                                    location=location,
                                    figure=figure.raw,
                                    snippet=_excerpt(text, figure),
                                    detail=(
                                        f"{figure.raw} is on file, but not in the text of the cited "
                                        f"source{'s' if len(with_text) > 1 else ''} {', '.join(with_text)}; "
                                        "check the citation"
                                    ),
                                    severity="P1",
                                    citations=list(cited_sources),
                                )
                            )
                        continue
                    result.unsupported += 1
                    result.findings.append(
                        FactFinding(
                            code="unsupported_figure",
                            section_id=section_id,
                            location=location,
                            figure=figure.raw,
                            snippet=_excerpt(text, figure),
                            detail=(
                                f"the figure {figure.raw} appears in none of the sources on file "
                                f"({corpus_label}); cite the source that carries it — adding it to the "
                                "package sources with its URL if it is missing — or state it as an "
                                "estimate with a [C#] calculation note whose inputs are sourced"
                            ),
                            severity="P0",
                            citations=list(cited),
                            headline=headline_section(section_id, block.get("component")),
                        )
                    )
    result.repair_feed, result.repair_feed_reason = should_enforce(
        corpus,
        result.coverage_pct,
        headline_unsupported=len(result.unsupported_headline_findings),
    )
    if honest is not None:
        try:
            result.tiers, result.figure_tiers = honest.finish()
        except Exception:  # noqa: BLE001 — report-only; must never sink the check
            logger.warning("fact check: honest tiers failed", exc_info=True)
            result.tiers, result.figure_tiers = {}, []
    return result


# ---- same-metric contradictions (report only, 2026-09-23) --------------------
#
# One memo, one metric, two values: "NextNav net loss $189.3M" in the
# comps table and "$111.9M net loss" in the valuation prose (Gemini v2,
# 2026-09-23). Figures are grouped by the metric word nearest them, the
# proper nouns in their sentence and the qualifiers around them (bear /
# base / bull, gross / net, trailing / forward, a year); two figures in
# one group that do not round to each other are a conflict. A table
# column that repeats one value is one value, never a conflict.

_QUALIFIER_WORDS = frozenset(
    {
        "bear", "base", "bull", "downside", "upside", "low", "high", "mid",
        "midpoint", "min", "max", "minimum", "maximum", "floor", "ceiling",
        "gross", "net", "trailing", "forward", "ltm", "ntm", "fy", "cumulative",
        "annual", "annualized", "annualised", "monthly", "quarterly", "primary",
        "secondary", "pre", "post", "prior", "current", "target", "implied",
        "fair", "entry", "exit", "headline", "adjusted", "recurring", "lifetime",
        "realizable", "realisable", "modeled", "modelled", "estimated", "q1",
        "q2", "q3", "q4", "h1", "h2", "per", "total", "average", "median",
        "filed", "issued", "granted", "pending", "fresh", "extension",
        "insider", "binding", "signed", "committed", "disclosed", "verified",
        "diluted", "undiluted", "ev", "equity", "enterprise", "cash", "burn",
    }
)
_MONTH_WORDS_SET = frozenset(
    {
        "january", "february", "march", "april", "may", "june", "july",
        "august", "september", "october", "november", "december",
    }
)
_YEAR_WORD_RE = re.compile(r"(?<!\d)(?:19|20)\d{2}(?!\d)")
_PROPER_NOUN_RE = re.compile(r"\b[A-Z][A-Za-z0-9&'\-]{2,}\b")
_SENTENCE_SPLIT_RE = re.compile(r"(?<=[.!?;])\s+")
_NEGATIVE_MONEY_RE = re.compile(r"(?<=\s)[-\u2212\u2013](?=\$)")
_METRIC_WINDOW = 60
MAX_METRIC_CONFLICTS = 20
# More distinct values than this for one metric is a series (scenario
# ARRs, a market forecast by year), not a contradiction.
_MAX_CONFLICT_VALUES = 3
# Precision guards (live ZaiNar run 2026-09-23 raised 20 "conflicts", nearly
# all two different quantities tied to one nearby metric word — "$10M" and
# "$1.0B" both read as "valuation" in "$10M ... at a $1.0B valuation"):
# - the metric word must sit this close to the figure (characters between
#   the two spans) for the pairing to count;
# - two readings of one metric that genuinely conflict are close in size
#   (NextNav's loss as $189.3M and $111.9M); values further apart than this
#   ratio are different quantities, not a contradiction;
# - both values inside one sentence are a range or a comparison the
#   sentence itself draws ("10% to 20% dilution"), never a conflict.
_CONFLICT_METRIC_MAX_GAP = 28
_CONFLICT_MAX_RATIO = 5.0


@dataclass
class _MetricMention:
    figure: Figure
    metric: str
    entities: frozenset[str]
    qualifiers: frozenset[str]
    years: frozenset[str]
    location: str
    section_id: str
    text: str

    @property
    def key(self) -> tuple[str, str, frozenset[str], frozenset[str]]:
        return (self.figure.klass, self.metric, self.entities, self.qualifiers)


def _nearest_metric_word(text: str, figure: Figure) -> tuple[str, int, int] | None:
    """The metric word (``_METRIC_WORDS`` stem) nearest the figure within
    the window, with its span; None when nothing names the metric."""
    best: tuple[int, str, int, int] | None = None
    lo = max(0, figure.start - _METRIC_WINDOW)
    hi = min(len(text), figure.end + _METRIC_WINDOW)
    for match in _CONTEXT_WORD_RE.finditer(text, lo, hi):
        stem = _stem(match.group(0).lower())
        if stem not in _METRIC_WORDS:
            continue
        if match.start() >= figure.end:
            # "$120M ARR", "$4.02M in trailing revenue": the metric follows.
            distance = match.start() - figure.end
        else:
            # "loss of $189.3M" names it just before; a metric word further
            # back usually belongs to another figure in the sentence.
            distance = figure.start - match.end()
            if distance > 6:
                distance += 8
        if best is None or distance < best[0]:
            best = (distance, stem, match.start(), match.end())
    return (best[1], best[2], best[3]) if best else None


def _sentence_around(text: str, figure: Figure) -> str:
    start = 0
    for match in _SENTENCE_SPLIT_RE.finditer(text):
        if match.end() <= figure.start:
            start = match.end()
        else:
            return text[start:match.start()]
    return text[start:]


def _nearest_year(text: str, figure: Figure, *, radius: int = 40) -> str | None:
    best: tuple[int, str] | None = None
    for match in _YEAR_WORD_RE.finditer(text, max(0, figure.start - radius), min(len(text), figure.end + radius)):
        if figure.start <= match.start() < figure.end:
            continue
        distance = min(abs(match.start() - figure.end), abs(figure.start - match.end()))
        if best is None or distance < best[0]:
            best = (distance, match.group(0))
    return best[1] if best else None


def _entity_words(text: str, company_tokens: set[str]) -> frozenset[str]:
    """Mixed-case proper nouns (a company, a person, a product) — not
    all-caps tickers and acronyms, months, metric words or qualifiers."""
    return frozenset(
        word.lower().removesuffix("'s")
        for word in _PROPER_NOUN_RE.findall(text)
        if not word.isupper()
        and word.lower() not in company_tokens
        and _stem(word.lower()) not in _METRIC_WORDS
        and word.lower() not in _CONTEXT_STOPWORDS
        and word.lower() not in _QUALIFIER_WORDS
        and word.lower() not in _MONTH_WORDS_SET
        and not _YEAR_WORD_RE.fullmatch(word)
    )


def _table_strings(block: dict):
    """``(path, text, row_context)`` for a table: each cell with its row's
    label cell and its column header as context (a scenario table's
    "Bear" / "Base" / "Bull" live in the label cell, not the value)."""
    headers = [_loc(h) or (str(h) if isinstance(h, str) else "") for h in block.get("headers") or []]
    for r_index, row in enumerate(block.get("rows") or []):
        cells = row.get("cells") if isinstance(row, dict) else row
        if not isinstance(cells, (list, tuple)):
            continue
        texts = [_loc(cell) or (str(cell) if isinstance(cell, str) else "") for cell in cells]
        label = texts[0] if texts else ""
        for c_index, text in enumerate(texts):
            if not text:
                continue
            header = headers[c_index] if c_index < len(headers) else ""
            context = " ".join(part for part in (label if c_index else "", header) if part)
            yield f"rows[{r_index}][{c_index}]", text, context


_RANGE_JOIN_RE = re.compile(r"^\s*(?:to|and|through|or|-|–|—|~)\s*$", re.IGNORECASE)
_SCENARIO_WORDS = ("bear", "base", "bull", "downside", "upside")
_SCENARIO_WORD_RE = re.compile(r"\b(bear|base|bull|downside|upside)\b", re.IGNORECASE)


def _range_endpoint_ids(text: str, figures: list[Figure]) -> set[int]:
    """ids of figures that are an endpoint of a stated range ("$5.0M to
    $10.0M", "75%-85%", "between 10% and 20%"): the two figures sit next
    to each other with only a range word between them."""
    ordered = sorted(figures, key=lambda f: f.start)
    ids: set[int] = set()
    for left, right in zip(ordered, ordered[1:]):
        if left.klass != right.klass:
            continue
        gap = text[left.end:right.start] if 0 <= left.end <= right.start <= len(text) else ""
        if _RANGE_JOIN_RE.match(gap):
            ids.add(id(left))
            ids.add(id(right))
    return ids


def _sentence_scenarios(text: str, figure: Figure) -> frozenset[str]:
    """Scenario words (bear/base/bull...) in the sentence that holds the
    figure: "In the bull case, ... a 5.10x net MOIC" is the bull reading,
    whatever sits in the few characters beside the number."""
    start = 0
    for part in _SENTENCE_SPLIT_RE.split(text):
        end = start + len(part)
        if start <= figure.start <= end:
            return frozenset(m.group(1).lower() for m in _SCENARIO_WORD_RE.finditer(part))
        start = end + 1
    return frozenset()


def _metric_mentions(package: dict, *, company_tokens: set[str] | None = None) -> list[_MetricMention]:
    company_tokens = {t.lower() for t in (company_tokens or set())}
    mentions: list[_MetricMention] = []
    for s_index, section in enumerate(package.get("sections") or []):
        if not isinstance(section, dict):
            continue
        section_id = str(section.get("id") or f"sections[{s_index}]")
        for b_index, block in enumerate(section.get("blocks") or []):
            if not isinstance(block, dict) or block.get("type") == "heading":
                continue
            base = f"sections[{s_index}].blocks[{b_index}]"
            if block.get("type") == "table":
                units = list(_table_strings(block))
            else:
                units = [(path, text, "") for path, text in _iter_block_strings(block)]
            for path, raw_text, context in units:
                # Figure offsets are on the citation-stripped text. A
                # negative money figure ("-$111.9M net loss") is one the
                # extractor skips as an identifier; here the sign is a
                # space of the same width, so offsets still line up.
                raw_text = _NEGATIVE_MONEY_RE.sub(" ", raw_text)
                text = _strip_citations(raw_text)
                figures = extract_figures(raw_text)
                ranged = _range_endpoint_ids(text, figures)
                for figure in figures:
                    if id(figure) in ranged:
                        # An endpoint of "X to Y" is half of one stated
                        # range, not a reading of its own.
                        continue
                    nearest = _nearest_metric_word(text, figure)
                    if nearest and max(nearest[1] - figure.end, figure.start - nearest[2]) > _CONFLICT_METRIC_MAX_GAP:
                        # A metric word this far away describes some other
                        # figure in the sentence, not this one.
                        nearest = None
                    metric = nearest[0] if nearest else None
                    # Qualifiers sit on the figure or its metric word ("net
                    # loss", "trailing revenue", "base case"): the span from
                    # a little before the first of the two to the last.
                    span_lo = min(figure.start, nearest[1] if nearest else figure.start)
                    span_hi = max(figure.end, nearest[2] if nearest else figure.end)
                    qualifier_text = text[max(0, span_lo - 24):span_hi + 6]
                    if not metric and context:
                        for match in _CONTEXT_WORD_RE.finditer(context):
                            stem = _stem(match.group(0).lower())
                            if stem in _METRIC_WORDS:
                                metric = stem
                                break
                    if not metric:
                        continue
                    window = text[max(0, figure.start - _METRIC_WINDOW):figure.end + _METRIC_WINDOW]
                    scope = f"{window} {context}"
                    qualifiers = frozenset(
                        word.lower()
                        for word in _CONTEXT_WORD_RE.findall(f"{qualifier_text} {context}")
                        if word.lower() in _QUALIFIER_WORDS
                    ) | _sentence_scenarios(text, figure)
                    year = _nearest_year(text, figure) or (
                        (_YEAR_WORD_RE.findall(context) or [None])[0] if context else None
                    )
                    mentions.append(
                        _MetricMention(
                            figure=figure,
                            metric=metric,
                            entities=_entity_words(scope, company_tokens),
                            qualifiers=qualifiers,
                            years=frozenset({year} if year else ()),
                            location=f"{base}.{path}" if path else base,
                            section_id=section_id,
                            text=text,
                        )
                    )
    return mentions


def _same_value(a: Figure, b: Figure) -> bool:
    """Two figures that state one value (one may be the other rounded)."""
    if a.klass != b.klass:
        return False
    if math.isclose(a.value, b.value, rel_tol=1e-9, abs_tol=1e-12):
        return True
    lo_a, hi_a = _significant(a.mantissa)
    lo_b, hi_b = _significant(b.mantissa)
    for sig in range(min(lo_a, lo_b), max(hi_a, hi_b) + 1):
        if _rounds_to(a.value, b.value, sig) or _rounds_to(b.value, a.value, sig):
            return True
    return False


def _package_from(package_or_docx: Any) -> dict | None:
    if isinstance(package_or_docx, dict):
        return package_or_docx
    path = Path(str(package_or_docx))
    if path.suffix.lower() == ".json":
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return None
        return payload if isinstance(payload, dict) else None
    if path.suffix.lower() == ".docx":
        try:
            from docx import Document  # type: ignore
        except ImportError:
            return None
        try:
            document = Document(str(path))
        except Exception:  # noqa: BLE001
            return None
        blocks: list[dict] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                blocks.append({"type": "paragraph", "text": {"en": text}})
        for table in document.tables:
            rows = [[{"en": cell.text.strip()} for cell in row.cells] for row in table.rows]
            blocks.append({"type": "table", "rows": rows})
        return {"sections": [{"id": "document", "blocks": blocks}]}
    return None


def metric_conflicts(
    package_or_docx: Any, *, company_tokens: set[str] | None = None
) -> list[dict]:
    """Different values for one metric within one memo:
    ``[{metric, values: [...], locations: [...]}]``. Report-only; never
    raises (an unreadable input is an empty list)."""
    try:
        package = _package_from(package_or_docx)
        if not isinstance(package, dict):
            return []
        if company_tokens is None:
            company = package.get("company") if isinstance(package.get("company"), dict) else {}
            name = company.get("name")
            name = name.get("en") if isinstance(name, dict) else name
            company_tokens = {
                word.lower()
                for word in re.findall(r"[A-Za-z][A-Za-z0-9&'\-]+", str(name or ""))
                if len(word) > 2
            }
        groups: dict[tuple, list[_MetricMention]] = {}
        for mention in _metric_mentions(package, company_tokens=company_tokens):
            groups.setdefault(mention.key, []).append(mention)
        out: list[dict] = []
        for key, mentions in groups.items():
            distinct: list[_MetricMention] = []
            for mention in mentions:
                if any(_same_value(mention.figure, other.figure) for other in distinct):
                    continue
                distinct.append(mention)
            if len(distinct) < 2:
                continue
            # A year distinguishes two readings of one metric (2024 vs 2025
            # revenue); a figure with no year conflicts with any year.
            conflicting: list[_MetricMention] = []
            for mention in distinct:
                if any(
                    mention.years and other.years and mention.years.isdisjoint(other.years)
                    for other in conflicting
                ):
                    continue
                conflicting.append(mention)
            if len(conflicting) < 2 or len(conflicting) > _MAX_CONFLICT_VALUES:
                continue
            # One sentence stating both values draws its own range or
            # comparison; only readings in different places can disagree.
            if len({m.location for m in conflicting}) < 2:
                continue
            magnitudes = [abs(m.figure.value) for m in conflicting if m.figure.value]
            if len(magnitudes) >= 2 and max(magnitudes) / min(magnitudes) > _CONFLICT_MAX_RATIO:
                continue
            metric, qualifiers, entities = key[1], key[3], key[2]
            label = " ".join(sorted(qualifiers) + [metric])
            if entities:
                label = f"{label} ({', '.join(sorted(entities))})"
            out.append(
                {
                    "metric": label,
                    "values": [m.figure.raw for m in conflicting],
                    "locations": [m.location for m in conflicting],
                    "sections": sorted({m.section_id for m in conflicting}),
                    "snippets": [_excerpt(m.text, m.figure, width=50) for m in conflicting],
                }
            )
            if len(out) >= MAX_METRIC_CONFLICTS:
                break
        return out
    except Exception:  # noqa: BLE001 — report-only
        logger.warning("metric conflicts failed", exc_info=True)
        return []


# ---- quote-backed claims (report only, 2026-09-23) ----------------------------

_QUOTE_DASHES = "\u2010\u2011\u2012\u2013\u2014\u2015\u2212"
# Every quotation mark — curly or straight, single or double — becomes one
# apostrophe: a page's straight quotes and a model's curly ones are the
# same words.
_QUOTE_MARKS = {mark: "'" for mark in "\u2018\u2019\u201a\u201c\u201d\u201e\u00ab\u00bb\""}
_QUOTE_MIN_CHARS = 12


def normalize_quote(text: Any) -> str:
    """Whitespace collapsed, quotation marks and dashes unified, lower
    case — what a verbatim quote and the page it came from agree on."""
    out = str(text or "")
    for mark, plain in _QUOTE_MARKS.items():
        out = out.replace(mark, plain)
    for dash in _QUOTE_DASHES:
        out = out.replace(dash, "-")
    out = out.replace("\u00a0", " ").replace("\u2026", "...")
    return re.sub(r"\s+", " ", out).strip().lower()


def _run_company_id(run_dir: Path) -> str | None:
    try:
        payload = json.loads((run_dir / "logs" / "run_inputs.json").read_text(encoding="utf-8"))
        if isinstance(payload, dict):
            for key in ("company_id", "company", "slug"):
                value = payload.get(key)
                if isinstance(value, str) and value.strip():
                    return value.strip()
    except (OSError, ValueError):
        pass
    # data/memos/<company>/<date>__<time>__<company>__memo-run
    if run_dir.parent.name and run_dir.parent.name != "memos":
        return run_dir.parent.name
    parts = run_dir.name.split("__")
    if len(parts) >= 4 and parts[2]:
        return parts[2]
    return None


_ELISION_RE = re.compile(r"\s*(?:\.\s*\.\s*\.|\[\s*\.\.\.\s*\]|\[\s*…\s*\])\s*")


def _elided_quote_in(quote: str, page: str) -> bool:
    """A quote that skips words with "..." matches when every fragment
    appears on the page, in order, and at least one fragment is
    ``_QUOTE_MIN_CHARS`` long (so "ZaiNar... today announced ..." counts
    but a quote made only of scraps does not). Both inputs are already
    normalised (``normalize_quote``)."""
    fragments = [part.strip(" ,;:") for part in _ELISION_RE.split(quote)]
    fragments = [part for part in fragments if part]
    if len(fragments) < 2 or max(len(part) for part in fragments) < _QUOTE_MIN_CHARS:
        return False
    position = 0
    for part in fragments:
        found = page.find(part, position)
        if found < 0:
            return False
        position = found + len(part)
    return True


def _cached_source_text(run_dir: Path, url: str, company_id: str | None, run_texts: dict[str, str]) -> str:
    canon = source_cache.canonical_url(url)
    if canon and canon in run_texts:
        return run_texts[canon]
    if company_id and url:
        try:
            record = source_cache.find_by_url(company_id, url)
            if record is not None:
                return source_cache.source_text(company_id, record)
        except Exception:  # noqa: BLE001
            return ""
    return ""


def check_quotes(run_dir: Path | str, *, company_id: str | None = None) -> dict:
    """Match every recorded evidence quote (``logs/evidence_quotes.json``,
    ``[{claim, url, quote, pass_id}]``) against the cached text of its
    page: the run's frozen sources first, then the company's source cache.
    Normalised substring match. Returns ``{checked, matched, unmatched:
    [{claim, url, quote, pass_id, reason}], uncached: [...]}``; never
    raises. ``unmatched`` holds only quotes whose page IS cached and does
    not contain them — the suspicious case. A quote whose page was never
    cached (a paywalled article, a page fetched outside the cache) cannot be
    checked either way and goes to ``uncached``: live on 2026-09-23 twelve
    of a Gemini run's "not found" quotes were real sentences from pages the
    cache simply did not hold, including the CEO's own words in The
    Information."""
    result: dict[str, Any] = {"checked": 0, "matched": 0, "unmatched": [], "uncached": []}
    try:
        run_dir = Path(run_dir)
        path = run_dir / "logs" / "evidence_quotes.json"
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return result
        if isinstance(payload, dict):
            payload = payload.get("quotes") or []
        if not isinstance(payload, list):
            return result
        company_id = company_id or _run_company_id(run_dir)
        run_texts: dict[str, str] = {}
        base = run_dir / source_cache.SOURCES_DIRNAME
        for row in source_cache.run_manifest(run_dir):
            canon = source_cache.canonical_url(row.get("url"))
            if canon and row.get("file") and canon not in run_texts:
                try:
                    run_texts[canon] = (base / str(row["file"])).read_text(encoding="utf-8")
                except OSError:
                    continue
        normalized_cache: dict[str, str] = {}
        for item in payload:
            if not isinstance(item, dict):
                continue
            quote = normalize_quote(item.get("quote"))
            url = str(item.get("url") or "").strip()
            if len(quote) < _QUOTE_MIN_CHARS:
                continue
            result["checked"] += 1
            entry = {
                "claim": str(item.get("claim") or ""),
                "url": url,
                "quote": str(item.get("quote") or ""),
                "pass_id": item.get("pass_id"),
            }
            if url not in normalized_cache:
                normalized_cache[url] = normalize_quote(
                    _cached_source_text(run_dir, url, company_id, run_texts)
                )
            page = normalized_cache[url]
            if not page:
                entry["reason"] = "no cached text for this URL"
                result["uncached"].append(entry)
                continue
            if quote in page or _elided_quote_in(quote, page):
                result["matched"] += 1
            else:
                entry["reason"] = "quote not found in the cached page text"
                result["unmatched"].append(entry)
        return result
    except Exception:  # noqa: BLE001 — report-only
        logger.warning("evidence quote check failed", exc_info=True)
        return result


# ---- honest tiers (report only) ----------------------------------------------

# Sections a reader takes as the memo's headline claims.
_HEADLINE_SECTION_RE = re.compile(r"exec|summary|highlight|recommend|decision|verdict", re.I)
# Words around a number that say nothing about which metric it is.
_CONTEXT_STOPWORDS = {
    "the", "and", "for", "with", "that", "this", "from", "are", "was", "were",
    "has", "have", "had", "its", "our", "their", "into", "not", "but", "than",
    "then", "about", "approximately", "approx", "roughly", "around", "over",
    "under", "more", "less", "least", "most", "nearly", "almost", "some",
    "only", "also", "which", "while", "will", "would", "could", "should",
    "may", "might", "per", "each", "every", "year", "years", "month", "months",
    "quarter", "quarters", "billion", "million", "thousand", "trillion",
    "percent", "times", "since", "after", "before", "during", "through",
    "between", "against", "both", "such", "these", "those", "other", "onto",
    "upon", "already", "still", "even", "just", "like", "much", "many",
    "company", "companies", "firm", "bsh", "usd", "est", "estimated",
    "estimate", "reported", "reports", "report", "said", "says", "according",
    "total", "value", "figure", "number", "level", "rate", "base", "case",
    "they", "them", "there", "here", "what", "when", "where", "who", "how",
    "all", "any", "one", "two", "three", "new", "now", "yet", "out",
    "his", "her", "him", "she", "you", "your", "can", "did", "does", "being",
    "been", "via", "within", "without", "across", "because", "whose",
}
# Words that name WHICH metric a number is (stems, see ``_stem``). When the
# memo's phrasing around a figure has any of these, only these count as its
# context; a verb like "reached" sits next to every kind of number.
_METRIC_WORDS = {
    "arr", "mrr", "revenue", "sale", "booking", "billing", "growth", "grew",
    "margin", "customer", "client", "user", "subscriber", "account", "logo",
    "employee", "headcount", "staff", "engineer", "valuation", "valued",
    "tam", "sam", "som", "market", "cagr", "ebitda", "ebit", "profit",
    "income", "loss", "cash", "burn", "runway", "funding", "raised", "raise",
    "round", "investment", "price", "share", "stake", "ownership", "pipeline",
    "contract", "backlog", "gmv", "nrr", "ndr", "grr", "churn", "retention",
    "patent", "unit", "deployment", "site", "country", "hospital", "pilot",
    "order", "capex", "opex", "debt", "dividend", "eps", "earning", "multiple",
    "ev", "run-rate", "capacity", "production", "volume", "shipment",
    "download", "install", "installation", "location", "store", "vehicle",
    "device", "sensor", "partner", "investor", "fund", "fee", "cost",
    "spend", "spending", "budget", "salary", "compensation", "arpu", "acv",
    "tcv", "ltv", "cac", "payback", "yield", "return", "irr", "moic",
}
_CONTEXT_WORD_RE = re.compile(r"[A-Za-z][A-Za-z&'-]+")
_NUMBER_TOKEN_RE = re.compile(r"(?<![\w.])\d[\d,]*(?:\.\d+)?")


def _stem(word: str) -> str:
    word = word.strip("'-").lower()
    if len(word) > 4 and word.endswith("ies"):
        return word[:-3] + "y"
    if len(word) > 3 and word.endswith("s") and not word.endswith("ss"):
        return word[:-1]
    return word


def _context_words(text: str) -> set[str]:
    words: set[str] = set()
    for word in _CONTEXT_WORD_RE.findall(text):
        lowered = word.lower()
        if len(lowered) < 3 and not (len(word) == 2 and word.isupper()):
            continue
        stem = _stem(lowered)
        if lowered in _CONTEXT_STOPWORDS or stem in _CONTEXT_STOPWORDS:
            continue
        words.add(stem)
    return words


def _number_key(raw: str) -> str:
    key = raw.replace(",", "")
    if "." in key:
        key = key.rstrip("0").rstrip(".") or "0"
    return key


class _HonestTally:
    """The report-only split of every checked figure (see the module
    docstring). Uses only the strict tier indexes and the subject's own
    domains; the enforced buckets are computed before it and never read
    it, so nothing here can change what the repair enforces."""

    def __init__(self, package: dict, corpus: Corpus, strict_sources: dict[str, TextIndex]):
        self.corpus = corpus
        tiers = getattr(corpus, "tier_index", None) or {}
        self.independent: TextIndex | None = tiers.get("independent")
        self.company: TextIndex | None = tiers.get("company")
        self.registry: TextIndex | None = tiers.get("registry")
        hosts = set(getattr(corpus, "company_hosts", None) or ())
        self.company_source_ids: set[str] = set()
        for source in package.get("sources") or []:
            if not isinstance(source, dict):
                continue
            if _on_hosts(_host_of(source.get("url")), hosts):
                self.company_source_ids.add(str(source.get("id") or ""))
        self.strict_sources = strict_sources
        self.counts = {
            "verified": 0,
            "found_elsewhere": 0,
            "in_context": 0,
            "derived": 0,
            "company_reported": 0,
            "company_reported_headline": 0,
            "registry_only": 0,
            "not_traced": 0,
        }
        self.rows: list[dict] = []
        self._context_queue: list[tuple[Figure, set[str]]] = []

    def classify(
        self,
        figure: Figure,
        *,
        bucket: str,
        text: str,
        section_id: str,
        location: str,
        cited_sources: list[str],
    ) -> str:
        if bucket == "derived":
            tier = "derived"
        elif bucket == "unsupported":
            # The loose match found nothing; the strict tiers cannot either.
            tier = "not_traced"
        else:
            cited_strict = [sid for sid in cited_sources if sid in self.strict_sources]
            if any(
                sid not in self.company_source_ids and self.strict_sources[sid].supports(figure)
                for sid in cited_strict
            ):
                tier = "verified"
            elif self.independent is not None and self.independent.supports(figure):
                tier = "found_elsewhere"
            elif any(
                sid in self.company_source_ids and self.strict_sources[sid].supports(figure)
                for sid in cited_strict
            ) or (self.company is not None and self.company.supports(figure)):
                tier = "company_reported"
            elif self.registry is not None and self.registry.supports(figure):
                tier = "registry_only"
            else:
                tier = "not_traced"
        self.counts[tier] += 1
        headline = bool(_HEADLINE_SECTION_RE.search(section_id))
        if tier == "company_reported" and headline:
            self.counts["company_reported_headline"] += 1
        if tier == "found_elsewhere":
            keywords = self._keywords(text, figure)
            if keywords:
                self._context_queue.append((figure, keywords))
        if tier in {"company_reported", "registry_only", "not_traced"}:
            self.rows.append(
                {
                    "section_id": section_id,
                    "location": location,
                    "figure": figure.raw,
                    "snippet": _excerpt(text, figure),
                    "tier": tier,
                    "headline": headline,
                    "unsupported": bucket == "unsupported",
                }
            )
        return tier

    def _keywords(self, text: str, figure: Figure) -> set[str]:
        window = text[max(0, figure.start - 80):min(len(text), figure.end + 50)]
        tokens = getattr(self.corpus, "company_tokens", None) or set()
        words = {word for word in _context_words(window) if word not in tokens}
        return (words & _METRIC_WORDS) or words

    def _count_in_context(self) -> int:
        if not self._context_queue:
            return 0
        wanted = {_number_key(figure.mantissa) for figure, _ in self._context_queue}
        positions: dict[str, list[tuple[int, int, int]]] = {}
        texts = self.corpus.tier_texts_of("independent")
        for text_index, text in enumerate(texts):
            for match in _NUMBER_TOKEN_RE.finditer(text):
                key = _number_key(match.group(0))
                if key in wanted:
                    rows = positions.setdefault(key, [])
                    if len(rows) < MAX_CONTEXT_POSITIONS:
                        rows.append((text_index, match.start(), match.end()))
        count = 0
        for figure, keywords in self._context_queue:
            for text_index, start, end in positions.get(_number_key(figure.mantissa), ()):
                text = texts[text_index]
                window = text[max(0, start - CONTEXT_WINDOW_CHARS):end + CONTEXT_WINDOW_CHARS]
                if keywords & _context_words(window):
                    count += 1
                    break
        return count

    def finish(self) -> tuple[dict[str, Any], list[dict]]:
        self.counts["in_context"] = self._count_in_context()
        tiers: dict[str, Any] = dict(self.counts)
        tiers["basis"] = "tiered"
        return tiers, self.rows


def _excerpt(text: str, figure: Figure, *, width: int = 70) -> str:
    lo = max(0, figure.start - width)
    hi = min(len(text), figure.end + width)
    piece = " ".join(text[lo:hi].split())
    return (("…" if lo > 0 else "") + piece + ("…" if hi < len(text) else ""))[:220]


def check_memo_run(
    *,
    run_dir: Path,
    package: dict,
    company_id: str,
    research_dir: Path | None = None,
    session_dir: Path | None = None,
) -> FactCheckResult:
    """The gate's entry point: build the corpus for this run and check.

    Also audits the package's source URLs (report only, written into
    ``to_dict()["source_urls"]``). ``session_dir`` is the Memo Studio
    session the run was built from; when omitted it is read from the run
    manifest."""
    try:
        corpus = build_corpus(company_id, run_dir=run_dir, research_dir=research_dir)
        source_texts = _source_texts(package, company_id, run_dir)
        result = check_package(package, corpus, source_texts=source_texts)
    except Exception as exc:  # noqa: BLE001
        logger.warning("memo fact check failed", exc_info=True)
        result = FactCheckResult()
        result.error = f"{type(exc).__name__}: {exc}"
        return result
    try:
        result.source_urls = audit_source_urls(
            package, company_id=company_id, run_dir=run_dir, session_dir=session_dir
        )
    except Exception:  # noqa: BLE001 — report-only
        logger.warning("source URL audit failed", exc_info=True)
    return result


# ---- comparisons that state the wrong direction (2026-09-23) ----------------
#
# "about $597M — still below the August 2021 Series A post-money of
# $579.37M" (Claude's IC memo, ZaiNar 2026-09-23) and "The $1B+ mark sits in
# the top fifth of that band" of $315M-$1.4B (it sits at 63%) are arithmetic
# a reader checks in a second and a model gets wrong. Only a comparison whose
# subject is the first figure is read: between it and the word only a dash,
# a comma or a linking word may stand ("$597M — still below"), so "$2.36B on
# under $4.0M of revenue" (where "under" qualifies the second figure) is not
# a comparison of the two.

_COMPARATOR_RE = re.compile(
    r"\b(?P<cmp>below|under|less than|lower than|beneath|above|over|more than|higher than|greater than|exceeds|exceeding)\b",
    re.I,
)
_LINK_WORDS = frozenset(
    "is are was were sits sit stands stand remains remain still lies lie falls fall lands land "
    "comes come in well just slightly far only which that it and so already even".split()
)
_BELOW_WORDS = frozenset({"below", "under", "less than", "lower than", "beneath"})
_SENTENCE_BREAK_RE = re.compile(r"[.;:!?](?:\s|$)|\n")
_BAND_RE = re.compile(r"\b(?P<side>top|bottom|upper|lower)\s+(?P<part>fifth|quarter|third|half|tenth)\b", re.I)
_BAND_PARTS = {"tenth": 0.1, "fifth": 0.2, "quarter": 0.25, "third": 1 / 3, "half": 0.5}


def _money(figures: list[Figure]) -> list[Figure]:
    return [f for f in figures if f.klass == "amount" and re.match(rf"\s*{_CURRENCY}", f.raw)]


_CURRENCY_TOKEN_RE = re.compile(rf"\s*({_CURRENCY})", re.I)
_US_DOLLAR_TOKENS = {"$", "US$", "U.S.$", "USD"}


def _currency_of(figure: Figure) -> str:
    """The figure's currency, normalised ("$", "US$" and "USD" are one)."""
    match = _CURRENCY_TOKEN_RE.match(figure.raw)
    if not match:
        return ""
    token = re.sub(r"\s+", "", match.group(1)).upper()
    return "USD" if token in _US_DOLLAR_TOKENS else token


def _negative(text: str, figure: Figure) -> bool:
    """A minus sign or an accounting parenthesis directly before the
    figure ("−$5M", "($5M)") — a sign the parser does not carry. A range's
    dash ("$315M–$1.4B") follows a digit or a scale letter, not a space."""
    before = text[max(0, figure.start - 3):figure.start]
    return bool(re.search(r"(?:^|[\s(\[:,])[-−–—(]\s?\Z", before))


# A tail that makes the second figure something other than the comparison's
# object: a comparison of another thing ("over the last 12 months vs $12M",
# "under the new CEO, from $6M"), or a part of the first figure ("less than
# $3M of it", "more than half of the $8M total").
_TAIL_OTHER_OBJECT_RE = re.compile(
    r"\b(?:vs\.?|versus|compared|against|from|with|half|third|quarter|fifth|tenth|"
    r"twice|double|triple|times)\b",
    re.I,
)
_TIME_OR_PEOPLE_RE = re.compile(
    r"\b(?:years?|months?|quarters?|weeks?|days?|decades?|period|horizon|hold|holding|"
    r"life|lifetime|term|time|next|last|past|ceo|cfo|founders?|leadership|management)\b",
    re.I,
)
_PART_OF_RE = re.compile(r"\s+of\s+(?:it|which|them|that|this|those|these)\b", re.I)
_RANGE_JOIN_RE = re.compile(r"\s?(?:–|—|-|to)\s?|\s+and\s+", re.I)
# A band's own clause ends at a sentence break or a table-cell rule; its
# subject stands within this many characters of the band phrase ("The $1B+
# mark sits in the top fifth").
_BAND_SCOPE_BREAK_RE = re.compile(r"[.;:!?](?:\s|$)|\n|\|")
_BAND_SUBJECT_MAX_CHARS = 100
_NEGATION_BEFORE_RE = re.compile(r"\b(?:not|never|no longer|isn't|aren't|wasn't)\b[^.;:!?]{0,30}\Z", re.I)


def _money_ranges(text: str, money: list[Figure]) -> list[tuple[Figure, Figure]]:
    """Adjacent money figures written as a range: "$315M–$1.4B", "$315M to
    $1.4B", "$1.2–1.5B" (the parser shares the scale), "between $315M and
    $1.4B" — same currency, low below high."""
    out = []
    for low, high in zip(money, money[1:]):
        joined = text[low.end:high.start]
        if not _RANGE_JOIN_RE.fullmatch(joined):
            continue
        if "and" in joined.lower() and not re.search(r"\bbetween\s+\Z", text[max(0, low.start - 12):low.start], re.I):
            continue
        if _currency_of(low) != _currency_of(high) or high.value <= low.value:
            continue
        out.append((low, high))
    return out


def comparison_errors(text: str) -> list[dict]:
    """Comparisons in ``text`` whose direction the figures contradict:
    ``[{kind, snippet, detail}]`` — "comparison" (X below / above Y) and
    "band" (X in the top / bottom part of a stated range). Conservative:
    only two amounts in the same currency, with the first figure the
    comparison's subject ("$597M — still below $579M"), never a difference
    ("sits $165M above the $835M bound"), a part ("less than $3M of it"),
    another clause's figure or a negative."""
    cleaned = _strip_citations(text or "")
    figures = extract_figures(cleaned)
    money = _money(figures)
    out: list[dict] = []
    for match in _COMPARATOR_RE.finditer(cleaned):
        before = [f for f in figures if f.end <= match.start()]
        after = [f for f in figures if f.start >= match.end()]
        if not before or not after:
            continue
        subject, other = before[-1], after[0]
        if subject not in money or other not in money:
            continue
        if _currency_of(subject) != _currency_of(other):
            continue
        if _negative(cleaned, subject) or _negative(cleaned, other):
            continue
        gap = cleaned[subject.end:match.start()]
        # Nothing between the figure and the word: the figure is a
        # difference ("sits $165M above the bound"), not the subject.
        if not gap.strip():
            continue
        words = re.findall(r"[A-Za-z]+", gap)
        if _SENTENCE_BREAK_RE.search(gap) or any(w.lower() not in _LINK_WORDS for w in words) or len(words) > 4:
            continue
        tail = cleaned[match.end():other.start]
        # The second figure must be the comparison's object, not the next
        # clause's ("sits $165M above the upper bound and nearly doubles the
        # $505M midpoint" compares nothing to $505M).
        if (
            _SENTENCE_BREAK_RE.search(tail)
            or "," in tail
            or len(re.findall(r"[A-Za-z0-9]+", tail)) > 12
            or re.search(r"\b(?:and|or|but|while|whereas)\b", tail, re.I)
            or _TAIL_OTHER_OBJECT_RE.search(tail)
            or _PART_OF_RE.match(cleaned[other.end:other.end + 12])
        ):
            continue
        comparator = match.group("cmp").lower()
        # "over" / "under" also mean time and people ("over three years",
        # "under the new CEO").
        if comparator in {"over", "under"} and _TIME_OR_PEOPLE_RE.search(tail):
            continue
        # Equal values are no error.
        if math.isclose(subject.value, other.value, rel_tol=1e-6):
            continue
        below = comparator in _BELOW_WORDS
        wrong = subject.value > other.value if below else subject.value < other.value
        if wrong:
            out.append(
                {
                    "kind": "comparison",
                    "snippet": " ".join(cleaned[max(0, subject.start - 40):other.end + 20].split()),
                    "detail": (
                        f"{subject.raw} is {'above' if below else 'below'} {other.raw}, "
                        f"not {comparator} it"
                    ),
                }
            )
    ranges = _money_ranges(cleaned, money)
    for match in _BAND_RE.finditer(cleaned):
        sentence_start = max(
            (m.end() for m in _BAND_SCOPE_BREAK_RE.finditer(cleaned, 0, match.start())), default=0
        )
        sentence_end_match = _BAND_SCOPE_BREAK_RE.search(cleaned, match.end())
        sentence_end = sentence_end_match.start() if sentence_end_match else len(cleaned)
        if _NEGATION_BEFORE_RE.search(cleaned[sentence_start:match.start()]):
            continue
        # The range the band names: one right after it in the same sentence
        # ("in the top fifth of $315M–$1.4B"), else the last one before it.
        following = [
            r for r in ranges
            if match.end() <= r[0].start <= min(sentence_end, match.end() + 60)
        ]
        preceding = [r for r in ranges if r[1].end <= match.start()]
        chosen = following[0] if following else (preceding[-1] if preceding else None)
        if chosen is None:
            continue
        low_fig, high_fig = chosen
        subjects = [
            f for f in money
            if sentence_start <= f.start and f.end <= match.start()
            and match.start() - f.end <= _BAND_SUBJECT_MAX_CHARS
            and f is not low_fig and f is not high_fig
            and _currency_of(f) == _currency_of(low_fig)
        ]
        if not subjects:
            continue
        subject = subjects[-1]
        low, high, value = low_fig.value, high_fig.value, subject.value
        position = (value - low) / (high - low)
        part = _BAND_PARTS[match.group("part").lower()]
        side = match.group("side").lower()
        inside = position >= 1 - part if side in {"top", "upper"} else position <= part
        if not inside:
            out.append(
                {
                    "kind": "band",
                    "snippet": " ".join(cleaned[max(0, min(low_fig.start, subject.start)):match.end() + 20].split())[:220],
                    "detail": (
                        f"{subject.raw} sits at {position:.0%} of the {low_fig.raw}–"
                        f"{high_fig.raw} range, not in its {side} {match.group('part').lower()}"
                    ),
                }
            )
    return out


# ---- calculation inputs cited to a source that does not carry them ----------
#
# A calculation note's inputs are what every derived figure rests on, so an
# input is where an invented number does the most damage: ZaiNar 2026-09-23
# (Gemini, v2) pinned "Estimated Baseline Contracted Revenue = $30M [S2]",
# where S2 is a paywalled headline carrying no revenue figure at all, and
# the 33.3x entry multiple, the $400M-$750M fair value and the risk ratings
# were all computed from it. The prose check could not see it: $30M occurs
# elsewhere in the corpus, attached to something else, so the loose match
# counted it as "supported". An input passes when a source it cites carries
# the value, or when some text on file carries the same value within
# CONTEXT_WINDOW_CHARS of a word from the input's own name ("revenue",
# "Skyhook"); an input that fails both is stated as what it is — our
# assumption — and the run says so.

CALCULATION_INPUTS_FILENAME = "calculation_inputs.json"
_CALC_INPUT_REF_RE = re.compile(r"\bS\d+\b")


def _input_keywords(name: str, value: str, company_tokens: set[str]) -> set[str]:
    words = _context_words(f"{name} {value}")
    return {w for w in words if w not in company_tokens and not any(ch.isdigit() for ch in w)}


# A scale heading over a column of bare numbers: "(Billions)", "(USD
# millions)", "in $ billions" — never any scale word nearby ("$5 million
# last year, with $30 average revenue per user" scales nothing).
_HEADER_SCALE_RE = re.compile(
    r"\(\s*(?:in\s+)?(?:(?:USD|US\$|\$)\s*)?(billion|bn|million|mn|thousand)s?\s*(?:USD)?\s*\)"
    r"|\bin\s+(?:USD\s+|US\$\s*|\$\s*)?(billion|million|thousand)s\b",
    re.I,
)


def _number_spellings(figure: Figure) -> set[str]:
    """The digit strings a text may write the figure's value with: as
    written, and at each scale ("$10,000,000" is "10" million)."""
    spellings = {figure.mantissa.replace(",", "")}
    if figure.klass == "amount":
        for scale in (1e3, 1e6, 1e9):
            scaled = figure.value / scale
            if scaled >= 1:
                spellings.add(f"{scaled:.6f}".rstrip("0").rstrip("."))
    return {s for s in spellings if s and any(ch.isdigit() and ch != "0" for ch in s)}


def _window_figures(flat: str, start: int, end: int, figure: Figure) -> list[Figure]:
    """The figures at one occurrence. A bare number under a scale heading
    ("Market Size (Billions) ~USD 18.31") takes the heading's scale."""
    found = extract_figures(flat[max(0, start - 8):end + 16])
    if figure.klass != "amount" or figure.value < 1e5:
        return found
    heading = [
        next(group for group in groups if group)
        for groups in _HEADER_SCALE_RE.findall(flat[max(0, start - 60):start])
    ]
    if not heading:
        return found
    factor = _SCALES.get(heading[-1].lower().rstrip("s"), 1.0)
    return found + [
        Figure(raw=f.raw, value=f.value * factor, klass="amount", start=f.start, end=f.end, mantissa=f.mantissa)
        for f in found
        if f.klass == "amount" and f.value < 1e5
    ]


def _value_in_context(figure: Figure, keywords: set[str] | None, texts: list[str]) -> bool:
    """Some text carries the figure's value (same number and scale, within
    rounding) with one of ``keywords`` inside the context window; with
    ``keywords`` None, anywhere (a source the input cites)."""
    if keywords is not None and not keywords:
        return False
    spellings = _number_spellings(figure)
    for text in texts:
        flat = text.replace(",", "") if "," in text else text
        for needle in spellings:
            # Only the leading edge is a boundary: "18.3" may be the start
            # of "18.31", which rounds to it (the parse decides). Positions
            # inside another number never count against the cap.
            pattern = re.compile(r"(?<![\d.])" + re.escape(needle))
            for seen, hit in enumerate(pattern.finditer(flat)):
                if seen >= MAX_CONTEXT_POSITIONS:
                    break
                start, end = hit.start(), hit.end()
                window = flat[max(0, start - CONTEXT_WINDOW_CHARS):end + CONTEXT_WINDOW_CHARS]
                if keywords is not None and not keywords & _context_words(window):
                    continue
                for candidate in _window_figures(flat, start, end, figure):
                    single: dict[str, list[float]] = {"amount": [], "pct": [], "mult": []}
                    single[candidate.klass] = [candidate.value]
                    if _value_supported(figure, single):
                        return True
    return False


def unsourced_calculation_inputs(
    package: dict,
    corpus: Corpus,
    source_texts: dict[str, str],
) -> list[dict]:
    """Calculation inputs whose ref names sources ("S2") that do not carry
    the value, and that no text on file carries in context either.
    ``[{calc_id, input_index, name, value, ref}]``; an input with no figure,
    or whose ref is an assumption or another note, is never listed."""
    out: list[dict] = []
    source_texts = source_texts or {}
    indexes = {sid: TextIndex.of(text) for sid, text in source_texts.items()}
    texts: list[str] | None = None
    tokens = set(getattr(corpus, "company_tokens", None) or ())
    for calc in package.get("calculations") or []:
        if not isinstance(calc, dict):
            continue
        for index, item in enumerate(calc.get("inputs") or []):
            if not isinstance(item, dict):
                continue
            ref = str(item.get("ref") or "")
            cited = _CALC_INPUT_REF_RE.findall(ref)
            # An input that is also another note's result is derived there.
            if not cited or re.search(r"\bC\d+\b", ref):
                continue
            value = _loc(item.get("value"))
            figures = extract_figures(value)
            if not figures:
                continue
            # Carried by a cited source: as indexed, or as a bare number
            # under the source's own scale heading ("(Billions) ~USD 18.31").
            missing = [
                figure
                for figure in figures
                if not any(sid in indexes and indexes[sid].supports(figure) for sid in cited)
                and not _value_in_context(
                    figure, None, [source_texts[sid] for sid in cited if source_texts.get(sid)]
                )
            ]
            if not missing:
                continue
            name = _loc(item.get("name"))
            keywords = _input_keywords(name, value, tokens)
            if texts is None:
                texts = [text for _, text in corpus.texts] + [text for _, _, text in corpus.tier_texts]
            if all(_value_in_context(figure, keywords, texts) for figure in missing):
                continue
            out.append(
                {
                    "calc_id": str(calc.get("id") or ""),
                    "input_index": index,
                    "name": name,
                    "value": value,
                    "ref": ref,
                    "cited_text_chars": {sid: len(source_texts.get(sid) or "") for sid in cited},
                }
            )
    return out


def relabel_unsourced_calculation_inputs(package: dict, entries: list[dict]) -> int:
    """Set each listed input's ref to "assumption" (the renderer prints it as
    our assumption in both languages). Returns how many were changed."""
    by_id = {
        str(calc.get("id") or ""): calc
        for calc in package.get("calculations") or []
        if isinstance(calc, dict)
    }
    changed = 0
    for entry in entries:
        calc = by_id.get(str(entry.get("calc_id") or ""))
        inputs = calc.get("inputs") if isinstance(calc, dict) else None
        index = entry.get("input_index")
        if not isinstance(inputs, list) or not isinstance(index, int) or not 0 <= index < len(inputs):
            continue
        item = inputs[index]
        if isinstance(item, dict) and str(item.get("ref") or "") == entry.get("ref"):
            item["ref"] = "assumption"
            changed += 1
    return changed


def final_check(
    *,
    run_dir: Path,
    package: dict,
    company_id: str,
    research_dir: Path | None = None,
    session_dir: Path | None = None,
) -> tuple[FactCheckResult, list[dict]]:
    """The check of the package as delivered, after every repair: calculation
    inputs no source carries are relabelled as assumptions (in place), then
    every figure is checked. Returns the result and the relabelled inputs.
    Never raises."""
    relabelled: list[dict] = []
    try:
        corpus = build_corpus(company_id, run_dir=run_dir, research_dir=research_dir)
        source_texts = _source_texts(package, company_id, run_dir)
        if not corpus.thin:
            relabelled = unsourced_calculation_inputs(package, corpus, source_texts)
            relabel_unsourced_calculation_inputs(package, relabelled)
        result = check_package(package, corpus, source_texts=source_texts)
    except Exception as exc:  # noqa: BLE001
        logger.warning("final memo fact check failed", exc_info=True)
        result = FactCheckResult()
        result.error = f"{type(exc).__name__}: {exc}"
        return result, relabelled
    try:
        result.source_urls = audit_source_urls(
            package, company_id=company_id, run_dir=run_dir, session_dir=session_dir
        )
    except Exception:  # noqa: BLE001 — report-only
        logger.warning("source URL audit failed", exc_info=True)
    return result, relabelled


_TIER_LINES = (
    ("verified", "verified in a cited source that is not the subject's own site"),
    ("found_elsewhere", "found in independent evidence on file"),
    ("in_context", "  of which the source names the same metric near the number"),
    ("derived", "derived in a calculation note"),
    ("company_reported", "only on the subject company's own pages"),
    ("company_reported_headline", "  of which in a headline section"),
    ("registry_only", "only in registry fields with source_refs"),
    ("not_traced", "not traced to evidence"),
)


def render_markdown_report(result: FactCheckResult, *, attempt: int | None = None) -> str:
    lines = ["# Memo fact check", ""]
    if attempt is not None:
        lines.append(f"Attempt: {attempt}")
    lines.extend(
        [
            f"Status: {result.status}",
            f"Figures checked: {result.checked}",
            f"- verified in the cited source: {result.verified}",
            f"- supported by a source on file: {result.supported}",
            f"- derived in a calculation note: {result.derived}",
            f"- unsupported: {result.unsupported}",
            f"- citations without text on file: {result.unverifiable_citations}",
            f"Coverage: {result.coverage_pct if result.coverage_pct is not None else 'n/a'}%",
            f"Corpus: {result.corpus_summary}",
            f"Enforced (fed to repair): {'yes' if result.repair_feed else 'no'} — {result.repair_feed_reason}",
        ]
    )
    if result.error:
        lines.append(f"Error: {result.error}")
    if result.tiers:
        lines.extend(["", "## What the figures rest on (report only, not enforced)"])
        for key, label in _TIER_LINES:
            if key in result.tiers:
                lines.append(f"- {label}: {result.tiers[key]}")
        flagged = [row for row in result.figure_tiers if row.get("tier") == "company_reported" and row.get("headline")]
        for row in flagged[:12]:
            lines.append(
                f"  - headline figure {row['figure']} in {row['section_id']} rests only on the "
                f"company's own pages: \"{row['snippet']}\""
            )
    if result.source_urls:
        lines.extend(["", "## Source URLs (report only)"])
        lines.extend(_source_url_summary_lines(result.source_urls))
    if result.corpus_texts:
        lines.extend(["", "## Sources on file"])
        for row in result.corpus_texts[:80]:
            lines.append(f"- {row['label']} ({row['chars']:,} chars)")
        if len(result.corpus_texts) > 80:
            lines.append(f"- … and {len(result.corpus_texts) - 80} more")
    lines.extend(["", "## Findings"])
    if not result.findings:
        lines.append("None.")
    for finding in result.findings:
        lines.append(
            f"- [{finding.severity}] {finding.code} — {finding.figure} in {finding.section_id} "
            f"({finding.location}): \"{finding.snippet}\" — {finding.detail}"
        )
    return "\n".join(lines) + "\n"


# ---- deterministic URL attachment ------------------------------------------


def _artifact_url_candidates(run_dir: Path | None) -> list[tuple[str, str]]:
    """``(source name, url)`` pairs the analysis passes recorded.

    ``_write_fast_pass_outputs`` nests the pass output under ``data``
    (``{"pass_id", "status", …, "data": {"supporting_evidence": […]}}``);
    reading ``supporting_evidence`` at the top level attached nothing. A
    file without ``data`` is read as the bare pass output."""
    if run_dir is None:
        return []
    fast_dir = Path(run_dir) / "analysis" / "fast"
    out: list[tuple[str, str]] = []
    try:
        files = sorted(fast_dir.glob("*.json"))
    except OSError:
        return []
    for path in files:
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        if not isinstance(payload, dict):
            continue
        data = payload.get("data", payload)
        if not isinstance(data, dict):
            continue  # a failed pass: "data": null
        for item in data.get("supporting_evidence") or []:
            if not isinstance(item, dict):
                continue
            url = item.get("url")
            if source_cache.canonical_url(url):
                out.append((str(item.get("source") or ""), str(url).strip()))
    return out


_URL_IN_TEXT_RE = re.compile(r"https?://[^\s<>()\[\]{}\"'`|]+", re.I)


def _analysis_markdown_urls(run_dir: Path | None) -> set[str]:
    """Seen-keys of every URL written into the run's analysis artifacts
    (``analysis/*.md`` — the legacy pipeline's passes write prose there)."""
    if run_dir is None:
        return set()
    out: set[str] = set()
    try:
        files = sorted((Path(run_dir) / "analysis").glob("*.md"))
    except OSError:
        return out
    for path in files:
        try:
            text = path.read_text(encoding="utf-8")
        except OSError:
            continue
        for match in _URL_IN_TEXT_RE.finditer(text):
            key = _seen_key(match.group(0).rstrip(".,;:"))
            if key:
                out.add(key)
    return out


_SESSION_ID_LINE_RE = re.compile(r"^- analysis_session_id:\s*(\S+)\s*$", re.M)


def _session_dir_for_run(run_dir: Path | None, company_id: str) -> Path | None:
    """The Memo Studio session a run was prepared from, read from the
    ``analysis_session_id`` line ``memo_prep`` writes into
    ``logs/run_manifest.md``; None when there is none on disk."""
    if run_dir is None:
        return None
    try:
        text = (Path(run_dir) / "logs" / "run_manifest.md").read_text(encoding="utf-8")
    except OSError:
        return None
    match = _SESSION_ID_LINE_RE.search(text)
    if not match:
        return None
    try:
        from . import serena_analysis

        path = serena_analysis.session_dir(company_id, match.group(1))
    except (ValueError, ImportError):
        return None
    return path if path.is_dir() else None


def _serena_url_candidates(session_dir: Path | None) -> list[dict]:
    """URLs a Memo Studio session recorded, with their excerpts: every
    ``source_traces`` row and every ``supporting_evidence`` /
    ``contradicting_evidence`` row (the strategic risk map's and the
    research tasks') whose ``url`` or ``locator`` is a web page.

    Rows are ``{"name", "url", "excerpt", "origin"}``. ``name`` is the
    row's own title (source traces carry one; risk evidence does not), so
    title matching runs on it and the URL's words — never on the excerpt,
    whose many words would let a weak title match."""
    if session_dir is None:
        return []
    root = Path(session_dir)
    documents: list[Any] = []
    session_file = root / "session.yaml"
    try:
        if session_file.is_file():
            session = yaml.safe_load(session_file.read_text(encoding="utf-8"))
            if isinstance(session, dict) and isinstance(session.get("artifacts"), dict):
                documents.append(session["artifacts"])
        if not documents:
            for path in sorted(root.glob("*.yaml")):
                documents.append(yaml.safe_load(path.read_text(encoding="utf-8")))
    except (OSError, yaml.YAMLError):
        logger.warning("Serena session unreadable for URL candidates: %s", root, exc_info=True)
        return []
    out: list[dict] = []
    seen: set[tuple[str, str]] = set()

    def take(row: Any, origin: str) -> None:
        if not isinstance(row, dict):
            return
        url = ""
        for key in ("url", "locator"):
            value = str(row.get(key) or "").strip()
            if value.lower().startswith(("http://", "https://")) and source_cache.canonical_url(value):
                url = value
                break
        if not url:
            return
        excerpt = " ".join(str(row.get("excerpt") or "").split())[:300]
        key = (_seen_key(url) or url, excerpt[:120])
        if key in seen:
            return
        seen.add(key)
        name = str(row.get("title") or row.get("source_title") or row.get("source") or "").strip()
        out.append({"name": name[:200], "url": url, "excerpt": excerpt, "origin": origin})

    def walk(node: Any) -> None:
        if isinstance(node, dict):
            for key, value in node.items():
                if key == "source_traces" and isinstance(value, list):
                    for row in value:
                        take(row, "Serena source trace")
                elif key in {"supporting_evidence", "contradicting_evidence"} and isinstance(value, list):
                    for row in value:
                        take(row, "Serena evidence locator")
                walk(value)
        elif isinstance(node, list):
            for item in node:
                walk(item)

    for document in documents:
        walk(document)
    return out


def _url_tokens(name: str, url: str) -> set[str]:
    return source_cache.significant_tokens(name) | source_cache.significant_tokens(
        re.sub(r"[/._-]+", " ", url.split("://", 1)[-1])
    )


def _best_url_match(title: str, candidates: list[tuple[str, str]], *, min_score: float = 0.6) -> tuple[str, str, float] | None:
    query = source_cache.significant_tokens(title)
    if len(query) < 2:
        return None
    best: tuple[str, str, float] | None = None
    for name, url in candidates:
        tokens = _url_tokens(name, url)
        if not tokens:
            continue
        score = len(query & tokens) / len(query)
        if score >= min_score and (best is None or score > best[2]):
            best = (name, url, score)
    return best


def _url_match_ties(title: str, candidates: list[tuple[str, str]], best: tuple[str, str, float]) -> list[str]:
    """Other pages that match ``title`` exactly as well as ``best`` does."""
    query = source_cache.significant_tokens(title)
    if not query:
        return []
    seen_keys = {_seen_key(best[1])}
    ties: list[str] = []
    for name, url in candidates:
        key = _seen_key(url)
        if not key or key in seen_keys:
            continue
        tokens = _url_tokens(name, url)
        if tokens and math.isclose(len(query & tokens) / len(query), best[2]):
            seen_keys.add(key)
            ties.append(url)
    return ties


def _seen_key(url: Any) -> str | None:
    """A comparable key for "the same page": the canonical URL without its
    scheme or a leading ``www.`` (so ``http://www.x.com/a`` and
    ``https://x.com/a`` are one page)."""
    canon = source_cache.canonical_url(url)
    if not canon:
        return None
    parts = urlsplit(canon)
    host = (parts.netloc or "").lower()
    if host.startswith("www."):
        host = host[4:]
    return f"{host}{parts.path or '/'}" + (f"?{parts.query}" if parts.query else "")


def _is_bare_homepage(url: Any) -> bool:
    canon = source_cache.canonical_url(url)
    if not canon:
        return False
    parts = urlsplit(canon)
    return parts.path in ("", "/") and not parts.query


# Files the pipeline writes into a research folder by itself, from the web:
# the known-sources digest and the tracked-news digest. They can vouch for a
# figure, but they are not private material — a source drawn from them came
# from a page a reader can open, so it has a URL to carry.
_WEB_DIGEST_FILES = {"known_sources.md", "recent_news.md"}


def private_material_on_file(
    company_id: str, *, research_dir: Path | None = None
) -> list[str]:
    """What the firm holds on this company that a reader cannot browse to:
    research-folder documents (not the web digests the pipeline writes
    there), founder updates and KPI rows, call transcripts and reference
    calls. One label per kind found; empty means there is nothing private
    on file, so no source in the memo can honestly be one.

    The renderer's URL rule exempts a source whose class or title says it is
    private. Nothing checked that claim, and live on 2026-09-21 a RadixArk
    Gemini memo — no research folder, no uploads, nothing private at all —
    classed three public sources as "internal document" and passed, because
    the rule's own error message told it that was the way out.
    """
    found: list[str] = []
    if research_dir is None:
        try:
            research_dir = research_store.RESEARCH_ROOT / company_paths.storage_key(company_id)
        except ValueError:
            research_dir = None
    root = Path(research_dir) if research_dir is not None else None
    if root is not None and root.is_dir():
        try:
            for path in sorted(p for p in root.iterdir() if p.is_file()):
                name = path.name
                if (
                    name.startswith("index.yaml")
                    or name.endswith(".progress.jsonl")
                    or name in _WEB_DIGEST_FILES
                ):
                    continue
                if path.stat().st_size > 0:
                    found.append(f"research file: {name}")
        except OSError:
            logger.warning("private material: could not list %s", root, exc_info=True)
    try:
        from . import portfolio

        if portfolio.has_record(company_id):
            record = portfolio.get_portfolio(company_id)
            if record.get("updates") or record.get("kpis"):
                found.append("founder updates / KPI rows")
    except Exception:  # noqa: BLE001
        logger.debug("private material: no portfolio record for %s", company_id, exc_info=True)
    try:
        from . import transcripts

        if any(t.get("company_id") == company_id for t in transcripts.all_transcripts()):
            found.append("call transcripts")
    except Exception:  # noqa: BLE001
        logger.debug("private material: transcripts unavailable", exc_info=True)
    try:
        from . import ic_room

        if ic_room.list_reference_calls(company_id).get("items"):
            found.append("reference calls")
    except Exception:  # noqa: BLE001
        logger.debug("private material: reference calls unavailable", exc_info=True)
    return found


def stamp_private_material(
    package: dict, *, company_id: str, research_dir: Path | None = None
) -> list[str]:
    """Record on the package envelope whether the firm holds anything private
    on this company, so the renderer's URL rule can tell an honest "internal
    document" from a public page that dropped its URL. Returns the labels."""
    found = private_material_on_file(company_id, research_dir=research_dir)
    if isinstance(package, dict):
        run = package.get("run")
        if not isinstance(run, dict):
            run = {}
            package["run"] = run
        run["private_material_on_file"] = bool(found)
    return found


# ---- the run's private inventory (G1) ----------------------------------------
#
# ``private_material_on_file`` answers one package-wide yes/no, and it said
# yes for things the writer never saw (a deselected document) or that are
# not private at all (the user-edited fact ledger, the pipeline's own
# decision record). The inventory names each private item this run may
# actually read, so a URL-less "private" source can be checked against it
# one by one (:func:`inventory_match`).

# Digests the pipeline or the user writes into the research folder. They are
# not private material a source can be: the fact ledger holds public facts,
# the decision record is the pipeline's own, and the other two are web
# digests (claude_runner.MEMO_*_FILENAME).
_NON_PRIVATE_DIGESTS = {
    "fact_ledger.md",
    "decision_record.md",
    "recent_news.md",
    "known_sources.md",
}
MAX_PRIVATE_INVENTORY = 150
_BSH_TAG_RE = re.compile(r"\b(?:bsh|diligence)\b", re.I)
# A tag that names demo or placeholder data is not diligence, whatever else
# it says ("Demo placeholder …, not diligence").
_NOT_DILIGENCE_RE = re.compile(
    r"\b(?:demo|mock|mockup|placeholder|sample|fixture|unverified|not\s+(?:bsh\s+)?diligence)\b",
    re.I,
)


def _diligence_tagged(ref: Any) -> bool:
    if not isinstance(ref, dict):
        return False
    tag = " ".join(
        str(ref.get(key) or "") for key in ("source_class", "label", "title", "origin")
    )
    return bool(_BSH_TAG_RE.search(tag)) and not _NOT_DILIGENCE_RE.search(tag)


def _entry_title(entry: dict) -> str:
    filename = str(entry.get("filename") or entry.get("stored_name") or "").strip()
    label = str(entry.get("label") or "").strip()
    if label and filename and label.lower() != filename.lower():
        return f"{label} — {filename}"[:200]
    return (label or filename)[:200]


def private_inventory(
    run_dir: Path | None,
    company: dict | str,
    *,
    research_dir: Path | None = None,
    extra_items: list[dict] | None = None,
) -> list[dict]:
    """The private items this run may read, as ``{"id", "kind", "title",
    "ref"}`` rows:

    - research-folder documents the run's listing offers
      (``claude_runner._research_file_listing`` rules): uploads
      (``research_document``), document analyses (``research_analysis``,
      plus one ``research_document`` row per document an analysis distils,
      whose ``ref`` is the analysis file), and files dropped into the
      folder outside the upload index (``research_file``). The run's
      ``logs/evidence_selection.json``
      (``claude_runner.memo_evidence_selection``) is honoured: a deselected
      analysis and the documents behind it are left out. The digests
      ``fact_ledger.md``, ``decision_record.md``, ``recent_news.md`` and
      ``known_sources.md`` never count;
    - registry metrics whose ``source_refs`` are tagged BSH / diligence
      (``registry_metric``) — not ones labelled demo, placeholder or
      unverified;
    - ``extra_items``: rows the caller stages itself (the call and founder
      digests, once they exist), passed through with the same four keys.

    ``company`` is the registry record or the company id."""
    if isinstance(company, dict):
        record: dict = company
        company_id = str(company.get("id") or "")
    else:
        company_id = str(company or "")
        record = storage.get_company(company_id) or {}
    items: list[dict] = []
    seen_ids: set[str] = set()

    def add(item_id: str, kind: str, title: str, ref: str) -> None:
        if not item_id or item_id in seen_ids or not str(title or "").strip():
            return
        seen_ids.add(item_id)
        items.append({"id": item_id, "kind": kind, "title": str(title)[:200], "ref": str(ref)[:300]})

    if research_dir is None and company_id:
        try:
            research_dir = research_store.RESEARCH_ROOT / company_paths.storage_key(company_id)
        except ValueError:
            research_dir = None
    root = Path(research_dir) if research_dir is not None else None
    if root is not None and root.is_dir():
        from . import claude_runner

        selected = claude_runner.memo_evidence_selection(Path(run_dir)) if run_dir is not None else None
        try:
            entries = research_store.list_files(company_id) if company_id else []
        except Exception:  # noqa: BLE001
            logger.warning("private inventory: research index unreadable for %s", company_id, exc_info=True)
            entries = []
        by_stored = {str(e.get("stored_name") or ""): e for e in entries if isinstance(e, dict) and e.get("stored_name")}
        by_id = {str(e.get("id") or ""): e for e in entries if isinstance(e, dict)}
        covered: set[str] = set()
        dropped: set[str] = set()
        for entry in by_id.values():
            target = entry.get("analysis_of")
            if not target:
                continue
            target = str(target)
            if selected is not None and str(entry.get("id")) not in selected:
                dropped.add(str(entry.get("id")))
            if target.startswith("fld"):
                covered |= {str(m.get("id")) for m in by_id.values() if m.get("folder_id") == target}
            else:
                covered.add(target)
        try:
            names = sorted(
                p.name
                for p in root.iterdir()
                if p.is_file()
                and not p.name.startswith("index.yaml")
                and not p.name.endswith(".progress.jsonl")
            )
        except OSError:
            logger.warning("private inventory: could not list %s", root, exc_info=True)
            names = []
        for name in names:
            if name in _NON_PRIVATE_DIGESTS:
                continue
            try:
                if (root / name).stat().st_size <= 0:
                    continue
            except OSError:
                continue
            entry = by_stored.get(name)
            if entry is None:
                add(f"file:{name}", "research_file", name, name)
                continue
            entry_id = str(entry.get("id") or "")
            if entry_id in covered or entry_id in dropped:
                continue
            target = entry.get("analysis_of")
            if not target:
                add(entry_id, "research_document", _entry_title(entry), name)
                continue
            add(entry_id, "research_analysis", _entry_title(entry), name)
            target = str(target)
            if target.startswith("fld"):
                members = [m for m in by_id.values() if m.get("folder_id") == target]
                folder_name = next((m.get("folder_name") for m in members if m.get("folder_name")), None)
                if folder_name:
                    add(f"{entry_id}:{target}", "research_document", str(folder_name), name)
                for member in members:
                    add(f"{entry_id}:{member.get('id')}", "research_document", _entry_title(member), name)
            elif target in by_id:
                add(f"{entry_id}:{target}", "research_document", _entry_title(by_id[target]), name)

    metrics = record.get("metrics") if isinstance(record, dict) else None
    for index, metric in enumerate(metrics if isinstance(metrics, list) else []):
        if not isinstance(metric, dict):
            continue
        refs = [ref for ref in metric.get("source_refs") or [] if _diligence_tagged(ref)]
        if not refs:
            continue
        label = str(metric.get("label") or metric.get("name") or f"metric {index + 1}").strip()
        value = str(metric.get("value") or "").strip()
        ref_label = str(refs[0].get("label") or refs[0].get("title") or refs[0].get("source_class") or "").strip()
        title = f"{ref_label}: {label} {value}".strip() if ref_label else f"{label} {value}".strip()
        add(f"registry:metrics[{index}]", "registry_metric", title, f"companies.yaml#{company_id}/metrics[{index}]")

    for item in extra_items or []:
        if isinstance(item, dict):
            add(
                str(item.get("id") or ""),
                str(item.get("kind") or "staged_digest"),
                str(item.get("title") or ""),
                str(item.get("ref") or ""),
            )
    return items[:MAX_PRIVATE_INVENTORY]


def stamp_private_inventory(package: dict, inventory: list[dict]) -> list[dict]:
    """Record the run's private inventory on the package envelope
    (``package["run"]["private_inventory"]``), next to the older
    ``private_material_on_file`` boolean. Returns what was stamped."""
    rows = [
        {key: str(item.get(key) or "") for key in ("id", "kind", "title", "ref")}
        for item in inventory or []
        if isinstance(item, dict)
    ][:MAX_PRIVATE_INVENTORY]
    if isinstance(package, dict):
        run = package.get("run")
        if not isinstance(run, dict):
            run = {}
            package["run"] = run
        run["private_inventory"] = rows
    return rows


def _title_tokens(value: Any) -> set[str]:
    text = str(value or "")
    text = re.sub(r"^[0-9a-f]{12}__", "", text)  # research_store's "<id>__" prefix
    text = re.sub(r"\.(?:pdf|pptx?|docx?|txt|md|png|jpe?g|gif|webp|xlsx?|csv)\b", " ", text, flags=re.I)
    return source_cache.significant_tokens(re.sub(r"[_.]+", " ", text))


def inventory_match(source: dict, inventory: list[dict] | None) -> dict | None:
    """The inventory item a URL-less private-class source names, or None.

    A ``private_ref`` on the source must equal an item's ``id`` or ``ref``
    (or its whole title). Otherwise the source title must name the item:
    every significant word of the item's title appears in the source title
    (at least two words), or at least 70% of the source title's words are
    the item's. Pure — the renderer may import it."""
    if not isinstance(source, dict) or not inventory:
        return None
    private_ref = str(source.get("private_ref") or "").strip().lower()
    if private_ref:
        for item in inventory:
            if not isinstance(item, dict):
                continue
            if private_ref in {
                str(item.get("id") or "").strip().lower(),
                str(item.get("ref") or "").strip().lower(),
                " ".join(str(item.get("title") or "").lower().split()),
            }:
                return item
    title = _loc(source.get("title")) or str(source.get("title") or "")
    query = _title_tokens(title)
    if len(query) < 2:
        return None
    best: tuple[float, dict] | None = None
    for item in inventory:
        if not isinstance(item, dict):
            continue
        tokens = _title_tokens(item.get("title"))
        if not tokens:
            continue
        if len(tokens) >= 2 and tokens <= query:
            score = 1.0
        else:
            score = len(query & tokens) / len(query)
        if score >= 0.7 and (best is None or score > best[0]):
            best = (score, item)
    return best[1] if best else None


def attach_source_urls(
    package: dict,
    *,
    company_id: str,
    run_dir: Path | None,
    session_dir: Path | None = None,
    attempt: int | None = None,
    write_report: bool = True,
) -> list[str]:
    """Fill in ``url`` on package sources that lack one, from what this
    run's analysis passes recorded, the company's source cache and the
    Memo Studio session's locators. Deterministic and logged: returns one
    line per attachment. Only a strong title match attaches, so a wrong URL
    is far rarer than a missing one — and a missing one fails validation,
    where the repair can still supply it.

    Candidates, in order: the analysis passes' ``supporting_evidence``
    (score ≥ 0.6), the known-pages list of the source cache (≥ 0.7), then
    the Serena session's source traces and evidence locators (≥ 0.7, and
    refused when two different pages tie — a wrong link is worse than
    none). ``session_dir`` defaults to the session named in the run
    manifest.

    With a ``run_dir`` (and ``write_report``), the URL status after
    attachment — every URL seen or unseen, bare homepages reused — is
    appended to ``logs/source_urls.md`` (see :func:`audit_source_urls`) and
    stamped on ``package["run"]["source_url_status"]``
    (:func:`stamp_source_url_status`). Nothing is ever removed from the
    package."""
    if not isinstance(package, dict):
        return []
    sources = package.get("sources")
    if not isinstance(sources, list):
        return []
    artifact_candidates = _artifact_url_candidates(run_dir)
    try:
        page_candidates = [
            (str(page.get("title") or ""), str(page.get("url") or ""))
            for page in source_cache.known_pages(company_id, limit=200)
        ]
    except ValueError:
        page_candidates = []
    session = session_dir if session_dir is not None else _session_dir_for_run(run_dir, company_id)
    serena_rows = _serena_url_candidates(session)
    serena_candidates = [(row["name"], row["url"]) for row in serena_rows]
    excerpts = {row["url"]: row["excerpt"] for row in serena_rows}
    notes: list[str] = []
    attachments: list[dict] = []
    refused: list[dict] = []
    origins = (
        ("analysis artifact", artifact_candidates, 0.6, False),
        ("known source", page_candidates, 0.7, False),
        ("Serena session locator", serena_candidates, 0.7, True),
    )
    for source in sources:
        if not isinstance(source, dict):
            continue
        if str(source.get("url") or "").strip():
            continue
        title = _loc(source.get("title")) or str(source.get("title") or "")
        if not title.strip():
            continue
        for origin, candidates, min_score, refuse_ties in origins:
            match = _best_url_match(title, candidates, min_score=min_score)
            if match is None:
                continue
            ties = _url_match_ties(title, candidates, match)
            if ties and refuse_ties:
                refused.append(
                    {
                        "id": str(source.get("id") or ""),
                        "title": title[:120],
                        "origin": origin,
                        "urls": [match[1], *ties][:5],
                        "score": round(match[2], 2),
                    }
                )
                break
            name, url, score = match
            source["url"] = url
            # Risk-evidence locators carry no title; show their excerpt
            # (display only — matching never reads it).
            shown = name or (excerpts.get(url) or "")[:80]
            notes.append(
                f"{source.get('id')} ← {url} (matched {origin} '{shown[:80]}', score {score:.2f})"
            )
            attachment = {
                "id": str(source.get("id") or ""),
                "url": url,
                "origin": origin,
                "matched": shown[:120],
                "score": round(score, 2),
            }
            if ties:
                attachment["tied_with"] = ties[:5]
            if url in excerpts and excerpts[url]:
                attachment["excerpt"] = excerpts[url]
            attachments.append(attachment)
            break
    if run_dir is not None and write_report:
        try:
            audit = audit_source_urls(
                package,
                company_id=company_id,
                run_dir=run_dir,
                session_dir=session,
                attachments=attachments,
                refused=refused,
            )
            write_source_url_report(run_dir, audit, attempt=attempt)
            stamp_source_url_status(package, audit)
        except Exception:  # noqa: BLE001 — the report must never sink the run
            logger.warning("source URL report failed", exc_info=True)
    return notes


def stamp_source_url_status(package: dict, audit: dict) -> dict:
    """Record the URL audit on the package envelope as
    ``package["run"]["source_url_status"]`` — ``{"checked_at", "unseen":
    {source id: url}, "homepage_reused": {source id: url}}`` — so the
    renderer can print an unseen source's title unlinked. A consumer should
    apply an entry only while the source still carries that same URL (a
    later repair may have changed it). Optional field: packages without it
    render as before. Returns what was stamped."""
    rows = [row for row in (audit or {}).get("sources") or [] if isinstance(row, dict) and row.get("url")]
    status = {
        "checked_at": (audit or {}).get("checked_at"),
        "unseen": {str(row["id"]): str(row["url"]) for row in rows if not row.get("seen")},
        "homepage_reused": {str(row["id"]): str(row["url"]) for row in rows if row.get("homepage_reused")},
    }
    if isinstance(package, dict):
        run = package.get("run")
        if not isinstance(run, dict):
            run = {}
            package["run"] = run
        run["source_url_status"] = status
    return status


# The signals that count a URL as seen: retrieval fetched it (the company's
# source cache, this run's manifest) or an upstream artifact recorded it
# (the analysis passes, the Memo Studio session). A page that appeared only
# as a link in search results is reported, but it was never read.
_SEEN_SIGNALS = ("source_cache", "run_manifest", "analysis_artifacts", "serena_locators")
MAX_URL_AUDIT_ROWS = 120


def audit_source_urls(
    package: dict,
    *,
    company_id: str,
    run_dir: Path | None,
    session_dir: Path | None = None,
    attachments: list[dict] | None = None,
    refused: list[dict] | None = None,
) -> dict:
    """Seen/unseen status of every package source URL. Report only: the
    package is never edited, because a dropped URL would fail the
    renderer's source-URL rule (``BSH_MEMO_SOURCE_URL_REQUIRED``).

    A URL is ``seen`` when the company's source cache fetched it
    (``source_cache.find_by_url``), this run's ``sources/manifest.jsonl``
    lists it, the analysis artifacts recorded it, or the Memo Studio
    session's locators carry it. ``seen_in`` names every signal that
    matched, plus ``search_results`` when a search or grounding pass only
    listed it. A bare homepage (no path) cited by two or more sources is
    flagged in ``homepage_reused``: one homepage cannot be the source of
    several different facts."""
    sources = package.get("sources") if isinstance(package, dict) else None
    fetched: set[str] = set()
    search_links: set[str] = set()
    try:
        for row in source_cache.list_sources(company_id):
            key = _seen_key(row.get("canonical_url") or row.get("url"))
            if key and row.get("kind") == "web_fetch":
                fetched.add(key)
            for link in row.get("links") or []:
                link_key = _seen_key(link.get("url")) if isinstance(link, dict) else None
                if link_key:
                    search_links.add(link_key)
    except ValueError:
        pass
    manifest: set[str] = set()
    if run_dir is not None:
        for row in source_cache.run_manifest(run_dir):
            key = _seen_key(row.get("url"))
            if key:
                manifest.add(key)
    artifacts = {key for _, url in _artifact_url_candidates(run_dir) if (key := _seen_key(url))}
    artifacts |= _analysis_markdown_urls(run_dir)
    session = session_dir if session_dir is not None else _session_dir_for_run(run_dir, company_id)
    serena = {key for row in _serena_url_candidates(session) if (key := _seen_key(row["url"]))}
    signals = {
        "source_cache": fetched,
        "run_manifest": manifest,
        "analysis_artifacts": artifacts,
        "serena_locators": serena,
        "search_results": search_links,
    }
    rows: list[dict] = []
    homepages: dict[str, list[str]] = {}
    for index, source in enumerate(sources if isinstance(sources, list) else []):
        if not isinstance(source, dict):
            continue
        source_id = str(source.get("id") or f"sources[{index}]")
        url = str(source.get("url") or "").strip()
        if not url:
            rows.append({"id": source_id, "url": None, "seen": None, "seen_in": []})
            continue
        key = _seen_key(url)
        seen_in = [name for name, keys in signals.items() if key and key in keys]
        row = {
            "id": source_id,
            "url": url,
            "seen": any(name in _SEEN_SIGNALS for name in seen_in),
            "seen_in": seen_in,
        }
        if _is_bare_homepage(url):
            row["bare_homepage"] = True
            homepages.setdefault(key or url, []).append(source_id)
        rows.append(row)
    reused = [{"url": next(r["url"] for r in rows if r["id"] == ids[0]), "ids": ids} for ids in homepages.values() if len(ids) >= 2]
    reused_ids = {source_id for item in reused for source_id in item["ids"]}
    for row in rows:
        if row["id"] in reused_ids:
            row["homepage_reused"] = True
    with_url = [row for row in rows if row["url"]]
    return {
        "checked_at": datetime.now(timezone.utc).isoformat(),
        "with_url": len(with_url),
        "seen": sum(1 for row in with_url if row["seen"]),
        "unseen": sum(1 for row in with_url if not row["seen"]),
        "no_url": len(rows) - len(with_url),
        "unseen_ids": [row["id"] for row in with_url if not row["seen"]],
        "homepage_reused": reused,
        "serena_session": Path(session).name if session else None,
        "signal_counts": {name: len(keys) for name, keys in signals.items()},
        "sources": rows[:MAX_URL_AUDIT_ROWS],
        "attachments": list(attachments or []),
        "refused": list(refused or []),
    }


def _source_url_summary_lines(audit: dict) -> list[str]:
    lines = [
        f"Sources with a URL: {audit.get('with_url', 0)} — seen {audit.get('seen', 0)}, "
        f"unseen {audit.get('unseen', 0)}; without a URL: {audit.get('no_url', 0)}.",
    ]
    for item in audit.get("homepage_reused") or []:
        lines.append(
            f"- bare homepage {item.get('url')} is cited by {len(item.get('ids') or [])} sources "
            f"({', '.join(item.get('ids') or [])}): one homepage cannot back several facts"
        )
    for row in audit.get("sources") or []:
        if row.get("url") and not row.get("seen"):
            weak = " (listed by a search, never fetched)" if "search_results" in (row.get("seen_in") or []) else ""
            lines.append(
                f"- {row.get('id')} UNSEEN{weak}: {row.get('url')} — not in the source cache, the run "
                "manifest, the analysis artifacts or the Memo Studio session"
            )
    return lines


def write_source_url_report(run_dir: Path, audit: dict, *, attempt: int | None = None) -> Path | None:
    """Append one URL-status section to ``<run_dir>/logs/source_urls.md``.
    Returns the path, or None when it could not be written."""
    path = Path(run_dir) / "logs" / "source_urls.md"
    header = "## URL check" + (f" — attempt {attempt}" if attempt is not None else "")
    lines = [f"{header} ({audit.get('checked_at', '')[:19]}Z)"]
    lines.extend(_source_url_summary_lines(audit))
    for row in audit.get("sources") or []:
        if row.get("url") and row.get("seen"):
            lines.append(f"- {row.get('id')} seen ({', '.join(row.get('seen_in') or [])}): {row.get('url')}")
    for item in audit.get("attachments") or []:
        line = (
            f"- attached {item.get('id')} ← {item.get('url')} (from {item.get('origin')} "
            f"'{item.get('matched', '')}', score {item.get('score')})"
        )
        if item.get("tied_with"):
            line += f"; tied with {', '.join(item['tied_with'])} — the first was kept"
        if item.get("excerpt"):
            line += f"; excerpt: \"{item['excerpt'][:200]}\""
        lines.append(line)
    for item in audit.get("refused") or []:
        lines.append(
            f"- not attached {item.get('id')} '{item.get('title')}': {item.get('origin')} pages tie "
            f"at score {item.get('score')} ({', '.join(item.get('urls') or [])}) — a wrong link is worse than none"
        )
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("a", encoding="utf-8") as handle:
            handle.write("\n".join(lines) + "\n\n")
    except OSError:
        logger.warning("source URL report write failed: %s", path, exc_info=True)
        return None
    return path


# ---- the reader-facing summary --------------------------------------------------


def summarize_fact_check(fact_check: dict | None) -> dict:
    """What a reader may be told about a memo's numbers, from a stored fact
    check (``logs/fact_check.json`` / the report's ``memo_fact_check``).

    Buckets partition ``checked``: ``verified`` + ``found_elsewhere`` +
    ``derived`` + ``company_reported`` + ``registry_only`` + ``not_traced``
    (``in_context`` is a subset of ``found_elsewhere``). ``unsupported`` and
    ``p0_count`` are the enforced gate's own counts (``unsupported`` is a
    subset of ``not_traced``). ``coverage_pct`` is the share found in
    sources on file (verified, found elsewhere, company-reported or
    derived) and is None whenever ``thin_corpus`` is true — a thin corpus
    is "not checkable", never a bare percentage.

    ``status``: ``not_run`` (no check on record — e.g. a memo that predates
    2026-09-20), ``error``, ``no_figures``, ``not_checkable`` (thin
    corpus), or the gate's own ``pass`` / ``warn`` / ``fail``. ``basis`` is
    ``tiered`` for checks that recorded the honest tiers and ``legacy`` for
    older ones, where ``found_elsewhere`` still includes registry-only and
    company-reported figures and those two read None (unknown)."""
    summary: dict[str, Any] = {
        "checked": 0,
        "verified": 0,
        "found_elsewhere": 0,
        "in_context": 0,
        "derived": 0,
        "company_reported": 0,
        "company_reported_headline": 0,
        "registry_only": 0,
        "not_traced": 0,
        "unsupported": 0,
        "coverage_pct": None,
        "thin_corpus": False,
        "p0_count": 0,
        "status": "not_run",
        "basis": None,
        "urls": None,
    }
    if not isinstance(fact_check, dict) or not fact_check:
        return summary

    def count(key: str) -> int:
        try:
            return max(0, int(fact_check.get(key) or 0))
        except (TypeError, ValueError):
            return 0

    checked = count("checked")
    thin = bool(fact_check.get("thin_corpus"))
    summary.update(
        {
            "checked": checked,
            "unsupported": count("unsupported"),
            "thin_corpus": thin,
            "p0_count": count("p0_count"),
        }
    )
    tiers = fact_check.get("tiers")
    if isinstance(tiers, dict) and tiers.get("basis") == "tiered":
        summary["basis"] = "tiered"
        for key in (
            "verified",
            "found_elsewhere",
            "in_context",
            "derived",
            "company_reported",
            "company_reported_headline",
            "registry_only",
            "not_traced",
        ):
            try:
                summary[key] = max(0, int(tiers.get(key) or 0))
            except (TypeError, ValueError):
                summary[key] = 0
        traced = (
            summary["verified"]
            + summary["found_elsewhere"]
            + summary["company_reported"]
            + summary["derived"]
        )
        coverage = round(traced / checked * 100) if checked else None
    else:
        summary["basis"] = "legacy"
        summary.update(
            {
                "verified": count("verified"),
                "found_elsewhere": count("supported"),
                "in_context": None,
                "derived": count("derived"),
                "company_reported": None,
                "company_reported_headline": None,
                "registry_only": None,
                "not_traced": count("unsupported"),
            }
        )
        coverage = fact_check.get("coverage_pct")
        if not isinstance(coverage, (int, float)) or isinstance(coverage, bool):
            coverage = None
    raw_status = str(fact_check.get("status") or "")
    if fact_check.get("error") or raw_status == "error":
        status = "error"
    elif not checked:
        status = "no_figures"
    elif thin:
        status = "not_checkable"
    else:
        status = raw_status if raw_status in {"pass", "warn", "fail"} else "warn"
    summary["status"] = status
    summary["coverage_pct"] = None if (thin or not checked or status == "error") else coverage
    urls = fact_check.get("source_urls")
    if isinstance(urls, dict):
        summary["urls"] = {
            "with_url": int(urls.get("with_url") or 0),
            "seen": int(urls.get("seen") or 0),
            "unseen": int(urls.get("unseen") or 0),
            "homepage_reused": len(urls.get("homepage_reused") or []),
        }
    return summary


# ---- company-level view (the research desk card) ---------------------------


def check_company(company_id: str) -> dict:
    """The latest memo package for ``company_id`` checked against everything
    on file — the payload behind the research desk's numbers card."""
    from . import comps

    path, package = comps._latest_memo_package(company_id)
    base = {
        "company_id": company_id,
        "memo_package": str(path) if path else None,
        "checked": 0,
        "supported": 0,
        "unsupported": 0,
        "coverage_pct": None,
        "findings": [],
        "sources": [],
    }
    if not package or path is None:
        base["note"] = "No memo package on record"
        return base
    run_dir = path.parent.parent if path.parent.name == "logs" else path.parent
    if _is_buffett_package(package):
        return _check_buffett_company(base, package, run_dir=run_dir, company_id=company_id)
    result = check_memo_run(run_dir=run_dir, package=package, company_id=company_id)
    titles = {
        str(section.get("id")): _loc(section.get("title")) or str(section.get("id"))
        for section in package.get("sections") or []
        if isinstance(section, dict)
    }
    base.update(
        {
            "checked": result.checked,
            "supported": result.checked - result.unsupported,
            "verified": result.verified,
            "derived": result.derived,
            "unsupported": result.unsupported,
            "coverage_pct": result.coverage_pct,
            "status": result.status,
            "thin_corpus": result.thin_corpus,
            "evidence_chars": result.evidence_chars,
            "findings": [
                {
                    "section": titles.get(f.section_id, f.section_id),
                    "number": f.figure,
                    "excerpt": f.snippet,
                    "code": f.code,
                    "looked_for": sorted(numbers_lint._digits_forms(f.figure.rstrip("+")))[:4],
                }
                for f in result.findings
                if f.code == "unsupported_figure"
            ],
            "sources": [row["label"] for row in result.corpus_texts][:60],
            # Report-only honest split (see ``summarize_fact_check``).
            "tiers": dict(result.tiers),
            "summary": summarize_fact_check(result.to_dict()),
            "note": (
                "A figure counts as supported when a source on file carries the same value "
                "in any spelling, or a calculation note derives it."
                + (" The corpus is thin: add a fact ledger, documents or run a memo so retrieval is cached." if result.thin_corpus else "")
            ),
            "error": result.error,
        }
    )
    return base


def _is_buffett_package(package: dict) -> bool:
    """A Buffett-method package carries Markdown, not the late-stage
    section/block schema that ``check_memo_run`` walks."""
    if package.get("kind") == "buffett_investment_memo":
        return True
    return bool(package.get("markdown_en")) and not package.get("sections")


def _check_buffett_company(base: dict, package: dict, *, run_dir: Path, company_id: str) -> dict:
    """The numbers card for a company whose latest memo is Buffett-method:
    the same report-only trace the Buffett run records, recomputed against
    what is on file now. Writes nothing."""
    from . import buffett_checks, buffett_memo_renderer

    base["memo_kind"] = "buffett_investment_memo"
    try:
        validated = buffett_memo_renderer.validate_package(package)
    except Exception as exc:  # noqa: BLE001 — an unreadable package has nothing to trace
        base["note"] = f"The Buffett-method memo on record could not be read: {exc}"
        base["error"] = str(exc)
        return base
    try:
        payload, _report = buffett_checks.fact_check(
            validated, company_id=company_id, run_dir=run_dir
        )
    except Exception as exc:  # noqa: BLE001
        base["note"] = "The number check could not run on this memo."
        base["error"] = str(exc)
        return base
    checked = int(payload.get("checked") or 0)
    unsupported = int(payload.get("unsupported") or 0)
    base.update(
        {
            "checked": checked,
            "supported": checked - unsupported,
            "verified": payload.get("verified"),
            "derived": payload.get("derived"),
            "unsupported": unsupported,
            "coverage_pct": payload.get("coverage_pct"),
            "status": payload.get("status"),
            "thin_corpus": payload.get("thin_corpus"),
            "evidence_chars": payload.get("evidence_chars"),
            "findings": [
                {
                    "section": finding.get("section_id"),
                    "number": finding.get("figure"),
                    "excerpt": finding.get("snippet"),
                    "code": finding.get("code"),
                    "looked_for": sorted(
                        numbers_lint._digits_forms(str(finding.get("figure") or "").rstrip("+"))
                    )[:4],
                }
                for finding in payload.get("findings") or []
                if finding.get("code") == "unsupported_figure"
            ],
            "sources": [row.get("label") for row in payload.get("corpus") or [] if isinstance(row, dict)][:60],
            "tiers": dict(payload.get("tiers") or {}),
            "summary": summarize_fact_check(payload),
            "note": (
                "Buffett-method memo: a figure counts as supported when a source on file carries it; "
                "report only."
                + (" The corpus is thin: add a fact ledger or documents so there is something to check against." if payload.get("thin_corpus") else "")
            ),
            "error": payload.get("error"),
        }
    )
    return base
