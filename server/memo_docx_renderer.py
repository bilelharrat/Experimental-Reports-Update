"""Parameterized DOCX renderer for BSH investment memos.

Claude owns memo judgment and writes ``memo_package.json``. This module owns
all DOCX construction so memo runs do not generate bespoke Python/JS renderers.

The renderer also owns the document's shape: page one is a masthead, the
cover facts and a contents list linked to ``bsh_sec_<n>`` bookmarks on real
Word headings; core sections print their canonical titles numbered
"I. II." / "一、二、", extra sections follow them, and all sources material
lands in one sources section at the back, numbered or not as the structure
profile says (see ``_document_plan``).
"""
from __future__ import annotations

import argparse
import contextvars
import datetime as _dt
import io
import json
import logging
import math
import os
import re
from decimal import Decimal, InvalidOperation
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import itertools

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Twips

from server import memo_structure, source_tiers

logger = logging.getLogger(__name__)

SCHEMA_VERSION = 1
# Stamped into every rendered file (core properties identifier/version),
# the validation reports and the run manifest, so a stored memo says which
# layout produced it and a re-render can tell when it is out of date.
RENDERER_VERSION = "2026-09-22.2"
# The page headers carry the review stamp from ``package["run"]["review"]``
# (DRAFT / Reviewed by … / WITHDRAWN); report_rerender reads this flag.
SUPPORTS_REVIEW_STAMP = True
FIRM_NAME = "Berkeley Summit House"

EN_FONT = "Arial"
ZH_FONT = "Microsoft YaHei"
NAVY = "1B2A4A"
TIFFANY = "0ABAB5"
PALE_TIFFANY = "E6F7F6"
PALE_GOLD = "F7F0D9"
WARM_GREY = "F3F3F3"
WHITE = "FFFFFF"
BLACK = "111111"
GREY = "666666"
BORDER = "BFBFBF"
# Severity tints for the leading token of a risk card's Likelihood /
# Rating ("High:", "8/10") and the review stamp in the page header.
RISK_HIGH = "A33A2B"
RISK_MEDIUM = "9A6A12"
RISK_LOW = "2E6B4F"
STAMP_DRAFT = "B45309"
STAMP_WITHDRAWN = "B91C1C"
STAMP_APPROVED = "2E6B4F"

# Typography: whole half-points (Word rounds anything else), a milder line
# spacing than a report template (these memos are long) and ~2.2 cm margins.
_TEXT_SIZE = 10.5
_TABLE_SIZE = 9
_LINE_SPACING = {"en": 1.15, "zh": 1.3}
_CELL_LINE_SPACING = {"en": 1.05, "zh": 1.15}
_MARGIN_CM = 2.2

GENERATED_RENDERER_SCRIPT_PATTERNS = (
    "build_memo*.py",
    "build_memos*.py",
    "generate_memo*.py",
    "generate_memos*.py",
    "render_memo*.py",
    "render_memos*.py",
    "build_memo*.js",
    "build_memos*.js",
    "generate_memo*.js",
    "generate_memos*.js",
    "render_memo*.js",
    "render_memos*.js",
)

# Structure-derived skeleton (server/memo_structure.py is the single
# source of truth; numbering is positional). The module constants are the
# late v1 defaults; every validation/render path resolves the package's own
# structure from its ``structure`` meta stamp via ``_structure_for``.
SECTION_TITLES = memo_structure.LATE.section_titles()
REQUIRED_SECTION_IDS = memo_structure.LATE.section_ids
REQUIRED_MEMO_COMPONENTS = memo_structure.LATE.components


def _structure_for(package: Any) -> memo_structure.MemoStructure:
    return memo_structure.for_package(package)
SUPPORTED_BLOCK_TYPES = {
    "heading",
    "paragraph",
    "bullets",
    "callout",
    "table",
    "chart",
    "spacer",
    # Optional "Terms used" block (component ``glossary``): items of
    # ``{"term": {en, zh}, "definition": {en, zh}}``. A package-level
    # ``glossary`` list of the same items renders after the executive
    # summary. Neither is required.
    "glossary",
}
GLOSSARY_TITLE = {"en": "Terms used", "zh": "术语说明"}
SUPPORTED_CHART_TYPES = {"bar", "grouped_bar", "hbar", "line", "pie"}
_SINGLE_SERIES_CHART_TYPES = {"bar", "hbar", "pie"}
# Top-level section titles ("I."–"X." or "一、"–"十、") are emitted automatically
# by ``_add_section`` from ``SECTION_TITLES``. A heading *block* that carries the
# same numbered prefix is a redundant restatement of the section title; rendering
# it produces a doubled section heading (and trips the Chinese-parity
# heading-count gate, since EN/ZH restatements are not always recognized
# symmetrically). Legitimate sub-headings are unnumbered, so this prefix is a
# safe signal to drop the block.
_NUMBERED_SECTION_HEADING_RE = re.compile(
    # longest-first so xi..xv match before their x/i prefixes
    r"^\s*(?:(?:xiii|xiv|xii|xv|xi|x|ix|viii|vii|vi|v|iv|iii|ii|i)\."
    r"|[一二三四五六七八九十]+[、.．])",
    re.IGNORECASE,
)
# Per-risk cards in `investment_risk`: a "Risk N: <one-line summary>" heading
# followed by a key_value table with these five row labels. Enforced at
# generation time only (english_package_validation_errors), so packages from
# runs that predate the format still re-render.
_RISK_CARD_HEADING_RE = re.compile(r"^\s*risk\s+(\d+)\s*[:：]\s*\S", re.IGNORECASE)
_RISK_RATING_VALUE_RE = re.compile(r"^\s*(10|[1-9])\s*/\s*10\b")
_RISK_LIKELIHOOD_VALUE_RE = re.compile(
    r"^\s*(high|medium|low|高|中|低)\s*[:：]\s*\S", re.IGNORECASE
)
_RISK_GENERIC_HEADINGS = {
    "commercial risk",
    "competition risk",
    "execution risk",
    "financing risk",
    "liquidity risk",
    "market risk",
    "regulatory risk",
    "technology risk",
    "valuation risk",
}
_RISK_GENERIC_WATCH_RE = re.compile(
    r"^\s*(?:monitor|track|watch|confirm|request|obtain|"
    r"cross[- ]check|verify|require|ensure|validate|ask for)\b",
    re.IGNORECASE,
)
_RISK_CARD_ROW_LABELS = (
    ("risk type", "Risk Type"),
    ("why it matters", "Why it matters"),
    ("what we watch", "What we watch"),
    ("likelihood", "Likelihood"),
    ("risk rating", "Risk Rating"),
)
# Risk Card v2 (structure-v2 family): Verdict and Impact rows open the
# card (founder feedback 2026-09-13 — say which aspect, how big, then
# explain), Mitigation joins it. "Why it matters" / "What we watch" keep
# their exact labels: the checks below index rows by those prefixes.
_RISK_CARD_ROW_LABELS_V2 = (
    ("risk type", "Risk Type"),
    ("verdict", "Verdict"),
    ("impact", "Impact"),
    ("why it matters", "Why it matters"),
    ("what we watch", "What we watch"),
    ("mitigation", "Mitigation"),
    ("likelihood", "Likelihood"),
    ("risk rating", "Risk Rating"),
)


def _risk_card_row_labels(
    structure: "memo_structure.MemoStructure",
) -> tuple[tuple[str, str], ...]:
    if structure.scorecard_weights():
        return _RISK_CARD_ROW_LABELS_V2
    return _RISK_CARD_ROW_LABELS


def _spelled_row_count(count: int) -> str:
    words = {5: "five", 6: "six", 7: "seven", 8: "eight", 9: "nine"}
    return words.get(count, str(count))
# The {section_id} placeholder is filled with the structure's risk-role
# section at check time, so the repair mapper can attribute the finding.
_RISK_CARD_FORMAT_HINT = (
    "section {section_id} must present risks as per-risk cards: "
    "{card_range} "
    "`heading` blocks titled 'Risk N: <one-line summary>', each immediately "
    "followed by a `table` block with component 'risk_register', layout "
    "'key_value', headers [], and exactly {row_count} two-cell rows labeled "
    "{row_list}, with 'Likelihood' written as 'High|Medium|Low: short "
    "reason' and 'Risk Rating' written as 'N/10: short reason', ordered "
    "highest rating first"
)
VALUATION_CONTENT_TERMS = (
    "model",
    "scenario",
    "range",
    "valuation",
    "revenue",
    "arr",
    "margin",
    "gross margin",
    "valuation sensitivity",
    "sensitivity",
    "multiple",
    "forecast",
    "收入",
    "估值",
    "情景",
    "区间",
    "毛利率",
    "利润率",
    "模型",
    "敏感性",
)
GENERIC_CONTENT_PATTERNS = (
    re.compile(r"^more diligence is needed\.?$", re.IGNORECASE),
    re.compile(r"^further diligence is needed\.?$", re.IGNORECASE),
    re.compile(r"^additional diligence is needed\.?$", re.IGNORECASE),
    re.compile(r"^to be determined\.?$", re.IGNORECASE),
    re.compile(r"^tbd\.?$", re.IGNORECASE),
)


class MemoRenderError(ValueError):
    """Raised when a memo package cannot be rendered."""


def find_generated_renderer_scripts(run_dir: Path | str) -> list[Path]:
    """Return forbidden per-run renderer scripts left by old memo workflows."""
    root = Path(run_dir)
    matches: list[Path] = []
    if not root.exists():
        return matches
    for pattern in GENERATED_RENDERER_SCRIPT_PATTERNS:
        matches.extend(path for path in root.rglob(pattern) if path.is_file())
    return sorted(set(matches))


def _read_package(path: Path | str) -> Any:
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise MemoRenderError(f"Invalid memo package JSON: {exc}") from exc


def load_package(path: Path | str, *, strict_sources: bool = True) -> dict:
    package = _read_package(path)
    validate_package(package, strict_sources=strict_sources)
    return package


def validate_package(package: Any, *, strict_sources: bool = True) -> list[str]:
    """Fail closed when Claude's memo package is structurally incomplete.

    ``strict_sources=False`` is for re-rendering a stored package that
    predates the source-URL rule: a missing web-source URL is then logged
    and returned as a warning instead of failing the render. New runs keep
    the default, and every other rule stays fail-closed either way.
    Returns the downgraded warnings (empty when strict)."""
    warnings: list[str] = []
    errors = _package_validation_errors(
        package, strict_sources=strict_sources, warnings=warnings
    )
    if errors:
        raise MemoRenderError("Invalid memo package: " + "; ".join(errors))
    for warning in warnings:
        logger.warning("memo package (legacy source rule): %s", warning)
    return warnings


def fill_blank_zh_placeholders(package: dict) -> dict:
    """Deep-copy a package and fill blank ``zh`` strings with a placeholder,
    so English-only (phase 3) output can pass through render-shaped checks."""
    import copy as _copy

    probe = _copy.deepcopy(package)

    def _fill_zh(node: Any) -> None:
        if isinstance(node, dict):
            if "en" in node and not str(node.get("zh") or "").strip():
                node["zh"] = "占位"
            for value in node.values():
                _fill_zh(value)
        elif isinstance(node, list):
            for value in node:
                _fill_zh(value)

    _fill_zh(probe)
    return probe


_SOURCE_KEY_SYNONYMS = {
    # Analysis-pass evidence vocabulary that keeps leaking into package
    # sources. Renaming is lossless, so repair it instead of failing the run.
    "label": "title",
    "source_class": "class",
    "evidence_class": "class",
    "source_type": "class",
    # A source's type IS its class ("corporate_filing",
    # "third_party_analyst_estimate"); six sources arrived with nothing else.
    "type": "class",
    "category": "class",
    "kind": "class",
    "detail": "treatment",
    "name": "title",
    # When the source was current. Live on 2026-09-20 every one of ten
    # sources dated itself under `date` and the attempt died ten times
    # over on "as_of is required".
    "date": "as_of",
    "as_of_date": "as_of",
}

# A block names its own fields as often as it names the schema's. Each of
# these cost a whole generation attempt on one live Gemini run (Koch, Inc.,
# 2026-09-18: three attempts, a different one of these killing each), so
# they are renamed rather than allowed to fail the memo.
_BLOCK_KEY_SYNONYMS = {
    # The block's type under another name; without it the block reads as an
    # untyped paragraph and fails as "text is required" (22 blocks in one run).
    "kind": "type",
    # The component slug under the analysis passes' word for it.
    "slug": "component",
    # A table's header row.
    "columns": "headers",
    "column_headers": "headers",
    # A chart's reading note under the analysis passes' word for it.
    "reading_note": "reading",
}

# A bullets or callout list under a key that echoes the block's own type.
_ITEMS_KEY_SYNONYMS = ("bullets", "points", "list_items", "entries")

_BLOCK_TYPE_SYNONYMS = {
    "bullet": "bullets",
    "bullet_list": "bullets",
    "bulleted_list": "bullets",
    "list": "bullets",
    "para": "paragraph",
    "text": "paragraph",
    "call_out": "callout",
    "box": "callout",
    # Seen from the per-section wave on Gemini: prose under its own name.
    "prose": "paragraph",
    "narrative": "paragraph",
    "body_text": "paragraph",
    "header": "heading",
    "subheading": "heading",
    "subheader": "heading",
    # Compound types the model invents for a titled paragraph, or a table
    # with its commentary in the same block; the prose half of a table is
    # split out into the paragraph that follows it before the main pass.
    "header_and_prose": "paragraph",
    "heading_and_prose": "paragraph",
    "table_and_prose": "table",
    "table_with_prose": "table",
}


def _repair_titleish(value: Any) -> str:
    """First-sentence-ish English text usable as a derived title."""
    text = _content_text(value)
    if not text:
        return ""
    sentence = re.split(r"(?<=[.!?;:])\s", text, maxsplit=1)[0].strip()
    words = sentence.split()
    if len(words) > 10:
        sentence = " ".join(words[:10]).rstrip(",;:") + "…"
    return sentence.rstrip(".")


def _repair_localized(node: dict, key: str, repairs: list[str], where: str) -> None:
    """Wrap a plain non-language-neutral string as {"en": ..., "zh": ""}."""
    value = node.get(key)
    if (
        isinstance(value, str)
        and value.strip()
        and not _is_language_neutral_text(value)
    ):
        node[key] = {"en": value, "zh": ""}
        repairs.append(f"{where}: wrapped plain string as bilingual en value")


def _decode_serialized_list(
    block: dict, key: str, repairs: list[str], where: str
) -> None:
    """Decode a list the model serialized as a JSON string inside the block
    — the escaping the section prompt warns against, seen live on a
    table's ``headers`` and ``rows``."""
    value = block.get(key)
    if isinstance(value, str) and value.lstrip().startswith("["):
        try:
            decoded = json.loads(value)
        except ValueError:
            return
        if isinstance(decoded, list):
            block[key] = decoded
            repairs.append(f"{where}.{key}: decoded a list serialized as a string")


# Prose a table block carries alongside its rows, and where it belongs.
_TABLE_PROSE_BEFORE = ("intro", "lead", "preface")
_TABLE_PROSE_AFTER = ("content", "commentary", "footnote", "note", "discussion")
_SERIALIZED_LIST_KEYS = ("headers", "columns", "rows", "items", "bullets", "paragraphs")
_CONTENT_KEYS = ("text", "content", "body", "paragraphs", "items", "bullets", "headers", "rows", "title", "level")


def _infer_block_type(block: dict) -> str:
    """The type a block's own fields imply — for a name the renderer does not
    know. The model names its block types freely (``content_block``,
    ``table_block``, ``table_and_prose`` …, a new one on each run); what it
    puts in them is stable."""
    if block.get("headers") or block.get("rows"):
        return "table"
    if isinstance(block.get("items"), list) or isinstance(block.get("bullets"), list):
        return "bullets"
    if block.get("paragraphs") is not None:
        return "paragraph"
    if block.get("level") is not None and not _content_text(
        block.get("content") or block.get("body")
    ):
        return "heading"
    if not any(block.get(key) for key in _CONTENT_KEYS):
        return "spacer"
    return "paragraph"


def _series_from_column_keyed_points(data: Any, where: str) -> list[dict] | None:
    """One series per measure, from points that name their own columns.

    The model writes a scenario chart the way it writes a table row —
    ``{"scenario": {en, zh}, "moic": 1.4, "irr": "7.0%"}`` — where the
    contract wants ``series`` of ``{label, points: [{x, y}]}``. The one
    non-numeric column is the x axis; every column that is a real number
    in EVERY point is a series, labelled by its own name, because here
    the column name is what the numbers are (unlike a flat ``{label,
    value}`` list, where the name carries nothing and the block title
    has to supply it). A column that is a number in some points and a
    string in others is not a series and is left out — ``irr: "7.0%"``
    is text, not a value this can plot.
    """
    if not isinstance(data, list) or not data:
        return None
    points = [point for point in data if isinstance(point, dict)]
    if len(points) != len(data):
        return None
    shared = set(points[0])
    for point in points[1:]:
        shared &= set(point)
    if not shared:
        return None

    def _is_number(value: Any) -> bool:
        return isinstance(value, (int, float)) and not isinstance(value, bool)

    # The x axis is the column the model wrote bilingually — a scenario
    # name, a year, a segment. The other text columns are formatted
    # numbers ("7.0%", "$90M"), which read as labels too, so asking only
    # "is it non-numeric" found three axes where there was one.
    localized = [
        key
        for key in shared
        if all(
            isinstance(point[key], dict)
            and isinstance(point[key].get("en"), str)
            and point[key]["en"].strip()
            for point in points
        )
    ]
    if len(localized) == 1:
        labels = localized
    else:
        labels = [
            key
            for key in shared
            if all(
                not _is_number(point[key]) and _content_text(point[key]).strip()
                for point in points
            )
        ]
    measures = [
        key for key in shared if all(_is_number(point[key]) for point in points)
    ]
    if len(labels) != 1 or not 1 <= len(measures) <= 4:
        return None
    x_key = labels[0]
    return [
        {
            "label": str(measure).replace("_", " ").replace("-", " ").strip(),
            "points": [
                {"x": _content_text(point[x_key]), "y": point[measure]}
                for point in points
            ],
        }
        for measure in sorted(measures, key=lambda key: sorted(shared).index(key))
    ]


def _adopt_column_keyed_rows(block: dict, repairs: list[str], where: str) -> None:
    """Turn rows keyed by column name into rows of cells, in place.

    The column order comes from the headers' own ``key`` fields when they
    carry them, and otherwise from the first row's key order (JSON keeps
    it). Every row must answer the same columns: a table where they
    disagree is not one this can order, and is left for the main pass.
    """
    rows = block.get("rows")
    if not isinstance(rows, list) or not rows:
        return
    if not all(
        isinstance(row, dict) and "cells" not in row and row for row in rows
    ):
        return
    headers = block.get("headers")
    order = [
        str(header["key"])
        for header in (headers if isinstance(headers, list) else [])
        if isinstance(header, dict) and header.get("key")
    ]
    if order:
        # The headers declare the columns, so a row may carry a key they
        # do not name — live, a `deal_terms` row answered both "value"
        # and "detail" with the same sentence, and the table has two
        # columns. The extra key is surplus, not a lost column; dropping
        # the whole table over it would cost ten real rows.
        if any(not set(order) <= set(row) for row in rows):
            return
        extra = sorted({key for row in rows for key in row} - set(order))
    else:
        # Nothing declares the order but the rows themselves, so they all
        # have to agree on it.
        order = [str(key) for key in rows[0]]
        if not order or any(set(row) != set(order) for row in rows):
            return
        extra = []
    block["rows"] = [{"cells": [row[key] for key in order]} for row in rows]
    note = f" (ignored {', '.join(repr(k) for k in extra)})" if extra else ""
    repairs.append(
        f"{where}.rows: ordered {len(rows)} column-keyed rows into cells{note}"
    )


def _expand_block(block: dict, repairs: list[str], where: str) -> list[dict]:
    """Normalize one block before the main pass; return the block(s) it
    stands for.

    Decodes lists serialized as strings, renames the field synonyms, reads
    an unknown type off the block's fields, moves a table's intro and
    footnote prose into the paragraphs around it, expands a titled group
    of paragraphs into paragraphs, and drops an empty table. Each of these
    shapes cost a live generation attempt on a Gemini run.
    """
    for old_key, new_key in _BLOCK_KEY_SYNONYMS.items():
        if old_key in block and not block.get(new_key):
            block[new_key] = block.pop(old_key)
            repairs.append(f"{where}: renamed {old_key!r} to {new_key!r}")
    for key in _SERIALIZED_LIST_KEYS:
        _decode_serialized_list(block, key, repairs, where)
    raw_type = str(block.get("type") or "").strip().lower()
    kind = _BLOCK_TYPE_SYNONYMS.get(raw_type, raw_type)
    if raw_type and kind not in SUPPORTED_BLOCK_TYPES:
        kind = _infer_block_type(block)
        block["type"] = kind
        repairs.append(
            f"{where}.type: read unknown type {raw_type!r} as {kind!r} from its fields"
        )
    elif raw_type and kind != raw_type:
        block["type"] = kind
        repairs.append(f"{where}.type: normalized {raw_type!r} to {kind!r}")
    out: list[dict] = []
    if kind == "chart":
        # A series with no points is the same placeholder an empty table is:
        # the model announced a chart and then had nothing to put in it. The
        # table case has been dropped rather than fatal since the first
        # Gemini runs; this one had not, so on 2026-09-19 a RadixArk memo
        # that was otherwise finished — seven sections, 378 other defects
        # already absorbed — died on one chart nobody could have rendered.
        series = block.get("series")
        if isinstance(series, list):
            kept = [
                one
                for one in series
                if isinstance(one, dict)
                and isinstance(one.get("points"), list)
                and one.get("points")
            ]
            if len(kept) != len(series):
                repairs.append(
                    f"{where}.series: dropped "
                    f"{len(series) - len(kept)} series with no points"
                )
            if not kept:
                repairs.append(f"{where}: dropped empty chart")
                return []
            block["series"] = kept
    # A single-series chart whose points arrived as a flat `data` list —
    # {"label": {en, zh}, "value": n} per point — where the contract wants
    # `series` of {label, points:[{x, y}]}. The points map one to one; only
    # the legend label has to come from somewhere, and the block's own title
    # is where it comes from, as a callout's missing title already comes from
    # its body. Live on 2026-09-20 this cost a valuation chart its block.
    if kind == "chart" and not block.get("series"):
        data = block.get("data")
        points = [
            {"x": _content_text(point.get("label")), "y": point.get("value")}
            for point in data
            if isinstance(point, dict)
        ] if isinstance(data, list) else []
        usable = [
            point
            for point in points
            if str(point["x"]).strip()
            and not isinstance(point["y"], bool)
            and isinstance(point["y"], (int, float))
        ]
        if usable and len(usable) == len(points):
            block["series"] = [
                {
                    "label": _content_text(block.get("title")) or "Series 1",
                    "points": usable,
                }
            ]
            block.pop("data", None)
            repairs.append(
                f"{where}: turned a flat 'data' list of {len(usable)} points "
                "into one series"
            )
        else:
            derived = _series_from_column_keyed_points(data, where)
            if derived:
                block["series"] = derived
                block.pop("data", None)
                names = ", ".join(repr(one["label"]) for one in derived)
                repairs.append(
                    f"{where}: read {len(derived)} series ({names}) off "
                    "points that named their own columns"
                )

    # A chart with nothing left to plot is the placeholder an empty table
    # is, and dropping it is what already happens to a chart whose series
    # carry no points. Reaching validation without one is fatal — "series
    # must be a list of 1-4 series" ended a live attempt on 2026-09-20
    # after every section had already been written.
    if kind == "chart" and not block.get("series"):
        repairs.append(f"{where}: dropped a chart with no series to plot")
        return []

    # A bullets block with no items, carrying its content as `text`: that is
    # a paragraph, and typing it as one is what the block already is. Live on
    # 2026-09-20 six of these failed a run's validation as "items must be a
    # non-empty list" while their prose sat unread in `text`.
    if kind == "bullets" and not block.get("items"):
        if _content_text(block.get("text") or block.get("body")):
            block.pop("items", None)
            block["type"] = "paragraph"
            kind = "paragraph"
            repairs.append(
                f"{where}: retyped an itemless bullets block as the "
                "paragraph it already was"
            )
    if kind == "table":
        # A column-keyed table: `columns` of {key, label} (renamed to
        # `headers` above) and each row a dict keyed by those column keys
        # instead of a list of cells. Every cell is already a localized
        # object and in the right column — only the order has to be
        # recovered, and the headers carry it. Unrepaired, each row reads
        # as an empty row and the whole table is dropped: live on
        # 2026-09-20 this silently emptied the Key Metrics Snapshot of
        # twelve rows on RadixArk and the same table plus Deal Snapshot on
        # Databricks, which is why `company_team` was the short section on
        # every Gemini run.
        _adopt_column_keyed_rows(block, repairs, where)
        # A key-value table delivered under `items` instead of `rows`: every
        # pair is a two-cell row, which is the layout this component renders
        # anyway. Live on 2026-09-20 the deal_terms table arrived this way
        # with ten pairs in it — and on the run before, the same table was
        # dropped as "empty" while its content sat in `items` unread.
        if not block.get("rows") and isinstance(block.get("items"), list):
            pairs = [
                item
                for item in block["items"]
                if isinstance(item, dict) and set(item.keys()) == {"key", "value"}
            ]
            if pairs and len(pairs) == len(block["items"]):
                block["rows"] = [
                    {"cells": [pair["key"], pair["value"]]} for pair in pairs
                ]
                block.pop("items")
                block.setdefault("headers", [])
                repairs.append(
                    f"{where}: turned {len(pairs)} key/value items into "
                    "two-cell rows"
                )
        # A table with no rows is a table with nothing in it, even when it
        # announced its columns: the reader gets a heading and one bare
        # header row. Live on 2026-09-20 the Key Metrics Snapshot shipped
        # that way on RadixArk. Dropping it lets the component gate say
        # what is actually true — the memo is missing that component —
        # instead of passing on the declared slug of an empty block.
        if not block.get("rows"):
            repairs.append(f"{where}: dropped empty table")
            return []
        for key in _TABLE_PROSE_BEFORE:
            prose = block.pop(key, None)
            if isinstance(prose, (str, dict)) and _content_text(prose):
                out.append({"type": "paragraph", "text": prose})
                repairs.append(f"{where}: moved its {key!r} into a paragraph before the table")
            elif prose is not None:
                block[key] = prose
        out.append(block)
        for key in _TABLE_PROSE_AFTER:
            prose = block.pop(key, None)
            if isinstance(prose, (str, dict)) and _content_text(prose):
                out.append({"type": "paragraph", "text": prose})
                repairs.append(f"{where}: moved its {key!r} into a paragraph after the table")
            elif prose is not None:
                block[key] = prose
        return out
    paragraphs = block.get("paragraphs")
    if kind == "paragraph" and isinstance(paragraphs, list) and paragraphs:
        del block["paragraphs"]
        carried = {k: v for k, v in block.items() if k not in ("text", "content", "body")}
        for item in paragraphs:
            if isinstance(item, (str, dict)) and _content_text(item):
                para: dict = {"type": "paragraph", "text": item}
                if not out:
                    para = {**carried, **para}
                out.append(para)
        if out:
            repairs.append(f"{where}: expanded {len(out)} paragraph(s) it carried as a list")
            return out
    return [block]


def _adopt_items_synonym(block: dict, repairs: list[str], where: str) -> None:
    """Move a bullet list filed under the block's own name into ``items``.

    Live: a bullets block arrived as ``{"type": "bullets", "bullets": [...],
    "items": null}`` and failed as "items must be a non-empty list", which
    ended the run's last attempt. A populated ``items`` always wins.
    """
    current = block.get("items")
    if isinstance(current, list) and current:
        return
    for key in _ITEMS_KEY_SYNONYMS:
        candidate = block.get(key)
        if isinstance(candidate, list) and candidate:
            block["items"] = candidate
            block.pop(key, None)
            repairs.append(f"{where}: moved {key!r} into 'items'")
            return


_SOURCE_CLASS_SLUG_RE = re.compile(r"^[a-z][a-z0-9]*(?:[_-][a-z0-9]+)*$")


def _repair_source_class_and_treatment(
    source: dict, repairs: list[str], where: str
) -> None:
    """Fill a source's ``class`` and ``treatment`` from whichever one arrived.

    A source classes itself in exactly one of these two fields and leaves the
    other empty, and both halves are required. One live run (Koch, Inc.,
    2026-09-18) failed twice on this, once each way: every source came back
    with ``class: "company_reported"`` and no treatment, then on the next
    attempt with ``treatment: {"en": "company_reported"}`` and no class.

    A bare slug is a class, never a treatment sentence, so it is promoted and
    the treatment restated as the sentence that field is for. A treatment
    already written as prose is left exactly as the author wrote it.
    """
    class_text = _content_text(source.get("class")).strip()
    treatment = source.get("treatment")
    treatment_text = _content_text(treatment).strip()

    def _label(slug: str) -> str:
        return slug.replace("_", " ").replace("-", " ").strip()

    if not class_text and treatment_text and _SOURCE_CLASS_SLUG_RE.match(
        treatment_text
    ):
        label = _label(treatment_text)
        source["class"] = label
        sentence = f"Weighted as {label}."
        if isinstance(treatment, dict):
            treatment["en"] = sentence
        else:
            source["treatment"] = {"en": sentence, "zh": ""}
        repairs.append(
            f"{where}: read class {label!r} from a treatment holding only "
            "that slug, and restated the treatment"
        )
        return

    if class_text and not treatment_text:
        label = _label(class_text)
        if _SOURCE_CLASS_SLUG_RE.match(class_text):
            source["class"] = label
        source["treatment"] = {"en": f"Weighted as {label}.", "zh": ""}
        repairs.append(
            f"{where}: wrote the treatment its class {label!r} implies"
        )


# A localized value the model nested one level deeper than the contract, under
# a name for the column plus a key for the machine:
#     {"key": "metric", "label": {"en": "Metric", "zh": "指标"}}
# where the renderer wants the localized object itself. Both halves are there
# and already translated, so lifting the payload out loses nothing; the `key`
# is the model's own bookkeeping and no reader ever sees it. Live on
# 2026-09-19 and 2026-09-20 this one shape was 31 of 54 validation errors on
# a package and blocked three runs between them.
_LOCALIZED_PAYLOAD_KEYS = ("label", "header", "title", "text")


def _unwrapped_localized(item: Any) -> tuple[dict, str] | None:
    """The localized object nested inside ``item``, and the key it sat under."""
    if not isinstance(item, dict) or "en" in item or "zh" in item:
        return None
    found = [key for key in _LOCALIZED_PAYLOAD_KEYS if key in item]
    if len(found) != 1:
        return None
    key = found[0]
    payload = item[key]
    if isinstance(payload, dict) and ("en" in payload or "zh" in payload):
        return payload, key
    if (
        isinstance(payload, str)
        and payload.strip()
        and not _is_language_neutral_text(payload)
    ):
        return {"en": payload, "zh": ""}, key
    return None


def _joined_localized(item: Any, keys: tuple[str, str]) -> dict | None:
    """Two localized halves of one value, joined into the one the contract
    wants. ``None`` when the item is not that shape."""
    if not isinstance(item, dict) or "en" in item or "zh" in item:
        return None
    if set(item.keys()) != set(keys):
        return None
    first, second = (item.get(keys[0]), item.get(keys[1]))
    if not isinstance(first, dict) or not isinstance(second, dict):
        return None
    joined = {}
    for half in ("en", "zh"):
        parts = [
            str(first.get(half) or "").strip(),
            str(second.get(half) or "").strip(),
        ]
        joined[half] = " ".join(part for part in parts if part)
    return joined if joined.get("en") else None


def _text_of_block_list(item: Any) -> dict | None:
    """A bullet delivered as a list of paragraph blocks, joined into the one
    localized string a bullet is. ``None`` unless every entry is a block
    carrying a localized ``text``.

    Live on 2026-09-21 (RadixArk 2026-09-21__195426, attempt 2) twelve
    bullets across three thesis_market lists arrived as
    ``[[{"type": "paragraph", "text": {en, zh}}], ...]`` — every word
    present in both languages, one level too deep — and failed the attempt
    as "must be bilingual with en and zh"."""
    if not isinstance(item, list) or not item:
        return None
    texts: list[dict] = []
    for entry in item:
        text = entry.get("text") if isinstance(entry, dict) else None
        if not isinstance(text, dict) or not isinstance(text.get("en"), str):
            return None
        texts.append(text)
    joined = {
        half: " ".join(
            str(text.get(half) or "").strip()
            for text in texts
            if str(text.get(half) or "").strip()
        )
        for half in ("en", "zh")
    }
    return joined if joined["en"] else None


def _repair_localized_list(items: Any, repairs: list[str], where: str) -> list:
    if not isinstance(items, list):
        return items
    out = []
    for index, item in enumerate(items):
        if (
            isinstance(item, str)
            and item.strip()
            and not _is_language_neutral_text(item)
        ):
            out.append({"en": item, "zh": ""})
            repairs.append(f"{where}[{index}]: wrapped plain string as bilingual en value")
            continue
        flattened = _text_of_block_list(item)
        if flattened is not None:
            out.append(flattened)
            repairs.append(
                f"{where}[{index}]: joined the {len(item)} paragraph block(s) "
                "it carried into one bullet"
            )
            continue
        nested = _unwrapped_localized(item)
        if nested is not None:
            payload, key = nested
            out.append(payload)
            repairs.append(
                f"{where}[{index}]: lifted the localized value out of "
                f"{key!r}"
            )
            continue
        # A bullet delivered as a lead line plus a body. The contract wants
        # one localized string, and the renderer already bolds a bullet's
        # lead by splitting at its first ". " — so joining the halves is
        # both lossless and exactly the look the split was reaching for.
        joined = _joined_localized(item, ("title", "text"))
        if joined is not None:
            out.append(joined)
            repairs.append(
                f"{where}[{index}]: joined the 'title' and 'text' halves "
                "into one bullet"
            )
            continue
        out.append(item)
    return out


def repair_package_structure(package: Any) -> tuple[Any, list[str]]:
    """Deterministically fix mechanical package defects before validation.

    Only lossless, unambiguous repairs are applied — a missing callout title
    is derived from the callout's own body, plain strings in bilingual slots
    are wrapped as ``{"en": ..., "zh": ""}``, and analysis-pass source
    vocabulary (``label``/``source_class``/``detail``) is renamed to the
    renderer contract. Anything judgment-shaped is left for the validation
    feedback loop. Returns ``(repaired_copy, repair_descriptions)``; the
    input is never mutated and an empty repair list means the package was
    returned unchanged.
    """
    import copy as _copy

    if not isinstance(package, dict):
        return package, []
    repaired = _copy.deepcopy(package)
    repairs: list[str] = []

    if not str(repaired.get("schema_version") or "").strip():
        repaired["schema_version"] = SCHEMA_VERSION
        repairs.append(f"schema_version: defaulted to {SCHEMA_VERSION}")

    sections = repaired.get("sections")
    for s_index, section in enumerate(sections if isinstance(sections, list) else []):
        if not isinstance(section, dict):
            continue
        where_section = f"sections[{s_index}]"
        _repair_localized(section, "title", repairs, f"{where_section}.title")
        blocks = section.get("blocks")
        # An empty table — no headers, no rows — is a placeholder the model
        # never filled. Validation rejects the whole package over it and a
        # regeneration attempt follows; dropping the block is what a human
        # editor would do. Seen live from the per-section wave on Gemini.
        if isinstance(blocks, list):
            kept: list = []
            for b_index, block in enumerate(blocks):
                if not isinstance(block, dict):
                    kept.append(block)
                    continue
                kept.extend(
                    _expand_block(block, repairs, f"{where_section}.blocks[{b_index}]")
                )
            blocks[:] = kept
        for b_index, block in enumerate(blocks if isinstance(blocks, list) else []):
            if not isinstance(block, dict):
                continue
            where = f"{where_section}.blocks[{b_index}]"
            for old_key, new_key in _BLOCK_KEY_SYNONYMS.items():
                if old_key in block and not block.get(new_key):
                    block[new_key] = block.pop(old_key)
                    repairs.append(
                        f"{where}: renamed {old_key!r} to {new_key!r}"
                    )
            # Prose under `content`, with the block's `text` carrying what
            # is really its title (nine blocks of one live valuation
            # section). The prose is the text; the short line above it is
            # the title. A list under `content` is never prose.
            content = block.get("content")
            if isinstance(content, (str, dict)) and _content_text(content):
                if not _content_text(block.get("text")):
                    block["text"] = block.pop("content")
                    repairs.append(f"{where}: moved 'content' into 'text'")
                elif not _content_text(block.get("title")):
                    block["title"] = block.pop("text")
                    block["text"] = block.pop("content")
                    repairs.append(
                        f"{where}: moved 'content' into 'text' and the "
                        "line it carried as text into 'title'"
                    )
            raw_kind = str(block.get("type") or "paragraph").strip().lower()
            kind = _BLOCK_TYPE_SYNONYMS.get(raw_kind, raw_kind)
            if not str(block.get("type") or "").strip():
                # The validator reads an untyped block as a paragraph; say so
                # in the package rather than leave every reader to default it.
                block["type"] = kind
                repairs.append(f"{where}.type: untyped block read as {kind!r}")
            elif kind != str(block.get("type")):
                block["type"] = kind
                repairs.append(f"{where}.type: normalized {raw_kind!r} to {kind!r}")
            if kind == "paragraph" and not _content_text(
                block.get("text") or block.get("body")
            ):
                # Some generations put paragraph prose under `title`.
                if _content_text(block.get("title")):
                    block["text"] = block.pop("title")
                    repairs.append(f"{where}: moved paragraph title into text")
            if kind == "heading":
                _repair_localized(block, "text", repairs, f"{where}.text")
                _repair_localized(block, "title", repairs, f"{where}.title")
            elif kind == "paragraph":
                _repair_localized(block, "text", repairs, f"{where}.text")
                _repair_localized(block, "body", repairs, f"{where}.body")
                _repair_localized(block, "title", repairs, f"{where}.title")
            elif kind == "bullets":
                _adopt_items_synonym(block, repairs, where)
                block["items"] = _repair_localized_list(
                    block.get("items"), repairs, f"{where}.items"
                )
            elif kind == "callout":
                _repair_localized(block, "title", repairs, f"{where}.title")
                _repair_localized(block, "label", repairs, f"{where}.label")
                _repair_localized(block, "body", repairs, f"{where}.body")
                # The validator reads a callout's prose as `body or text`, so
                # prose left as a plain string under `text` fails as ".body
                # must be bilingual". That single unrepaired shape was what
                # sent a 5,000-word first draft into the regeneration retry.
                _repair_localized(block, "text", repairs, f"{where}.text")
                _adopt_items_synonym(block, repairs, where)
                block["items"] = _repair_localized_list(
                    block.get("items"), repairs, f"{where}.items"
                )
                if not _content_text(block.get("title") or block.get("label")):
                    derived = _repair_titleish(
                        block.get("body") or block.get("text")
                    ) or _repair_titleish(
                        next(iter(block.get("items") or []), None)
                    ) or "Key Consideration"
                    block["title"] = {"en": derived, "zh": ""}
                    repairs.append(
                        f"{where}.title: derived callout title from its content"
                    )
            elif kind == "chart":
                _repair_localized(block, "title", repairs, f"{where}.title")
                _repair_localized(block, "caption", repairs, f"{where}.caption")
                _repair_localized(block, "reading", repairs, f"{where}.reading")
            elif kind == "table":
                _repair_localized(block, "title", repairs, f"{where}.title")
                block["headers"] = _repair_localized_list(
                    block.get("headers"), repairs, f"{where}.headers"
                )
                rows = block.get("rows")
                # An empty row is the same placeholder as an empty table,
                # one level down ("rows[0] must contain cells" cost a live
                # attempt). Drop it rather than fail the package.
                if isinstance(rows, list):
                    kept_rows = [
                        row for row in rows
                        if (row.get("cells") if isinstance(row, dict) else row)
                    ]
                    if len(kept_rows) != len(rows):
                        repairs.append(
                            f"{where}.rows: dropped {len(rows) - len(kept_rows)} empty row(s)"
                        )
                        rows[:] = kept_rows
                for r_index, row in enumerate(rows if isinstance(rows, list) else []):
                    cells = row.get("cells") if isinstance(row, dict) else row
                    # A cell written as {"text": <localized>} carries its
                    # value one level down; the validator wants the localized
                    # object itself and reports "cells[0].en is required".
                    if isinstance(cells, list):
                        for c_index, cell in enumerate(cells):
                            if (
                                isinstance(cell, dict)
                                and "en" not in cell
                                and "zh" not in cell
                                and isinstance(cell.get("text"), (dict, str))
                                and len(cell) == 1
                            ):
                                cells[c_index] = cell["text"]
                                repairs.append(
                                    f"{where}.rows[{r_index}].cells[{c_index}]: "
                                    "unwrapped cell text"
                                )
                    fixed = _repair_localized_list(
                        cells, repairs, f"{where}.rows[{r_index}].cells"
                    )
                    if isinstance(row, dict):
                        row["cells"] = fixed
                    elif isinstance(rows, list):
                        rows[r_index] = fixed

    sources = repaired.get("sources")
    for index, source in enumerate(sources if isinstance(sources, list) else []):
        if not isinstance(source, dict):
            continue
        where = f"sources[{index}]"
        for old_key, new_key in _SOURCE_KEY_SYNONYMS.items():
            if old_key in source and not str(source.get(new_key) or "").strip():
                source[new_key] = source.pop(old_key)
                repairs.append(f"{where}: renamed {old_key!r} to {new_key!r}")
        _repair_source_class_and_treatment(source, repairs, where)
        if not str(source.get("id") or "").strip():
            source["id"] = f"S{index + 1}"
            repairs.append(f"{where}.id: assigned S{index + 1}")
        # `title` may legally stay a plain string; only `treatment` is
        # validated as a bilingual value.
        _repair_localized(source, "treatment", repairs, f"{where}.treatment")

    return repaired, repairs


def english_package_validation_errors(
    package: Any, *, editorial_risk_checks: bool = True
) -> list[str]:
    """Validate an English-only package (phase 3 output, ``zh`` still blank).

    Runs the full renderer validation against a copy whose empty ``zh``
    strings are placeholder-filled, so structural defects (wrong source
    vocabulary, missing sections, shallow content) surface at generation
    time — where the synthesis pass can retry with the errors fed back —
    instead of 20 minutes later at render time.

    ``editorial_risk_checks=False`` leaves out the risk-card checks that
    judge wording rather than shape — the economic-consequence vocabulary
    of "Why it matters", a "What we watch" that reads as a command or says
    too little, and the exact row count of a card — so a caller can treat
    them as quality findings (:func:`risk_card_quality_findings`) that go
    through the surgical repair and end as warnings, instead of forcing a
    regeneration. The default keeps every check blocking, as before.
    """
    if not isinstance(package, dict):
        return ["memo package must be a JSON object"]
    filled = fill_blank_zh_placeholders(package)
    errors = _package_validation_errors(filled)
    errors.extend(
        _risk_card_format_errors(
            filled, part="all" if editorial_risk_checks else "structural"
        )
    )
    errors.extend(_subsection_heading_errors(filled))
    errors.extend(_exec_summary_format_errors(filled))
    errors.extend(_chart_reading_errors(filled))
    errors.extend(_word_budget_errors(filled))
    return errors


def _section_en_word_count(section: dict) -> int:
    """All English words in a section's blocks — prose, bullets, table
    cells, chart captions alike. What the reader has to get through."""
    words = 0

    def walk(value: Any) -> None:
        nonlocal words
        if isinstance(value, dict):
            if "en" in value:
                words += len(str(value.get("en") or "").split())
            else:
                for item in value.values():
                    walk(item)
        elif isinstance(value, (list, tuple)):
            for item in value:
                walk(item)

    walk(section.get("blocks") or [])
    return words


def section_en_word_count(section: dict) -> int:
    """Public name for the count the compact-budget gate enforces, so other
    modules measure a section the way the renderer does rather than growing
    a second counter that drifts from this one."""
    return _section_en_word_count(section)


# How far over its target a section may run before the gate fires, when the
# profile does not set a per-section multiple of its own.
_BUDGET_GRACE = 1.10


def _word_budget_errors(package: dict) -> list[str]:
    """Generation-time gate for profiles that declare word budgets
    (compact only). Prose budgets alone failed twice live (5.2K and
    5.5K words against a 2.6-3.2K target); the gate is deterministic so
    the retry loop can enforce it.

    `budget_words` is the SOFT target the writer aims at; the gate fires
    only at `budget_words * budget_hard_multiple`. The two are separate
    because a section that stops mid-argument to respect a word count is
    worse than one that runs long: the point of the budget is to keep
    commentary out, not to truncate an answer. The multiples are per
    section (owner-set 2026-09-16) because a complete answer costs
    different amounts in different sections — a risk register that has
    genuinely found eight risks cannot be as short as a valuation
    summary, hence risks at 2x against valuation_returns at 1.3x.
    """
    sections = package.get("sections")
    if not isinstance(sections, list):
        return []
    structure = _structure_for(package)
    errors: list[str] = []
    by_id = {
        str(s.get("id") or ""): s for s in sections if isinstance(s, dict)
    }
    for sdef in structure.sections:
        if not sdef.budget_words:
            continue
        section = by_id.get(sdef.id)
        if section is None:
            continue
        count = _section_en_word_count(section)
        # Re-emitting a section has never reliably shortened one:
        # company_team went 1358 -> 1187 -> 1018 against a 750-word target
        # over three rounds on 2026-09-16. So the gate fires only at the
        # hard cap, where the section is long enough to be worth the trim
        # pass, and everything between the target and the cap is left alone.
        hard_cap = sdef.budget_words * (
            sdef.budget_hard_multiple or _BUDGET_GRACE
        )
        if count > hard_cap:
            errors.append(
                f"section {sdef.id} runs {count} English words against its "
                f"{sdef.budget_words}-word target and its "
                f"{int(hard_cap)}-word hard cap — this is the COMPACT memo: "
                "cut commentary (one bullet per point, one clause per "
                "judgment) until it is under the cap; never cut pinned "
                "facts, subsection headings, or the scorecard sentences"
            )
    return errors


def _chart_reading_errors(package: dict) -> list[str]:
    """Generation-time gate: every structure-v2 chart carries a `reading`
    note ("Higher is better", "Bars below 1.0x lose money") so the
    reader knows how to read it. Render-time validation stays lenient —
    packages from runs that predate the field still re-render."""
    structure = _structure_for(package)
    if not structure.scorecard_weights():
        return []
    errors: list[str] = []
    sections = package.get("sections")
    for section in sections if isinstance(sections, list) else []:
        if not isinstance(section, dict):
            continue
        section_id = str(section.get("id") or "")
        for index, block in enumerate(section.get("blocks") or []):
            if not isinstance(block, dict) or block.get("type") != "chart":
                continue
            if not _content_text(block.get("reading")):
                errors.append(
                    f"section {section_id} blocks[{index}]: chart must carry "
                    "a `reading` note (bilingual) telling the reader how to "
                    "read it — 'Higher is better', 'Bars below 1.0x lose "
                    "money'"
                )
    return errors


_SUBSECTION_NUMBER_RE = re.compile(r"^\s*(\d+)\s*[\.、]\s*")


def _subsection_heading_errors(package: dict) -> list[str]:
    """Generation-time gate for the fixed numbered subsections a
    structure-v2 profile declares per section.

    Each declared subsection must appear, in order, as a heading block
    whose text starts with its position number ("2. Market size").
    Extra headings between declared ones (risk cards, sub-headings the
    contract asks for) are allowed. Kept out of render-time validation
    so packages from runs that predate subsections still re-render.
    """
    sections = package.get("sections")
    if not isinstance(sections, list):
        return []
    structure = _structure_for(package)
    errors: list[str] = []
    by_id = {
        str(s.get("id") or ""): s for s in sections if isinstance(s, dict)
    }
    for sdef in structure.sections:
        if not sdef.subsections:
            continue
        section = by_id.get(sdef.id)
        if section is None:
            continue  # the missing-section error is raised elsewhere
        headings = [
            _content_text(b.get("text") or b.get("title"))
            for b in section.get("blocks") or []
            if isinstance(b, dict)
            and str(b.get("type") or "paragraph") == "heading"
        ]
        cursor = 0
        for number, sub in enumerate(sdef.subsections, start=1):
            expected = f"{number}. {sub.en}"
            found = False
            while cursor < len(headings):
                heading = headings[cursor]
                cursor += 1
                match = _SUBSECTION_NUMBER_RE.match(heading)
                if (
                    match
                    and int(match.group(1)) == number
                    and sub.en.lower()
                    in _SUBSECTION_NUMBER_RE.sub("", heading).strip().lower()
                ):
                    found = True
                    break
            if not found:
                errors.append(
                    f"section {sdef.id} must contain its fixed numbered "
                    f"subsection heading {expected!r} (a `heading` block, "
                    "level 2, in the declared order)"
                )
    return errors


def _exec_summary_format_errors(package: dict) -> list[str]:
    """Generation-time gate: a structure-v2 executive summary carries no
    tables — its job is the point, not the data. The snapshot tables
    live in the overview section the profile routes them to."""
    structure = _structure_for(package)
    if not structure.scorecard_weights():
        return []
    try:
        exec_id = structure.section_for_role("exec").id
    except KeyError:
        return []
    sections = package.get("sections")
    if not isinstance(sections, list):
        return []
    section = next(
        (
            s
            for s in sections
            if isinstance(s, dict) and str(s.get("id") or "") == exec_id
        ),
        None,
    )
    if section is None:
        return []
    routing = structure.component_section()
    homes = ", ".join(
        f"{slug} belongs in {routing[slug]}"
        for slug in ("deal_terms", "key_metrics_snapshot")
        if routing.get(slug)
    )
    errors: list[str] = []
    for index, block in enumerate(section.get("blocks") or []):
        if not isinstance(block, dict):
            continue
        if str(block.get("type") or "") == "table":
            errors.append(
                f"section {exec_id} blocks[{index}]: the executive summary "
                "must contain NO table blocks — state its numbers in "
                "interpreted prose"
                + (f" ({homes})" if homes else "")
            )
        if str(block.get("type") or "") == "bullets":
            for item_index, item in enumerate(block.get("items") or []):
                lead = _bullet_lead_sentence(_content_text(item))
                if (
                    lead
                    and len(lead.split()) < 4
                    # Digits or an inner period mean a figure or an
                    # abbreviation ("U.S."), not a bare topic label.
                    and "." not in lead
                    and not any(ch.isdigit() for ch in lead)
                ):
                    errors.append(
                        f"section {exec_id} blocks[{index}]"
                        f".items[{item_index}]: bullet opens with the topic "
                        f"label {lead!r} — open with a short claim that "
                        "carries the direction ('The price sits below every "
                        "disclosed peer — 13.8x vs a 21x median.'), never a "
                        "naked label"
                    )
    return errors


_BULLET_LEAD_SPLIT_RE = re.compile(r"(?<=[.!?。])\s+|(?<=[。])")


def _bullet_lead_sentence(text: str) -> str:
    """The words a bullet leads with — up to the first sentence break.
    A bare topic label ("Price. 13.8x vs 21x.") shows up here as a one-
    to three-word fragment; a colon lead ("Market size: the pool is
    $125B...") keeps its judgment in the same sentence and passes."""
    stripped = text.strip().lstrip("*_•- ").strip()
    if not stripped:
        return ""
    parts = _BULLET_LEAD_SPLIT_RE.split(stripped, maxsplit=1)
    return parts[0].strip().rstrip(".!?。*_").strip()


def risk_card_quality_findings(package: Any) -> list[str]:
    """The editorial half of the risk-card checks, as quality findings: a
    "Why it matters" row with no economic consequence (or too thin to name
    one), a "What we watch" row that is a diligence command or too thin to
    name a signal, and a card without exactly its declared rows. The
    renderer draws all of these; they are wording to repair, not a package
    to regenerate. Empty for packages whose profile has no risk cards."""
    if not isinstance(package, dict):
        return []
    return _risk_card_format_errors(
        fill_blank_zh_placeholders(package), part="editorial"
    )


def _risk_card_format_errors(package: dict, *, part: str = "all") -> list[str]:
    """Generation-time gate for the per-risk card format in investment_risk.

    Kept out of ``_package_validation_errors`` on purpose: render-time
    validation must keep accepting packages from runs that predate the card
    format, while the synthesis retry loop gets precise errors to fix.

    ``part``: "all" (default), "structural" (the card shape: heading then a
    key_value table, card count, empty headers, row labels in order,
    Likelihood / Rating forms, rating order, specific and distinct
    headings) or "editorial" (see :func:`risk_card_quality_findings`).
    """
    structural_errors: list[str] = []
    editorial_errors: list[str] = []

    def result() -> list[str]:
        if part == "structural":
            return structural_errors
        if part == "editorial":
            return editorial_errors
        return structural_errors + editorial_errors
    sections = package.get("sections")
    if not isinstance(sections, list):
        return []
    structure = _structure_for(package)
    if structure.risk_format != "cards":
        # Compact profiles present the pinned risks as verdict-lead
        # bullets; the pin-echo gate still enforces the risk list.
        return []
    try:
        risk_id = structure.section_for_role("risk").id
    except KeyError:
        return []
    row_labels = _risk_card_row_labels(structure)
    row_index = {prefix: i for i, (prefix, _label) in enumerate(row_labels)}
    row_list_text = " / ".join(label for _p, label in row_labels)
    section = next(
        (
            s
            for s in sections
            if isinstance(s, dict) and str(s.get("id") or "") == risk_id
        ),
        None,
    )
    if section is None:
        # The missing-section error is already raised by shared validation.
        return []
    errors = structural_errors
    blocks = [b for b in section.get("blocks") or [] if isinstance(b, dict)]
    cards: list[tuple[str, dict, str]] = []
    for index, block in enumerate(blocks):
        if str(block.get("type") or "paragraph") != "heading":
            continue
        heading_text = _content_text(block.get("text") or block.get("title"))
        if not _RISK_CARD_HEADING_RE.match(heading_text):
            continue
        location = f"{risk_id} blocks[{index}]"
        nxt = blocks[index + 1] if index + 1 < len(blocks) else None
        if not isinstance(nxt, dict) or str(nxt.get("type") or "") != "table":
            errors.append(
                f"{location}: risk heading {heading_text!r} must be "
                "immediately followed by its key_value risk card table"
            )
            continue
        cards.append((heading_text, nxt, f"{risk_id} blocks[{index + 1}]"))
    row_count_word = _spelled_row_count(len(row_labels))
    low, high = memo_structure.risk_count_bounds(structure)
    if not low <= len(cards) <= high:
        errors.append(
            _RISK_CARD_FORMAT_HINT.format(
                section_id=risk_id,
                card_range=f"{low}-{high}",
                row_count=row_count_word,
                row_list=", ".join(
                    f"'{label}'" for _p, label in row_labels
                ),
            )
        )
        errors.append(
            f"{risk_id}: risk register must contain {low}-{high} material "
            f"risk cards, found {len(cards)}"
        )
        return result()

    ratings: list[tuple[str, int]] = []
    seen_headings: set[str] = set()
    for heading_text, table, location in cards:
        summary = heading_text.split(":", 1)[-1].strip().lower()
        if summary in _RISK_GENERIC_HEADINGS:
            errors.append(
                f"{location}: risk heading {heading_text!r} is only a "
                "generic category; name the concrete failure mode"
            )
        normalized_summary = re.sub(r"[^a-z0-9]+", " ", summary).strip()
        if normalized_summary in seen_headings:
            errors.append(
                f"{location}: risk heading {heading_text!r} duplicates "
                "another risk card"
            )
        seen_headings.add(normalized_summary)
        if str(table.get("layout") or "").strip().lower() != "key_value":
            errors.append(
                f"{location}: risk card table must declare \"layout\": \"key_value\""
            )
        if [header for header in table.get("headers") or [] if _content_text(header)]:
            errors.append(f"{location}: risk card table must have empty headers")
        row_texts: list[tuple[str, str] | None] = []
        for row in table.get("rows") or []:
            cells = row.get("cells") if isinstance(row, dict) else row
            if isinstance(cells, (list, tuple)) and len(cells) == 2:
                row_texts.append((_content_text(cells[0]), _content_text(cells[1])))
            else:
                row_texts.append(None)
        if len(row_texts) != len(row_labels) or any(
            row is None for row in row_texts
        ):
            # Editorial: the card still draws; its rows need rewriting.
            editorial_errors.append(
                f"{location}: risk card table needs exactly "
                f"{row_count_word} two-cell rows "
                f"({row_list_text})"
            )
            continue
        for (prefix, label), row in zip(row_labels, row_texts):
            label_text, value_text = row  # type: ignore[misc]
            if not label_text.lower().startswith(prefix):
                errors.append(
                    f"{location}: row label {label_text!r} must be {label!r}"
                )
            elif not value_text.strip():
                errors.append(f"{location}: row {label!r} must not be empty")
        why_row = row_texts[row_index["why it matters"]]
        if why_row and why_row[0].lower().startswith("why it matters"):
            why_text = why_row[1].strip()
            if not _has_meaningful_text(why_text, min_chars=35, min_words=6):
                editorial_errors.append(
                    f"{location}: Why it matters must name the fact, failure "
                    "mode, and economic consequence"
                )
            # "pric" covers price/pricing; discount, fee, billing, profit,
            # burn and churn are the vocabulary of services and SaaS risks —
            # "15-25% productivity discounts on time-and-materials
            # contracts" is an economic consequence, and a live run failed
            # three retries and a repair pass on exactly that sentence.
            if not re.search(
                r"\b(?:revenue|margin|cash|valuation|pric|return|dilut|"
                r"capital|exit|control|loss|multiple|growth|financ|cost|"
                r"conversion|discount|fee|billing|profit|burn|churn)\w*\b",
                why_text,
                re.IGNORECASE,
            ):
                # The message is fed back to the retry and to the repair
                # pass, so it has to say what would satisfy it: a model that
                # believes it stated a consequence cannot act on a bare
                # "must state one".
                editorial_errors.append(
                    f"{location}: Why it matters must state an economic "
                    "consequence — name the effect with a financial term "
                    "(revenue, margin, cost, cash, pricing, valuation, "
                    "dilution or exit value), not only the operational cause"
                )
        watch_row = row_texts[row_index["what we watch"]]
        if watch_row and watch_row[0].lower().startswith("what we watch"):
            watch_text = watch_row[1].strip()
            if not _has_meaningful_text(watch_text, min_chars=18, min_words=3):
                editorial_errors.append(
                    f"{location}: What we watch must name an observable signal"
                )
            if _RISK_GENERIC_WATCH_RE.match(watch_text):
                editorial_errors.append(
                    f"{location}: What we watch must be a signal, not a "
                    "confirmation or diligence command"
                )
        likelihood_row = row_texts[row_index["likelihood"]]
        if (
            likelihood_row
            and likelihood_row[0].lower().startswith("likelihood")
            and not _RISK_LIKELIHOOD_VALUE_RE.match(likelihood_row[1])
        ):
            errors.append(
                f"{location}: Likelihood must be written as "
                "'High|Medium|Low: short reason'"
            )
        rating_row = row_texts[row_index["risk rating"]]
        if rating_row and rating_row[0].lower().startswith("risk rating"):
            match = _RISK_RATING_VALUE_RE.match(rating_row[1])
            if match:
                ratings.append((heading_text, int(match.group(1))))
            else:
                errors.append(
                    f"{location}: Risk Rating must be written as "
                    "'N/10: short reason' with N from 1-10"
                )
    for (prev_title, prev_rating), (title, rating) in zip(ratings, ratings[1:]):
        if rating > prev_rating:
            errors.append(
                f"{risk_id}: risk cards must be ordered by Risk Rating, "
                f"highest first — {title!r} ({rating}/10) is rated above "
                f"{prev_title!r} ({prev_rating}/10) but listed after it"
            )
            break
    return result()


def _package_validation_errors(
    package: Any,
    *,
    strict_sources: bool = True,
    warnings: list[str] | None = None,
) -> list[str]:
    errors: list[str] = []
    if not isinstance(package, dict):
        return ["memo package must be a JSON object"]
    structure = _structure_for(package)
    section_titles = structure.section_titles()
    floors = structure.content_floors()
    try:
        version = int(package.get("schema_version") or SCHEMA_VERSION)
    except (TypeError, ValueError):
        errors.append(
            f"unsupported schema_version {package.get('schema_version')!r}; expected {SCHEMA_VERSION}"
        )
        version = SCHEMA_VERSION
    if version != SCHEMA_VERSION:
        errors.append(
            f"unsupported schema_version {version}; expected {SCHEMA_VERSION}"
        )
    if not isinstance(package.get("company"), dict):
        errors.append("company must be an object")
    elif not str(package["company"].get("name") or "").strip():
        errors.append("company.name is required")
    sections = package.get("sections")
    if not isinstance(sections, list):
        errors.append("sections must be a list")
        sections = []
    sources = package.get("sources")
    if not isinstance(sources, list) or not sources:
        errors.append("sources must be a non-empty list")
    errors.extend(_calculation_errors(package))
    errors.extend(_citation_errors(package))
    if package.get("glossary") is not None:
        _validate_glossary_items(package.get("glossary"), "glossary", errors)

    by_id: dict[str, dict] = {}
    for index, section in enumerate(sections):
        location = f"sections[{index}]"
        if not isinstance(section, dict):
            errors.append(f"{location} must be an object")
            continue
        section_id = str(section.get("id") or "").strip()
        if section_id:
            by_id[section_id] = section
            if section_id not in section_titles and not _loc(
                section.get("title"), "en"
            ):
                errors.append(
                    f"{location} section id {section_id!r} must provide a "
                    "bilingual title or use a renderer-supported section id"
                )
        blocks = section.get("blocks")
        if not isinstance(blocks, list) or not blocks:
            errors.append(f"{location} blocks must be a non-empty list")
            blocks = []
        _validate_localized_value(
            section.get("title"),
            f"{location}.title",
            errors,
            required=False,
        )
        for block_index, block in enumerate(blocks):
            _validate_block(block, f"{location}.blocks[{block_index}]", errors)

    for section_id in structure.section_ids:
        if section_id not in by_id:
            errors.append(f"missing required section {section_id}")
        else:
            _validate_section_content_floor(
                section_id, by_id[section_id], errors, floors.get(section_id)
            )
    _validate_required_memo_components(package, errors, structure)

    if isinstance(sources, list):
        # Stamped by the pipeline before validation (``memo_fact_check.
        # stamp_private_material``). Absent — a stored package, a fixture —
        # the class-based exemption stands as it always did.
        run_meta = package.get("run") if isinstance(package.get("run"), dict) else {}
        private_on_file = run_meta.get("private_material_on_file")
        # The itemised version (``memo_fact_check.stamp_private_inventory``):
        # when present, a URL-less private source must name an item in it.
        inventory = run_meta.get("private_inventory")
        for index, source in enumerate(sources):
            _validate_source(
                source,
                f"sources[{index}]",
                errors,
                private_material=private_on_file if isinstance(private_on_file, bool) else None,
                strict_sources=strict_sources,
                warnings=warnings,
                private_inventory=inventory if isinstance(inventory, list) else None,
            )
    return errors


def _validate_required_memo_components(
    package: dict,
    errors: list[str],
    structure: memo_structure.MemoStructure | None = None,
) -> None:
    structure = structure or _structure_for(package)
    coverage = _memo_component_coverage(package, structure)
    for component in structure.components:
        component_id = str(component["id"])
        if not coverage.get(component_id):
            errors.append(
                f"missing required memo component {component_id}: "
                f"{component['label']}"
            )


def _memo_component_coverage(
    package: dict,
    structure: memo_structure.MemoStructure | None = None,
) -> dict[str, bool]:
    structure = structure or _structure_for(package)
    components = structure.components
    section_titles = structure.section_titles()
    coverage = {str(component["id"]): False for component in components}
    if isinstance(package.get("sources"), list) and package.get("sources"):
        coverage["source_index"] = True
    for section in package.get("sections") or []:
        if not isinstance(section, dict):
            continue
        section_title = _content_text(section.get("title")) or _section_title(
            section, "en", section_titles
        )
        previous_heading = ""
        for block in section.get("blocks") or []:
            if not isinstance(block, dict):
                continue
            for component_id in _declared_component_ids(block):
                if component_id in coverage:
                    coverage[component_id] = True
            signature = _block_signature_text(block, section_title, previous_heading)
            kind = str(block.get("type") or "paragraph")
            for component in components:
                component_id = str(component["id"])
                if coverage.get(component_id) or kind not in component["block_types"]:
                    continue
                if any(
                    re.search(pattern, signature, flags=re.IGNORECASE)
                    for pattern in component["patterns"]
                ):
                    coverage[component_id] = True
            if kind == "heading":
                previous_heading = _content_text(block.get("text") or block.get("title"))
    return coverage


def _declared_component_ids(block: dict) -> list[str]:
    value = block.get("component") or block.get("components")
    if isinstance(value, str):
        return [value.strip()]
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    return []


def _block_signature_text(block: dict, section_title: str, previous_heading: str) -> str:
    values = [section_title, previous_heading]
    for key in ("component", "title", "label", "text", "body", "caption"):
        value = block.get(key)
        if isinstance(value, list):
            values.extend(_content_text(item) for item in value)
        else:
            values.append(_content_text(value))
    if block.get("type") == "table":
        values.extend(_content_text(header) for header in block.get("headers") or [])
        for row in block.get("rows") or []:
            cells = row.get("cells") if isinstance(row, dict) else row
            if isinstance(cells, (list, tuple)):
                values.extend(_content_text(cell) for cell in cells)
    if block.get("type") in {"bullets", "callout"}:
        values.extend(_content_text(item) for item in block.get("items") or [])
    return " ".join(value for value in values if value)


def _validate_section_content_floor(
    section_id: str,
    section: dict,
    errors: list[str],
    floor: memo_structure.ContentFloor | None,
) -> None:
    score = _section_content_score(section)
    if score["real_blocks"] < 1:
        errors.append(f"section {section_id} must contain substantive memo content")
        return
    if floor is None:
        return
    if floor.min_real_blocks > 1 and score["real_blocks"] < floor.min_real_blocks:
        errors.append(
            f"section {section_id} must contain at least "
            f"{'two' if floor.min_real_blocks == 2 else floor.min_real_blocks} "
            "substantive content blocks"
        )
    if floor.bullets_or_prose:
        has_bullets = score["bullet_items"] >= 2
        has_table_or_callout_with_prose = (
            score["paragraphs"] >= 1
            and (score["tables"] + score["callouts"]) >= 1
        )
        if not has_bullets and not has_table_or_callout_with_prose:
            errors.append(
                f"section {section_id} must contain at least two substantive "
                "bullets or explanatory prose plus a substantive table/callout"
            )
    if floor.require_valuation_refs and score["valuation_refs"] < 1:
        errors.append(
            f"section {section_id} must reference scenario "
            "ranges, valuation, revenue, margins, or what moves the number"
        )


def _section_content_score(section: dict) -> dict[str, int]:
    score = {
        "real_blocks": 0,
        "paragraphs": 0,
        "bullet_items": 0,
        "tables": 0,
        "callouts": 0,
        "valuation_refs": 0,
    }
    for block in section.get("blocks") or []:
        block_score = _block_content_score(block)
        for key, value in block_score.items():
            score[key] += value
    return score


def _block_content_score(block: Any) -> dict[str, int]:
    score = {
        "real_blocks": 0,
        "paragraphs": 0,
        "bullet_items": 0,
        "tables": 0,
        "callouts": 0,
        "valuation_refs": 0,
    }
    if not isinstance(block, dict):
        return score

    kind = str(block.get("type") or "paragraph")
    if kind == "paragraph":
        text = block.get("text") or block.get("body")
        if _has_meaningful_text(text, min_chars=35, min_words=6):
            score["real_blocks"] = 1
            score["paragraphs"] = 1
            score["valuation_refs"] = int(_has_valuation_reference(text))
    elif kind == "bullets":
        item_count = sum(
            1
            for item in block.get("items") or []
            if _has_meaningful_text(item, min_chars=24, min_words=4)
        )
        if item_count:
            score["real_blocks"] = 1
            score["bullet_items"] = item_count
    elif kind == "callout":
        body = block.get("body") or block.get("text")
        items = block.get("items") or []
        item_count = sum(
            1
            for item in items
            if _has_meaningful_text(item, min_chars=24, min_words=4)
        )
        if _has_meaningful_text(body, min_chars=35, min_words=6) or item_count:
            score["real_blocks"] = 1
            score["callouts"] = 1
    elif kind == "table":
        if _table_has_real_body_row(block):
            score["real_blocks"] = 1
            score["tables"] = 1
            score["valuation_refs"] = int(_has_valuation_reference(_table_body_text(block)))
    elif kind == "chart":
        if any(
            isinstance(series, dict) and series.get("points")
            for series in block.get("series") or []
        ):
            score["real_blocks"] = 1
            score["tables"] = 1
            score["valuation_refs"] = int(
                _has_valuation_reference(
                    " ".join(
                        _content_text(block.get(key))
                        for key in ("title", "caption")
                    )
                )
            )
    return score


def _table_has_real_body_row(block: dict) -> bool:
    for row in block.get("rows") or []:
        cells = row.get("cells") if isinstance(row, dict) else row
        if not isinstance(cells, (list, tuple)):
            continue
        clean_cells = [_content_text(cell) for cell in cells]
        non_empty_cells = [cell for cell in clean_cells if cell]
        if any(
            _has_meaningful_text(cell, min_chars=12, min_words=2)
            for cell in non_empty_cells
        ):
            return True
        if len(non_empty_cells) >= 2:
            return True
    return False


def _table_body_text(block: dict) -> str:
    values: list[str] = []
    for row in block.get("rows") or []:
        cells = row.get("cells") if isinstance(row, dict) else row
        if isinstance(cells, (list, tuple)):
            values.extend(_content_text(cell) for cell in cells)
    return " ".join(value for value in values if value)


def _has_meaningful_text(
    value: Any,
    *,
    min_chars: int,
    min_words: int,
) -> bool:
    text = _content_text(value)
    if not text:
        return False
    lowered = text.lower().strip(" .")
    if any(pattern.fullmatch(lowered) for pattern in GENERIC_CONTENT_PATTERNS):
        return False
    words = re.findall(r"[\w$%][\w$%'-]*", text, flags=re.UNICODE)
    return len(text) >= min_chars or len(words) >= min_words


def _has_valuation_reference(value: Any) -> bool:
    text = _content_text(value).lower()
    return any(term in text for term in VALUATION_CONTENT_TERMS)


def _content_text(value: Any) -> str:
    return re.sub(r"\s+", " ", _loc(value, "en")).strip()


def _validate_block(block: Any, location: str, errors: list[str]) -> None:
    if not isinstance(block, dict):
        errors.append(f"{location} must be an object")
        return
    kind = str(block.get("type") or "paragraph")
    if kind not in SUPPORTED_BLOCK_TYPES:
        errors.append(f"{location}.type {kind!r} is unsupported")
        return
    if kind == "heading":
        _validate_localized_value(
            block.get("text") or block.get("title"),
            f"{location}.text",
            errors,
        )
    elif kind == "paragraph":
        _validate_localized_value(
            block.get("text") or block.get("body"),
            f"{location}.text",
            errors,
        )
    elif kind == "bullets":
        items = block.get("items")
        if not isinstance(items, list) or not items:
            errors.append(f"{location}.items must be a non-empty list")
            return
        for index, item in enumerate(items):
            _validate_localized_value(item, f"{location}.items[{index}]", errors)
    elif kind == "callout":
        _validate_localized_value(
            block.get("title") or block.get("label"),
            f"{location}.title",
            errors,
        )
        _validate_localized_value(
            block.get("body") or block.get("text"),
            f"{location}.body",
            errors,
            required=False,
        )
        items = block.get("items") or []
        if not isinstance(items, list):
            errors.append(f"{location}.items must be a list")
            return
        for index, item in enumerate(items):
            _validate_localized_value(item, f"{location}.items[{index}]", errors)
    elif kind == "chart":
        _validate_chart_block(block, location, errors)
    elif kind == "glossary":
        _validate_glossary_items(block.get("items") or block.get("terms"), location, errors)
    elif kind == "table":
        _validate_localized_value(
            block.get("title"),
            f"{location}.title",
            errors,
            required=False,
        )
        headers = block.get("headers") or []
        rows = block.get("rows") or []
        if not headers and not rows:
            errors.append(f"{location} table must include headers or rows")
        # A blank corner cell over the row labels ("", "2025", "2026") is
        # a normal table: the renderer draws it blank in both languages.
        corner_blank = (
            len(headers) > 1
            and _blank_header(headers[0])
            and not any(_blank_header(header) for header in headers[1:])
        )
        for index, header in enumerate(headers):
            if index == 0 and corner_blank:
                continue
            _validate_localized_value(header, f"{location}.headers[{index}]", errors)
        for row_index, row in enumerate(rows):
            cells = row.get("cells") if isinstance(row, dict) else row
            if not isinstance(cells, (list, tuple)) or not cells:
                errors.append(f"{location}.rows[{row_index}] must contain cells")
                continue
            for cell_index, cell in enumerate(cells):
                _validate_localized_value(
                    cell,
                    f"{location}.rows[{row_index}].cells[{cell_index}]",
                    errors,
                )


def _validate_glossary_items(items: Any, location: str, errors: list[str]) -> None:
    """A glossary (block ``items`` or the package-level list): every item a
    bilingual term with a bilingual definition."""
    if not isinstance(items, list) or not items:
        errors.append(f"{location}.items must be a non-empty list of glossary terms")
        return
    for index, item in enumerate(items):
        where = f"{location}.items[{index}]"
        if not isinstance(item, dict):
            errors.append(f"{where} must be an object with term and definition")
            continue
        _validate_localized_value(item.get("term"), f"{where}.term", errors)
        _validate_localized_value(item.get("definition"), f"{where}.definition", errors)


def glossary_items(value: Any) -> list[dict]:
    """The renderable ``{term, definition}`` items of a glossary block or
    the package-level list (anything malformed is skipped)."""
    items = value.get("items") or value.get("terms") if isinstance(value, dict) else value
    if not isinstance(items, list):
        return []
    return [
        item
        for item in items
        if isinstance(item, dict) and item.get("term") and item.get("definition")
    ]


def _chart_label_key(value: Any) -> str:
    """A chart label as the x-shape check compares it: a bilingual label by
    its English half, anything else as written."""
    if isinstance(value, dict):
        return str(value.get("en") or "").strip()
    return str(value if value is not None else "").strip()


def _blank_header(value: Any) -> bool:
    """A header cell with no English text (the English decides: the
    generation gate fills blank Chinese halves with a placeholder)."""
    if isinstance(value, dict):
        return not str(value.get("en") or "").strip()
    return not str(value if value is not None else "").strip()


def _validate_chart_block(block: dict, location: str, errors: list[str]) -> None:
    """Validate a ``chart`` block (structure-v2 chart slots).

    The text AROUND the image (title, caption, reading, unit) is bilingual
    like any other block text. Text INSIDE the image — series labels, x
    categories, reference-line labels — may be plain (one string serves
    both documents) or bilingual ``{"en","zh"}`` (each document draws its
    own; a blank Chinese half draws the English, it never fails the
    package). Optional, drawing-only fields are accepted and never block:
    ``reference_lines`` [{"y": number, "label"?}] (up to three dashed
    rules, e.g. 1.0x breakeven; a malformed rule is skipped) and
    ``"estimate": true`` on a point (drawn lighter / dashed).
    """
    chart_type = str(block.get("chart_type") or "").strip().lower()
    if chart_type not in SUPPORTED_CHART_TYPES:
        errors.append(
            f"{location}.chart_type must be one of "
            f"{sorted(SUPPORTED_CHART_TYPES)}"
        )
        return
    _validate_localized_value(block.get("title"), f"{location}.title", errors)
    _validate_localized_value(
        block.get("caption"), f"{location}.caption", errors, required=False
    )
    _validate_localized_value(
        block.get("reading"), f"{location}.reading", errors, required=False
    )
    _validate_localized_value(
        block.get("unit"), f"{location}.unit", errors, required=False,
        allow_plain=True,
    )
    series = block.get("series")
    if not isinstance(series, list) or not 1 <= len(series) <= 4:
        errors.append(f"{location}.series must be a list of 1-4 series")
        return
    if chart_type in _SINGLE_SERIES_CHART_TYPES and len(series) != 1:
        errors.append(
            f"{location}: chart_type {chart_type!r} takes exactly one series "
            "— use 'grouped_bar' for several"
        )
    x_shapes: list[tuple[str, ...]] = []
    total_points = 0
    for s_index, one in enumerate(series):
        where = f"{location}.series[{s_index}]"
        if not isinstance(one, dict):
            errors.append(f"{where} must be an object")
            continue
        if not _chart_label_key(one.get("label")):
            errors.append(f"{where}.label must be a non-empty plain string")
        points = one.get("points")
        if not isinstance(points, list) or not points:
            errors.append(f"{where}.points must be a non-empty list")
            continue
        xs: list[str] = []
        for p_index, point in enumerate(points):
            spot = f"{where}.points[{p_index}]"
            if not isinstance(point, dict):
                errors.append(f"{spot} must be an object with x and y")
                continue
            x = point.get("x")
            if not _chart_label_key(x):
                errors.append(f"{spot}.x must be a non-empty label")
            y = point.get("y")
            if isinstance(y, bool) or not isinstance(y, (int, float)):
                errors.append(f"{spot}.y must be a plain number")
            xs.append(_chart_label_key(x))
            total_points += 1
        x_shapes.append(tuple(xs))
    if len({shape for shape in x_shapes}) > 1:
        errors.append(
            f"{location}: every series must share the same x categories in "
            "the same order"
        )
    if total_points < 2 and not errors:
        errors.append(
            f"{location}: a chart needs at least two data points — use "
            "prose for a single number"
        )
    source_ids = block.get("source_ids")
    if source_ids is not None and (
        not isinstance(source_ids, list)
        or any(not str(item or "").strip() for item in source_ids)
    ):
        errors.append(f"{location}.source_ids must be a list of source ids")


_CALC_ID_RE = re.compile(r"^C\d+$")


def _calculation_errors(package: dict) -> list[str]:
    """`calculations` is optional; when present it is a list of notes with
    an id "C<n>", a formula and a result (any text field may be localized
    or plain), and ids are unique."""
    calculations = package.get("calculations")
    if calculations is None:
        return []
    if not isinstance(calculations, list):
        return ["calculations must be a list when present"]
    errors: list[str] = []
    seen: set[str] = set()
    for index, calc in enumerate(calculations):
        location = f"calculations[{index}]"
        if not isinstance(calc, dict):
            errors.append(f"{location} must be an object")
            continue
        calc_id = str(calc.get("id") or "").strip()
        if not _CALC_ID_RE.match(calc_id):
            errors.append(f"{location}.id must look like C1, C2, ...")
        elif calc_id in seen:
            errors.append(f"{location}.id {calc_id} is duplicated")
        seen.add(calc_id)
        for key in ("formula", "result"):
            if not _loc(calc.get(key), "en").strip():
                errors.append(f"{location}.{key} is required")
        inputs = calc.get("inputs")
        if inputs is not None and not isinstance(inputs, list):
            errors.append(f"{location}.inputs must be a list")
    return errors


def _iter_strings(value: Any):
    if isinstance(value, str):
        yield value
    elif isinstance(value, dict):
        for item in value.values():
            yield from _iter_strings(item)
    elif isinstance(value, list):
        for item in value:
            yield from _iter_strings(item)


def _citation_errors(package: dict) -> list[str]:
    """Every inline `[S#]` / `[C#]` must name an existing source or
    calculation note — a dangling link is a fabricated citation."""
    known: set[str] = set()
    for source in package.get("sources") or []:
        if isinstance(source, dict):
            known.add(str(source.get("id") or "").strip())
    for calc in package.get("calculations") or []:
        if isinstance(calc, dict):
            known.add(str(calc.get("id") or "").strip())
    errors: list[str] = []
    reported: set[str] = set()
    for index, section in enumerate(package.get("sections") or []):
        if not isinstance(section, dict):
            continue
        section_id = str(section.get("id") or f"sections[{index}]")
        for text in _iter_strings(section.get("blocks")):
            for citation_id in citation_ids(text):
                if citation_id in known or citation_id in reported:
                    continue
                reported.add(citation_id)
                errors.append(
                    f"{section_id}: unknown_citation [{citation_id}] — cite "
                    "only ids that exist in the package sources or the "
                    "pinned calculation notes"
                )
    return errors


# A source whose class or title says it came from somewhere a reader cannot
# browse to — diligence, an interview, a deck, the data room, an internal
# model — may omit `url`. Everything else was retrieved from the web and must
# carry the page it came from: four in five shipped citations had none, and a
# citation the reader cannot follow is not a citation.
_SOURCE_URL_OPTIONAL_RE = re.compile(
    r"\b(?:diligence|interview|reference calls?|call notes?|founder updates?|"
    r"management (?:call|meeting|interview|presentation)|data ?room|private|"
    r"internal|bsh|memo studio|uploads?|uploaded|file|deck|transcripts?|kpis?|"
    r"portfolio|investors?|intermediary|board|term sheet|cap table|model|"
    r"company[- _]?(?:reported|disclosure|documents?|materials|data|provided)|"
    r"company|proprietary|confidential|email|correspondence|survey|expert|primary)\b",
    re.I,
)


def _memo_source_url_required_enabled() -> bool:
    return os.environ.get("BSH_MEMO_SOURCE_URL_REQUIRED", "1") == "1"


def source_url_optional(source: dict) -> bool:
    """Whether this source may omit its URL: its class or title marks it as
    private material rather than a web page."""
    haystack = " ".join(
        [
            _loc(source.get("class"), "en") or str(source.get("class") or ""),
            _loc(source.get("title"), "en") or str(source.get("title") or ""),
        ]
    )
    return bool(_SOURCE_URL_OPTIONAL_RE.search(haystack))


# When a source was current. ``as_of`` is the legacy field; ``published_at``
# and ``data_period`` are the honest ones (YYYY, YYYY-MM, YYYY-MM-DD or
# "undated"), and any one of them dates the source.
_SOURCE_DATE_KEYS = ("as_of", "published_at", "data_period")


# A private class that claims the firm's own diligence ("BSH primary
# diligence", "internal diligence notes"): with nothing behind it, the
# figures it carries are the registry's, not the firm's (R4) — the fix is
# to drop them, not to find a URL. Other private-sounding classes ("internal
# document") were public pages in disguise: they keep the carry-the-URL
# message.
_BSH_DILIGENCE_CLASS_RE = re.compile(r"\b(?:bsh|diligence)\b", re.I)

# The class the sources contract prescribes for a registry value with no
# document behind it (claude_runner.MEMO_PACKAGE_SOURCES_CONTRACT). It says
# outright that a reader has nothing to open, so it needs no URL and claims
# no diligence — even when its title names the BSH registry (ZaiNar
# 2026-09-23__015923 failed on exactly this source).
_UNVERIFIED_REGISTRY_CLASS_RE = re.compile(r"\bunverified registry value\b", re.I)


def _inventory_titles(inventory: list, limit: int = 5) -> str:
    titles = [
        f"\u201c{str(item.get('title') or item.get('id') or '').strip()}\u201d"
        for item in inventory
        if isinstance(item, dict) and str(item.get("title") or item.get("id") or "").strip()
    ]
    shown = ", ".join(titles[:limit])
    if len(titles) > limit:
        shown += f" (+{len(titles) - limit} more)"
    return shown


def _validate_source(
    source: Any,
    location: str,
    errors: list[str],
    *,
    private_material: bool | None = None,
    strict_sources: bool = True,
    warnings: list[str] | None = None,
    private_inventory: list | None = None,
) -> None:
    """``private_material`` is whether the firm holds anything private on
    this company (None = unknown). When it is False, no source can honestly
    be private, so the class-based URL exemption does not apply — otherwise
    a public page only has to call itself an "internal document".

    ``private_inventory`` (``run.private_inventory``, optional) itemises that
    material; when given, a URL-less private-class source passes only when
    ``memo_fact_check.inventory_match`` finds the item it names (by title,
    or by ``private_ref``). Absent, the boolean rule above applies.

    ``strict_sources=False`` (re-rendering a stored package) reports a
    missing URL into ``warnings`` instead of ``errors``."""
    if not isinstance(source, dict):
        errors.append(f"{location} must be an object")
        return
    for key in ("id", "title", "class", "treatment"):
        if not str(source.get(key) or "").strip():
            errors.append(f"{location}.{key} is required")
    if not any(str(source.get(key) or "").strip() for key in _SOURCE_DATE_KEYS):
        errors.append(f"{location}.as_of is required")
    url_errors: list[str] = []
    url = source.get("url")
    claims_private = source_url_optional(source)
    source_class = _loc(source.get("class"), "en") or str(source.get("class") or "")
    claims_bsh = bool(
        _BSH_DILIGENCE_CLASS_RE.search(
            " ".join([source_class, _loc(source.get("title"), "en") or str(source.get("title") or "")])
        )
    )
    no_bsh_document = (
        f"{location}.url is required: its class {source_class!r} names BSH "
        "diligence, but no BSH document backs this; drop the source and "
        "remove or restate as unverified the figures that rest on it"
    )
    if url is not None and str(url).strip():
        if not isinstance(url, str) or not url.strip().startswith(
            ("http://", "https://")
        ):
            errors.append(f"{location}.url must be an http(s) URL when present")
    elif _UNVERIFIED_REGISTRY_CLASS_RE.search(source_class):
        pass
    elif (
        _memo_source_url_required_enabled()
        and claims_private
        and private_inventory is not None
    ):
        from server import memo_fact_check  # lazy: memo_fact_check imports widely

        if memo_fact_check.inventory_match(source, private_inventory) is None:
            if not private_inventory:
                url_errors.append(
                    no_bsh_document
                    if claims_bsh
                    else f"{location}.url is required: its class {source_class!r} "
                    "says it is private, but the firm holds nothing private on "
                    "this company, so it came from a page a reader can open. "
                    "Carry that page's URL from the analysis artifacts or the "
                    "known-sources list; do not reclassify it"
                )
            else:
                url_errors.append(
                    f"{location}.url is required: its class {source_class!r} says "
                    "it is private, but it names none of the firm's private "
                    "material on this company — "
                    f"{_inventory_titles(private_inventory)}. Cite the document "
                    "it came from by its exact title (or set private_ref to its "
                    "id); if it came from a web page, carry that page's URL; "
                    "otherwise drop the source and remove or restate as "
                    "unverified the figures that rest on it"
                )
    elif _memo_source_url_required_enabled() and claims_private and private_material is False:
        url_errors.append(
            no_bsh_document
            if claims_bsh
            else f"{location}.url is required: its class {source_class!r} says it is "
            "private, but the firm holds nothing private on this company — no "
            "research documents, founder updates, transcripts or reference "
            "calls — so it came from a page a reader can open. Carry that "
            "page's URL from the analysis artifacts or the known-sources list; "
            "do not reclassify it"
        )
    elif _memo_source_url_required_enabled() and not claims_private:
        # Only offer reclassifying when it could be true. With nothing
        # private on file it is the loophole, and it would cost a round to
        # be caught by the branch above (RadixArk 2026-09-21__195426).
        nothing_private = private_material is False or private_inventory == []
        way_out = (
            "; the firm holds nothing private on this company, so it cannot "
            "be an internal document. For a paid report (IDC, Gartner and the "
            "like) cite the publisher's own press release or summary page for "
            "the figure; only if no page carries it, drop the source and the "
            "claims that rest on it"
            if nothing_private
            else "; if it is really a private file, interview or internal "
            "document, say so in its class instead"
        )
        url_errors.append(
            f"{location}.url is required: a source of class {source_class!r} is "
            "web-retrieved, so carry the page URL the analysis artifacts or the "
            "known-sources list recorded" + way_out
        )
    if strict_sources:
        errors.extend(url_errors)
    elif warnings is not None:
        warnings.extend(url_errors)
    _validate_localized_value(
        source.get("class"),
        f"{location}.class",
        errors,
        allow_plain=True,
    )
    _validate_localized_value(
        source.get("treatment"),
        f"{location}.treatment",
        errors,
        allow_plain=False,
    )


_LATIN_WORD_RE = re.compile(r"[A-Za-z]{2,}")
_CJK_CHAR_RE = re.compile(r"[㐀-䶿一-鿿]")


def _is_language_neutral_text(text: str) -> bool:
    """True for plain strings safe to render identically in both locales.

    Dates, figures, percentages, tickers, and proper-noun names carry no
    translatable prose, so ``{"en": "2026-05", "zh": "2026-05"}`` would be
    pure ceremony. Anything containing ordinary lowercase English words
    ("120+ filed / 90+ issued") still requires an explicit bilingual pair.
    """
    stripped = text.strip()
    if not stripped or _CJK_CHAR_RE.search(stripped):
        return False
    words = _LATIN_WORD_RE.findall(stripped)
    if not words:
        return True
    if len(words) > 5:
        return False
    return all(word[0].isupper() for word in words)


def _validate_localized_value(
    value: Any,
    location: str,
    errors: list[str],
    *,
    required: bool = True,
    allow_plain: bool = False,
) -> None:
    if value is None or value == "":
        if required:
            errors.append(f"{location} is required")
        return
    if isinstance(value, dict):
        if not str(value.get("en") or "").strip():
            errors.append(f"{location}.en is required")
        if not str(value.get("zh") or "").strip():
            errors.append(f"{location}.zh is required")
        return
    if allow_plain or isinstance(value, (int, float)):
        return
    if isinstance(value, str) and _is_language_neutral_text(value):
        return
    errors.append(
        f"{location} must be bilingual with en and zh (plain strings are "
        "allowed only for language-neutral values such as dates, figures, "
        "or proper-noun names)"
    )


MEMO_LOCALES = ("en", "zh")
_VALIDATION_FILE_NAMES = {"en": "validation.txt", "zh": "validation_cn.txt"}


def _save_locale_document(
    payload: dict,
    locale: str,
    path: Path,
    zh_fallbacks: list[tuple[str, str, str]] | None,
) -> None:
    """Build one language's document and write it to ``path``; a Chinese
    build records every English fallback into ``zh_fallbacks``."""
    path.parent.mkdir(parents=True, exist_ok=True)
    sink = _ZH_FALLBACKS.set(zh_fallbacks if locale == "zh" else None)
    try:
        document = _build_document(payload, locale)
    finally:
        _ZH_FALLBACKS.reset(sink)
    document.save(path)


def render_memos(
    package: dict | Path | str,
    *,
    out_en: Path | str,
    out_zh: Path | str,
    validation_en: Path | str | None = None,
    validation_zh: Path | str | None = None,
    inventory_path: Path | str | None = None,
    manifest_path: Path | str | None = None,
    strict_sources: bool = True,
) -> dict:
    """Render English and Chinese DOCX files from one structured package.

    ``strict_sources=False`` re-renders a stored package that predates the
    source-URL rule: a missing web-source URL is logged as a warning (and
    returned under ``warnings``) instead of failing the render. New runs
    keep the default. One language at a time: :func:`render_memo_locale`."""
    payload = _read_package(package) if isinstance(package, (str, Path)) else package
    warnings = validate_package(payload, strict_sources=strict_sources)
    outputs = {
        "en": Path(out_en),
        "zh": Path(out_zh),
    }
    zh_fallbacks: list[tuple[str, str, str]] = []
    for locale, path in outputs.items():
        _save_locale_document(payload, locale, path, zh_fallbacks)

    run_dir = outputs["en"].parent.parent
    logs_dir = run_dir / "logs"
    logs_dir.mkdir(parents=True, exist_ok=True)
    validation_paths = {
        "en": Path(validation_en) if validation_en else logs_dir / "validation.txt",
        "zh": Path(validation_zh) if validation_zh else logs_dir / "validation_cn.txt",
    }
    for locale, path in validation_paths.items():
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(
            _validation_report(
                payload,
                outputs[locale],
                locale,
                zh_fallbacks=zh_fallbacks if locale == "zh" else None,
                source_warnings=warnings,
            ),
            encoding="utf-8",
        )
    if inventory_path:
        _write_inventory(Path(inventory_path), payload, outputs, validation_paths)
    else:
        _write_inventory(logs_dir / "file_inventory.md", payload, outputs, validation_paths)
    if manifest_path:
        _append_manifest(Path(manifest_path), payload, outputs, validation_paths)
    else:
        _append_manifest(logs_dir / "run_manifest.md", payload, outputs, validation_paths)
    return {
        "ok": True,
        "outputs": {locale: str(path) for locale, path in outputs.items()},
        "validation": {locale: str(path) for locale, path in validation_paths.items()},
        "renderer_version": RENDERER_VERSION,
        "warnings": list(warnings),
    }


def render_memo_locale(
    package: dict | Path | str,
    locale: str,
    out_path: Path | str,
    *,
    strict_sources: bool = True,
    validation_path: Path | str | None = None,
    inventory_path: Path | str | None = None,
    manifest_path: Path | str | None = None,
) -> dict:
    """Render ONE language of the memo — the pipeline's English-first
    delivery (``"en"`` before the Chinese exists) and "Retry Chinese"
    (``"zh"`` alone). The other language's file is never read or written.

    Validation is the pair render's for that language: ``"en"`` checks the
    package the way :func:`english_package_validation_errors` checks its
    structure — blank Chinese halves are allowed, nothing else is relaxed;
    ``"zh"`` checks the full bilingual package exactly as
    :func:`render_memos` does. ``strict_sources`` as in :func:`render_memos`.

    Side effects, the same the pair render has for that language: the docx
    at ``out_path``; ``logs/validation.txt`` (en) or ``logs/validation_cn.txt``
    (zh, with its "## Chinese fallbacks" list) under the run dir
    (``out_path``'s grandparent) unless ``validation_path`` is given;
    ``logs/file_inventory.md`` rewritten for this language with the other
    language's lines kept; one "## Analysis Finalization" block appended to
    ``logs/run_manifest.md``. Returns ``{"ok", "locale", "output",
    "outputs": {locale: path}, "validation": {locale: path},
    "renderer_version", "warnings"}`` plus ``"zh_fallback_count"`` for zh.
    Raises :class:`MemoRenderError` like :func:`render_memos`."""
    if locale not in MEMO_LOCALES:
        raise ValueError(f"locale must be one of {MEMO_LOCALES}, got {locale!r}")
    payload = _read_package(package) if isinstance(package, (str, Path)) else package
    if locale == "en" and isinstance(payload, dict):
        # The Chinese may not exist yet: validate the structure with blank
        # zh halves filled, render the package itself.
        warnings = validate_package(
            fill_blank_zh_placeholders(payload), strict_sources=strict_sources
        )
    else:
        warnings = validate_package(payload, strict_sources=strict_sources)
    output = Path(out_path)
    zh_fallbacks: list[tuple[str, str, str]] | None = [] if locale == "zh" else None
    _save_locale_document(payload, locale, output, zh_fallbacks)

    logs_dir = output.parent.parent / "logs"
    logs_dir.mkdir(parents=True, exist_ok=True)
    validation = (
        Path(validation_path)
        if validation_path
        else logs_dir / _VALIDATION_FILE_NAMES[locale]
    )
    validation.parent.mkdir(parents=True, exist_ok=True)
    validation.write_text(
        _validation_report(
            payload,
            output,
            locale,
            zh_fallbacks=zh_fallbacks,
            source_warnings=warnings,
        ),
        encoding="utf-8",
    )
    outputs = {locale: output}
    validation_paths = {locale: validation}
    _write_inventory(
        Path(inventory_path) if inventory_path else logs_dir / "file_inventory.md",
        payload,
        outputs,
        validation_paths,
        keep_other_locales=True,
    )
    _append_manifest(
        Path(manifest_path) if manifest_path else logs_dir / "run_manifest.md",
        payload,
        outputs,
        validation_paths,
    )
    result: dict[str, Any] = {
        "ok": True,
        "locale": locale,
        "output": str(output),
        "outputs": {locale: str(output)},
        "validation": {locale: str(validation)},
        "renderer_version": RENDERER_VERSION,
        "warnings": list(warnings),
    }
    if zh_fallbacks is not None:
        result["zh_fallback_count"] = len(dict.fromkeys(zh_fallbacks))
    return result


def pin_section_titles(package: dict, structure) -> list[str]:
    """Replace a core section's heading with the canonical title when the
    parity gate would not recognise it. Returns one note per change.

    The package's own heading normally wins, so a model that writes
    "公司概况与公司治理" where the profile says "公司概览" ships that heading —
    and the Chinese parity gate, which finds sections by full-line anchored
    patterns, then reports the section missing although every paragraph of
    it is there. A live Gemini run produced four such false "missing" P0s.
    Only headings the gate would reject are replaced: a recognised variant
    such as "II. Company Overview" keeps its numbering.

    The renderer itself now always prints the canonical title for a core
    section (it owns the numbering); this still normalizes the package for
    the callers that read titles from it.
    """
    titles = structure.section_titles()
    patterns = structure.parity_patterns()
    notes: list[str] = []
    for section in package.get("sections") or []:
        if not isinstance(section, dict):
            continue
        section_id = str(section.get("id") or "")
        if section_id not in titles or section_id not in patterns:
            continue
        for locale in ("en", "zh"):
            canonical = _raw_text(titles[section_id], locale)
            pattern = patterns[section_id].get(locale)
            if not canonical or pattern is None:
                continue
            own = _raw_text(section.get("title"), locale)
            if own and pattern.match(own):
                continue
            title = section.get("title")
            if not isinstance(title, dict):
                title = {"en": own, "zh": own} if own else {}
            title[locale] = canonical
            section["title"] = title
            notes.append(f"{section_id}.title.{locale}: {own!r} -> {canonical!r}")
    return notes


# ---- document plan ------------------------------------------------------------
#
# The renderer owns the document's shape. Core sections print their
# canonical structure titles, numbered by position ("IV. Investment Risk" /
# "四、投资风险"); sections the package adds beyond the profile are numbered
# after the core ones with their own titles (any numeral the model wrote is
# dropped); every sources-like section the model wrote folds into the ONE
# sources section; back matter follows its profile entry — ``numbered:
# true`` takes the next numeral after the body ("XIII. Sources, ..." /
# "十三、来源、..."), otherwise its own title — and the renderer's
# Calculation notes come last, unnumbered. Each level-1 heading carries the
# bookmark ``bsh_sec_<n>``
# — n is its order in the document, identical in the EN and ZH files — and
# each level-2 heading ``bsh_sub_<n>_<m>``, so the contents list, Word's
# navigation pane and the web viewer's outline all address the same ids.


SECTION_BOOKMARK_PREFIX = "bsh_sec_"
SUBSECTION_BOOKMARK_PREFIX = "bsh_sub_"


def section_bookmark(number: int) -> str:
    return f"{SECTION_BOOKMARK_PREFIX}{number}"


def subsection_bookmark(number: int, sub: int) -> str:
    return f"{SUBSECTION_BOOKMARK_PREFIX}{number}_{sub}"


@dataclass
class _PlanEntry:
    kind: str  # "section" | "sources" | "calculations" | "back_matter"
    number: int  # document order among level-1 headings -> bsh_sec_<number>
    title: dict[str, str]  # full heading, numeral included
    bare: dict[str, str]  # title alone (contents list, _toc_titles)
    label: dict[str, str]  # "IV." / "四、"; "" for unnumbered back matter
    section: dict | None = None
    section_id: str = ""
    folded: list[dict] = field(default_factory=list)


# Titles that name the sources index itself ("Sources", "Sources, Source
# Classes, and Disclosures", "Source Classes and Fact Index", "Fact Reference
# Index and Disclosures") — never a body topic that merely starts with the
# word ("Source of Funds", "Sources and Uses of Capital", "资金来源").
_SOURCES_TITLE_EN_RE = re.compile(
    r"^(?:sources?(?!\s*(?:and|&)\s*uses)(?=\s*(?:$|[,:;]|and\b|&))"
    r"|sources?\s+(?:classes|index|list)\b"
    r"|fact\s+(?:reference\s+)?index\b"
    r"|references\s*$)",
    re.IGNORECASE,
)
_SOURCES_TITLE_ZH_RE = re.compile(
    r"^(?:(?:资料|信息|参考)?来源(?:$|[、与及和：:]|清单|索引|类别)"
    r"|资料类别|来源类别|事实索引|参考资料)"
)
DEFAULT_SOURCES_TITLE = {
    "en": "Sources, Source Classes, and Fact Reference Index",
    "zh": "来源、来源类别与事实索引",
}


def _raw_text(value: Any, locale: str) -> str:
    """A localized value's text in ``locale`` with no English fallback (and
    no fallback bookkeeping) — for decisions, never for printing."""
    if isinstance(value, dict):
        return str(value.get(locale) or "")
    return str(value or "")


def _is_sources_like(section: dict, structure: memo_structure.MemoStructure) -> bool:
    """A model-written section that restates the sources index: its id
    starts with ``source``, its title reads as a sources/fact-index title,
    or it carries the ``source_index`` component. Folded into the one
    sources section so the memo never prints two."""
    section_id = str(section.get("id") or "").strip().lower()
    if section_id.startswith("source"):
        return True
    patterns = structure.parity_patterns().get("sources") or {}
    for locale, title_re in (("en", _SOURCES_TITLE_EN_RE), ("zh", _SOURCES_TITLE_ZH_RE)):
        title = _raw_text(section.get("title"), locale).strip()
        if not title:
            continue
        if memo_structure.heading_matches(patterns.get(locale), title):
            return True
        if title_re.match(memo_structure.strip_section_numeral(title)):
            return True
    return any(
        isinstance(block, dict) and "source_index" in _declared_component_ids(block)
        for block in section.get("blocks") or []
    )


def _document_plan(
    package: dict, structure: memo_structure.MemoStructure
) -> list[_PlanEntry]:
    sections = [s for s in package.get("sections") or [] if isinstance(s, dict)]
    core_ids = list(structure.section_ids)
    pseudo_ids = {ps.id for ps in structure.pseudo_sections if ps.id != "sources"}
    core: dict[str, dict] = {}
    extras: list[dict] = []
    folded: list[dict] = []
    back: list[dict] = []
    for section in sections:
        section_id = str(section.get("id") or "")
        if section_id in core_ids and section_id not in core:
            core[section_id] = section
        elif _is_sources_like(section, structure):
            folded.append(section)
        elif section_id in pseudo_ids:
            back.append(section)
        else:
            extras.append(section)

    entries: list[_PlanEntry] = []

    def _numbered(position: int, bare: dict[str, str], section: dict, section_id: str) -> None:
        entries.append(
            _PlanEntry(
                kind="section",
                number=len(entries) + 1,
                title={
                    locale: memo_structure.numbered_title(position, bare[locale], locale)
                    for locale in ("en", "zh")
                },
                bare=bare,
                label={
                    "en": f"{memo_structure.section_numeral(position, 'en')}.",
                    "zh": f"{memo_structure.section_numeral(position, 'zh')}、",
                },
                section=section,
                section_id=section_id,
            )
        )

    for index, section_id in enumerate(core_ids, start=1):
        section = core.get(section_id)
        if section is None:
            continue
        definition = structure.section(section_id)
        bare = {
            "en": definition.en_title if definition else section_id,
            "zh": definition.zh_title if definition else section_id,
        }
        _numbered(index, bare, section, section_id)
    position = len(core_ids)
    for section in extras:
        position += 1
        section_id = str(section.get("id") or "")
        fallback = section_id.replace("_", " ").strip().title() or "Section"
        en = memo_structure.strip_section_numeral(_raw_text(section.get("title"), "en"))
        zh = memo_structure.strip_section_numeral(_raw_text(section.get("title"), "zh"))
        if en and not zh:
            _note_zh_fallback({"en": en}, None)
        _numbered(position, {"en": en or fallback, "zh": zh or en or fallback}, section, section_id)

    def _back_matter(
        kind: str,
        bare: dict[str, str],
        numbered: bool,
        *,
        section: dict | None = None,
        section_id: str,
        folded_sections: list[dict] | None = None,
    ) -> None:
        # Back matter follows its own profile entry: ``numbered: true``
        # takes the next numeral after every body section, otherwise the
        # profile's title prints as written.
        nonlocal position
        if numbered:
            position += 1
            _numbered(position, bare, section, section_id)
            entry = entries[-1]
            entry.kind = kind
            entry.folded = list(folded_sections or [])
            return
        entries.append(
            _PlanEntry(
                kind=kind,
                number=len(entries) + 1,
                title=dict(bare),
                bare=dict(bare),
                label={"en": "", "zh": ""},
                section=section,
                section_id=section_id,
                folded=list(folded_sections or []),
            )
        )

    if package.get("sources") or folded:
        pseudo = structure.pseudo_section("sources")
        _back_matter(
            "sources",
            {"en": pseudo.en_title, "zh": pseudo.zh_title}
            if pseudo
            else dict(DEFAULT_SOURCES_TITLE),
            pseudo.numbered if pseudo else True,
            section_id="sources",
            folded_sections=folded,
        )
    for section in back:
        section_id = str(section.get("id") or "")
        pseudo = structure.pseudo_section(section_id)
        _back_matter(
            "back_matter",
            {"en": pseudo.en_title, "zh": pseudo.zh_title}
            if pseudo
            else {"en": section_id, "zh": section_id},
            bool(pseudo and pseudo.numbered),
            section=section,
            section_id=section_id,
        )
    if package.get("calculations"):
        # The renderer's own appendix behind every [C#] link: no profile
        # entry, so it keeps its plain title, last.
        _back_matter(
            "calculations",
            dict(CALCULATIONS_TITLE),
            False,
            section_id="calculations",
        )
    return entries


def _build_document(package: dict, locale: str) -> Document:
    structure = _structure_for(package)
    pin_section_titles(package, structure)
    plan = _document_plan(package, structure)
    document = Document()
    _configure_document(document, package, locale)
    _add_front_page(document, package, locale, plan)
    context = _RenderContext(package=package, structure=structure, locale=locale)
    for index, entry in enumerate(plan):
        if index:
            document.add_page_break()
        token = _RENDER_SECTION.set(entry.section_id or entry.kind)
        try:
            _add_plan_entry(document, entry, context)
        finally:
            _RENDER_SECTION.reset(token)
    return document


# ---- page setup, styles, metadata, header and footer -------------------------


def _text_width_twips(document: Document) -> int:
    section = document.sections[0]
    width = section.page_width - section.left_margin - section.right_margin
    return int(width / 635)


def _configure_document(document: Document, package: dict, locale: str) -> None:
    font = _font(locale)
    section = document.sections[0]
    section.top_margin = Cm(_MARGIN_CM)
    section.bottom_margin = Cm(_MARGIN_CM)
    section.left_margin = Cm(_MARGIN_CM)
    section.right_margin = Cm(_MARGIN_CM)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    section.different_first_page_header_footer = True

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(_TEXT_SIZE)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = _LINE_SPACING[locale]
    _set_rfonts(normal.element.get_or_add_rPr(), font, locale)
    _configure_heading_styles(document, locale)
    _configure_list_style(document, locale)
    _ensure_derived_styles(document)
    if locale == "zh":
        _set_zh_language_defaults(document)
    _set_core_properties(document, package, locale)
    _write_header_footer(document, package, locale)


_HEADING_LOOK = {
    # level: (size pt, color, space before pt, space after pt)
    1: (15, NAVY, 12, 5),
    2: (11.5, NAVY, 16, 4),
    3: (10.5, BLACK, 7, 3),
}


def _strip_theme_attributes(rpr: Any) -> None:
    """Word lets a theme font or theme color win over an explicit one, so
    a restyled built-in style must lose them."""
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is not None:
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
    color = rpr.find(qn("w:color"))
    if color is not None:
        for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
            if color.get(qn(attr)) is not None:
                del color.attrib[qn(attr)]


def _configure_heading_styles(document: Document, locale: str) -> None:
    """The built-in Heading 1/2/3 styles, restyled to the memo's look
    (Arial / Microsoft YaHei, navy, the renderer's sizes and spacing) with
    their outline levels, so Word's navigation pane and a PDF outline see
    real headings. Runs keep the same direct formatting, so the web viewer
    and Word render them identically."""
    font = _font(locale)
    for level, (size, color, before, after) in _HEADING_LOOK.items():
        try:
            style = document.styles[f"Heading {level}"]
        except KeyError:
            continue
        rpr = style.element.get_or_add_rPr()
        _strip_theme_attributes(rpr)
        style.font.name = font
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = RGBColor.from_string(color)
        _set_rfonts(rpr, font, locale)
        fmt = style.paragraph_format
        fmt.space_before = Pt(before)
        fmt.space_after = Pt(after)
        fmt.keep_with_next = True
        fmt.line_spacing = _LINE_SPACING[locale]
        ppr = style.element.get_or_add_pPr()
        outline = ppr.find(qn("w:outlineLvl"))
        if outline is None:
            outline = _insert_ordered(ppr, OxmlElement("w:outlineLvl"), _PPR_ORDER)
        outline.set(qn("w:val"), str(level - 1))


def _configure_list_style(document: Document, locale: str) -> None:
    try:
        style = document.styles["List Bullet"]
    except KeyError:
        return
    style.font.name = _font(locale)
    style.font.size = Pt(_TEXT_SIZE)
    _set_rfonts(style.element.get_or_add_rPr(), _font(locale), locale)
    ppr = style.element.get_or_add_pPr()
    contextual = ppr.find(qn("w:contextualSpacing"))
    if contextual is not None:
        ppr.remove(contextual)
    style.paragraph_format.space_after = Pt(3)
    style.paragraph_format.line_spacing = _LINE_SPACING[locale]


def _ensure_derived_styles(document: Document) -> None:
    styles = document.styles
    names = {style.name for style in styles}
    if memo_structure.DERIVED_PARAGRAPH_STYLE not in names:
        paragraph_style = styles.add_style(
            memo_structure.DERIVED_PARAGRAPH_STYLE, WD_STYLE_TYPE.PARAGRAPH
        )
        paragraph_style.base_style = styles["Normal"]
        paragraph_style.hidden = False
        paragraph_style.quick_style = False
    if memo_structure.DERIVED_TABLE_STYLE not in names:
        table_style = styles.add_style(
            memo_structure.DERIVED_TABLE_STYLE, WD_STYLE_TYPE.TABLE
        )
        try:
            table_style.base_style = styles["Normal Table"]
        except KeyError:
            pass


def _set_zh_language_defaults(document: Document) -> None:
    """Tag the Chinese file as zh-CN (python-docx ships eastAsia en-US and a
    Japanese theme-font language), and make Microsoft YaHei the East Asian
    default so a run without explicit fonts never falls back to the
    Japanese theme font. Narrow XML edits, Chinese document only."""
    doc_defaults = document.styles.element.find(qn("w:docDefaults"))
    rpr = None
    if doc_defaults is not None:
        rpr_default = doc_defaults.find(qn("w:rPrDefault"))
        if rpr_default is not None:
            rpr = rpr_default.find(qn("w:rPr"))
    if rpr is not None:
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is not None:
            if rfonts.get(qn("w:eastAsiaTheme")) is not None:
                del rfonts.attrib[qn("w:eastAsiaTheme")]
            rfonts.set(qn("w:eastAsia"), ZH_FONT)
        lang = rpr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            rpr.append(lang)
        lang.set(qn("w:eastAsia"), "zh-CN")
    theme_lang = document.settings.element.find(qn("w:themeFontLang"))
    if theme_lang is not None:
        theme_lang.set(qn("w:eastAsia"), "zh-CN")


def _date_value_for_properties(value: str) -> _dt.datetime:
    parsed = source_tiers.parse_partial_date(value)
    if parsed:
        parts = [int(part) for part in parsed.split("-")] + [1, 1]
        return _dt.datetime(parts[0], parts[1], parts[2], tzinfo=_dt.timezone.utc)
    return _dt.datetime.now(_dt.timezone.utc).replace(microsecond=0)


def _set_core_properties(document: Document, package: dict, locale: str) -> None:
    """File > Info metadata the recipient sees: the memo's title and date,
    the firm as author, the review state, and the renderer version."""
    properties = document.core_properties
    company = _company_name(package, locale)
    properties.title = (
        f"{company} — 投资备忘录" if locale == "zh" else f"{company} — Investment Memo"
    )
    properties.subject = _raw_text(
        (package.get("company") or {}).get("descriptor"), locale
    )[:255] if isinstance(package.get("company"), dict) else ""
    properties.author = FIRM_NAME
    properties.last_modified_by = FIRM_NAME
    properties.category = "Confidential"
    properties.comments = ""
    properties.keywords = ""
    properties.language = "zh-CN" if locale == "zh" else "en-US"
    written = _written_date(package)
    stamp = _date_value_for_properties(written)
    properties.created = stamp
    properties.modified = stamp
    properties.revision = 1
    properties.identifier = RENDERER_VERSION
    properties.version = RENDERER_VERSION
    state = _review_state(package)
    properties.content_status = {
        "approved": "Reviewed",
        "withdrawn": "Withdrawn",
    }.get(state, "Draft")


def _review(package: dict) -> dict:
    run = package.get("run") if isinstance(package.get("run"), dict) else {}
    review = run.get("review")
    return review if isinstance(review, dict) else {}


def _review_state(package: dict) -> str:
    state = str(_review(package).get("state") or "draft").strip().lower()
    return state if state in {"draft", "in_review", "approved", "withdrawn"} else "draft"


def _review_stamp(package: dict, locale: str) -> tuple[str, str]:
    """The header line that says whether anyone has reviewed this memo."""
    state = _review_state(package)
    if state == "approved":
        review = _review(package)
        name = str(review.get("reviewer") or "").strip()
        when = source_tiers.parse_partial_date(
            str(review.get("reviewed_at") or "")[:10]
        ) or ""
        if locale == "zh":
            text = "已审阅" + (f"：{name}" if name else "")
            if when:
                text += f"，{_format_date(when, 'zh')}"
        else:
            text = "Reviewed" + (f" by {name}" if name else "")
            if when:
                text += f", {when}"
        return text, STAMP_APPROVED
    if state == "withdrawn":
        return (
            "已撤回 — 请勿依据本备忘录"
            if locale == "zh"
            else "WITHDRAWN — do not rely on this memo"
        ), STAMP_WITHDRAWN
    return (
        "草稿 — AI 生成，未经审阅"
        if locale == "zh"
        else "DRAFT — AI-generated, not reviewed"
    ), STAMP_DRAFT


def _write_header_footer(document: Document, package: dict, locale: str) -> None:
    section = document.sections[0]
    stamp, stamp_color = _review_stamp(package, locale)
    company_line = (
        f"{_company_name(package, locale)} | BSH 机密投资备忘录"
        if locale == "zh"
        else f"{_company_name(package, locale)} | BSH Confidential Investment Memo"
    )
    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    header.paragraph_format.tab_stops.add_tab_stop(
        Twips(_text_width_twips(document)), WD_TAB_ALIGNMENT.RIGHT
    )
    _add_run(header, stamp, size=8, bold=True, color=stamp_color, locale=locale)
    tab = header.add_run()
    tab.add_tab()
    _style_run(tab, size=8, color=GREY, locale=locale)
    _add_run(header, company_line, size=8, color=GREY, locale=locale)

    first_header = section.first_page_header.paragraphs[0]
    first_header.paragraph_format.space_after = Pt(0)
    _add_run(first_header, stamp, size=8, bold=True, color=stamp_color, locale=locale)

    written = _written_date(package)
    if locale == "zh":
        confidential = " · ".join(
            part for part in ("BSH 机密", _format_date(written, "zh"), "中文") if part
        )
    else:
        confidential = " · ".join(
            part for part in ("BSH Confidential", written, "EN") if part
        )
    for footer in (section.footer, section.first_page_footer):
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        _add_page_x_of_y(paragraph, locale)
        _add_run(paragraph, confidential, size=8, color=GREY, locale=locale)


def _add_page_x_of_y(paragraph: Any, locale: str) -> Any:
    """"Page X of Y · " / "第 X 页，共 Y 页 · " as PAGE and NUMPAGES fields
    with cached results. Word and the Word-made PDF compute the numbers.
    The literal words sit in the same run as the field characters, so a
    viewer that cannot paginate (docx-preview drops field runs) prints no
    stray "Page  of " — only the confidentiality line after it."""
    run = paragraph.add_run()
    element = run._r

    def _text(value: str) -> None:
        node = OxmlElement("w:t")
        node.set(qn("xml:space"), "preserve")
        node.text = value
        element.append(node)

    def _field(instruction: str) -> None:
        element.append(_field_char("begin"))
        element.append(_instr_text(f" {instruction} "))
        element.append(_field_char("separate"))
        _text("1")
        element.append(_field_char("end"))

    if locale == "zh":
        _text("第 ")
        _field("PAGE")
        _text(" 页，共 ")
        _field("NUMPAGES")
        _text(" 页 · ")
    else:
        _text("Page ")
        _field("PAGE")
        _text(" of ")
        _field("NUMPAGES")
        _text(" · ")
    _style_run(run, size=8, color=GREY, locale=locale)
    return run


# ---- cover facts ---------------------------------------------------------------

_RUN_ID_DATE_RE = re.compile(r"^\s*(\d{4}-\d{2}-\d{2})(?!\d)")
_EN_MONTHS = (
    "Jan", "Feb", "Mar", "Apr", "May", "Jun",
    "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
)


def _run_meta(package: dict) -> dict:
    run = package.get("run")
    return run if isinstance(run, dict) else {}


def _written_date(package: dict) -> str:
    """The day the memo was written, as YYYY-MM-DD — never the raw run id
    (ZaiNar's cover once printed "2026-08-31__224828")."""
    run = _run_meta(package)
    for key in ("as_of", "generated_on", "run_date"):
        parsed = source_tiers.parse_partial_date(run.get(key))
        if parsed:
            return parsed
    match = _RUN_ID_DATE_RE.match(str(run.get("run_id") or ""))
    if match and source_tiers.parse_partial_date(match.group(1)):
        return match.group(1)
    return ""


def _evidence_through(package: dict, written: str) -> str:
    """The evidence cutoff: the one the package states (``run.evidence_cutoff``
    or its alias ``run.evidence_ceiling``), unless it is missing or merely
    repeats the run date — then the latest dated non-internal source (R21)."""
    run = _run_meta(package)
    stated = ""
    for value in (
        run.get("evidence_cutoff"),
        run.get("evidence_ceiling"),
    ):
        parsed = source_tiers.parse_partial_date(value)
        if parsed:
            stated = parsed
            break
    if stated and stated != written:
        return stated
    derived = source_tiers.evidence_cutoff(package.get("sources"), written or None)
    return derived or stated


def _format_date(value: str, locale: str, *, human: bool = False) -> str:
    """``2026-08-17`` as the memo prints it: ISO in English cover lines,
    ``17 Aug 2026`` / ``Aug 2026`` in the Sources table, ``2026年8月17日``
    in Chinese."""
    parsed = source_tiers.parse_partial_date(value)
    if not parsed:
        return str(value or "")
    parts = [int(part) for part in parsed.split("-")]
    if locale == "zh":
        text = f"{parts[0]}年"
        if len(parts) > 1:
            text += f"{parts[1]}月"
        if len(parts) > 2:
            text += f"{parts[2]}日"
        return text
    if not human:
        return parsed
    if len(parts) == 1:
        return str(parts[0])
    month = _EN_MONTHS[parts[1] - 1]
    if len(parts) == 2:
        return f"{month} {parts[0]}"
    return f"{parts[2]} {month} {parts[0]}"


def _company(package: dict) -> dict:
    company = package.get("company")
    return company if isinstance(company, dict) else {}


_PUBLIC_STATUS_RE = re.compile(
    r"^\s*(?:public(?:ly)?(?:\s+(?:listed|traded))?|listed)\b", re.IGNORECASE
)


def _is_public_company(company: dict) -> bool:
    if str(company.get("company_type") or "").strip().lower() == "public":
        return True
    if company.get("ticker") or company.get("exchange"):
        return True
    return bool(_PUBLIC_STATUS_RE.match(_raw_text(company.get("status"), "en")))


def _round_value(company: dict, locale: str) -> str:
    """The Round row: the package's own round; "No round on offer" when the
    package says there is none (the key present but empty) on a private
    company; nothing for a public company or a legacy package that never
    stated it — an absent key is unknown, and ZaiNar (no round key, an SPV
    on offer) shows why "absent" must not print as "none"."""
    if "round" in company:
        text = _loc(company.get("round"), locale).strip()
        if text:
            return text
        if not _is_public_company(company):
            return "暂无在售轮次" if locale == "zh" else "No round on offer"
    return ""


def _cover_fact_rows(package: dict, locale: str) -> list[tuple[str, str]]:
    company = _company(package)
    location = company.get("location") or company.get("hq") or company.get("headquarters")
    rows = [
        (_label("Stage", "阶段", locale), _loc(company.get("stage"), locale)),
        (_label("Sector", "行业", locale), _loc(company.get("sector"), locale)),
        (_label("Location", "地点", locale), _loc(location, locale)),
    ]
    if not _is_public_company(company):
        rows.append((_label("Round", "轮次", locale), _round_value(company, locale)))
    return [(label, value) for label, value in rows if str(value or "").strip()]


# ---- page one: masthead, cover facts, contents -------------------------------------


def _derived_paragraph(container: Any) -> Any:
    """A paragraph the docx gates skip (see memo_structure.DERIVED_*)."""
    paragraph = container.add_paragraph()
    try:
        paragraph.style = memo_structure.DERIVED_PARAGRAPH_STYLE
    except KeyError:
        pass
    return paragraph


def _mark_derived_table(table: Any) -> None:
    try:
        table.style = memo_structure.DERIVED_TABLE_STYLE
    except KeyError:
        pass


def _add_derived_title(document: Document, text: str, locale: str, *, size: float = 13) -> None:
    paragraph = _derived_paragraph(document)
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    _add_bottom_border(paragraph, TIFFANY)
    _add_run(paragraph, text, size=size, bold=True, color=NAVY, locale=locale)


def _add_derived_caption(document: Document, text: str, locale: str) -> None:
    paragraph = _derived_paragraph(document)
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    _add_run(paragraph, text, size=9.5, bold=True, color=NAVY, locale=locale)


def _add_front_page(
    document: Document, package: dict, locale: str, plan: list[_PlanEntry]
) -> None:
    """Page one: a two-line masthead, the company, the cover facts ("Written
    … · Evidence through …", the Stage/Sector/Location/Round table) and the
    contents list. The body starts on page two."""
    masthead = _derived_paragraph(document)
    masthead.paragraph_format.space_after = Pt(2)
    _add_run(masthead, "BERKELEY SUMMIT HOUSE", size=9, bold=True, color=TIFFANY, locale=locale)
    _add_run(
        masthead,
        "  ·  机密投资备忘录" if locale == "zh" else "  ·  Confidential Investment Memo",
        size=9,
        color=GREY,
        locale=locale,
    )
    _add_paragraph(
        document,
        _company_name(package, locale),
        size=22,
        bold=True,
        color=NAVY,
        after=2,
        locale=locale,
    )
    descriptor = _loc(_company(package).get("descriptor"), locale)
    if descriptor:
        _add_paragraph(document, descriptor, size=10.5, color=GREY, after=4, locale=locale)
    written = _written_date(package)
    evidence = _evidence_through(package, written)
    source_count = len([s for s in package.get("sources") or [] if isinstance(s, dict)])
    if locale == "zh":
        parts = [
            f"撰写于 {_format_date(written, 'zh')}" if written else "",
            f"证据截至 {_format_date(evidence, 'zh')}" if evidence else "",
            f"{source_count} 个来源" if source_count else "",
        ]
    else:
        parts = [
            f"Written {written}" if written else "",
            f"Evidence through {evidence}" if evidence else "",
            f"{source_count} source{'s' if source_count != 1 else ''}" if source_count else "",
        ]
    meta = " · ".join(part for part in parts if part)
    generated = generated_with_line(package, locale)
    if meta:
        line = _derived_paragraph(document)
        line.paragraph_format.space_after = Pt(2 if generated else 8)
        _add_run(line, meta, size=8.5, color=GREY, locale=locale)
    if generated:
        line = _derived_paragraph(document)
        line.paragraph_format.space_after = Pt(8)
        _add_run(line, generated, size=8.5, color=GREY, locale=locale)
    facts = _cover_fact_rows(package, locale)
    if facts:
        _add_key_value_table(document, facts, locale)
    _add_contents(document, plan, locale)
    document.add_page_break()


_ENGINE_LABELS = {"claude": "Claude", "gemini": "Gemini"}
_QUALITY_LABELS = {
    "best": {"en": "Best quality", "zh": "最佳质量"},
    "balanced": {"en": "Balanced quality", "zh": "均衡质量"},
    "economy": {"en": "Economy quality", "zh": "经济质量"},
}
_MODE_LABELS = {
    "compact": {"en": "compact", "zh": "精简版"},
    "full": {"en": "full", "zh": "完整版"},
}


def _template_label(generated: dict, structure: dict, locale: str) -> str:
    """"Standard memo (v1)" or "Founder's IC template (v2, compact)" — the
    template the run was written on. The run's own stamp wins; an older
    package is read from its structure block (the late v1 profile is the
    standard memo, anything else is the v2 family)."""
    template = str(generated.get("template") or "").strip().lower()
    stage = str(structure.get("stage") or "").strip().lower()
    version = structure.get("version")
    if not template and (stage or version is not None):
        template = "standard" if stage in ("", "late") and str(version) in ("1", "None") else "ic_v2"
    if template == "standard":
        return "标准备忘录（v1）" if locale == "zh" else "Standard memo (v1)"
    if template != "ic_v2":
        return ""
    mode = str((generated.get("structure") or {}).get("mode") or structure.get("mode") or "").strip().lower()
    if not mode and stage.endswith("_compact"):
        mode = "compact"
    mode_label = (_MODE_LABELS.get(mode) or {}).get(locale, "")
    if locale == "zh":
        return f"创始人投委会模板（v2{'，' + mode_label if mode_label else ''}）"
    return f"Founder's IC template (v2{', ' + mode_label if mode_label else ''})"


def generated_with_line(package: dict, locale: str) -> str:
    """Page one's provenance line: which engine and model wrote the memo,
    on which template, at which quality, from which code — so two memos
    on the same company can be told apart at a glance. Reads the run's
    ``generated_with`` stamp; a package without one still names its
    template. Empty when nothing is known."""
    run = package.get("run") if isinstance(package.get("run"), dict) else {}
    generated = run.get("generated_with") if isinstance(run.get("generated_with"), dict) else {}
    structure = package.get("structure") if isinstance(package.get("structure"), dict) else {}
    parts: list[str] = []
    engine = _ENGINE_LABELS.get(str(generated.get("engine") or "").strip().lower(), "")
    models = generated.get("models") if isinstance(generated.get("models"), dict) else {}
    writer = str(generated.get("writer_model") or models.get("SECTION") or models.get("ENGLISH") or "").strip()
    if engine or writer:
        who = f"{engine}（{writer}）" if locale == "zh" and engine and writer else (
            f"{engine} ({writer})" if engine and writer else (engine or writer)
        )
        parts.append(f"生成：{who}" if locale == "zh" else f"Generated with {who}")
    template = _template_label(generated, structure, locale)
    if template:
        parts.append(template)
    quality = (_QUALITY_LABELS.get(str(generated.get("quality") or "").strip().lower()) or {}).get(locale)
    if quality:
        parts.append(quality)
    translator = str(generated.get("translation_model") or models.get("TRANSLATION") or "").strip()
    if translator and locale == "zh":
        parts.append(f"中文翻译：{translator}")
    code = str(generated.get("code_version") or "").strip()
    if code:
        parts.append(f"代码 {code}" if locale == "zh" else f"code {code}")
    return " · ".join(parts)


def _add_contents(document: Document, plan: list[_PlanEntry], locale: str) -> None:
    """The contents list: one internal link per level-1 heading, to the
    ``bsh_sec_<n>`` bookmark on that heading. No page numbers (the web
    viewer cannot paginate) and no Word TOC field (it would render empty
    in the web viewer and make Word ask to update fields on open)."""
    _add_derived_title(document, "目录" if locale == "zh" else "Contents", locale, size=15)
    if not plan:
        return
    table = document.add_table(rows=len(plan), cols=2)
    _mark_derived_table(table)
    # Wide enough for the longest numeral ("XIII." / "十三、") in 10.5pt
    # bold, no wider: the titles sit close to their numbers.
    longest_label = max((_display_len(entry.label.get(locale) or "") for entry in plan), default=0)
    label_twips = max(int(Cm(1.2) / 635), longest_label * 120 + _CELL_PAD_TWIPS)
    total = _text_width_twips(document)
    for row, entry in zip(table.rows, plan):
        label_cell, title_cell = row.cells
        for cell in (label_cell, title_cell):
            _clear_cell_borders(cell)
        label = entry.label.get(locale) or ""
        _cell_text(label_cell, label, locale=locale, bold=True, color=NAVY, size=10.5)
        paragraph = title_cell.paragraphs[0]
        _cell_paragraph_format(paragraph, locale)
        paragraph.paragraph_format.space_after = Pt(4)
        _append_hyperlink(
            paragraph,
            entry.bare.get(locale) or entry.title.get(locale) or "",
            anchor=section_bookmark(entry.number),
            locale=locale,
            size=10.5,
            color=NAVY,
            underline=False,
        )
        label_cell.paragraphs[0].paragraph_format.space_after = Pt(4)
    _apply_table_widths(table, [label_twips, total - label_twips])


# ---- risk cards (read by the risk summary) -----------------------------------------


def _pair(value: Any) -> dict[str, str] | None:
    """A value as {"en", "zh"}: a localized dict keeps both halves, a plain
    language-neutral string serves both; empty -> None."""
    if isinstance(value, dict):
        en = str(value.get("en") or "").strip()
        zh = str(value.get("zh") or "").strip()
        if not en and not zh:
            return None
        return {"en": en or zh, "zh": zh}
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        text = str(value)
        return {"en": text, "zh": text}
    text = str(value or "").strip()
    return {"en": text, "zh": text} if text else None


def _row_cells(row: Any) -> list[Any]:
    cells = row.get("cells") if isinstance(row, dict) else row
    return list(cells) if isinstance(cells, (list, tuple)) else []


_RISK_TITLE_PREFIX_RE = re.compile(r"^\s*(?:risk|风险)\s*\d+\s*[:：]\s*", re.IGNORECASE)
_RISK_TABLE_TITLE_RE = re.compile(r"^\s*(?:risk|风险)\s*(\d+)\s*$", re.IGNORECASE)
_RISK_LIKELIHOOD_TOKEN_RE = re.compile(
    r"^\s*((?:high|medium|low|高|中|低)\s*[:：])", re.IGNORECASE
)


def _risk_cards(section: dict | None) -> list[dict]:
    """The per-risk cards of a risk section: each "Risk N: …" heading block
    followed by its key-value table, with the table's rows by EN label."""
    if not isinstance(section, dict):
        return []
    blocks = [b for b in section.get("blocks") or [] if isinstance(b, dict)]
    cards: list[dict] = []
    for index, block in enumerate(blocks):
        if str(block.get("type") or "paragraph") != "heading":
            continue
        heading = block.get("text") or block.get("title")
        match = _RISK_CARD_HEADING_RE.match(_raw_text(heading, "en"))
        if not match or index + 1 >= len(blocks):
            continue
        table = blocks[index + 1]
        if str(table.get("type") or "") != "table":
            continue
        rows: dict[str, Any] = {}
        for row in table.get("rows") or []:
            cells = _row_cells(row)
            if len(cells) == 2:
                rows[_raw_text(cells[0], "en").strip().lower()] = cells[1]
        cards.append(
            {
                "block": block,
                "number": int(match.group(1)),
                "heading": heading,
                "rows": rows,
            }
        )
    return cards


def _card_row(card: dict, prefix: str) -> Any:
    for label, value in card["rows"].items():
        if label.startswith(prefix):
            return value
    return None


def _card_title(card: dict) -> dict[str, str] | None:
    pair = _pair(card["heading"])
    if not pair:
        return None
    return {
        locale: _RISK_TITLE_PREFIX_RE.sub("", pair[locale], count=1).strip()
        for locale in ("en", "zh")
    }


def _rating_token(value: Any) -> str:
    for locale in ("en", "zh"):
        match = _RISK_RATING_VALUE_RE.match(_raw_text(value, locale))
        if match:
            return re.sub(r"\s+", "", match.group(0))
    return ""


def _likelihood_token(value: Any, locale: str) -> str:
    match = _RISK_LIKELIHOOD_TOKEN_RE.match(_raw_text(value, locale))
    return match.group(1).rstrip(":：").strip() if match else ""


# ---- optional package fields the renderer reads ---------------------------------
#
# Every field below is OPTIONAL; a package without them renders as before.
#
# package["run"]["review"] — the human review state (G2), printed in the page
# headers only, never in the body:
#   {"state": "draft"|"in_review"|"approved"|"withdrawn",
#    "reviewer": str, "reviewed_at": ISO date or datetime}
# package["run"]["checks"] — what the gates found, printed as one provenance
# line at the top of the sources section:
#   {"fact_check": {"verified": int, "found_elsewhere": int, "derived": int,
#                   "not_traced": int, "thin_corpus": bool},
#    "gates": {"quality": str, "chinese_parity": str, "fact_check": str}}
#                   # "passed" | "warnings" | "failed" | "not_run"
# package["run"]["as_of"] / ["generated_on"] / ["run_date"] — the day the memo
# was written (the run id's date part otherwise); ["evidence_cutoff"] (alias
# ["evidence_ceiling"]) — the evidence ceiling.
# package["run"]["private_inventory"] — the firm's private material on this
# company (memo_fact_check.private_inventory); when present, a URL-less
# private-class source must match an item in it.
# package["run"]["source_url_status"] — {"unseen": {id: url}, ...}
# (memo_fact_check.stamp_source_url_status): a source whose URL no retrieval
# ever saw prints its title unlinked while its URL is still that one.
# package["company"]["round"] — present but empty/None means "no round on
# offer" (printed for a private company); an ABSENT key means unknown and the
# row is left out. ["location"] falls back to ["hq"]; ["name"] may be
# {"en","zh"}; ["website"] (or ["domain"]) marks the company's own pages A-.
# package["sources"][i] — besides the legacy "as_of": optional
# "published_at", "data_period" (YYYY | YYYY-MM | YYYY-MM-DD | "undated")
# and "retrieved_at" (filled in Python from the source cache, never by the
# model); see server/source_tiers.py.


def _count(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return None


_GATE_WORDS = {
    "passed": ("passed", "通过"),
    "pass": ("passed", "通过"),
    "ok": ("passed", "通过"),
    "true": ("passed", "通过"),
    "warnings": ("warnings", "有警告"),
    "warning": ("warnings", "有警告"),
    "complete_with_warnings": ("warnings", "有警告"),
    "failed": ("failed", "未通过"),
    "fail": ("failed", "未通过"),
    "false": ("failed", "未通过"),
    "skipped": ("not run", "未执行"),
    "not_run": ("not run", "未执行"),
}


def _provenance_line(package: dict, locale: str) -> str:
    """One line of what the gates found, from the optional
    ``run.checks`` block (fact-check counts and gate results)."""
    checks = _run_meta(package).get("checks")
    if not isinstance(checks, dict):
        return ""
    zh = locale == "zh"
    parts: list[str] = []
    fact = checks.get("fact_check") if isinstance(checks.get("fact_check"), dict) else {}
    counts = {
        key: _count(fact.get(key))
        for key in ("verified", "found_elsewhere", "derived", "not_traced")
    }
    known = {key: value for key, value in counts.items() if value is not None}
    if known:
        total = sum(known.values())
        verified = known.get("verified", 0)
        elsewhere = known.get("found_elsewhere", 0)
        untraced = known.get("not_traced", 0)
        derived = known.get("derived", 0)
        if zh:
            text = (
                f"共 {total} 个数字：{verified} 个已核实 · {elsewhere} 个在其他来源找到"
                f" · {untraced} 个未能溯源"
            )
            if derived:
                text += f" · {derived} 个为推算"
        else:
            text = (
                f"{total} figures: {verified} verified · {elsewhere} found elsewhere"
                f" · {untraced} not traced"
            )
            if derived:
                text += f" · {derived} derived"
        if fact.get("thin_corpus") is True:
            text += "（可供核对的来源较少）" if zh else " (few sources on file to check against)"
        parts.append(text)
    gates = checks.get("gates") if isinstance(checks.get("gates"), dict) else {}
    gate_parts = []
    for key, en_name, zh_name in (
        ("quality", "quality", "质量检查"),
        ("chinese_parity", "Chinese parity", "中文一致性"),
        ("fact_check", "fact check", "事实核查"),
    ):
        if key not in gates:
            continue
        raw = gates.get(key)
        word = _GATE_WORDS.get(str(raw).strip().lower())
        shown = (word[1] if zh else word[0]) if word else str(raw)
        gate_parts.append(f"{zh_name} {shown}" if zh else f"{en_name} {shown}")
    if gate_parts:
        parts.append(("检查：" if zh else "Checks: ") + " · ".join(gate_parts))
    return " · ".join(parts)


# ---- sections and blocks -----------------------------------------------------------


@dataclass
class _RenderContext:
    package: dict
    structure: memo_structure.MemoStructure
    locale: str
    section_number: int = 0
    sub_count: int = 0
    previous_heading_en: str = ""
    # The last heading printed in this locale (the section title, then
    # each sub-heading), as ``_title_key`` reads it: a heading block or a
    # table title that says the same words again is not printed twice.
    last_heading_key: str = ""
    _tiers: dict[str, str] | None = None

    def next_sub_bookmark(self) -> str | None:
        if not self.section_number:
            return None
        self.sub_count += 1
        return subsection_bookmark(self.section_number, self.sub_count)

    def source_tiers(self) -> dict[str, str]:
        if self._tiers is None:
            self._tiers = _source_tier_map(self.package)
        return self._tiers


def _source_tier_map(package: dict) -> dict[str, str]:
    """Source id -> reliability tier, for every source with a web URL."""
    company = _company(package)
    domain = company.get("website") or company.get("domain") or company.get("url")
    tiers: dict[str, str] = {}
    for source in package.get("sources") or []:
        if not isinstance(source, dict):
            continue
        url = str(source.get("url") or "").strip()
        source_id = str(source.get("id") or "").strip()
        if source_id and url.startswith(("http://", "https://")):
            tiers[source_id] = source_tiers.tier_for_url(url, domain)
    return tiers


def _add_plan_entry(document: Document, entry: _PlanEntry, context: _RenderContext) -> None:
    locale = context.locale
    context.section_number = entry.number
    context.sub_count = 0
    context.previous_heading_en = ""
    context.last_heading_key = _title_key(entry.bare.get(locale) or entry.bare.get("en") or "")
    _add_heading(
        document,
        entry.title.get(locale) or entry.title.get("en") or "",
        level=1,
        locale=locale,
        bookmark=section_bookmark(entry.number),
    )
    if entry.kind == "sources":
        _add_sources_section(document, context.package, locale, entry=entry, context=context)
    elif entry.kind == "calculations":
        _add_calculations_section(document, context.package, locale, heading=False)
    elif entry.section is not None:
        _add_section_blocks(document, entry.section, context)
        if entry.section_id == _exec_section_id(context.structure):
            # The package-level glossary reads once, right after the
            # executive summary, before the first section that uses the
            # terms without re-defining them.
            _add_glossary(document, glossary_items(context.package.get("glossary")), locale)


def _exec_section_id(structure: memo_structure.MemoStructure) -> str:
    try:
        return structure.section_for_role("exec").id
    except KeyError:
        return ""


def _add_section(
    document: Document,
    section: dict,
    locale: str,
    section_titles: dict[str, dict[str, str]] | None = None,
) -> None:
    """A section with its (own or structure) title — kept for callers that
    render one section outside a document plan."""
    title = _section_title(section, locale, section_titles)
    if title:
        _add_heading(document, title, level=1, locale=locale)
    context = _RenderContext(
        package={"sections": [section]},
        structure=memo_structure.LATE,
        locale=locale,
    )
    _add_section_blocks(document, section, context)


def _add_section_blocks(document: Document, section: dict, context: _RenderContext) -> None:
    structure = context.structure
    section_id = str(section.get("id") or "")
    try:
        risk_id = structure.section_for_role("risk").id
    except KeyError:
        risk_id = ""
    first_card = None
    cards: list[dict] = []
    if section_id == risk_id and structure.risk_format == "cards":
        try:
            cards = _risk_cards(section)
        except Exception:  # noqa: BLE001 — the summary is optional
            logger.exception("risk summary could not be built")
            cards = []
        if cards:
            first_card = cards[0]["block"]
    blocks = [block for block in section.get("blocks") or [] if isinstance(block, dict)]
    merged = merged_heading_titles(blocks)
    for index, block in enumerate(blocks):
        if first_card is not None and block is first_card:
            _add_risk_summary(document, cards, context)
            first_card = None
        heading_override, suppress_title = merged.get(index, (None, False))
        _add_block(
            document,
            block,
            context.locale,
            context,
            heading_override=heading_override,
            suppress_title=suppress_title,
        )


_TITLE_KEY_RE = re.compile(r"[^0-9a-z\u3400-\u4dbf\u4e00-\u9fff]+")


def _title_key(text: str) -> str:
    """A heading or table title for duplicate detection: lower case,
    citations and punctuation dropped, one space between words."""
    text = _CITATION_RE.sub(" ", str(text or ""))
    return _TITLE_KEY_RE.sub(" ", text.lower()).strip()


def merged_heading_titles(blocks: list[dict]) -> dict[int, tuple[Any, bool]]:
    """Where a subsection heading is immediately followed by a table whose
    title is the same words (case / punctuation aside) or a longer form of
    them ("Revenue Picture", then "Revenue picture — what is disclosed and
    what it will carry"), render ONE heading — the longer one — and no
    table title. Returns block index -> (heading text override or None,
    suppress this table's title)."""
    out: dict[int, tuple[Any, bool]] = {}
    for index, block in enumerate(blocks[:-1]):
        if str(block.get("type") or "paragraph") != "heading":
            continue
        nxt = blocks[index + 1]
        if str(nxt.get("type") or "paragraph") != "table" or not nxt.get("title"):
            continue
        heading_value = block.get("text") or block.get("title")
        heading_key = _title_key(_raw_text(heading_value, "en") if isinstance(heading_value, dict) else heading_value)
        title_key = _title_key(_raw_text(nxt["title"], "en") if isinstance(nxt["title"], dict) else nxt["title"])
        if not heading_key or not title_key:
            continue
        shorter, longer = sorted((heading_key, title_key), key=len)
        same = heading_key == title_key
        extends = (
            len(shorter.split()) >= 2
            and longer != shorter
            and longer.startswith(shorter + " ")
        )
        if not (same or extends):
            continue
        override = nxt["title"] if (not same and longer == title_key) else None
        out[index] = (override, False)
        out[index + 1] = (None, True)
    return out


def _is_numbered_section_heading(text: str) -> bool:
    """True when a heading block restates a top-level numbered section title.

    Numbered section titles are rendered automatically by :func:`_add_section`,
    so a heading block carrying such a prefix is a duplicate and must be
    dropped. See :data:`_NUMBERED_SECTION_HEADING_RE`.
    """
    return bool(text and _NUMBERED_SECTION_HEADING_RE.match(text))


def _add_block(
    document: Document,
    block: dict,
    locale: str,
    context: _RenderContext | None = None,
    *,
    heading_override: Any = None,
    suppress_title: bool = False,
) -> None:
    kind = str(block.get("type") or "paragraph")
    if kind == "heading":
        value = heading_override if heading_override is not None else (block.get("text") or block.get("title"))
        en_text = _raw_text(value, "en") if isinstance(value, dict) else str(value or "")
        if _is_numbered_section_heading(en_text):
            # Redundant restatement of the renderer's own section title — skip
            # in both languages (decided on the English, so EN and ZH match).
            return
        text = _loc(value, locale)
        # A roman or Chinese section numeral the model put on a sub-heading
        # is not the document's numbering; the renderer owns that.
        text = _NUMBERED_SECTION_HEADING_RE.sub("", text, count=1).strip() or text
        if context is not None:
            context.previous_heading_en = en_text
            key = _title_key(text)
            if key and key == context.last_heading_key:
                # The same words as the heading just printed (the section
                # title restated, or a sub-heading the translation
                # collapsed into the previous one): print once.
                return
            context.last_heading_key = key
        try:
            level = int(block.get("level") or 2)
        except (TypeError, ValueError):
            level = 2
        # Level 1 belongs to the renderer's own section titles.
        level = max(2, min(3, level))
        bookmark = context.next_sub_bookmark() if context and level == 2 else None
        _add_heading(document, text, level=level, locale=locale, bookmark=bookmark)
    elif kind == "bullets":
        bold_lead = str(block.get("component") or "") in _BOLD_LEAD_COMPONENTS
        for item in block.get("items") or []:
            _add_bullet(
                document, _loc(item, locale), locale=locale, bold_lead=bold_lead
            )
    elif kind == "table":
        title = None if suppress_title else block.get("title")
        if title and not _is_repeated_risk_title(title, context):
            text = _loc(title, locale)
            key = _title_key(text)
            if context is not None and key and key == context.last_heading_key:
                text = ""
            elif context is not None:
                context.last_heading_key = key
            if text:
                _add_heading(document, text, level=3, locale=locale)
        _add_table(document, block, locale, context)
    elif kind == "chart":
        _add_chart(document, block, locale)
    elif kind == "callout":
        _add_callout(document, block, locale)
    elif kind == "glossary":
        _add_glossary(document, glossary_items(block), locale)
    elif kind == "spacer":
        document.add_paragraph()
    else:
        text = _loc(block.get("text") or block.get("body") or "", locale)
        if text:
            _add_paragraph(document, text, locale=locale)


def _add_glossary(document: Document, items: list[dict], locale: str) -> None:
    """The "Terms used" block: a small heading and a two-column term /
    definition table. Nothing when there are no items."""
    rows = [
        (_loc(item.get("term"), locale), _loc(item.get("definition"), locale))
        for item in items
    ]
    rows = [(term, definition) for term, definition in rows if term and definition]
    if not rows:
        return
    _add_heading(document, GLOSSARY_TITLE[locale], level=3, locale=locale)
    _add_key_value_table(document, rows, locale)


def _is_repeated_risk_title(title: Any, context: _RenderContext | None) -> bool:
    """A risk card's table titled "Risk 3" right under its own "Risk 3: …"
    heading says the same thing twice; decided on the English so both
    documents drop it together."""
    if context is None:
        return False
    match = _RISK_TABLE_TITLE_RE.match(_raw_text(title, "en") if isinstance(title, dict) else str(title or ""))
    if not match:
        return False
    heading = _RISK_CARD_HEADING_RE.match(context.previous_heading_en or "")
    return bool(heading and heading.group(1) == match.group(1))


def _add_risk_summary(document: Document, cards: list[dict], context: _RenderContext) -> None:
    """A one-glance table of every risk card (# | risk | area | likelihood |
    rating), built from the cards themselves in both languages, so the EN
    and ZH table counts stay equal."""
    locale = context.locale
    zh = locale == "zh"
    rows = []
    for card in cards:
        title = _card_title(card) or {"en": "", "zh": ""}
        area = _card_row(card, "risk type")
        likelihood = _card_row(card, "likelihood")
        rows.append(
            [
                str(card["number"]),
                title.get(locale) or title.get("en") or "",
                _loc(area, locale) if area is not None else "—",
                _likelihood_token(likelihood, locale) or "—",
                _rating_token(_card_row(card, "risk rating")) or "—",
            ]
        )
    if not rows:
        return
    _add_derived_caption(document, "风险一览" if zh else "Risk summary", locale)
    _add_grid_table(
        document,
        ["#", "风险", "领域", "可能性", "评分"] if zh else ["#", "Risk", "Area", "Likelihood", "Rating"],
        rows,
        locale,
        derived=True,
        severity_col=4,
        likelihood_col=3,
    )


def _sources_restatement(block: dict, structure: memo_structure.MemoStructure) -> bool:
    """A heading inside a folded sources section that repeats the sources
    section's own title ("Sources", "VI. Sources, …")."""
    if str(block.get("type") or "") != "heading":
        return False
    value = block.get("text") or block.get("title")
    patterns = structure.parity_patterns().get("sources") or {}
    for locale in ("en", "zh"):
        text = _raw_text(value, locale).strip() if isinstance(value, dict) else str(value or "").strip()
        if not text:
            continue
        if _is_numbered_section_heading(text):
            return True
        if memo_structure.heading_matches(patterns.get(locale), text):
            return True
    return False


_SOURCE_CLASS_LABELS: dict[str, tuple[str, str]] = {
    "company reported": ("Company-reported", "公司自述"),
    "company announcement": ("Company announcement", "公司公告"),
    "company first party announcement": ("Company announcement", "公司公告"),
    "company disclosure": ("Company disclosure", "公司披露"),
    "company materials": ("Company materials", "公司材料"),
    "company website": ("Company website", "公司官网"),
    "investor materials": ("Investor materials", "投资者材料"),
    "public filings": ("Public filings", "公开申报文件"),
    "public filing": ("Public filing", "公开申报文件"),
    "regulatory filing": ("Regulatory filing", "监管申报文件"),
    "regulatory filings": ("Regulatory filings", "监管申报文件"),
    "court record": ("Court record", "法院记录"),
    "court records": ("Court records", "法院记录"),
    "press reporting": ("Press reporting", "媒体报道"),
    "business press": ("Business press", "财经媒体"),
    "third party press": ("Third-party press", "第三方媒体报道"),
    "independent secondary": ("Independent secondary reporting", "独立第三方报道"),
    "party affiliated secondary": ("Party-affiliated secondary source", "关联方二手资料"),
    "third party attestation": ("Third-party attestation", "第三方证言"),
    "third party market data": ("Third-party market data", "第三方市场数据"),
    "market data": ("Market data", "市场数据"),
    "third party analyst estimate": ("Third-party analyst estimate", "第三方分析师估算"),
    "analyst research": ("Analyst research", "分析师研究"),
    "industry research": ("Industry research", "行业研究"),
    "syndicated market research": ("Syndicated market research", "第三方行业研究报告"),
    "data vendor": ("Data vendor", "数据服务商"),
    "aggregator": ("Aggregator write-up", "第三方汇总文章"),
    "bsh primary diligence": ("BSH primary diligence", "BSH 一手尽调"),
    "bsh diligence": ("BSH diligence", "BSH 尽调"),
    "internal model": ("Internal model", "内部模型"),
    "bsh internal model": ("BSH internal model", "BSH 内部模型"),
    "internal record": ("Internal record", "内部记录"),
    "bsh internal record": ("BSH internal record", "BSH 内部记录"),
    "company registry": ("Company registry", "公司登记信息"),
    "founder update": ("Founder update", "创始人更新"),
    "founder updates": ("Founder updates", "创始人更新"),
    "management interview": ("Management interview", "管理层访谈"),
    "reference call": ("Reference call", "背景调查访谈"),
    "reference calls": ("Reference calls", "背景调查访谈"),
    "expert call": ("Expert call", "专家访谈"),
    "data room": ("Data room", "资料室文件"),
    # Classes the memo prompts prescribe (the sources contract and the
    # passes' source-hunt vocabulary), as live packages wrote them
    # (2026-09-23 ZaiNar runs). Each gets ONE Chinese label; the Chinese
    # memo never falls back to a generic class for these.
    "unverified registry value no document on file": (
        "Unverified registry value (no document on file)",
        "未经核实的登记值（无文件佐证）",
    ),
    "unverified registry value": (
        "Unverified registry value",
        "未经核实的登记值（无文件佐证）",
    ),
    "secondary press report original article not retrieved": (
        "Secondary press report, original article not retrieved",
        "二手媒体报道（未获取原文）",
    ),
    "secondary press report": ("Secondary press report", "二手媒体报道（未获取原文）"),
    "government program page": ("Government program page", "政府项目页面"),
    "government page": ("Government page", "政府项目页面"),
    "company reported press release": ("Company-reported press release", "公司新闻稿"),
    "company press release": ("Company press release", "公司新闻稿"),
    "press release": ("Press release", "公司新闻稿"),
    "independent trade press": ("Independent trade press", "独立行业媒体"),
    "trade press": ("Trade press", "独立行业媒体"),
    "tier 1 press": ("Tier-1 press", "一线媒体报道"),
    "tier 1 news": ("Tier-1 news", "一线媒体报道"),
    "third party company profile": ("Third-party company profile", "第三方公司资料"),
    "company profile": ("Company profile", "第三方公司资料"),
    "competitor company materials": ("Competitor company materials", "竞争对手公司材料"),
    "competitor materials": ("Competitor materials", "竞争对手公司材料"),
    "third party private market data": ("Third-party private-market data", "第三方私募市场数据"),
    "private market data": ("Private-market data", "第三方私募市场数据"),
    "third party market research": ("Third-party market research", "第三方市场研究"),
    "market research": ("Market research", "第三方市场研究"),
    "public patent office record": ("Public patent-office record", "公开专利记录"),
    "patent office record": ("Patent-office record", "公开专利记录"),
    "patent record": ("Patent record", "公开专利记录"),
    "crowdsourced org chart aggregator": ("Crowdsourced org-chart aggregator", "众包组织架构汇总站"),
    "org chart aggregator": ("Org-chart aggregator", "众包组织架构汇总站"),
    "investor intermediary": ("Investor/intermediary", "投资方/中介资料"),
    "investor or intermediary": ("Investor/intermediary", "投资方/中介资料"),
    "independent secondary research": ("Independent secondary research", "独立第三方研究"),
    "independent secondary press": ("Independent secondary press", "独立第三方报道"),
    "low reliability aggregator commentary": ("Low-reliability aggregator commentary", "低可靠性汇总评论"),
    "aggregator commentary": ("Aggregator commentary", "低可靠性汇总评论"),
    "company reported deck": ("Company-reported deck", "公司自述材料（路演材料）"),
    "bsh reference call": ("BSH reference call", "BSH 背景调查访谈"),
}


def _source_class_key(raw: str) -> str:
    """A source class as the label map keys it: lower case, punctuation
    and hyphens/underscores/slashes as spaces, one space between words —
    so "Third-Party Market Data", "third_party market data" and
    "third-party market data" are one class."""
    key = re.sub(r"[\s_\-/,;:()\[\]{}\"'“”‘’.]+", " ", str(raw or "").lower())
    return re.sub(r"\s+", " ", key).strip()


def _known_source_class(raw: str) -> tuple[str, str] | None:
    """The canonical (en, zh) labels for a class the map knows, matched on
    the English text: exactly, or by the longest known class the text
    starts with ("third-party market data (aggregator)" is third-party
    market data with a qualifier)."""
    key = _source_class_key(raw)
    if not key:
        return None
    known = _SOURCE_CLASS_LABELS.get(key)
    if known:
        return known
    best: tuple[str, str] | None = None
    best_len = 0
    for candidate, labels in _SOURCE_CLASS_LABELS.items():
        if len(candidate) > best_len and (key + " ").startswith(candidate + " "):
            best, best_len = labels, len(candidate)
    return best


def canonical_source_class_zh(raw: str) -> str | None:
    """The one Chinese label for a prompt-prescribed source class written
    in English, or None when the class is not one the renderer knows."""
    known = _known_source_class(raw)
    return known[1] if known else None


# "registry" alone is a public register (a patent office, a companies
# house), never BSH's own: ZaiNar 2026-09-23 labelled an entrepreneurloop
# blog post classed "Patent office registry" as BSH 内部资料. Our registry's
# classes ("unverified registry value") are in the map above.
_SOURCE_CLASS_ZH_FALLBACKS = (
    (re.compile(r"\b(?:bsh|internal|diligence)\b", re.I), "BSH 内部资料"),
    (re.compile(r"\b(?:filing|regulat|court|exchange|government|registry|register|patent)", re.I), "公开申报与法律文件"),
    (re.compile(r"\b(?:company|first[- ]party|issuer|corporate|announcement)\b", re.I), "公司披露"),
    (re.compile(r"\b(?:aggregator|blog|commentary|forum|social)\b", re.I), "第三方汇总与评论"),
    # Research before media: "analyst report", "industry report" and
    # "academic journal article" are research, not press coverage.
    (re.compile(r"\b(?:research|analyst|academic|industry report|estimate|market data|data)\b", re.I), "研究与市场数据"),
    (re.compile(r"\b(?:press|news|media|journal(?:ism)?|report(?:ing|er)?)\b", re.I), "媒体报道"),
)


def _source_class_label(source: dict, locale: str) -> str:
    """A source's class as a reader-facing label in the memo's language:
    a bilingual class as written; a known slug from the map; an unknown
    slug as readable English, and in Chinese a generic class — never a raw
    slug in the Chinese memo."""
    value = source.get("class") if source.get("class") is not None else source.get("type")
    raw = _raw_text(value, "en").strip() if isinstance(value, dict) else str(value or "").strip()
    known = _known_source_class(raw) if raw else None
    if isinstance(value, dict) and _raw_text(value, locale).strip():
        # A bilingual class: the Chinese half is the model's translation,
        # and for a class the map knows the canonical label wins over it
        # (one class, one label, every memo — the ZH run of 2026-09-23
        # rendered the same class three ways).
        if locale == "zh" and known:
            return known[1]
        return _raw_text(value, locale).strip()
    if not raw:
        return ""
    if known:
        if locale == "zh":
            return known[1]
        return known[0] if _source_class_key(raw) in _SOURCE_CLASS_LABELS else raw[:1].upper() + raw[1:]
    if locale == "zh":
        if _CJK_CHAR_RE.search(raw):
            return raw
        for pattern, label in _SOURCE_CLASS_ZH_FALLBACKS:
            if pattern.search(raw.replace("_", " ").replace("-", " ")):
                return label
        return "其他来源"
    if _SOURCE_CLASS_SLUG_RE.match(raw):
        return " ".join(word.capitalize() for word in re.split(r"[_\-]+", raw))
    return raw[:1].upper() + raw[1:]


def _source_dates_text(source: dict, run_date: str, locale: str) -> str:
    """One stacked date cell: "Aug 2026 · covers Jul 2026 · retrieved
    1 Sep 2026" / "2026年8月发布 · 覆盖2026年7月 · 2026年9月1日检索"."""
    normalized = source_tiers.normalize_source_dates(source, run_date)
    published = normalized.get("published_at")
    period = normalized.get("data_period")
    retrieved = normalized.get("retrieved_at")
    parts: list[str] = []
    if locale == "zh":
        if published == source_tiers.UNDATED:
            parts.append("未注明日期")
        elif published:
            parsed = source_tiers.parse_partial_date(published)
            parts.append(f"{_format_date(parsed, 'zh')}发布" if parsed else str(published))
        if period and period != source_tiers.UNDATED:
            parsed = source_tiers.parse_partial_date(period)
            parts.append(f"覆盖{_format_date(parsed, 'zh')}" if parsed else f"覆盖{period}")
        if retrieved:
            parts.append(f"{_format_date(retrieved, 'zh')}检索")
    else:
        if published == source_tiers.UNDATED:
            parts.append("Undated")
        elif published:
            parsed = source_tiers.parse_partial_date(published)
            parts.append(_format_date(parsed, "en", human=True) if parsed else str(published))
        if period and period != source_tiers.UNDATED:
            parsed = source_tiers.parse_partial_date(period)
            parts.append(
                f"covers {_format_date(parsed, 'en', human=True)}" if parsed else f"covers {period}"
            )
        if retrieved:
            parts.append(f"retrieved {_format_date(retrieved, 'en', human=True)}")
    return " · ".join(parts)


def _add_sources_section(
    document: Document,
    package: dict,
    locale: str,
    section_titles: dict[str, dict[str, str]] | None = None,
    *,
    entry: _PlanEntry | None = None,
    context: _RenderContext | None = None,
) -> None:
    """The one sources section: any sources-like prose and tables the model
    wrote (class weighting, disclosures), then the renderer's source list —
    ID, linked title with its domain, class, reliability tier, treatment and
    one stacked date cell — and the tier legend."""
    if entry is None:
        _add_heading(
            document,
            _section_title({"id": "sources"}, locale, section_titles),
            level=1,
            locale=locale,
        )
    provenance = _provenance_line(package, locale)
    if provenance:
        # What the gates found (the optional ``run.checks``), where a
        # reader checking the sources looks for it. Renderer text: the
        # gates skip it.
        line = _derived_paragraph(document)
        line.paragraph_format.space_after = Pt(6)
        _add_run(line, provenance, size=8.5, color=GREY, locale=locale)
    structure = context.structure if context else _structure_for(package)
    folded = entry.folded if entry is not None else []
    for section in folded:
        token = _RENDER_SECTION.set(str(section.get("id") or "sources"))
        try:
            for block in section.get("blocks") or []:
                if isinstance(block, dict) and not _sources_restatement(block, structure):
                    _add_block(document, block, locale, context)
        finally:
            _RENDER_SECTION.reset(token)
    if folded:
        _add_heading(
            document,
            "来源清单" if locale == "zh" else "Source list",
            level=2,
            locale=locale,
            bookmark=context.next_sub_bookmark() if context else None,
        )
    sources = [s for s in package.get("sources") or [] if isinstance(s, dict)]
    if not sources:
        return
    company = _company(package)
    company_domain = company.get("website") or company.get("domain") or company.get("url")
    written = _written_date(package)
    unseen = _unseen_source_urls(package)
    headers = (
        ["编号", "来源", "类别", "等级", "处理方式", "日期"]
        if locale == "zh"
        else ["ID", "Source", "Class", "Tier", "Treatment", "Dates"]
    )
    rows: list[list[str]] = []
    for source in sources:
        url = str(source.get("url") or "").strip()
        linked = url.startswith(("http://", "https://"))
        title = _loc(source.get("title") or source.get("id"), locale, rewrite=False)
        domain = source_tiers.display_domain(url) if linked else ""
        if linked and unseen.get(str(source.get("id") or "").strip()) == url:
            domain = " · ".join(part for part in (domain, _UNSEEN_LINK_NOTE[locale]) if part)
        rows.append(
            [
                str(source.get("id") or "").strip(),
                f"{title}\n{domain}" if domain else title,
                _source_class_label(source, locale),
                source_tiers.tier_for_url(url, company_domain) if linked else "—",
                _loc(source.get("treatment"), locale),
                _source_dates_text(source, written, locale),
            ]
        )
    total = _text_width_twips(document)
    # ID and tier are a few characters; the stacked date cell reads best at
    # a fixed width rather than squeezed by the prose columns around it.
    fixed = {
        0: int(Cm(1.25) / 635),
        3: int(Cm(1.15) / 635),
        5: int(Cm(2.6 if locale == "zh" else 2.4) / 635),
    }
    widths = _plan_column_widths(headers, rows, total, fixed=fixed)
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _fill_header_row(table.rows[0], headers, locale)
    for index, (source, values) in enumerate(zip(sources, rows), start=1):
        cells = table.rows[index].cells
        for col, value in enumerate(values):
            cell = cells[col]
            if index % 2 == 1:
                _shade_cell(cell, "FAFAFA")
            _set_cell_borders(cell)
            if col == 1:
                continue
            _cell_text(cell, value, locale=locale, align="center" if col == 3 else None)
        source_id = values[0]
        if source_id:
            _add_bookmark(cells[0].paragraphs[0], citation_anchor(source_id))
        # Title cell: the linked title, then the domain in small grey type so
        # a printed or PDF copy stays traceable. A URL no retrieval ever saw
        # (run.source_url_status) prints unlinked and says so — while the
        # source still carries that same URL.
        title_cell = cells[1]
        paragraph = _reset_cell(title_cell)
        _cell_paragraph_format(paragraph, locale)
        url = str(source.get("url") or "").strip()
        title = _loc(source.get("title") or source_id, locale, rewrite=False)
        if url.startswith(("http://", "https://")):
            not_seen = bool(source_id) and unseen.get(source_id) == url
            if not_seen:
                _add_run(paragraph, title, size=_TABLE_SIZE, locale=locale)
            else:
                _append_hyperlink(paragraph, title, url=url, locale=locale, size=_TABLE_SIZE)
            domain = source_tiers.display_domain(url)
            if domain or not_seen:
                line = title_cell.add_paragraph()
                _cell_paragraph_format(line, locale)
                note = _UNSEEN_LINK_NOTE[locale] if not_seen else ""
                _add_run(
                    line,
                    " · ".join(part for part in (domain, note) if part),
                    size=8,
                    color=GREY,
                    locale=locale,
                )
        else:
            _add_run(paragraph, title, size=_TABLE_SIZE, locale=locale)
    _apply_table_widths(table, widths)
    _set_row_flags(table.rows[0], header=True)
    for row, values in zip(list(table.rows)[1:], rows):
        if _row_is_short(values, widths):
            _set_row_flags(row, cant_split=True)
    legend = _derived_paragraph(document)
    legend.paragraph_format.space_before = Pt(4)
    _add_run(legend, source_tiers.TIER_LEGEND[locale], size=8, color=GREY, locale=locale)


_UNSEEN_LINK_NOTE = {"en": "link not verified", "zh": "链接未经核实"}


def _unseen_source_urls(package: dict) -> dict[str, str]:
    """``run.source_url_status["unseen"]`` — source id -> the URL no
    retrieval of this run ever saw (memo_fact_check.stamp_source_url_status);
    empty when the run carries no status."""
    status = _run_meta(package).get("source_url_status")
    unseen = status.get("unseen") if isinstance(status, dict) else None
    if not isinstance(unseen, dict):
        return {}
    return {
        str(key).strip(): str(value).strip()
        for key, value in unseen.items()
        if str(key).strip() and str(value or "").strip()
    }


CALCULATIONS_TITLE = {"en": "Calculation notes", "zh": "计算说明"}


# What an input's ``ref`` says when it is not a source or calculation id.
_CALCULATION_REF_WORDS = {
    "assumption": {"en": "assumption", "zh": "假设"},
    "assumed": {"en": "assumption", "zh": "假设"},
    "estimate": {"en": "estimate", "zh": "估算"},
    "our estimate": {"en": "our estimate", "zh": "自行估算"},
    "derived": {"en": "derived", "zh": "推算"},
    "calculated": {"en": "derived", "zh": "推算"},
    # BSH's own check and post-money, from the deal record (memo_returns'
    # "BSH proceeds by case" note).
    "deal terms": {"en": "deal terms", "zh": "交易条款"},
}


def _calculation_inputs_text(calculation: dict, locale: str = "en") -> str:
    """"name = value [S3]; …" for the Inputs cell. A ``ref`` that names
    sources or calculation notes becomes their citation link; any other
    ref ("assumption") reads as a plain note in the memo's language, never
    as a bracket token. ``name`` may be bilingual (``{"en","zh"}``)."""
    parts = []
    for item in calculation.get("inputs") or []:
        if not isinstance(item, dict):
            continue
        name = _loc(item.get("name"), locale)
        value = _loc(item.get("value"), locale)
        ref = str(item.get("ref") or "").strip().strip("[]").strip()
        ref_note = ""
        if ref:
            ids = citation_ids(f"[{ref}]")
            if ids:
                ref_note = f" [{', '.join(ids)}]"
            else:
                word = _CALCULATION_REF_WORDS.get(ref.lower())
                shown = word[locale] if word else ref
                ref_note = f"（{shown}）" if locale == "zh" else f" ({shown})"
        parts.append(f"{name} = {value}{ref_note}")
    return "; ".join(parts)


def _add_calculations_section(
    document: Document, package: dict, locale: str, *, heading: bool = True
) -> None:
    """The appendix behind every `[C#]` link: one bookmarked row per
    pinned calculation note — what was computed, from which inputs, the
    arithmetic, the result, and what it means."""
    calculations = [
        c for c in package.get("calculations") or [] if isinstance(c, dict)
    ]
    if not calculations:
        return
    if heading:
        _add_heading(document, CALCULATIONS_TITLE[locale], level=1, locale=locale)
    headers = (
        ["ID", "What", "Inputs", "Formula", "Result", "Meaning"]
        if locale == "en"
        else ["编号", "计算内容", "输入", "公式", "结果", "含义"]
    )
    rows = [
        [
            str(calc.get("id") or ""),
            _loc(calc.get("label"), locale),
            _calculation_inputs_text(calc, locale),
            _loc(calc.get("formula"), locale),
            _loc(calc.get("result"), locale),
            _loc(calc.get("meaning"), locale),
        ]
        for calc in calculations
    ]
    table = _add_grid_table(document, headers, rows, locale, fixed={0: int(Cm(1.25) / 635)})
    for index, calc in enumerate(calculations):
        calc_id = str(calc.get("id") or "").strip()
        if calc_id:
            _add_bookmark(
                table.rows[index + 1].cells[0].paragraphs[0],
                citation_anchor(calc_id),
            )


# ---- headings, paragraphs, bullets ---------------------------------------------------


def _add_heading(
    document: Document,
    text: str,
    *,
    level: int,
    locale: str,
    bookmark: str | None = None,
) -> Any:
    """A real Word heading (Heading 1/2/3 style, outline level) with the
    renderer's look also set on the run, so Word, the Word-made PDF and the
    web viewer agree. No upper-casing: the heading reads as written."""
    if not text:
        return None
    level = max(1, min(3, int(level)))
    try:
        paragraph = document.add_paragraph(style=f"Heading {level}")
    except KeyError:
        paragraph = document.add_paragraph()
    size, color, before, after = _HEADING_LOOK[level]
    paragraph.paragraph_format.keep_with_next = True
    # Level-2 subsection headings get real air above them so the fixed
    # numbered subsections read as visual units, not a wall of text.
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    if level == 1:
        _add_bottom_border(paragraph, TIFFANY)
    _add_run(paragraph, text, bold=True, size=size, color=color, locale=locale)
    if bookmark:
        _add_bookmark(paragraph, bookmark)
    return paragraph


def _add_paragraph(
    document_or_cell: Any,
    text: str,
    *,
    locale: str = "en",
    size: float = _TEXT_SIZE,
    bold: bool = False,
    color: str = BLACK,
    align: str = "left",
    before: float = 0,
    after: float = 5,
) -> Any:
    paragraph = document_or_cell.add_paragraph()
    paragraph.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = _LINE_SPACING.get(locale, 1.15)
    _add_run(paragraph, text, size=size, bold=bold, color=color, locale=locale)
    return paragraph


def _spacer(document: Document, after: float = 4) -> None:
    document.add_paragraph().paragraph_format.space_after = Pt(after)


# Bullet blocks whose items open with a verdict headline (the executive
# summary's investment highlights and key risks): the lead sentence is
# rendered bold, the way the fund's own LP deck sets its highlights.
_BOLD_LEAD_COMPONENTS = frozenset({"investment_highlights", "key_risks"})
_LEAD_SENTENCE_RE = re.compile(r"^(.+?(?:[.!?](?=\s)|[。！？]))(.*)$", re.DOTALL)
# The docx has no markdown. Writers told a headline "is rendered bold"
# sometimes wrap it in "**...**" themselves (live compact run,
# 2026-09-14: the asterisks printed literally in both languages). A
# marked span renders as a bold run; a stray marker is dropped.
_MD_BOLD_RE = re.compile(r"\*\*(.+?)\*\*", re.DOTALL)
_MD_BOLD_LEAD_RE = re.compile(r"^\s*\*\*(.+?)\*\*(.*)$", re.DOTALL)


def strip_markdown_bold(text: str) -> str:
    """Drop markdown bold markers ("**Acme leads.**" -> "Acme leads.")."""
    if "**" not in text:
        return text
    return _MD_BOLD_RE.sub(r"\1", text).replace("**", "")


def split_lead_sentence(text: str) -> tuple[str, str]:
    """Split ``text`` into (lead sentence, remainder). The lead ends at
    the first sentence-final mark followed by whitespace (so "$1.5T" and
    "vs." inside a sentence do not split it) or at a CJK full stop. A
    single sentence with no remainder is entirely the lead. A lead the
    writer wrapped in markdown bold is taken as written, markers dropped."""
    marked = _MD_BOLD_LEAD_RE.match(text)
    if marked:
        lead = marked.group(1).strip()
        rest = strip_markdown_bold(marked.group(2))
        if rest[:1].isalnum() and lead.endswith((".", "!", "?")):
            rest = f" {rest}"
        return lead, rest
    text = strip_markdown_bold(text)
    match = _LEAD_SENTENCE_RE.match(text)
    if match:
        return match.group(1), match.group(2)
    return text, ""


# A real bullet: Word numbering with a "•" in Arial and a hanging indent,
# so wrapped lines align under the text, not under a typed character.
_BULLET_NUM_ID = "7301"


def _bullet_num_id(part: Any) -> str | None:
    try:
        numbering = part.numbering_part.element
    except Exception:  # noqa: BLE001 — no numbering part: plain paragraph
        return None
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) == _BULLET_NUM_ID:
            return _BULLET_NUM_ID
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), _BULLET_NUM_ID)
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (
        ("w:start", "1"),
        ("w:numFmt", "bullet"),
        ("w:lvlText", "•"),
        ("w:lvlJc", "left"),
    ):
        node = OxmlElement(tag)
        node.set(qn("w:val"), value)
        level.append(node)
    ppr = OxmlElement("w:pPr")
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "340")
    indent.set(qn("w:hanging"), "227")
    ppr.append(indent)
    level.append(ppr)
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), EN_FONT)
    fonts.set(qn("w:hint"), "default")
    rpr.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TIFFANY)
    rpr.append(color)
    level.append(rpr)
    abstract.append(level)
    nums = numbering.findall(qn("w:num"))
    if nums:
        nums[0].addprevious(abstract)
    else:
        numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), _BULLET_NUM_ID)
    abstract_id = OxmlElement("w:abstractNumId")
    abstract_id.set(qn("w:val"), _BULLET_NUM_ID)
    num.append(abstract_id)
    numbering.append(num)
    return _BULLET_NUM_ID


def _make_bullet(paragraph: Any) -> None:
    num_id = _bullet_num_id(paragraph.part)
    if num_id is None:
        paragraph.paragraph_format.left_indent = Cm(0.6)
        paragraph.paragraph_format.first_line_indent = Cm(-0.4)
        return
    numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    numbering.get_or_add_ilvl().val = 0
    numbering.get_or_add_numId().val = int(num_id)


def _add_bullet(
    document: Document, text: str, *, locale: str, bold_lead: bool = False
) -> Any:
    try:
        paragraph = document.add_paragraph(style="List Bullet")
    except KeyError:
        paragraph = document.add_paragraph()
    _make_bullet(paragraph)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = _LINE_SPACING.get(locale, 1.15)
    if bold_lead:
        lead, rest = split_lead_sentence(str(text or ""))
        _add_run(paragraph, lead, bold=True, locale=locale)
        if rest:
            _add_run(paragraph, rest, locale=locale)
        return paragraph
    _add_run(paragraph, str(text or ""), locale=locale)
    return paragraph


def _add_chart(document: Document, block: dict, locale: str) -> None:
    """Render a ``chart`` block: localized heading, one shared PNG (all
    image-internal text is English/neutral), localized caption. When the
    chart cannot be drawn (matplotlib missing or a render fault), fall
    back to the series as a small table so the data always ships."""
    title = _loc(block.get("title"), locale)
    if title:
        _add_heading(document, title, level=3, locale=locale)
    if locale == "zh":
        # A bilingual label inside the image with no Chinese half draws the
        # English; record it like any other fallback (validation_cn.txt).
        labels = [one.get("label") for one in block.get("series") or [] if isinstance(one, dict)]
        labels += [
            point.get("x")
            for one in block.get("series") or []
            if isinstance(one, dict)
            for point in one.get("points") or []
            if isinstance(point, dict)
        ]
        labels += [rule.get("label") for rule in block.get("reference_lines") or [] if isinstance(rule, dict)]
        for label in labels:
            if isinstance(label, dict) and not str(label.get("zh") or "").strip():
                _note_zh_fallback(label, label.get("zh"))
    png: bytes | None = None
    try:
        from server import memo_charts

        png = memo_charts.chart_png(block, locale)
    except Exception:
        logger.warning("chart could not be drawn; printing its data as a table", exc_info=True)
        png = None
    if png:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.add_run().add_picture(io.BytesIO(png), width=Inches(6.0))
    else:
        _add_table(document, _chart_fallback_table(block, locale), locale)
    caption = _loc(block.get("caption"), locale)
    reading = _loc(block.get("reading"), locale)
    if reading:
        reading_text = (
            f"Reading: {reading}" if locale == "en" else f"读法：{reading}"
        )
        caption = f"{reading_text}  {caption}".strip() if caption else reading_text
    source_ids = ", ".join(
        str(item).strip()
        for item in block.get("source_ids") or []
        if str(item or "").strip()
    )
    if source_ids:
        sources_text = (
            f"Sources: {source_ids}" if locale == "en" else f"来源：{source_ids}"
        )
        caption = f"{caption}  {sources_text}".strip() if caption else sources_text
    if caption:
        _add_paragraph(
            document, caption, locale=locale, size=8.5, color=GREY, after=8
        )


def _chart_fallback_table(block: dict, locale: str = "en") -> dict:
    """The chart's series as a plain table block, labels in ``locale``."""
    series = [s for s in block.get("series") or [] if isinstance(s, dict)]
    first = series[0] if series else {}
    xs = [
        _loc(p.get("x"), locale) if isinstance(p.get("x"), dict) else str(p.get("x"))
        for p in first.get("points") or []
        if isinstance(p, dict)
    ]
    unit = _loc(block.get("unit"), locale).strip()
    headers: list[Any] = [unit or "", *xs]
    rows = []
    for one in series:
        label = one.get("label")
        cells: list[Any] = [_loc(label, locale) if isinstance(label, dict) else str(label or "")]
        for point in one.get("points") or []:
            value = point.get("y") if isinstance(point, dict) else ""
            cells.append(str(value))
        rows.append(cells)
    return {"type": "table", "headers": headers, "rows": rows}


class _CellParagraphs:
    """Hands out a cell's paragraphs in order, starting with the one every
    cell already has, so a cell never opens with a blank line."""

    def __init__(self, cell: Any, locale: str) -> None:
        self.cell = cell
        self.locale = locale
        self.used = False

    def next(self, *, after: float = 2) -> Any:
        if not self.used:
            self.used = True
            paragraph = self.cell.paragraphs[0]
        else:
            paragraph = self.cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(after)
        paragraph.paragraph_format.line_spacing = _LINE_SPACING.get(self.locale, 1.15)
        return paragraph


def _add_callout(document: Document, block: dict, locale: str) -> None:
    title = _loc(block.get("title") or block.get("label"), locale)
    body = _loc(block.get("body") or block.get("text"), locale)
    items = [_loc(item, locale) for item in block.get("items") or [] if _loc(item, locale)]
    tone = str(block.get("tone") or "info")
    fill = PALE_GOLD if tone in {"warning", "critical", "risk"} else PALE_TIFFANY
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _shade_cell(cell, fill)
    _set_cell_borders(cell, left=(TIFFANY, "24"))
    paragraphs = _CellParagraphs(cell, locale)
    if title:
        _add_run(paragraphs.next(after=2), title, bold=True, color=NAVY, locale=locale)
    if body:
        _add_run(paragraphs.next(after=2), body, locale=locale)
    for item in items:
        paragraph = paragraphs.next(after=1)
        _make_bullet(paragraph)
        _add_run(paragraph, item, locale=locale)
    _apply_table_widths(table, [_text_width_twips(document)])
    _spacer(document, 4)


# ---- tables -------------------------------------------------------------------------

# Column-width planning, in twips. ~0.5em of 9pt Arial per character; CJK
# characters count double. Word's default cell margins add 0.08in a side.
_CHAR_TWIPS = 92
_CELL_PAD_TWIPS = 230
_KEY_VALUE_LABEL_TWIPS = int(Cm(4.2) / 635)
_FACTS_LABEL_TWIPS = int(Cm(3.2) / 635)
_SHORT_ROW_LINES = 8
_CJK_WIDE_RE = re.compile(r"[ᄀ-ᅟ⺀-꓏가-힣豈-﫿︰-﹏＀-｠￠-￦]")


def _display_len(text: str) -> int:
    text = str(text or "")
    return len(text) + len(_CJK_WIDE_RE.findall(text))


def _longest_word(text: str) -> int:
    words = re.split(r"\s+", str(text or "").strip())
    longest = 0
    for word in words:
        if _CJK_WIDE_RE.search(word):
            # Chinese breaks between characters; its "word" is one glyph.
            longest = max(longest, 2)
        else:
            longest = max(longest, len(word))
    return longest


def _lines(text: str, width_chars: float) -> int:
    width = max(1.0, width_chars)
    total = 0
    for part in str(text or "").split("\n"):
        length = _display_len(part)
        total += max(1, math.ceil(length * 1.08 / width)) if length else 1
    return max(1, total)


_NUMERIC_CELL_RE = re.compile(
    r"^[\s~≈<>≥≤±+\-−–—$€£¥%x×.,:/()\dKMBTkmbn约亿万倍至千元美]+$"
)
_DATE_CELL_RE = re.compile(
    r"^\s*\d{4}(?:[-/.]\d{1,2}){0,2}(?:\s*[–—-]\s*\d{4}(?:[-/.]\d{1,2}){0,2})?\s*$"
    r"|^\s*\d{4}\s*-?\s*Q[1-4]\s*$",
    re.IGNORECASE,
)


def _is_numeric_cell(text: str) -> bool:
    # "1.5x [C2]" is a number with its citation link, not prose.
    value = _CITATION_RE.sub("", str(text or "")).strip()
    return bool(
        value
        and len(value) <= 24
        and any(ch.isdigit() for ch in value)
        and _NUMERIC_CELL_RE.match(value)
        and not _DATE_CELL_RE.match(value)
    )


def _numeric_columns(rows: list[list[str]], col_count: int) -> set[int]:
    numeric: set[int] = set()
    for col in range(1, col_count):
        values = [row[col].strip() for row in rows if col < len(row) and str(row[col]).strip()]
        if values and sum(1 for value in values if _is_numeric_cell(value)) >= 0.7 * len(values):
            numeric.add(col)
    return numeric


def _plan_column_widths(
    headers: list[str],
    rows: list[list[str]],
    total_twips: int,
    *,
    fixed: dict[int, int] | None = None,
) -> list[int]:
    """Column widths (twips, summing to ``total_twips``) chosen from the
    content: every column starts wide enough for its longest word (numbers
    and dates unbroken), then width goes, a few characters at a time, to the
    column where it saves the most wrapped lines across the table; width
    left once nothing wraps less goes to the columns carrying the most
    text."""
    col_count = max([len(headers), *(len(row) for row in rows)] or [1])
    fixed = {col: width for col, width in (fixed or {}).items() if col < col_count}
    free = [col for col in range(col_count) if col not in fixed]
    if not free:
        widths = [fixed.get(col, 0) for col in range(col_count)]
        widths[-1] += total_twips - sum(widths)
        return widths
    numeric = _numeric_columns(rows, col_count)
    grid = [list(headers) + [""] * (col_count - len(headers))]
    grid += [[row[col] if col < len(row) else "" for col in range(col_count)] for row in rows]
    row_count = len(grid)
    # Per column, per row: the display length of each line of the cell (the
    # header is bold, so it counts a little wider).
    lengths = [
        [
            [_display_len(part) * (1.1 if row == 0 else 1.0) for part in str(grid[row][col]).split("\n")]
            for row in range(row_count)
        ]
        for col in range(col_count)
    ]
    budget = max(
        len(free) * 3,
        (total_twips - sum(fixed.values()) - len(free) * _CELL_PAD_TWIPS) // _CHAR_TWIPS,
    )
    width: dict[int, float] = {}
    ceiling: dict[int, float] = {}
    for col in free:
        body = [grid[row][col] for row in range(1, row_count)]
        longest = max((_longest_word(value) for value in [grid[0][col], *body]), default=3)
        full = max((length for row in range(1, row_count) for length in lengths[col][row]), default=3)
        date_column = bool(body) and all(
            _DATE_CELL_RE.match(str(value)) or not str(value).strip() for value in body
        )
        if col in numeric or date_column:
            # The numbers and dates stay unbroken; a prose cell that shares
            # the column ("4.9x growth × 0.36x multiple") may wrap.
            whole = max(
                (
                    _display_len(str(value))
                    for value in body
                    if date_column or _is_numeric_cell(str(value))
                ),
                default=3,
            )
            floor = min(max(whole, longest, 3), 22)
        else:
            floor = min(max(longest, 3), 18)
        width[col] = floor
        ceiling[col] = max(floor, full + 1, min(max(lengths[col][0]) + 1, 40))
    fixed_chars = {
        col: max(1.0, (value - _CELL_PAD_TWIPS) / _CHAR_TWIPS) for col, value in fixed.items()
    }
    used = sum(width.values())
    if used > budget:
        scale = budget / used
        for col in free:
            width[col] = max(3.0, width[col] * scale)

    def _column_lines(col: int, chars: float) -> list[int]:
        chars = max(1.0, chars)
        return [
            sum(max(1, math.ceil(length * 1.08 / chars)) if length else 1 for length in lengths[col][row])
            for row in range(row_count)
        ]

    current = {
        col: _column_lines(col, width[col] if col in width else fixed_chars[col])
        for col in range(col_count)
    }
    heights = [max(current[col][row] for col in range(col_count)) for row in range(row_count)]
    remaining = budget - sum(width.values())
    step = max(1, int(remaining // 80))
    while remaining >= step:
        best_col, best_gain, best_lines = None, 0, None
        for col in free:
            if width[col] >= ceiling[col]:
                continue
            candidate = _column_lines(col, width[col] + step)
            gain = 0
            for row in range(row_count):
                others = max(
                    (current[other][row] for other in range(col_count) if other != col),
                    default=1,
                )
                gain += heights[row] - max(others, candidate[row])
            if gain > best_gain:
                best_col, best_gain, best_lines = col, gain, candidate
        if best_col is None:
            break
        width[best_col] += step
        current[best_col] = best_lines
        remaining -= step
        heights = [max(current[col][row] for col in range(col_count)) for row in range(row_count)]
    if remaining > 0:
        volume = {
            col: sum(sum(lengths[col][row]) for row in range(1, row_count)) + 1 for col in free
        }
        spread = sum(volume.values())
        for col in free:
            width[col] += remaining * volume[col] / spread
    widths = [
        fixed[col] if col in fixed else int(width[col] * _CHAR_TWIPS) + _CELL_PAD_TWIPS
        for col in range(col_count)
    ]
    if sum(widths) > total_twips:
        # More columns than the page holds at their minimums: shrink them
        # all in proportion rather than squeeze one below zero.
        scale = total_twips / sum(widths)
        widths = [max(120, int(width * scale)) for width in widths]
    widest = max(range(col_count), key=lambda col: widths[col])
    widths[widest] += total_twips - sum(widths)
    return widths


def _apply_table_widths(table: Any, widths: list[int]) -> None:
    """Fixed layout with the planned widths written to BOTH the grid
    (``gridCol``, which the web viewer honours) and every cell (``tcW``,
    which Word honours), and the table width equal to their sum."""
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    for column, width in zip(table.columns, widths):
        column.width = Twips(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Twips(width)


def _set_row_flags(row: Any, *, cant_split: bool = False, header: bool = False) -> None:
    """``cantSplit`` keeps a short row on one page; ``tblHeader`` repeats a
    header row at the top of every page the table spans."""
    tr_pr = row._tr.get_or_add_trPr()
    for tag, wanted in (("w:cantSplit", cant_split), ("w:tblHeader", header)):
        if wanted and tr_pr.find(qn(tag)) is None:
            _insert_ordered(tr_pr, OxmlElement(tag), _TRPR_ORDER)


def _row_is_short(values: list[str], widths: list[int]) -> bool:
    for text, width in zip(values, widths):
        chars = max(1.0, (width - _CELL_PAD_TWIPS) / _CHAR_TWIPS)
        if _lines(text, chars) > _SHORT_ROW_LINES:
            return False
    return True


def _cell_paragraph_format(paragraph: Any, locale: str) -> None:
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = _CELL_LINE_SPACING.get(locale, 1.05)


def _fill_header_row(row: Any, headers: list[str], locale: str, numeric: set[int] | None = None) -> None:
    for col, (cell, value) in enumerate(zip(row.cells, headers)):
        _shade_cell(cell, NAVY)
        _set_cell_borders(cell, color=WHITE)
        _cell_text(
            cell,
            value,
            locale=locale,
            bold=True,
            color=WHITE,
            align="right" if numeric and col in numeric else None,
        )


def _clear_cell_borders(cell: Any) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:tcBorders"))
    if existing is not None:
        tc_pr.remove(existing)
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "nil")
        borders.append(border)
    _insert_ordered(tc_pr, borders, _TCPR_ORDER)


def _severity_color(token: str) -> str | None:
    text = token.strip().lower()
    match = re.match(r"^(10|[1-9])\s*/\s*10", text)
    if match:
        score = int(match.group(1))
        return RISK_HIGH if score >= 8 else RISK_MEDIUM if score >= 5 else RISK_LOW
    if text.startswith(("high", "高")):
        return RISK_HIGH
    if text.startswith(("medium", "中")):
        return RISK_MEDIUM
    if text.startswith(("low", "低")):
        return RISK_LOW
    return None


def _add_grid_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    locale: str,
    *,
    derived: bool = False,
    fixed: dict[int, int] | None = None,
    severity_col: int | None = None,
    likelihood_col: int | None = None,
    marked_rows: set[int] | None = None,
) -> Any:
    """A header-row table with content-planned widths, a repeating header
    row, unsplittable short rows and right-aligned number columns."""
    col_count = max([len(headers), *(len(row) for row in rows)] or [1])
    table = document.add_table(rows=(1 if headers else 0) + len(rows), cols=col_count)
    if derived:
        _mark_derived_table(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    numeric = _numeric_columns(rows, col_count)
    marked_rows = marked_rows or set()
    offset = 0
    if headers:
        padded = list(headers) + [""] * (col_count - len(headers))
        _fill_header_row(table.rows[0], padded, locale, numeric)
        offset = 1
    shown_rows: list[list[str]] = []
    for index, row in enumerate(rows):
        values = [row[col] if col < len(row) else "" for col in range(col_count)]
        if index in marked_rows and values:
            values[0] = f"{values[0]} †"
        shown_rows.append(values)
        cells = table.rows[index + offset].cells
        for col, value in enumerate(values):
            cell = cells[col]
            if index % 2 == 1:
                _shade_cell(cell, "FAFAFA")
            _set_cell_borders(cell)
            tint = None
            if col in (severity_col, likelihood_col):
                tint = _severity_color(str(value))
            _cell_text(
                cell,
                value,
                locale=locale,
                bold=bool(tint),
                color=tint or BLACK,
                align="right" if col in numeric else None,
            )
    widths = _plan_column_widths(
        list(headers), shown_rows, _text_width_twips(document), fixed=fixed
    )
    _apply_table_widths(table, widths)
    if headers:
        _set_row_flags(table.rows[0], header=True)
    for row, values in zip(list(table.rows)[offset:], shown_rows):
        if _row_is_short(values, widths):
            _set_row_flags(row, cant_split=True)
    _spacer(document, 4)
    return table


_KEY_METRICS_FOOTNOTE = {
    "en": "† Single third-party estimate, not corroborated by a filing or the company.",
    "zh": "† 单一第三方估算，未经监管申报文件或公司披露佐证。",
}


def _tier_c_rows(block: dict, context: _RenderContext | None) -> set[int]:
    """Key Metrics Snapshot rows whose cited [S#] sources are all tier C —
    decided on the English cells, so both documents mark the same rows."""
    if context is None or "key_metrics_snapshot" not in _declared_component_ids(block):
        return set()
    return _tier_c_row_indexes(block, context.source_tiers())


def _tier_c_row_indexes(block: dict, tiers: dict[str, str]) -> set[int]:
    if not tiers:
        return set()
    marked: set[int] = set()
    for index, row in enumerate(block.get("rows") or []):
        text = " ".join(_raw_text(cell, "en") if isinstance(cell, dict) else str(cell or "") for cell in _row_cells(row))
        cited = [item for item in citation_ids(text) if item.startswith("S")]
        if cited and all(tiers.get(item) == source_tiers.TIER_C for item in cited):
            marked.add(index)
    return marked


def _add_key_value_card_table(document: Document, rows: list[list[str]], locale: str) -> None:
    """Render a ``layout: "key_value"`` block: label column left, prose right.

    This is the risk-card layout — a shaded bold label column (4.2 cm) so
    the reader scans Risk Type / Why it matters / What we watch /
    Likelihood / Risk Rating at a glance. Only the leading assessment
    token ("High:", "8/10") is bold and tinted, not the whole sentence.
    """
    table = document.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    total = _text_width_twips(document)
    widths = [_KEY_VALUE_LABEL_TWIPS, total - _KEY_VALUE_LABEL_TWIPS]
    for idx, (label, value) in enumerate(rows):
        cells = table.rows[idx].cells
        _shade_cell(cells[0], WARM_GREY)
        _set_cell_borders(cells[0])
        _set_cell_borders(cells[1])
        _cell_text(cells[0], label, locale=locale, bold=True, color=NAVY)
        text = str(value or "")
        rating = _RISK_RATING_VALUE_RE.match(text)
        likelihood = _RISK_LIKELIHOOD_TOKEN_RE.match(text)
        token_end = rating.end() if rating else likelihood.end(1) if likelihood else 0
        if token_end:
            token = text[:token_end]
            paragraph = _reset_cell(cells[1])
            _cell_paragraph_format(paragraph, locale)
            _add_run(
                paragraph,
                token,
                bold=True,
                color=_severity_color(token) or NAVY,
                size=_TABLE_SIZE,
                locale=locale,
            )
            if text[token_end:]:
                _add_run(paragraph, text[token_end:], size=_TABLE_SIZE, locale=locale)
        else:
            _cell_text(cells[1], text, locale=locale)
    _apply_table_widths(table, widths)
    for row, values in zip(table.rows, rows):
        if _row_is_short(list(values), widths):
            _set_row_flags(row, cant_split=True)
    _spacer(document, 4)


def _add_key_value_table(document: Document, rows: list[tuple[str, str]], locale: str) -> None:
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    total = _text_width_twips(document)
    for idx, (key, value) in enumerate(rows):
        cells = table.rows[idx].cells
        _shade_cell(cells[0], WARM_GREY)
        _set_cell_borders(cells[0])
        _set_cell_borders(cells[1])
        _cell_text(cells[0], key, locale=locale, bold=True, color=NAVY)
        _cell_text(cells[1], str(value), locale=locale)
    _apply_table_widths(table, [_FACTS_LABEL_TWIPS, total - _FACTS_LABEL_TWIPS])
    _spacer(document, 4)


def _add_table(
    document: Document,
    block: dict,
    locale: str,
    context: _RenderContext | None = None,
) -> None:
    headers = [_loc(header, locale) for header in block.get("headers") or []]
    rows = [_row_values(row, locale) for row in block.get("rows") or []]
    if not rows and not headers:
        return
    if (
        str(block.get("layout") or "").strip().lower() == "key_value"
        and not headers
        and rows
        and all(len(row) == 2 for row in rows)
    ):
        _add_key_value_card_table(document, rows, locale)
        return
    marked = _tier_c_rows(block, context)
    _add_grid_table(document, headers, rows, locale, marked_rows=marked)
    if marked:
        footnote = _derived_paragraph(document)
        footnote.paragraph_format.space_after = Pt(6)
        _add_run(footnote, _KEY_METRICS_FOOTNOTE[locale], size=8, color=GREY, locale=locale)


def _reset_cell(cell: Any) -> Any:
    """Empty a cell (keeping its properties) down to one bare paragraph —
    ``cell.text = ""`` would leave an empty run in front of the content."""
    cell._tc.clear_content()
    cell._tc.add_p()
    return cell.paragraphs[0]


def _cell_text(
    cell: Any,
    text: Any,
    *,
    locale: str,
    bold: bool = False,
    color: str = BLACK,
    size: float = _TABLE_SIZE,
    align: str | None = None,
) -> None:
    paragraph = _reset_cell(cell)
    _cell_paragraph_format(paragraph, locale)
    if align == "right":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "center":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(paragraph, _loc(text, locale), locale=locale, bold=bold, color=color, size=size)


# Inline citations: `[S3]`, `[C2]`, `[S3, C2]` — ids of package sources
# (S) and calculation notes (C). The renderer turns each id into a
# superscript link to its bookmarked row in the Sources table or the
# Calculation notes appendix, so a reader can click through to where a
# number came from and how it was computed (founder feedback 2026-09-13).
_CITATION_RE = re.compile(r"\[((?:[SC]\d+)(?:\s*,\s*[SC]\d+)*)\]")
_BOOKMARK_IDS = itertools.count(9000)
CITATION_LINK_COLOR = "1F4E79"


def citation_ids(text: Any) -> list[str]:
    """Every citation id in ``text`` in order (with repeats)."""
    ids: list[str] = []
    for group in _CITATION_RE.findall(str(text or "")):
        ids.extend(part.strip() for part in group.split(","))
    return ids


def citation_anchor(citation_id: str) -> str:
    return f"src_{citation_id}" if citation_id.startswith("S") else f"calc_{citation_id}"


def _add_bookmark(paragraph: Any, name: str) -> None:
    """Bookmark the whole paragraph. The start goes after ``w:pPr`` (which
    must stay the paragraph's first child), the end after the content."""
    bookmark_id = str(next(_BOOKMARK_IDS))
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bookmark_id)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bookmark_id)
    p = paragraph._p
    p.insert(1 if p.pPr is not None else 0, start)
    p.append(end)


def _append_hyperlink(
    paragraph: Any,
    text: str,
    *,
    anchor: str | None = None,
    url: str | None = None,
    locale: str = "en",
    size: float = _TEXT_SIZE,
    superscript: bool = False,
    color: str | None = None,
    underline: bool = True,
    bold: bool = False,
) -> Any:
    """Append a run wrapped in ``w:hyperlink`` — internal (``anchor`` to a
    bookmark) or external (``url`` via a document relationship)."""
    hyperlink = OxmlElement("w:hyperlink")
    if url:
        r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        hyperlink.set(qn("r:id"), r_id)
    elif anchor:
        hyperlink.set(qn("w:anchor"), anchor)
    run = paragraph.add_run(str(text or ""))
    _style_run(run, size=size, bold=bold, color=color or CITATION_LINK_COLOR, locale=locale)
    if superscript:
        run.font.superscript = True
    elif underline:
        run.font.underline = True
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)
    return run


def _markdown_segments(text: str) -> list[tuple[str, bool]]:
    """``"A **bold** word"`` -> [("A ", False), ("bold", True), (" word",
    False)]; a stray marker is dropped, never printed."""
    if "**" not in text:
        return [(text, False)] if text else []
    segments: list[tuple[str, bool]] = []
    position = 0
    for match in _MD_BOLD_RE.finditer(text):
        if match.start() > position:
            segments.append((text[position : match.start()].replace("**", ""), False))
        segments.append((match.group(1).replace("**", ""), True))
        position = match.end()
    if position < len(text):
        segments.append((text[position:].replace("**", ""), False))
    return [(segment, bold) for segment, bold in segments if segment]


def _add_run(
    paragraph: Any,
    text: str,
    *,
    size: float = _TEXT_SIZE,
    bold: bool = False,
    color: str = BLACK,
    locale: str = "en",
) -> Any:
    """Append ``text`` as styled runs: a ``**marked**`` span renders bold
    (never as asterisks), and inline ``[S#]``/``[C#]`` citations become
    superscript links. Returns the last run."""
    segments = _markdown_segments(str(text or ""))
    if not segments:
        run = paragraph.add_run("")
        _style_run(run, size=size, bold=bold, color=color, locale=locale)
        return run
    last: Any = None
    for segment, marked in segments:
        last = _add_run_segment(
            paragraph, segment, size=size, bold=bold or marked, color=color, locale=locale
        )
    return last


def _add_run_segment(
    paragraph: Any,
    text: str,
    *,
    size: float,
    bold: bool,
    color: str,
    locale: str,
) -> Any:
    if "[" not in text or not _CITATION_RE.search(text):
        run = paragraph.add_run(text)
        _style_run(run, size=size, bold=bold, color=color, locale=locale)
        return run
    last: Any = None
    position = 0
    for match in _CITATION_RE.finditer(text):
        if match.start() > position:
            last = paragraph.add_run(text[position : match.start()])
            _style_run(last, size=size, bold=bold, color=color, locale=locale)
        ids = [part.strip() for part in match.group(1).split(",")]
        for index, citation_id in enumerate(ids):
            label = citation_id if index == len(ids) - 1 else f"{citation_id},"
            last = _append_hyperlink(
                paragraph,
                label,
                anchor=citation_anchor(citation_id),
                locale=locale,
                size=size,
                superscript=True,
            )
        position = match.end()
    if position < len(text):
        last = paragraph.add_run(text[position:])
        _style_run(last, size=size, bold=bold, color=color, locale=locale)
    return last


def _style_run(
    run: Any,
    *,
    size: float = _TEXT_SIZE,
    bold: bool = False,
    color: str = BLACK,
    locale: str = "en",
) -> None:
    font = _font(locale)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    _set_rfonts(run._element.get_or_add_rPr(), font, locale)


def _set_rfonts(rpr: Any, font: str, locale: str) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    if locale == "zh":
        rfonts.set(qn("w:eastAsia"), ZH_FONT)
        rfonts.set(qn("w:ascii"), EN_FONT)
        rfonts.set(qn("w:hAnsi"), EN_FONT)
        rfonts.set(qn("w:cs"), EN_FONT)
    else:
        rfonts.set(qn("w:ascii"), font)
        rfonts.set(qn("w:hAnsi"), font)
        rfonts.set(qn("w:cs"), font)


def _add_page_number(paragraph: Any, locale: str) -> None:
    if locale == "zh":
        _add_run(paragraph, "第 ", locale=locale, size=8.5, color=GREY)
    else:
        _add_run(paragraph, "Page ", locale=locale, size=8.5, color=GREY)
    run = paragraph.add_run()
    for element in (
        _field_char("begin"),
        _instr_text("PAGE"),
        _field_char("end"),
    ):
        run._r.append(element)
    _style_run(run, locale=locale, size=8.5, color=GREY)
    if locale == "zh":
        _add_run(paragraph, " 页", locale=locale, size=8.5, color=GREY)


def _field_char(kind: str) -> Any:
    element = OxmlElement("w:fldChar")
    element.set(qn("w:fldCharType"), kind)
    return element


def _instr_text(value: str) -> Any:
    element = OxmlElement("w:instrText")
    element.set(qn("xml:space"), "preserve")
    element.text = value
    return element


# WordprocessingML child order for the property containers this module
# edits by hand; Word expects schema order, so new children are inserted in
# place rather than appended.
_PPR_ORDER = (
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
    "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
    "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
    "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
    "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
    "pPrChange",
)
_TCPR_ORDER = (
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
    "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
    "headers", "cellIns", "cellDel", "cellMerge", "tcPrChange",
)
_TRPR_ORDER = (
    "cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter",
    "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden",
    "ins", "del", "trPrChange",
)


def _insert_ordered(parent: Any, child: Any, order: tuple[str, ...]) -> Any:
    name = child.tag.rsplit("}", 1)[-1]
    later = {qn(f"w:{tag}") for tag in order[order.index(name) + 1 :]}
    for existing in parent:
        if existing.tag in later:
            existing.addprevious(child)
            return child
    parent.append(child)
    return child


def _shade_cell(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = _insert_ordered(tc_pr, OxmlElement("w:shd"), _TCPR_ORDER)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def _set_cell_borders(
    cell: Any,
    *,
    color: str = BORDER,
    size: str = "4",
    left: tuple[str, str] | None = None,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:tcBorders"))
    if existing is not None:
        tc_pr.remove(existing)
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        if side == "left" and left:
            border.set(qn("w:sz"), left[1])
            border.set(qn("w:color"), left[0])
        else:
            border.set(qn("w:sz"), size)
            border.set(qn("w:color"), color)
        borders.append(border)
    _insert_ordered(tc_pr, borders, _TCPR_ORDER)


def _add_bottom_border(paragraph: Any, color: str) -> None:
    p_pr = paragraph._element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = _insert_ordered(p_pr, OxmlElement("w:pBdr"), _PPR_ORDER)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def _validation_report(
    package: dict,
    output_path: Path,
    locale: str,
    *,
    zh_fallbacks: list[tuple[str, str, str]] | None = None,
    source_warnings: list[str] | None = None,
) -> str:
    document = Document(output_path)
    table_count = len(document.tables)
    paragraph_count = sum(1 for paragraph in document.paragraphs if paragraph.text.strip())
    callout_count = sum(
        1
        for section in package.get("sections") or []
        for block in section.get("blocks", [])
        if isinstance(block, dict) and block.get("type") == "callout"
    )
    structure = _structure_for(package)
    coverage = _memo_component_coverage(package, structure)
    coverage_lines = [
        f"- {component['id']}: {'present' if coverage[str(component['id'])] else 'missing'}"
        for component in structure.components
    ]
    lines = [
        "# Memo Renderer Validation",
        "",
        f"- output: `{output_path}`",
        f"- locale: {locale}",
        f"- schema_version: {package.get('schema_version')}",
        f"- renderer_version: {RENDERER_VERSION}",
        f"- paragraphs: {paragraph_count}",
        f"- tables: {table_count}",
        f"- declared_callouts: {callout_count}",
        "",
        "## Content Coverage",
        "",
        *coverage_lines,
        "",
    ]
    if source_warnings:
        lines += [
            "## Source warnings (legacy package, rule downgraded)",
            "",
            *[f"- {warning}" for warning in source_warnings],
            "",
        ]
    if zh_fallbacks is not None:
        unique: list[tuple[str, str, str]] = list(dict.fromkeys(zh_fallbacks))
        lines += [
            "## Chinese fallbacks",
            "",
            f"- zh_fallback_count: {len(unique)}",
        ]
        for section_id, kind, english in unique:
            code = (
                "zh_fallback_to_english"
                if kind == "english_printed"
                else "zh_blank_translation"
            )
            snippet = re.sub(r"\s+", " ", english).strip()
            if len(snippet) > 160:
                snippet = snippet[:157] + "…"
            lines.append(f"- P1 {code} · {section_id}: “{snippet}”")
        lines.append("")
    lines += ["- status: passed", ""]
    return "\n".join(lines)


_INVENTORY_ENTRY_RE = re.compile(r"^- (memo|validation)_([a-z]{2}): `")


def _write_inventory(
    path: Path,
    package: dict,
    outputs: dict[str, Path],
    validation_paths: dict[str, Path],
    *,
    keep_other_locales: bool = False,
) -> None:
    """``logs/file_inventory.md``: the rendered files. A one-language render
    (``keep_other_locales``) keeps the other language's lines from the
    inventory already on disk, so English-first delivery and "Retry
    Chinese" leave both files listed."""
    entries: dict[tuple[str, str], str] = {}
    if keep_other_locales and path.exists():
        for line in path.read_text(encoding="utf-8").splitlines():
            match = _INVENTORY_ENTRY_RE.match(line)
            if match and match.group(2) not in outputs:
                entries[(match.group(1), match.group(2))] = line
    for locale, output in outputs.items():
        entries[("memo", locale)] = f"- memo_{locale}: `{output}`"
    for locale, validation in validation_paths.items():
        entries[("validation", locale)] = f"- validation_{locale}: `{validation}`"
    locales = list(MEMO_LOCALES) + sorted(
        {locale for _kind, locale in entries} - set(MEMO_LOCALES)
    )
    lines = ["# Memo File Inventory", ""]
    lines.append(f"- company: {_company_name(package)}")
    lines.append(f"- schema_version: {package.get('schema_version')}")
    lines.append(f"- renderer_version: {RENDERER_VERSION}")
    for kind in ("memo", "validation"):
        for locale in locales:
            if (kind, locale) in entries:
                lines.append(entries[(kind, locale)])
    lines.append("")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def _append_manifest(
    path: Path,
    package: dict,
    outputs: dict[str, Path],
    validation_paths: dict[str, Path],
) -> None:
    lines = [
        "",
        "## Analysis Finalization",
        "",
        "- renderer: server.memo_docx_renderer",
        f"- renderer_version: {RENDERER_VERSION}",
        f"- memo_package_schema_version: {package.get('schema_version')}",
    ]
    for locale, output in outputs.items():
        lines.append(f"- memo_{locale}: `{output}`")
    for locale, validation in validation_paths.items():
        lines.append(f"- validation_{locale}: `{validation}`")
    lines.append("- validation_status: passed")
    lines.append("")
    path.parent.mkdir(parents=True, exist_ok=True)
    existing = path.read_text(encoding="utf-8") if path.exists() else ""
    path.write_text(existing.rstrip() + "\n" + "\n".join(lines), encoding="utf-8")


def _has_section(package: dict, section_id: str) -> bool:
    return any(
        isinstance(section, dict) and section.get("id") == section_id
        for section in package.get("sections") or []
    )


def _section_title(
    section: dict,
    locale: str,
    section_titles: dict[str, dict[str, str]] | None = None,
) -> str:
    titles = SECTION_TITLES if section_titles is None else section_titles
    return _loc(section.get("title"), locale) or _loc(
        titles.get(str(section.get("id") or ""), ""), locale
    )


def _toc_titles(package: dict, locale: str) -> list[str]:
    """The level-1 titles in document order, without their numerals — the
    entries of the contents list."""
    plan = _document_plan(package, _structure_for(package))
    return [entry.bare.get(locale) or entry.bare.get("en") or "" for entry in plan]


# Placeholders the memo prompts prescribe for what the IC fills in later.
# The English is fixed by the prompts; the Chinese is fixed here so every
# memo prints the same words (the model's own renderings varied:
# "[待定]", "[由投委会决定]", the English left as is).
PLACEHOLDER_LABELS: dict[str, dict[str, str]] = {
    "[TO BE DETERMINED BY IC]": {"en": "[TO BE DETERMINED BY IC]", "zh": "【待投委会确定】"},
    "[date to set]": {"en": "[date to set]", "zh": "【日期待定】"},
    "[owner to assign]": {"en": "[owner to assign]", "zh": "【负责人待定】"},
}
# A translator may prefix its own rendering and keep the English after a
# slash ("[由投委会决定 / TO BE DETERMINED BY IC]", Gemini, ZaiNar
# 2026-09-23); the whole bracket is the placeholder.
_PLACEHOLDER_RE = re.compile(
    r"[\[【]\s*(?:[^\[\]【】/]{1,20}/\s*)?(?:to be determined by ic|tbd by ic|date to set|owner to assign)\s*[\]】]",
    re.IGNORECASE,
)
# Owner roles that appear alone in a cell of the follow-ups / action table.
OWNER_ROLE_LABELS: dict[str, str] = {
    "deal lead": "交易负责人",
    "partner": "合伙人",
    "legal": "法务",
    "finance": "财务",
    "deal lead/legal/finance": "交易负责人/法务/财务",
    "deal lead / legal / finance": "交易负责人/法务/财务",
}
_OWNER_ROLE_SPLIT_RE = re.compile(r"\s*[/、,，;；]\s*")


def _placeholder_zh(match: re.Match) -> str:
    inner = re.sub(r"\s+", " ", match.group(0).lower())
    if "date to set" in inner:
        return PLACEHOLDER_LABELS["[date to set]"]["zh"]
    if "owner to assign" in inner:
        return PLACEHOLDER_LABELS["[owner to assign]"]["zh"]
    return PLACEHOLDER_LABELS["[TO BE DETERMINED BY IC]"]["zh"]


# The founder's template fixes three whole lines around those placeholders
# ("Proposed amount: [TO BE DETERMINED BY IC]", "Allocation: ...",
# "Strategy: [重仓 / 跟投 / 卡位 — IC to select]"). A translator told to keep
# placeholders verbatim keeps the whole line (ZaiNar 2026-09-23: the Chinese
# recommendation box read "Proposed amount: 【待投委会确定】"). A label is
# rendered only when the rest of the string is the placeholder itself.
_STRATEGY_PLACEHOLDER_RE = re.compile(
    r"[\[【]\s*重仓\s*/\s*跟投\s*/\s*卡位\s*[—–-]+\s*(?:IC to select|由?投委会(?:选定|选择|决定))\s*[\]】]",
    re.IGNORECASE,
)
_STRATEGY_PLACEHOLDER_ZH = "【重仓 / 跟投 / 卡位 — 投委会选定】"
_FIXED_LINE_LABELS_ZH = (
    (re.compile(r"^(\s*)proposed amount\s*[:：]\s*(?=【待投委会确定】\s*$)", re.IGNORECASE), "拟投金额："),
    (re.compile(r"^(\s*)allocation\s*[:：]\s*(?=【待投委会确定】\s*$)", re.IGNORECASE), "额度分配："),
    (re.compile(r"^(\s*)strategy\s*[:：]\s*(?=【重仓 / 跟投 / 卡位 — 投委会选定】\s*$)", re.IGNORECASE), "投资策略："),
)


def apply_zh_placeholders(text: str) -> str:
    """The Chinese memo's fixed rendering of the prompts' placeholders,
    wherever they appear in a Chinese string (a cell, or mid-sentence),
    and of the template lines that consist of a label and a placeholder."""
    if not text or "[" not in text and "【" not in text:
        return text
    text = _STRATEGY_PLACEHOLDER_RE.sub(_STRATEGY_PLACEHOLDER_ZH, text)
    text = _PLACEHOLDER_RE.sub(_placeholder_zh, text)
    for pattern, label in _FIXED_LINE_LABELS_ZH:
        text = pattern.sub(lambda m, label=label: m.group(1) + label, text)
    return text


# Money in the Chinese memo keeps its English form ($30M, $1.35B) — the
# fixed number convention every translation prompt carries. A translator
# that writes 万/亿美元 anyway mixes the two styles in one memo (Gemini,
# ZaiNar 2026-09-23: "$30M" in one section, "3000万美元" in the next), so
# the renderer converts, exactly (Decimal): 3000万美元 → $30M, 5.7937 亿美元
# → $579.37M, "4 亿至 7.5 亿美元" → "$400M 至 $750M". Chinese-numeral
# amounts (数亿美元, 十亿美元) and small plain amounts ("每 1 美元") stay.
_ZH_NUM = r"\d+(?:,\d{3})*(?:\.\d+)?"
_ZH_MONEY_UNITS = {"万亿": Decimal(10) ** 12, "亿": Decimal(10) ** 8, "万": Decimal(10) ** 4}
_ZH_MONEY_RANGE_RE = re.compile(
    rf"(?<![\d.$])({_ZH_NUM})\s*(万亿|亿|万)?\s*(至|到|~|～|-|–|—)\s*({_ZH_NUM})\s*(万亿|亿|万)\s*美元"
)
_ZH_MONEY_RE = re.compile(rf"(?<![\d.$])({_ZH_NUM})\s*(万亿|亿|万)?\s*美元")


def _usd_english_form(amount: Decimal) -> str:
    for scale, suffix in ((Decimal(10) ** 12, "T"), (Decimal(10) ** 9, "B"), (Decimal(10) ** 6, "M"), (Decimal(10) ** 3, "K")):
        if amount >= scale:
            return f"${_plain_decimal(amount / scale)}{suffix}"
    return f"${_plain_decimal(amount)}"


def _plain_decimal(value: Decimal) -> str:
    text = format(value.normalize(), "f")
    return text.rstrip("0").rstrip(".") if "." in text else text


def _zh_amount(number: str, unit: str | None) -> Decimal:
    return Decimal(number.replace(",", "")) * _ZH_MONEY_UNITS.get(unit or "", Decimal(1))


def normalize_zh_money(text: str) -> str:
    """万/亿美元 amounts in a Chinese string rewritten in the memo's English
    money form (see the note above)."""
    if not text or "美元" not in text:
        return text

    def spaced(match: re.Match, rewritten: str) -> str:
        # A half-width space between the amount and a Chinese character or
        # a word that follows ("$5B GPS", never "$5BGPS").
        after = match.string[match.end():match.end() + 1]
        return rewritten + (" " if after and (_CJK_CHAR_RE.match(after) or after.isalnum()) else "")

    def range_sub(match: re.Match) -> str:
        unit = match.group(5)
        low = _zh_amount(match.group(1), match.group(2) or unit)
        high = _zh_amount(match.group(4), unit)
        return spaced(match, f"{_usd_english_form(low)} {match.group(3)} {_usd_english_form(high)}")

    def single_sub(match: re.Match) -> str:
        number, unit = match.group(1), match.group(2)
        amount = _zh_amount(number, unit)
        if not unit and "." not in number and amount < 10:
            return match.group(0)
        return spaced(match, _usd_english_form(amount))

    try:
        text = _ZH_MONEY_RANGE_RE.sub(range_sub, text)
        return _ZH_MONEY_RE.sub(single_sub, text)
    except (InvalidOperation, ValueError):
        return text


def zh_render_text(text: str) -> str:
    """Every fixed rewrite a Chinese string gets on its way to the page:
    the money form, then the placeholders and template lines."""
    return apply_zh_placeholders(normalize_zh_money(text))


def owner_role_zh(text: str) -> str | None:
    """The Chinese label for an owner-role cell ("deal lead",
    "legal/finance") — None when the cell is anything else."""
    key = re.sub(r"\s+", " ", str(text or "").strip().lower())
    if not key:
        return None
    if key in OWNER_ROLE_LABELS:
        return OWNER_ROLE_LABELS[key]
    parts = [part for part in _OWNER_ROLE_SPLIT_RE.split(key) if part]
    if len(parts) > 1 and all(part in OWNER_ROLE_LABELS for part in parts):
        return "/".join(OWNER_ROLE_LABELS[part] for part in parts)
    return None


def _row_values(row: Any, locale: str) -> list[str]:
    if isinstance(row, dict):
        if isinstance(row.get("cells"), list):
            values = [_loc(value, locale) for value in row["cells"]]
        else:
            values = [_loc(value, locale) for value in row.values()]
    elif isinstance(row, (list, tuple)):
        values = [_loc(value, locale) for value in row]
    else:
        values = [_loc(row, locale)]
    if locale == "zh":
        values = [owner_role_zh(value) or value for value in values]
    return values


# Where the Chinese render printed English (or nothing) because a value's
# ``zh`` half was missing or blank. Collected only while the ZH document is
# being built and written into validation_cn.txt, so a silent fallback is
# visible instead of shipping unnoticed.
_ZH_FALLBACKS: contextvars.ContextVar[list | None] = contextvars.ContextVar(
    "memo_zh_fallbacks", default=None
)
_RENDER_SECTION: contextvars.ContextVar[str] = contextvars.ContextVar(
    "memo_render_section", default="front_matter"
)


def _note_zh_fallback(value: dict, selected: Any) -> None:
    sink = _ZH_FALLBACKS.get()
    if sink is None:
        return
    english = value.get("en")
    if english is None or not str(english).strip():
        return
    kind = "english_printed" if selected is None else "blank_zh"
    sink.append((_RENDER_SECTION.get(), kind, str(english)))


def _loc(value: Any, locale: str, *, rewrite: bool = True) -> str:
    """``value`` in ``locale`` (English fallback). Chinese text gets the
    fixed rewrites (zh_render_text) unless ``rewrite`` is False — a cited
    source's title is quoted as published, never re-worded."""
    if value is None:
        return ""
    if isinstance(value, dict):
        selected = value.get(locale)
        if locale == "zh" and (selected is None or not str(selected).strip()):
            _note_zh_fallback(value, selected)
        if selected is None:
            selected = value.get("en")
        if selected is None:
            selected = next(iter(value.values()), "")
        text = str(selected or "")
        return zh_render_text(text) if locale == "zh" and rewrite else text
    return zh_render_text(str(value)) if locale == "zh" and rewrite else str(value)


def _company_name(package: dict, locale: str = "en") -> str:
    """The company name for ``locale``: a bilingual ``company.name`` gives
    the Chinese memo its Chinese name; a plain string serves both."""
    company = package.get("company") or {}
    name = company.get("name") or company.get("display_name") or "Company"
    if isinstance(name, dict):
        return _raw_text(name, locale).strip() or _raw_text(name, "en").strip() or "Company"
    return str(name)


def _font(locale: str) -> str:
    return ZH_FONT if locale == "zh" else EN_FONT


def _label(en: str, zh: str, locale: str) -> str:
    return zh if locale == "zh" else en


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Render BSH memo DOCX files from memo_package.json")
    parser.add_argument("--package", required=True, type=Path)
    parser.add_argument("--out-en", required=True, type=Path)
    parser.add_argument("--out-zh", required=True, type=Path)
    parser.add_argument("--validation-en", type=Path)
    parser.add_argument("--validation-zh", type=Path)
    parser.add_argument("--inventory", type=Path)
    parser.add_argument("--manifest", type=Path)
    args = parser.parse_args(argv)
    render_memos(
        args.package,
        out_en=args.out_en,
        out_zh=args.out_zh,
        validation_en=args.validation_en,
        validation_zh=args.validation_zh,
        inventory_path=args.inventory,
        manifest_path=args.manifest,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
