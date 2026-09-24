from __future__ import annotations

import json
import zipfile
from pathlib import Path

from docx import Document

from server import api as api_module
from server import (
    buffett_checks,
    buffett_memo_analysis,
    buffett_memo_renderer,
    claude_runner,
    internal_memo_renderer,
    memo_prep,
    storage,
)


def _pad(text: str, *, minimum: int = 850) -> str:
    filler = (
        " Owner's earnings, as we figure them, are the cash that can be taken "
        "out of the business without impairing its competitive position. We "
        "would rather be approximately right than precisely wrong, and we will "
        "not pay today for a story that has to come true next year. "
    )
    while len(text) < minimum:
        text += filler
    return text


def _english_memo(*, decision: str = "Pass", extra_viii: str = "") -> str:
    return _pad(
        f"""# Investment Memorandum — Example Co

BSH Research · Owner's analysis (Buffett method) · prices as of 2026-08-20

## I. Investment Decision

**{decision}.** We would not buy the whole company at the last disclosed price.
The economics are intelligible, but the price offers no margin of safety.

## II. The Business

Example Co sells industrial software to factories. Customers pay annually.
The product is useful, not magical.

## III. Circle of Competence

We understand how a factory pays for software that reduces downtime. This sits
inside our circle.

## IV. Economic Characteristics

Returns on tangible capital look decent where disclosed, but incremental
reinvestment needs are not proven. We will not capitalize hoped-for scale.

## V. Durable Competitive Advantage

Switching costs exist once the software is on the floor, yet a well-capitalized
rival can still win the next plant. The moat is narrow.

## VI. Management and Capital Allocation

Management appears able. There is no long record of owner-oriented capital
allocation, so we assume average until proven otherwise.

## VII. Financial Position

The balance sheet is not stretched. Leverage is not the issue; price is.

## VIII. Intrinsic Value and Margin of Safety

We would pay a price that capitalizes current owner's earnings at a conservative
rate. Today's quote is closer to a wonderful price for a fair business than a
fair price for a wonderful business.{extra_viii}

## IX. What Can Go Permanently Wrong

**Price.** The last round values the business ahead of proven owner's earnings.
If growth slows, the multiple can contract and permanent loss starts with the
price we pay, not with the product.

**Competition.** A well-capitalized rival can win the next plant at a lower
price. Switching costs may not protect the installed base, so revenue and the
ten-year return can fall together.

**Key person.** The founders still carry important operating knowledge. If
they leave before that knowledge is institutionalized, execution can weaken
and the business may need more capital than its economics justify.

**Dilution.** Growth spending may not earn its cost of capital. New equity
would then transfer more of the future owner's earnings to investors without
creating equal business value.

## X. Conclusion

{decision}. We would be glad to own the whole business only at a lower price,
and only if the installed base keeps renewing for the next decade.
"""
    )


def _chinese_memo(*, decision_zh: str = "放弃") -> str:
    return _pad(
        f"""# 投资备忘录 — Example Co

BSH 研究 · 巴菲特方法所有者分析 · 价格截至 2026 年 8 月 20 日

## 一、投资决策

**{decision_zh}。** 按最近一轮披露的价格，我们不会买下整家公司。生意能看懂，但没有安全边际。

## 二、这家企业

Example Co 向工厂出售工业软件，客户按年付费。产品有用，但并不神奇。

## 三、能力圈

我们理解工厂为何为减少停机的软件付钱。这在我们的能力圈内。

## 四、经济特征

有披露的有形资本回报看起来尚可，但再投资需求未经证明。我们不会把明年的故事资本化。

## 五、持久竞争优势

上线后有一定转换成本，但有资本的对手仍可拿下下一家工厂。护城河很窄。

## 六、管理层与资本配置

管理层看起来能干。目前没有长期的所有者导向资本配置记录，因此先按平均水平假设。

## 七、财务状况

资产负债表并不紧。问题不在杠杆，而在价格。

## 八、内在价值与安全边际

我们会按保守利率资本化当前股东盈余（owner's earnings）。今天的价格更像是给平庸生意付了美妙价格。

## 九、可能造成永久损失的因素

**价格。** 上一轮定价高于已证明的股东盈余。如果增长放缓，估值倍数可能收缩，永久损失首先来自买入价格，而不是产品本身。

**竞争。** 资金充足的对手可以用更低价格拿下下一家工厂。转换成本未必能保护现有客户，收入和十年回报可能同时下降。

**关键人物。** 创始人仍掌握重要运营知识。如果知识尚未制度化就离开，执行可能变弱，企业需要的资本可能超过其经济性所能支持的水平。

**稀释。** 增长投入可能赚不到资本成本。此时新股权会把更多未来股东盈余转给新投资者，却没有创造等量的企业价值。

## 十、结论

{decision_zh}。只有价格更低、且现有客户在未来十年持续续约，我们才乐于拥有整家公司。
""",
        minimum=900,
    )


def _package(**overrides) -> dict:
    package = {
        "schema_version": 1,
        "kind": "buffett_investment_memo",
        "company_name": "Example Co",
        "decision": "Pass",
        "buy_price": "$8 a share",
        "markdown_en": _english_memo(),
        "markdown_zh": _chinese_memo(),
    }
    package.update(overrides)
    return package


def _valuation_fields() -> dict:
    # The skill's worked example: every figure reconciles.
    return {
        "pass_kind": "price",
        "currency": "USD",
        "price": 150.0,
        "price_date": "2026-09-21",
        "ust10y": 4.25,
        "ust10y_date": "2026-09-21",
        "hurdle": 10.0,
        "g": 3.0,
        "value_basis": "per_share",
        "earnings": 1000,
        "d_and_a": 200,
        "maintenance_capex": 250,
        "working_capital": -10,
        "owner_earnings": 940,
        "diluted_shares": 100,
        "owner_earnings_per_share": 9.40,
        "multiple": 14,
        "adjustments": [{"label": "net cash", "label_zh": "净现金", "per_share": 2.0}],
        "value_low": 118,
        "value_central": 134,
        "value_high": 150,
        "required_mos_pct": 15,
        "buy_price_value": 114,
        "mos_pct": -11.9,
        "buy_price": "$114 a share",
        "buy_price_zh": "每股 114 美元",
    }


def _docx_part(path: Path, name: str) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read(name).decode("utf-8")


def _docx_parts_text(path: Path, prefix: str) -> str:
    with zipfile.ZipFile(path) as archive:
        return "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith(prefix)
        )


def _render(tmp_path: Path, package: dict, **kwargs) -> tuple[dict, Path, Path]:
    run_dir = tmp_path / "run"
    (run_dir / "logs").mkdir(parents=True, exist_ok=True)
    (run_dir / "memo").mkdir(exist_ok=True)
    (run_dir / "logs" / "memo_package.json").write_text(
        json.dumps(package, ensure_ascii=False), encoding="utf-8"
    )
    en_docx = run_dir / "memo" / "en.docx"
    zh_docx = run_dir / "memo" / "zh.docx"
    rendered = buffett_memo_renderer.render_package(
        buffett_memo_renderer.load_package(run_dir),
        run_dir=run_dir,
        memo_paths={"en": en_docx, "zh": zh_docx},
        **kwargs,
    )
    return rendered, en_docx, zh_docx


def test_report_types_include_buffett_memo():
    assert memo_prep.BUFFETT_REPORT_TYPE in api_module.REPORT_TYPES
    assert memo_prep.REPORT_TYPE in api_module.REPORT_TYPES
    assert memo_prep.is_memo_report_type(memo_prep.BUFFETT_REPORT_TYPE)
    assert memo_prep.is_buffett_kind(memo_prep.BUFFETT_KIND)
    assert not memo_prep.is_buffett_kind(memo_prep.LATESTAGE_KIND)


def test_buffett_filename_and_run_suffix(tmp_path, monkeypatch):
    monkeypatch.setattr(memo_prep, "MEMOS_ROOT", tmp_path / "memos")
    name = memo_prep._memo_filename("Acme", "2026-08-20__120000", "en", buffett=True)
    assert name == "Acme - Buffett-Method Memo - 2026-08-20__120000.docx"
    zh = memo_prep._memo_filename("Acme", "2026-08-20__120000", "zh", buffett=True)
    assert zh == "Acme - 巴菲特方法备忘录 - 2026-08-20__120000.docx"
    run_dir = memo_prep._make_run_dir(
        "acme", "2026-08-20__120000", run_suffix="buffett-memo-run"
    )
    assert run_dir.name.endswith("__buffett-memo-run")
    assert (run_dir / "memo").is_dir()


def test_buffett_filename_is_one_path_component_for_edgar_names():
    for language in ("en", "zh"):
        name = memo_prep._memo_filename(
            "Occidental Petroleum Corp /De/", "2026-08-25__002314", language, buffett=True
        )
        assert "/" not in name and "\\" not in name
        assert Path(name).name == name
        assert name.startswith("Occidental Petroleum Corp - ")


def _prompt(tmp_path: Path, **kwargs) -> str:
    companies = tmp_path / "companies.yaml"
    companies.write_text("companies: []\n", encoding="utf-8")
    defaults = dict(
        run_dir=tmp_path,
        company_name="Example Co",
        company_slug="example-co",
        run_id="2026-08-20__120000",
        companies_yaml_path=companies,
        memo_paths={
            "en": str(tmp_path / "memo" / "en.docx"),
            "zh": str(tmp_path / "memo" / "zh.docx"),
        },
        today="2026-09-22",
    )
    defaults.update(kwargs)
    return claude_runner._build_buffett_investment_memo_prompt(**defaults)


def test_buffett_prompt_is_bsh_research_not_buffett(tmp_path):
    prompt = _prompt(
        tmp_path,
        scope_check={
            "outcome": "warn",
            "classification": "early-stage",
            "reason": "Latest funding round 'series a' is early-stage.",
        },
        warnings=["Latest funding round 'series a' is early-stage."],
    )
    assert "Warren Buffett" in prompt
    assert "You are BSH Research" in prompt
    assert "The author is BSH Research, applying Buffett's owner framework" in prompt
    assert "The author is Warren Buffett" not in prompt
    assert "as Warren Buffett." not in prompt
    assert "no Omaha dateline" in prompt
    assert "Buy / Pass / Too Hard" in prompt
    assert "Do not adopt BSH LP sell-side voice" in prompt
    assert "Today is 2026-09-22." in prompt
    # No venture stage note for an owner's analysis; a Security block instead.
    assert "Stage note from prep" not in prompt
    assert "early-stage" not in prompt
    assert "## Security" in prompt
    assert "name the listed parent" in prompt
    # The research floor replaced "use the registry and independent reasoning".
    assert "## Research before writing (required)" in prompt
    assert "independent reasoning" not in prompt
    assert "6–12" in prompt
    assert "owner's earnings" in prompt.lower()
    assert "bsh-buffett-investment-memo-v1" in prompt
    assert "Too Hard" in prompt
    assert "3–6 distinct ways capital can be lost forever" in prompt
    assert "permanent economic consequence" in prompt
    assert "must not use risk ratings" in prompt
    # The skill is loaded without its front matter.
    assert 'description: "BSH Research' not in prompt
    assert "Never state Buffett's or Berkshire Hathaway's holdings" in prompt


def test_buffett_prompt_carries_pinned_inputs_and_provenance(tmp_path):
    research_dir = tmp_path / "research"
    research_dir.mkdir()
    (research_dir / "fact_ledger.md").write_text("- 2026-06-30: revenue $120M\n", encoding="utf-8")
    (research_dir / "known_sources.md").write_text(
        "- Annual report — https://example.com/ar\n", encoding="utf-8"
    )
    prompt = _prompt(
        tmp_path,
        research_dir=research_dir,
        security={"name": "Coca-Cola", "ticker": "KO", "exchange": "NYSE", "status": "public"},
        market_inputs={
            "ticker": "KO",
            "price": 88.61,
            "currency": "USD",
            "price_as_of": "2026-09-22",
            "price_source": "cnbc",
            "ust10y": 4.955,
            "ust10y_as_of": "2026-09-22",
            "ust10y_source": "cnbc",
        },
    )
    assert "- Ticker: KO" in prompt
    assert "Share price: KO USD 88.61 as of 2026-09-22 (source: cnbc)" in prompt
    assert "10-year US Treasury yield: 4.955% as of 2026-09-22" in prompt
    assert "Hurdle r = max(4.955% + 3%, 10%) = 10.00%" in prompt
    assert "## Curated fact ledger" in prompt and "revenue $120M" in prompt
    assert "## Known sources" in prompt and "https://example.com/ar" in prompt
    assert str(research_dir) in prompt


def test_buffett_prompt_unpinned_inputs_ask_for_model_sourced_dated_figures(tmp_path):
    prompt = _prompt(tmp_path, market_inputs={"ticker": "KO", "price": None, "ust10y": None})
    assert "not available for KO at run start" in prompt
    assert "model-sourced" in prompt
    assert "10-year US Treasury yield: not available at run start" in prompt


def test_buffett_prompt_asks_for_local_language_research(tmp_path):
    prompt = _prompt(
        tmp_path,
        company_name="CEInet Data",
        security={"name": "CEInet Data", "hq": "Beijing, China", "status": "subsidiary"},
    )
    assert "search in Simplified Chinese as well as English" in prompt
    generic = _prompt(tmp_path, security={"name": "Example Co", "hq": "Austin, Texas"})
    assert "search in its local" in generic


def test_buffett_runner_allows_web_tools_counts_lookups_and_captures_sources(
    tmp_path, monkeypatch
):
    run_dir = tmp_path / "run"
    (run_dir / "logs").mkdir(parents=True)
    companies = tmp_path / "companies.yaml"
    companies.write_text("companies: []\n", encoding="utf-8")
    events = [
        {
            "type": "assistant",
            "message": {
                "content": [
                    {"type": "tool_use", "id": "t1", "name": "WebSearch", "input": {"query": "example co revenue"}},
                    {"type": "tool_use", "id": "t2", "name": "WebFetch", "input": {"url": "https://example.com/about"}},
                    {"type": "tool_use", "id": "t3", "name": "Read", "input": {"file_path": "x"}},
                ]
            },
        },
        {
            "type": "user",
            "message": {
                "content": [
                    {
                        "type": "tool_result",
                        "tool_use_id": "t2",
                        "content": "Example Co makes industrial software. Revenue was $120 million in 2025.",
                    }
                ]
            },
        },
        {"type": "result", "subtype": "success", "total_cost_usd": 0.5, "duration_ms": 1000},
    ]

    class FakeProc:
        stdout = iter(json.dumps(event) for event in events)
        stderr = iter(())
        returncode = 0
        pid = 12345

        def wait(self, timeout=None):
            return self.returncode

        def poll(self):
            return self.returncode

    captured: dict = {}

    def fake_popen(args, **kwargs):
        captured["cmd"] = args
        return FakeProc()

    monkeypatch.setattr(claude_runner, "is_available", lambda: True)
    monkeypatch.setattr(claude_runner, "claude_path", lambda: "claude")
    monkeypatch.setattr(claude_runner.subprocess, "Popen", fake_popen)
    research_dir = buffett_memo_analysis._research_dir("example-co")
    claude_runner.register_memo_run_source_capture(
        run_dir, company_id="example-co", run_id="r1", research_dir=research_dir
    )

    result = claude_runner.run_buffett_investment_memo(
        run_dir=run_dir,
        company_name="Example Co",
        company_slug="example-co",
        run_id="r1",
        companies_yaml_path=companies,
        memo_paths={"en": str(run_dir / "en.docx"), "zh": str(run_dir / "zh.docx")},
        research_dir=research_dir,
        today="2026-09-22",
    )

    assert result["ok"] is True
    assert result["web_lookups"] == 2
    assert result["retrieved_sources"] == 1
    allowed = captured["cmd"][captured["cmd"].index("--allowedTools") + 1]
    assert "WebSearch" in allowed and "WebFetch" in allowed
    manifest = (run_dir / "sources" / "manifest.jsonl").read_text(encoding="utf-8")
    assert "https://example.com/about" in manifest


def test_buffett_package_renders_docx_with_bsh_frame(tmp_path):
    rendered, en_docx, zh_docx = _render(tmp_path, _package(run={"run_id": "2026-08-20__120000"}))
    assert rendered["ok"] is True
    assert rendered["decision"] == "Pass"
    assert en_docx.exists() and zh_docx.exists()
    assert (tmp_path / "run" / "memo" / "buffett_memo.en.md").exists()

    english = Document(str(en_docx))
    text = "\n".join(p.text for p in english.paragraphs)
    assert "Investment Decision" in text
    assert "Pass" in text
    # **bold** is a bold run, not stripped markers.
    decision_para = next(p for p in english.paragraphs if p.text.startswith("Pass."))
    assert decision_para.runs[0].bold is True
    assert "**" not in text
    # Decision box under the title.
    box = english.tables[0]
    assert box.cell(0, 0).text == "Call" and box.cell(0, 1).text == "Pass"
    assert box.cell(1, 0).text == "Buy price" and box.cell(1, 1).text == "$8 a share"
    header = _docx_parts_text(en_docx, "word/header")
    assert "Example Co | BSH Research — Owner's analysis (Buffett method)" in header
    assert "DRAFT — AI-generated, not reviewed" in header
    footer = _docx_parts_text(en_docx, "word/footer")
    assert "BSH Confidential" in footer
    assert "not written or endorsed by Warren Buffett or Berkshire Hathaway" in footer
    assert "NUMPAGES" in footer and " PAGE " in footer
    props = english.core_properties
    assert props.author == "Berkeley Summit House"
    assert props.title == "Example Co — Buffett-Method Memo"
    assert props.created.date().isoformat() == "2026-08-20"  # the run's date
    assert "BSH Confidential · 2026-08-20 · EN" in footer
    styles = _docx_part(en_docx, "word/styles.xml")
    assert 'w:color w:val="1B2A4A"' in styles
    assert 'w:ascii="Arial"' in styles

    chinese = Document(str(zh_docx))
    assert chinese.tables[0].cell(0, 0).text == "结论"
    assert chinese.tables[0].cell(0, 1).text == "放弃"
    # An English-only buy price (legacy package) reads in Chinese units.
    assert chinese.tables[0].cell(1, 1).text == "8 美元（详见第一节）"
    zh_header = _docx_parts_text(zh_docx, "word/header")
    assert "Example Co | BSH 研究 — 巴菲特方法所有者分析" in zh_header
    assert "草稿 — AI 生成，未经审阅" in zh_header
    zh_footer = _docx_parts_text(zh_docx, "word/footer")
    assert "BSH 机密" in zh_footer and "第 " in zh_footer and "页，共 " in zh_footer
    assert "并非由沃伦·巴菲特或伯克希尔·哈撒韦公司撰写或认可" in zh_footer
    zh_styles = _docx_part(zh_docx, "word/styles.xml")
    assert 'w:eastAsia="Microsoft YaHei"' in zh_styles
    assert 'w:eastAsia="zh-CN"' in zh_styles
    assert "eastAsiaTheme" not in zh_styles.split("<w:latentStyles")[0].split("</w:docDefaults>")[0]
    assert 'w:eastAsia="zh-CN"' in _docx_part(zh_docx, "word/settings.xml")
    assert chinese.core_properties.language == "zh-CN"


def test_buffett_render_draws_arithmetic_sources_and_price_rows(tmp_path):
    package = _package(
        **_valuation_fields(),
        markdown_en=_english_memo(
            extra_viii=(
                " The hurdle is 10%. We value the business at $118 to $150 a share, "
                "centred on $134; with a 15% required margin the buy price is $114, "
                "against $150 on 2026-09-21 [1]."
            )
        ),
        sources=[{"n": 1, "title": "Example Co annual report 2025"}],
    )
    run_dir = tmp_path / "run"
    (run_dir / "sources").mkdir(parents=True)
    (run_dir / "sources" / "manifest.jsonl").write_text(
        json.dumps(
            {
                "at": "2026-09-21T10:00:00+00:00",
                "tool": "WebFetch",
                "url": "https://example.com/annual-report-2025",
                "title": "Example Co Annual Report 2025",
                "file": "abc.txt",
            }
        )
        + "\n",
        encoding="utf-8",
    )
    rendered, en_docx, zh_docx = _render(tmp_path, package)
    assert rendered["pass_kind"] == "price"
    assert rendered["call_label"] == {
        "en": "Pass at today's price — buy at or below $114",
        "zh": "暂不买入（买入价 ≤ 114 美元）",
    }
    assert rendered["sources"][0]["url"] == "https://example.com/annual-report-2025"

    english = Document(str(en_docx))
    box = {row.cells[0].text: row.cells[1].text for row in english.tables[0].rows}
    assert box["Call"] == "Pass at today's price — buy at or below $114"
    assert box["Price & date"] == "$150 (2026-09-21)"
    assert box["Value range"] == "$118 – $150 (central $134)"
    assert box["Margin of safety"] == "−11.9% at today's price; 15% required"
    text = "\n".join(p.text for p in english.paragraphs)
    assert "The arithmetic" in text
    arithmetic = {row.cells[0].text: row.cells[1].text for row in english.tables[1].rows}
    assert arithmetic["= Owner's earnings"] == "$940 million"
    assert arithmetic["Owner's earnings per share"] == "$9.40"
    assert arithmetic["Buy price = central × (1 − required)"] == "$114"
    assert "Sources" in text
    assert "[1] Example Co annual report 2025 — https://example.com/annual-report-2025 (accessed 2026-09-21)" in text
    # The arithmetic sits inside section VIII, before section IX.
    headings = [p.text for p in english.paragraphs if p.style.name.startswith("Heading")]
    assert headings.index("The arithmetic") < headings.index("IX. What Can Go Permanently Wrong")

    chinese = Document(str(zh_docx))
    zh_box = {row.cells[0].text: row.cells[1].text for row in chinese.tables[0].rows}
    assert zh_box["结论"] == "暂不买入（买入价 ≤ 114 美元）"
    assert zh_box["买入价"] == "每股 114 美元"
    assert zh_box["股价与日期"] == "150 美元（2026-09-21）"
    zh_text = "\n".join(p.text for p in chinese.paragraphs)
    assert "估值算术" in zh_text and "资料来源" in zh_text
    zh_arithmetic = {row.cells[0].text: row.cells[1].text for row in chinese.tables[1].rows}
    assert zh_arithmetic["= 股东盈余"] == "9.4 亿美元"
    assert zh_arithmetic["调整项：净现金"] == "每股 +2 美元"

    record = buffett_memo_renderer.report_fields(rendered)
    assert record["decision"] == "Pass"
    assert record["pass_kind"] == "price"
    assert record["buy_price"] == "$114 a share"
    assert record["buffett_valuation"]["price"] == 150.0
    assert record["buffett_valuation"]["value_central"] == 134.0
    assert record["buffett_valuation"]["price_date"] == "2026-09-21"


def test_review_stamp_wording():
    assert buffett_memo_renderer.review_stamp(None)["en"] == "DRAFT — AI-generated, not reviewed"
    approved = buffett_memo_renderer.review_stamp(
        {"state": "approved", "reviewer": "Ann Lee", "reviewed_at": "2026-09-20T10:00:00Z"}
    )
    assert approved["en"] == "Reviewed by Ann Lee, 2026-09-20"
    assert approved["zh"] == "已审阅：Ann Lee，2026-09-20"
    withdrawn = buffett_memo_renderer.review_stamp({"state": "withdrawn"})
    assert withdrawn["en"] == "WITHDRAWN — do not rely on this memo"
    assert withdrawn["zh"] == "已撤回 — 请勿依据本备忘录"
    assert buffett_memo_renderer.review_stamp({"state": "in_review"})["state"] == "draft"


def test_render_uses_review_from_package_or_argument(tmp_path):
    package = _package(run={"review": {"state": "withdrawn"}})
    _rendered, en_docx, _zh = _render(tmp_path, package)
    assert "WITHDRAWN — do not rely on this memo" in _docx_parts_text(en_docx, "word/header")
    rendered, en_docx, zh_docx = _render(
        tmp_path,
        package,
        review={"state": "approved", "reviewer": "Ann Lee", "reviewed_at": "2026-09-20"},
    )
    assert rendered["review_stamp"]["state"] == "approved"
    assert "Reviewed by Ann Lee, 2026-09-20" in _docx_parts_text(en_docx, "word/header")
    assert "已审阅：Ann Lee，2026-09-20" in _docx_parts_text(zh_docx, "word/header")


def test_legacy_package_rerenders_purely(tmp_path):
    legacy = _package(
        company_name="Occidental Petroleum Corp /De/",
        markdown_en=_english_memo().replace(
            "# Investment Memorandum — Example Co\n\nBSH Research · Owner's analysis (Buffett method) · prices as of 2026-08-20",
            "# Investment Memorandum — Occidental Petroleum Corp /De/\n\nOmaha, August 2026",
        ),
    )
    first, en_docx, _ = _render(tmp_path, legacy)
    document_xml = _docx_part(en_docx, "word/document.xml")
    header_xml = _docx_parts_text(en_docx, "word/header")
    again, en_docx, _ = _render(tmp_path, legacy)
    assert _docx_part(en_docx, "word/document.xml") == document_xml
    assert _docx_parts_text(en_docx, "word/header") == header_xml
    assert first["pass_kind"] is None and first["fields"] == {}
    assert first["call_label"] == {"en": "Pass", "zh": "放弃"}
    assert first["company_name"] == "Occidental Petroleum Corp"
    assert "Occidental Petroleum Corp | BSH Research" in header_xml
    title = Document(str(en_docx)).paragraphs[0].text
    assert title == "Investment Memorandum — Occidental Petroleum Corp"
    # The stored memo text is not rewritten.
    assert "/De/" in (tmp_path / "run" / "memo" / "buffett_memo.en.md").read_text(encoding="utf-8")


def test_buffett_package_rejects_bsh_structure():
    package = {
        "schema_version": 1,
        "kind": "buffett_investment_memo",
        "company_name": "Example Co",
        "decision": "Buy",
        "markdown_en": "# Not a Buffett memo\n\nWe recommend participating.\n",
        "markdown_zh": "# 不是巴菲特备忘录\n\n我们建议参与。\n",
    }
    try:
        buffett_memo_renderer.validate_package(package)
    except buffett_memo_renderer.BuffettMemoPackageError as exc:
        message = str(exc)
        assert "too short" in message or "missing heading" in message
    else:
        raise AssertionError("expected BuffettMemoPackageError")


def test_malformed_structured_fields_are_dropped_not_errors():
    validated = buffett_memo_renderer.validate_package(
        _package(price="n/a", price_date="soon", pass_kind="maybe", sources="nope", hurdle=0.1)
    )
    assert "price" not in validated["fields"]
    assert "price_date" not in validated["fields"]
    assert validated["pass_kind"] is None
    assert validated["sources"] == []
    assert validated["fields"]["hurdle"] == 10.0
    nested = buffett_memo_renderer.validate_package(
        _package(valuation={"price": "$92.10", "value_central": 80, "currency": "usd"})
    )
    assert nested["fields"]["price"] == 92.1
    assert nested["fields"]["currency"] == "USD"
    assert nested["fields"]["value_basis"] == "per_share"


def test_check_valuation_reconciles_the_arithmetic():
    clean = _package(**_valuation_fields())
    assert buffett_memo_renderer.check_valuation(clean) == [] or all(
        finding["severity"] == "P2" for finding in buffett_memo_renderer.check_valuation(clean)
    )

    def codes(**changes) -> set[str]:
        fields = {**_valuation_fields(), **changes}
        return {f["code"] for f in buffett_memo_renderer.check_valuation(_package(**fields))}

    assert "owner_earnings_mismatch" in codes(owner_earnings=1200)
    assert "owner_earnings_no_d_and_a" in codes(d_and_a=None, owner_earnings=750, working_capital=None)
    assert "intrinsic_value_mismatch" in codes(multiple=20)
    assert "buy_price_above_mos_rule" in codes(buy_price_value=130)
    assert "buy_above_buy_price" in codes(decision="Buy", pass_kind=None)
    assert "pass_without_buy_price" in codes(buy_price_value=None, buy_price="")
    assert "mos_mismatch" in codes(mos_pct=25)
    assert "hurdle_mismatch" in codes(hurdle=8)
    assert "long_run_growth_above_cap" in codes(g=6)
    assert "value_range_inverted" in codes(value_low=160)
    for finding in buffett_memo_renderer.check_valuation(_package(**{**_valuation_fields(), "owner_earnings": 1200})):
        assert finding["severity"] in {"P1", "P2"}
    # Legacy packages carry no fields: nothing to check.
    assert buffett_memo_renderer.check_valuation(_package()) == []


def test_valuation_figures_must_appear_in_section_viii():
    findings = buffett_memo_renderer.check_valuation(_package(**_valuation_fields()))
    missing = [f for f in findings if f["code"] == "valuation_figure_not_in_section_viii"]
    assert missing and "value_central" in missing[0]["snippet"]
    stated = _package(
        **_valuation_fields(),
        markdown_en=_english_memo(
            extra_viii=(
                " The hurdle is 10%. Value runs $118 to $150 a share, centred on $134. "
                "With a 15% required margin the buy price is $114; the price was $150."
            )
        ),
    )
    codes = {f["code"] for f in buffett_memo_renderer.check_valuation(stated)}
    assert "valuation_figure_not_in_section_viii" not in codes


def test_voice_checks_flag_the_buffett_persona():
    legacy_en = _english_memo().replace(
        "BSH Research · Owner's analysis (Buffett method) · prices as of 2026-08-20",
        "Omaha, August 2026",
    ).replace(
        "The balance sheet is not stretched.",
        "Berkshire's 400 million shares are not for sale, and we were adding shares in late 2024. "
        "Charlie and I agreed on that. The balance sheet is not stretched.",
    )
    legacy_zh = _chinese_memo().replace(
        "BSH 研究 · 巴菲特方法所有者分析 · 价格截至 2026 年 8 月 20 日", "奥马哈，2026年8月"
    ).replace("资产负债表并不紧。", "伯克希尔的股份我们持有多年。资产负债表并不紧。")
    findings = buffett_checks.voice_findings(legacy_en, legacy_zh)
    codes = [f.code for f in findings]
    assert codes.count("persona_omaha_dateline") == 2
    assert "persona_berkshire_first_person" in codes
    assert "persona_buffett_first_person" in codes
    assert any(f.location == "markdown_zh" and f.code == "persona_berkshire_first_person" for f in findings)
    assert all(f.severity in {"P1", "P2"} for f in findings)

    third_party = _english_memo().replace(
        "The balance sheet is not stretched.",
        "Berkshire Hathaway reported owning 400 million shares at June 30 [3]. "
        "We would own the business at a lower price than Berkshire paid. "
        "If we owned the whole company, we would want the debt gone. The balance sheet is not stretched.",
    )
    assert buffett_checks.voice_findings(third_party, _chinese_memo()) == []

    # BSH has no position: a first-person holding is the persona even
    # without Berkshire's name in the sentence.
    holdings = _english_memo().replace(
        "The balance sheet is not stretched.",
        "We hold warrants on another 83.9 million shares. In January we bought the chemicals unit. "
        "The balance sheet is not stretched.",
    )
    zh_holdings = _chinese_memo().replace("资产负债表并不紧。", "我们仍然持有该公司 85 亿美元的优先股。资产负债表并不紧。")
    hits = [
        f for f in buffett_checks.voice_findings(holdings, zh_holdings)
        if f.code == "persona_berkshire_first_person"
    ]
    assert len([f for f in hits if f.location == "markdown_en"]) == 2
    assert len([f for f in hits if f.location == "markdown_zh"]) == 1


def test_stock_phrases_are_flagged():
    text = _english_memo().replace(
        "The moat is narrow.",
        "The moat is narrow, not a toll bridge. If the market closed tomorrow for ten years we would not mind.",
    )
    zh = _chinese_memo().replace("护城河很窄。", "护城河很窄，不是收费桥。")
    codes = [f.code for f in buffett_checks.stock_phrase_findings(text, zh)]
    assert codes.count("stock_phrase_toll_bridge") == 2
    assert "stock_phrase_market_closed" in codes
    zh_closed = _chinese_memo().replace("护城河很窄。", "假如市场明天关门，我们睡得着吗？")
    zh_codes = [
        f.code
        for f in buffett_checks.stock_phrase_findings(_english_memo(), zh_closed)
        if f.location == "markdown_zh"
    ]
    assert zh_codes == ["stock_phrase_market_closed"]


def test_assumption_only_prices_must_be_labelled():
    unlabelled = buffett_memo_renderer.validate_package(
        _package(
            valuation_basis="assumed",
            markdown_en=_english_memo().replace(
                "We would not buy the whole company", "We would pay $75 million for the whole company"
            ),
        )
    )
    findings = buffett_checks.price_basis_findings(unlabelled)
    assert any(f.code == "assumed_price_unlabelled" and f.severity == "P1" for f in findings)
    labelled = buffett_memo_renderer.validate_package(
        _package(
            valuation_basis="assumed",
            markdown_en=_english_memo().replace(
                "We would not buy the whole company",
                "We would pay $75 million (illustrative, built from assumptions) for the whole company",
            ),
        )
    )
    assert not [
        f for f in buffett_checks.price_basis_findings(labelled) if f.location.startswith("I.")
    ]
    heuristic = buffett_memo_renderer.validate_package(
        _package(
            markdown_en=_english_memo().replace(
                "We would not buy the whole company", "We would pay $75 million for the whole company"
            ),
        )
    )
    assert [f.code for f in buffett_checks.price_basis_findings(heuristic)] == ["price_basis_unstated"]


def test_number_parity_compares_units_across_languages():
    en = _english_memo().replace(
        "The balance sheet is not stretched.",
        "Revenue was $14.3 billion, the margin 28.6%, and 106 million shares trade at 32 times earnings.",
    )
    zh_ok = _chinese_memo().replace(
        "资产负债表并不紧。", "收入为 143 亿美元，利润率 28.6%，1.06 亿股，估值 32 倍。"
    )
    validated = buffett_memo_renderer.validate_package(_package(markdown_en=en, markdown_zh=zh_ok))
    parity = buffett_checks.number_parity(validated)
    assert not [f for f in parity["findings"] if f["code"] == "zh_number_not_in_en"]
    assert parity["p0_count"] == 0 and parity["status"] == "passed"

    zh_bad = _chinese_memo().replace("资产负债表并不紧。", "收入为 150 亿美元，利润率 28.6%。")
    parity = buffett_checks.number_parity(
        buffett_memo_renderer.validate_package(_package(markdown_en=en, markdown_zh=zh_bad))
    )
    wrong = [f for f in parity["findings"] if f["code"] == "zh_number_not_in_en"]
    assert wrong and "150" in wrong[0]["snippet"]
    missing = [f for f in parity["findings"] if f["code"] == "en_numbers_missing_in_zh"]
    assert missing and "$14.3 billion" in missing[0]["snippet"]


def _corpus(texts: list[str]):
    from server import memo_fact_check

    corpus = memo_fact_check.Corpus()
    for index, text in enumerate(texts):
        corpus.add(f"doc {index}", text)
    return corpus.finish()


def test_fact_check_traces_figures_to_sources(tmp_path, monkeypatch):
    from server import memo_fact_check

    memo = _english_memo().replace(
        "The balance sheet is not stretched.",
        "Revenue was $120 million [1]. Net cash is $45 million. Debt is $999 million.",
    )
    validated = buffett_memo_renderer.validate_package(
        _package(**_valuation_fields(), markdown_en=memo, sources=[{"n": 1, "title": "Annual report"}])
    )
    corpus = _corpus(["Net cash stood at $45 million at year end. " * 80])
    monkeypatch.setattr(memo_fact_check, "build_corpus", lambda *args, **kwargs: corpus)
    monkeypatch.setattr(
        buffett_checks, "_source_texts_by_n", lambda *args, **kwargs: {1: "Revenue was $120 million in 2025."}
    )
    payload, report = buffett_checks.fact_check(validated, company_id="example-co", run_dir=tmp_path)
    assert payload["verified"] >= 1
    assert payload["supported"] >= 1
    assert payload["unsupported"] >= 1
    unsupported = [f for f in payload["findings"] if f["code"] == "unsupported_figure"]
    assert any(f["figure"] == "$999 million" for f in unsupported)
    assert unsupported[0]["section_id"] == "VII. Financial Position"
    assert payload["p0_count"] == len(unsupported)
    assert payload["repair_feed"] is False
    assert payload["memo_kind"] == "buffett_investment_memo"
    assert "# Memo fact check" in report

    thin = _corpus(["short"])
    monkeypatch.setattr(memo_fact_check, "build_corpus", lambda *args, **kwargs: thin)
    payload, _ = buffett_checks.fact_check(validated, company_id="example-co", run_dir=tmp_path)
    assert payload["thin_corpus"] is True and payload["p0_count"] == 0


def test_web_lookup_count_reads_current_and_archived_streams(tmp_path):
    logs = tmp_path / "logs"
    logs.mkdir()
    assert buffett_checks.web_lookup_count(tmp_path) is None
    event = {"type": "claude_action", "action": "tool_use", "tool": "WebSearch", "preview": "q"}
    (logs / "stream.before_resume.20260101T000000Z.jsonl").write_text(
        json.dumps(event) + "\n", encoding="utf-8"
    )
    (logs / "stream.jsonl").write_text(
        json.dumps({**event, "tool": "Read"}) + "\n" + json.dumps({**event, "tool": "WebFetch"}) + "\n",
        encoding="utf-8",
    )
    assert buffett_checks.web_lookup_count(tmp_path) == 2


def test_pin_market_inputs_pins_price_and_yield(tmp_path, monkeypatch):
    from server import live_quotes

    (tmp_path / "logs").mkdir()
    monkeypatch.setattr(
        live_quotes,
        "fetch_quotes",
        lambda tickers: {
            "quotes": {
                "KO": {"last_price": 88.61, "currency": "USD", "as_of": "2026-09-22T20:10:00Z", "source": "cnbc"},
                "US10Y": {"last_price": 4.955, "as_of": "2026-09-22T21:05:00Z", "source": "cnbc"},
            }
        },
    )
    monkeypatch.setattr(buffett_memo_analysis, "_yahoo_tnx", lambda: (_ for _ in ()).throw(AssertionError))
    inputs = buffett_memo_analysis._pin_market_inputs({"ticker": "ko"}, tmp_path)
    assert inputs["status"] == "pinned"
    assert inputs["price"] == 88.61 and inputs["price_as_of"] == "2026-09-22"
    assert inputs["ust10y"] == 4.955 and inputs["ust10y_as_of"] == "2026-09-22"
    assert buffett_memo_analysis._load_market_inputs(tmp_path)["price"] == 88.61


def test_pin_market_inputs_records_unavailable_quotes(tmp_path, monkeypatch):
    from server import live_quotes

    (tmp_path / "logs").mkdir()

    def boom(*_args, **_kwargs):
        raise RuntimeError("offline")

    monkeypatch.setattr(live_quotes, "fetch_quotes", boom)
    monkeypatch.setattr(buffett_memo_analysis, "_yahoo_tnx", boom)
    inputs = buffett_memo_analysis._pin_market_inputs({"ticker": "KO"}, tmp_path)
    assert inputs["status"] == "unavailable"
    assert inputs["price"] is None and inputs["ust10y"] is None
    findings = buffett_checks.market_input_findings(inputs)
    assert findings and findings[0].code == "inputs_not_pinned"
    private = buffett_memo_analysis._pin_market_inputs({}, tmp_path)
    assert private["ticker"] is None and private["status"] == "unavailable"


def test_pin_market_inputs_uses_the_tnx_fallback(tmp_path, monkeypatch):
    from server import live_quotes

    (tmp_path / "logs").mkdir()
    monkeypatch.setattr(live_quotes, "fetch_quotes", lambda tickers: {"quotes": {}})
    monkeypatch.setattr(
        buffett_memo_analysis,
        "_yahoo_tnx",
        lambda: {"last_price": 4.2, "as_of": "2026-09-21T20:00:00Z", "source": "yahoo ^TNX"},
    )
    inputs = buffett_memo_analysis._pin_market_inputs({}, tmp_path)
    assert inputs["status"] == "pinned"
    assert inputs["ust10y"] == 4.2 and inputs["ust10y_source"] == "yahoo ^TNX"


def _finalize_env(tmp_path, monkeypatch, package: dict, *, report_extra: dict | None = None):
    from server import memo_fact_check

    data_root = tmp_path / "data"
    monkeypatch.setattr(memo_prep, "DATA_DIR", data_root)
    monkeypatch.setattr(memo_fact_check, "build_corpus", lambda *a, **k: _corpus(["x"]))
    storage._write_yaml(
        storage.COMPANIES_FILE, [{"id": "example-co", "name": "Example Co", "status": "private"}]
    )
    run_dir = data_root / "memos" / "example-co" / "2026-08-20__120000__example-co__buffett-memo-run"
    (run_dir / "logs").mkdir(parents=True)
    (run_dir / "memo").mkdir()
    (run_dir / "logs" / "memo_package.json").write_text(
        json.dumps(package, ensure_ascii=False), encoding="utf-8"
    )
    report = storage.create_report(
        company_id="example-co", report_type=memo_prep.BUFFETT_REPORT_TYPE, audience="Internal"
    )
    rel = run_dir.relative_to(data_root.parent)
    fields = {
        "kind": memo_prep.BUFFETT_KIND,
        "status": "analyzing",
        "run_dir": str(rel),
        "run_id": "2026-08-20__120000",
        "memo_files": [
            {"language": "en", "path": str(rel / "memo" / "en.docx")},
            {"language": "zh", "path": str(rel / "memo" / "zh.docx")},
        ],
    }
    fields.update(report_extra or {})
    storage.update_report(report["id"], **fields)
    return storage.get_report(report["id"]), run_dir


class _Stream:
    def __init__(self):
        self.events: list[dict] = []

    def emit(self, type_, **fields):
        self.events.append({"type": type_, **fields})


def test_finalize_persists_call_fields_and_check_results(tmp_path, monkeypatch):
    package = _package(
        **_valuation_fields(),
        markdown_en=_english_memo(
            extra_viii=(
                " The hurdle is 10%. Value runs $118 to $150 a share, centred on $134. "
                "With a 15% required margin the buy price is $114; the price was $150."
            )
        ),
    )
    report, run_dir = _finalize_env(tmp_path, monkeypatch, package)
    stream = _Stream()
    ok = buffett_memo_analysis._finalize_from_package(
        report_id=report["id"],
        report=report,
        run_dir=run_dir,
        stream=stream,
        result={"cost_usd": 1.0, "duration_ms": 10, "web_lookups": 7},
    )
    assert ok is True
    saved = storage.get_report(report["id"])
    assert saved["status"] == "complete"
    assert saved["decision"] == "Pass"
    assert saved["pass_kind"] == "price"
    assert saved["buy_price"] == "$114 a share"
    assert saved["buy_price_zh"] == "每股 114 美元"
    assert saved["call_label"]["zh"] == "暂不买入（买入价 ≤ 114 美元）"
    assert saved["buffett_valuation"]["buy_price_value"] == 114.0
    assert saved["web_lookups"] == 7
    assert saved["memo_quality_lint"]["p0_count"] == 0
    assert saved["memo_chinese_parity"]["status"] == "passed"
    assert saved["memo_fact_check"]["memo_kind"] == "buffett_investment_memo"
    assert saved["quality_warnings"] is None
    assert saved["renderer_version"] == buffett_memo_renderer.RENDERER_VERSION
    assert (run_dir / "logs" / "buffett_checks.md").exists()
    assert (run_dir / "logs" / "fact_check.json").exists()
    assert stream.events[-1]["type"] == "done" and stream.events[-1]["status"] == "complete"


def test_finalize_warns_on_zero_lookups_and_the_buffett_persona(tmp_path, monkeypatch):
    legacy = _package(
        markdown_en=_english_memo().replace(
            "BSH Research · Owner's analysis (Buffett method) · prices as of 2026-08-20",
            "Omaha, August 2026",
        )
    )
    report, run_dir = _finalize_env(tmp_path, monkeypatch, legacy)
    ok = buffett_memo_analysis._finalize_from_package(
        report_id=report["id"],
        report=report,
        run_dir=run_dir,
        stream=_Stream(),
        result={"web_lookups": 0},
    )
    assert ok is True
    saved = storage.get_report(report["id"])
    assert saved["status"] == "complete_with_warnings"
    assert saved["stage"] == "Memo ready (quality warnings)"
    assert "Written without any external lookups" in saved["quality_warnings"]
    assert "未进行任何外部检索" in saved["quality_warnings_zh"]
    assert any("Omaha" in warning for warning in saved["quality_warnings"])
    assert saved["pass_kind"] is None and saved["buffett_valuation"] is None
    assert saved["call_label"] == {"en": "Pass", "zh": "放弃"}


def test_recover_leaves_complete_with_warnings_alone(tmp_path, monkeypatch):
    report, _run_dir = _finalize_env(
        tmp_path, monkeypatch, _package(), report_extra={"status": "complete_with_warnings"}
    )
    called = []
    monkeypatch.setattr(
        buffett_memo_analysis, "_finalize_from_package", lambda **kwargs: called.append(kwargs)
    )
    assert buffett_memo_analysis.recover_stale_reports() == 0
    assert called == []
    assert storage.get_report(report["id"])["status"] == "complete_with_warnings"


def test_internal_memo_default_render_is_unchanged(tmp_path):
    md_path = tmp_path / "internal.md"
    md_path.write_text(
        "# Internal Diligence Memo\n\n## Recommendation\n\n**Proceed** with a staged allocation. "
        + "Internal use only. " * 20,
        encoding="utf-8",
    )
    out = tmp_path / "internal.docx"
    internal_memo_renderer.render_internal_memo(md_path, out)
    document = Document(str(out))
    body = next(p for p in document.paragraphs if p.text.startswith("Proceed"))
    assert all(run.bold is None for run in body.runs)
    assert document.core_properties.author == "python-docx"
    assert document.styles["Normal"].font.name == "Aptos"
    assert "NUMPAGES" not in _docx_parts_text(out, "word/")
    assert not any(name.startswith("word/header") for name in zipfile.ZipFile(out).namelist())


def test_buffett_bootstrap_dispatches_to_buffett_worker(tmp_path, monkeypatch):
    started: list[tuple[str, str]] = []
    monkeypatch.setattr(storage, "DATA_DIR", tmp_path)
    monkeypatch.setattr(storage, "COMPANIES_FILE", tmp_path / "companies.yaml")
    monkeypatch.setattr(storage, "REPORTS_DIR", tmp_path / "reports")
    monkeypatch.setattr(memo_prep, "DATA_DIR", tmp_path)
    monkeypatch.setattr(memo_prep, "MEMOS_ROOT", tmp_path / "memos")
    monkeypatch.setattr(
        memo_prep,
        "SETTINGS_FILE",
        tmp_path / "settings" / "serena_background.md",
    )
    monkeypatch.setattr(memo_prep, "COMPANIES_FILE", tmp_path / "companies.yaml")
    memo_prep.SETTINGS_FILE.parent.mkdir(parents=True)
    memo_prep.SETTINGS_FILE.write_text("background\n", encoding="utf-8")
    storage._write_yaml(
        storage.COMPANIES_FILE,
        [
            {
                "id": "example-co",
                "name": "Example Co /DE/",
                "status": "private",
                "description": "A software company.",
                "latest_funding": {"round": "Series A"},
            }
        ],
    )
    monkeypatch.setattr(
        buffett_memo_analysis,
        "start_analysis",
        lambda report_id: started.append(("buffett", report_id)),
    )
    from server import memo_analysis

    monkeypatch.setattr(
        memo_analysis,
        "start_analysis",
        lambda report_id: started.append(("latestage", report_id)),
    )

    result = memo_prep.bootstrap_memo_run(
        "example-co",
        report_type=memo_prep.BUFFETT_REPORT_TYPE,
    )
    assert result["failed"] is False
    report = storage.get_report(result["report_id"])
    assert report["kind"] == memo_prep.BUFFETT_KIND
    assert report["report_type"] == memo_prep.BUFFETT_REPORT_TYPE
    assert report["skill"] == memo_prep.BUFFETT_SKILL_NAME
    assert started == [("buffett", report["id"])]
    assert "buffett-memo-run" in str(result["run_dir"])
    manifest = Path(result["manifest_path"]).read_text(encoding="utf-8")
    assert "bsh-buffett-investment-memo-v1" in manifest
    assert "serena_background.md" not in manifest
    # Owner's analysis: no venture stage warning, the clean display name.
    assert report["scope_check"]["outcome"] == "pass"
    assert report["scope_check"]["classification"] == "private"
    assert report["company_name"] == "Example Co"
    assert all("/" not in Path(entry["path"]).name for entry in report["memo_files"])
    assert Path(report["memo_files"][0]["path"]).name.startswith("Example Co - Buffett-Method Memo - ")
