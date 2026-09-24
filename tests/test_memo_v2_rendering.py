"""The renderer on the memo every run now writes (v2 is on by default).

Renders the realistic late_v2 and late_compact packages from
``memo_v2_fixture`` — twelve / seven sections with their numbered
subsections, key-metrics, deal, scenario and scorecard tables, inline
[S#]/[C#] citations, calculation notes, 8-row risk cards, charts and a
sources list with URLs — strict and lenient, in both languages, and
checks what a reader and the gates see.
"""
from __future__ import annotations

import re

import pytest
from docx import Document
from docx.oxml.ns import qn

import memo_v2_fixture as fx
from server import (
    memo_chinese_parity,
    memo_docx_renderer,
    memo_quality_lint,
    memo_structure,
)

PACKAGES = {
    "late_v2": fx.late_v2_package,
    "late_compact": fx.late_compact_package,
}


@pytest.fixture(scope="module")
def rendered(tmp_path_factory):
    out = {}
    for name, build in PACKAGES.items():
        package = build()
        root = tmp_path_factory.mktemp(name)
        result = memo_docx_renderer.render_memos(
            package,
            out_en=root / "memo" / "en.docx",
            out_zh=root / "memo" / "zh.docx",
        )
        out[name] = {
            "package": package,
            "structure": memo_structure.for_package(package),
            "result": result,
            "en_path": root / "memo" / "en.docx",
            "zh_path": root / "memo" / "zh.docx",
            "en": Document(root / "memo" / "en.docx"),
            "zh": Document(root / "memo" / "zh.docx"),
            "root": root,
        }
    return out


def _h1(doc):
    return [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]


def _names(doc, tag="w:bookmarkStart"):
    return [el.get(qn("w:name")) for el in doc.element.body.iter(qn(tag))]


def _anchors(doc):
    return [
        h.get(qn("w:anchor"))
        for h in doc.element.body.iter(qn("w:hyperlink"))
        if h.get(qn("w:anchor"))
    ]


def _table(doc, first_cell):
    return next(t for t in doc.tables if t.rows[0].cells[0].text == first_cell)


@pytest.mark.parametrize("name", list(PACKAGES))
def test_the_v2_fixture_passes_the_generation_gates(name):
    package = PACKAGES[name]()
    assert memo_docx_renderer.english_package_validation_errors(package) == []
    assert memo_docx_renderer.validate_package(package) == []


@pytest.mark.parametrize("name", list(PACKAGES))
def test_strict_and_lenient_renders_agree(name, rendered, tmp_path):
    lenient = memo_docx_renderer.render_memos(
        PACKAGES[name](),
        out_en=tmp_path / "memo" / "en.docx",
        out_zh=tmp_path / "memo" / "zh.docx",
        strict_sources=False,
    )
    assert lenient["ok"] and lenient["warnings"] == []
    assert rendered[name]["result"]["warnings"] == []
    for locale in ("en", "zh"):
        assert _h1(Document(lenient["outputs"][locale])) == _h1(rendered[name][locale])


def test_late_v2_numbers_twelve_sections_then_the_sources(rendered):
    doc = rendered["late_v2"]
    assert _h1(doc["en"]) == [
        "I. Executive Summary",
        "II. Company Overview & Stage",
        "III. Market & Industry Analysis",
        "IV. Product, Business Model & Unit Economics",
        "V. Competitive Landscape",
        "VI. Moat & Defensibility",
        "VII. Financial Analysis",
        "VIII. Team & Governance",
        "IX. Valuation Analysis",
        "X. Return & Exit Analysis",
        "XI. Investment Risk",
        "XII. Investment Decision",
        "XIII. Sources, Source Classes, and Fact Reference Index",
        "Calculation notes",
    ]
    assert _h1(doc["zh"]) == [
        "一、执行摘要",
        "二、公司概况与发展阶段",
        "三、市场与行业分析",
        "四、产品、商业模式与单位经济",
        "五、竞争格局",
        "六、护城河",
        "七、财务分析",
        "八、团队与治理",
        "九、估值分析",
        "十、回报测算与退出分析",
        "十一、风险分析",
        "十二、最终投资决定",
        "十三、来源、来源类别与事实索引",
        "计算说明",
    ]


def test_late_compact_numbers_seven_sections_then_the_sources(rendered):
    doc = rendered["late_compact"]
    assert _h1(doc["en"])[0] == "I. Executive Summary"
    assert _h1(doc["en"])[6] == "VII. Investment Decision"
    assert _h1(doc["en"])[7:] == [
        "VIII. Sources, Source Classes, and Fact Reference Index",
        "Calculation notes",
    ]
    assert _h1(doc["zh"])[7:] == ["八、来源、来源类别与事实索引", "计算说明"]


@pytest.mark.parametrize("name", list(PACKAGES))
def test_every_heading_bookmark_contents_link_and_citation_resolves(name, rendered):
    doc = rendered[name]
    en, zh = doc["en"], doc["zh"]
    assert [n for n in _names(en) if n.startswith("bsh_")] == [n for n in _names(zh) if n.startswith("bsh_")]
    for d in (en, zh):
        marks = set(_names(d))
        anchors = _anchors(d)
        assert not [a for a in anchors if a not in marks], "every internal link lands"
        sections = [n for n in _names(d) if n.startswith("bsh_sec_")]
        assert sections == [f"bsh_sec_{n}" for n in range(1, len(_h1(d)) + 1)]
        assert [a for a in anchors if a.startswith("bsh_sec_")] == sections
    # every inline citation the package writes is a clickable link to its row
    cited = set()
    for section in doc["package"]["sections"]:
        for text in memo_docx_renderer._iter_strings(section["blocks"]):
            cited.update(memo_docx_renderer.citation_ids(text))
    assert cited, "the fixture cites"
    for d in (en, zh):
        linked = {a for a in _anchors(d) if a.startswith(("src_", "calc_"))}
        assert {memo_docx_renderer.citation_anchor(c) for c in cited} <= linked
        marks = set(_names(d))
        for source in doc["package"]["sources"]:
            assert f"src_{source['id']}" in marks
        for calc in doc["package"]["calculations"]:
            assert f"calc_{calc['id']}" in marks


def test_page_one_is_the_cover_and_the_contents_on_v2(rendered):
    doc = rendered["late_v2"]["en"]
    body = list(doc.element.body.iterchildren())
    first_break = next(
        i for i, el in enumerate(body)
        if el.tag == qn("w:p") and any(br.get(qn("w:type")) == "page" for br in el.iter(qn("w:br")))
    )
    tables = [el for el in body[:first_break] if el.tag == qn("w:tbl")]
    assert len(tables) == 2  # cover facts, contents
    contents = tables[1]
    assert len(contents.findall(qn("w:tr"))) == 14
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Written 2026-09-18 · Evidence through 2026-09-12 · 16 sources" in text
    assert "2026-09-18__101500" not in text


@pytest.mark.parametrize("name", list(PACKAGES))
def test_charts_are_drawn_and_captioned_once_in_each_language(name, rendered):
    doc = rendered[name]
    charts = [
        block
        for section in doc["package"]["sections"]
        for block in section["blocks"]
        if block.get("type") == "chart"
    ]
    for locale, reading, sources in (("en", "Reading: ", "Sources: "), ("zh", "读法：", "来源：")):
        d = doc[locale]
        assert len(d.inline_shapes) == len(charts)
        body = list(d.element.body.iterchildren())
        pictures = [i for i, el in enumerate(body) if el.find(".//" + qn("w:drawing")) is not None]
        for index, chart in zip(pictures, charts):
            caption = "".join(t.text or "" for t in body[index + 1].iter(qn("w:t")))
            assert caption.startswith(reading), caption
            assert chart["reading"][locale] in caption
            assert chart["caption"][locale] in caption
            assert sources in caption


@pytest.mark.parametrize("name, cards", [("late_v2", 6), ("late_compact", 5)])
def test_the_risk_summary_reads_the_eight_row_cards(name, cards, rendered):
    doc = rendered[name]
    for locale, header, first in (
        ("en", ["#", "Risk", "Area", "Likelihood", "Rating"], ["1", "The entry price already assumes the 2030 plan is delivered", "Valuation & exit", "High", "9/10"]),
        ("zh", ["#", "风险", "领域", "可能性", "评分"], ["1", "入场价格已经假设 2030 年计划全部兑现", "估值与退出", "高", "9/10"]),
    ):
        summary = next(t for t in doc[locale].tables if [c.text for c in t.rows[0].cells] == header)
        assert len(summary.rows) == cards + 1
        assert [c.text for c in summary.rows[1].cells] == first
        assert summary._tbl.tblPr.find(qn("w:tblStyle")).get(qn("w:val")) == memo_structure.DERIVED_TABLE_STYLE_ID
        ratings = [int(row.cells[4].text.split("/")[0]) for row in summary.rows[1:]]
        assert ratings == sorted(ratings, reverse=True)
    # each card keeps its eight rows, label column first
    card = next(t for t in doc["en"].tables if t.rows[0].cells[0].text == "Risk Type")
    assert [row.cells[0].text for row in card.rows] == [
        "Risk Type", "Verdict", "Impact", "Why it matters", "What we watch",
        "Mitigation", "Likelihood", "Risk Rating",
    ]


@pytest.mark.parametrize("name", list(PACKAGES))
def test_scorecard_and_scenario_tables_have_sane_widths(name, rendered):
    doc = rendered[name]["en"]
    section = doc.sections[0]
    text_width = int((section.page_width - section.left_margin - section.right_margin) / 635)
    scorecard = _table(doc, "Dimension")
    grid = [int(g.get(qn("w:w"))) for g in scorecard._tbl.find(qn("w:tblGrid"))]
    assert abs(sum(grid) - text_width) <= 2
    dimension, weight, score, why = grid
    assert weight < 1300 and score < 1300  # two-digit numbers, not prose columns
    assert why == max(grid) and dimension > weight
    for row in scorecard.rows[1:]:
        assert row.cells[2].paragraphs[0].alignment == 2  # numbers right-aligned
    scenario = _table(doc, "Scenario")
    grid = [int(g.get(qn("w:w"))) for g in scenario._tbl.find(qn("w:tblGrid"))]
    assert len(grid) == 9 and min(grid) >= 600
    # "1.5x [C2]" is a number with its citation: right-aligned like the rest
    moic = [c.text for c in scenario.columns[7].cells]
    assert moic[2].startswith("1.5x")
    assert scenario.rows[2].cells[7].paragraphs[0].alignment == 2


@pytest.mark.parametrize("name", list(PACKAGES))
def test_sources_section_lists_links_tiers_and_dates(name, rendered):
    doc = rendered[name]
    for locale, header in (
        ("en", ["ID", "Source", "Class", "Tier", "Treatment", "Dates"]),
        ("zh", ["编号", "来源", "类别", "等级", "处理方式", "日期"]),
    ):
        table = _table(doc[locale], header[0])
        assert [c.text for c in table.rows[0].cells] == header
        rows = {row.cells[0].text: [c.text for c in row.cells] for row in table.rows[1:]}
        assert len(rows) == len(doc["package"]["sources"])
        assert rows["S2"][3] == "A"  # sec.gov
        assert rows["S1"][3] == "A-"  # the company's own site
        assert rows["S3"][3] == "B"  # Reuters
        assert rows["S5"][3] == "—"  # a data-room deck has no public link
        assert rows["S3"][1].endswith("reuters.com")
    en_rows = {row.cells[0].text: row.cells[5].text for row in _table(doc["en"], "ID").rows[1:]}
    assert en_rows["S1"] == "4 Aug 2026 · retrieved 15 Sep 2026"
    assert en_rows["S3"] == "12 Mar 2026 · covers 2025"
    assert en_rows["S13"] == "Undated"
    # the run's checks open the sources section (renderer text the gates skip)
    paragraphs = [p for p in doc["en"].paragraphs if p.text.strip()]
    heading = next(i for i, p in enumerate(paragraphs) if p.text.endswith("Sources, Source Classes, and Fact Reference Index"))
    assert paragraphs[heading + 1].text.startswith("88 figures: 61 verified")


@pytest.mark.parametrize("name", list(PACKAGES))
def test_calculation_notes_read_in_the_memo_language(name, rendered):
    doc = rendered[name]
    calcs_en = next(t for t in doc["en"].tables if [c.text for c in t.rows[0].cells][:2] == ["ID", "What"])
    calcs_zh = next(t for t in doc["zh"].tables if [c.text for c in t.rows[0].cells][:2] == ["编号", "计算内容"])
    c2_en = [c.text for c in calcs_en.rows[2].cells]
    c2_zh = [c.text for c in calcs_zh.rows[2].cells]
    assert "dilution to exit = 15% (assumption)" in c2_en[2]
    assert "（假设）" in c2_zh[2] and "[assumption]" not in c2_zh[2]
    grid = [int(g.get(qn("w:w"))) for g in calcs_en._tbl.find(qn("w:tblGrid"))]
    # a long result does not starve the inputs and formula columns
    assert grid[2] > grid[4] and grid[3] > grid[4]


@pytest.mark.parametrize("name", list(PACKAGES))
def test_stamp_footer_and_file_metadata(name, rendered):
    doc = rendered[name]
    for locale, stamp, footer_tail in (
        ("en", "DRAFT — AI-generated, not reviewed", "BSH Confidential · 2026-09-18 · EN"),
        ("zh", "草稿 — AI 生成，未经审阅", "BSH 机密 · 2026年9月18日 · 中文"),
    ):
        section = doc[locale].sections[0]
        assert section.first_page_header.paragraphs[0].text == stamp
        assert section.header.paragraphs[0].text.startswith(stamp)
        footer = section.footer.paragraphs[0]
        assert footer.text.endswith(footer_tail)
        assert [i.text.strip() for i in footer._p.iter(qn("w:instrText"))] == ["PAGE", "NUMPAGES"]
        assert doc[locale].core_properties.version == memo_docx_renderer.RENDERER_VERSION


@pytest.mark.parametrize("name", list(PACKAGES))
def test_the_gates_pass_the_v2_memo(name, rendered):
    doc = rendered[name]
    lint = memo_quality_lint.lint_memo_docx(doc["en_path"], doc["structure"])
    assert not lint.p0_findings, [f.to_dict() for f in lint.p0_findings]
    parity = memo_chinese_parity.lint_chinese_memo_pair(
        doc["en_path"], doc["zh_path"], doc["structure"], package=doc["package"]
    )
    assert parity.findings == [], [f.to_dict() for f in parity.findings]
    shape = memo_chinese_parity._extract_docx_shape(doc["zh_path"], "zh", doc["structure"].parity_patterns())
    assert shape.section_ids.count("sources") == 1
    assert set(doc["structure"].section_ids) <= set(shape.section_ids)
    # no Chinese slot fell back to English
    report = (doc["root"] / "logs" / "validation_cn.txt").read_text(encoding="utf-8")
    assert "- zh_fallback_count: 0" in report


def test_the_chinese_memo_keeps_dollar_figures_as_written(rendered):
    doc = rendered["late_v2"]["zh"]
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    assert "投后估值 $2.4B" in text and "$180M" in text
    # no native-unit conversion (owner decision): "$180M" never becomes
    # "1.8 亿美元", nor "$2.4B" "24 亿美元"
    assert not re.search(r"1\.8\s*亿美元|24\s*亿美元", text)
