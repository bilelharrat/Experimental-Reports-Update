"""Shared fixtures for the Console test suite."""
from __future__ import annotations

import os
import sys
import threading
import time
from pathlib import Path

import pytest

_LIVE_DATA_DIR = os.path.realpath(Path(__file__).resolve().parent.parent / "data")
_WRITE_FLAGS = os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_APPEND | os.O_TRUNC


def _in_live_data_dir(path) -> bool:
    if path is None or isinstance(path, int):
        return False
    try:
        raw = os.fsdecode(path)
    except TypeError:
        return False
    if "data" not in raw:
        return False
    real = os.path.realpath(raw)
    return real == _LIVE_DATA_DIR or real.startswith(_LIVE_DATA_DIR + os.sep)


def _guard_live_data_dir(event, args):
    """Fail any write, rename, delete or new directory under the live ``data/``
    directory: a store path that escaped ``_isolate_data_dir`` raises here
    instead of silently touching real data."""
    if event == "open":
        path, mode, flags = args
        writes = (isinstance(mode, str) and any(c in mode for c in "wax+")) or (
            isinstance(flags, int) and bool(flags & _WRITE_FLAGS)
        )
        targets = (path,) if writes else ()
    elif event in {"os.rename", "os.replace", "shutil.move"}:
        targets = args[:2]
    elif event == "shutil.copyfile":
        targets = args[1:2]
    elif event in {"os.remove", "os.rmdir", "shutil.rmtree"}:
        targets = args[:1]
    elif event == "os.mkdir":
        targets = tuple(p for p in args[:1] if _in_live_data_dir(p) and not os.path.exists(os.fsdecode(p)))
    else:
        return
    for target in targets:
        if _in_live_data_dir(target):
            raise PermissionError(f"test touched the live data directory: {os.fsdecode(target)}")


sys.addaudithook(_guard_live_data_dir)


@pytest.fixture(autouse=True)
def _isolate_data_dir(monkeypatch, tmp_path):
    """Redirect every module-level data path to a per-test tmp dir.

    The live ``data/`` directory used to be the default for any test that
    forgot to monkeypatch a root — ``data/companies.yaml.broken_by_tests``
    was the scar. Modules freeze path constants at import time, so each one
    must be re-pointed individually; keep this list in sync when adding a
    store with its own module-level path constant. Tests that need their
    own layout simply monkeypatch over these (their patch runs later).
    """
    from server import (
        annotation_store,
        auth_store,
        cache,
        console_store,
        desk_store,
        external_store,
        files_store,
        hormuz_store,
        market_brief,
        memo_editor_store,
        memo_prep,
        news_brief,
        product_store,
        research_store,
        serena_analysis,
        stock_research,
        storage,
        weekly_stocks,
    )

    data_root = tmp_path / "data"
    monkeypatch.setattr(storage, "DATA_DIR", data_root)
    monkeypatch.setattr(storage, "COMPANIES_FILE", data_root / "companies.yaml")
    monkeypatch.setattr(storage, "REPORTS_DIR", data_root / "reports")
    monkeypatch.setattr(storage, "THREADS_DIR", data_root / "threads")
    monkeypatch.setattr(cache, "CACHE_ROOT", data_root / "cache")
    monkeypatch.setattr(auth_store, "USERS_FILE", data_root / "users.json")
    monkeypatch.setattr(auth_store, "SESSIONS_FILE", data_root / "sessions.json")
    monkeypatch.setattr(external_store, "EXTERNAL_ROOT", data_root / "external")
    monkeypatch.setattr(files_store, "UPLOADS_ROOT", data_root / "uploads")
    monkeypatch.setattr(research_store, "RESEARCH_ROOT", data_root / "research")
    monkeypatch.setattr(console_store, "CONSOLES_ROOT", data_root / "consoles")
    monkeypatch.setattr(memo_editor_store, "EDITOR_ROOT", data_root / "memo_editor")
    monkeypatch.setattr(product_store, "SETTINGS_ROOT", data_root / "settings")
    monkeypatch.setattr(memo_prep, "MEMOS_ROOT", data_root / "memos")
    monkeypatch.setattr(
        memo_prep, "SETTINGS_FILE", data_root / "settings" / "serena_background.md"
    )
    monkeypatch.setattr(memo_prep, "COMPANIES_FILE", data_root / "companies.yaml")
    monkeypatch.setattr(hormuz_store, "DATA_DIR", data_root)
    monkeypatch.setattr(
        hormuz_store,
        "SOURCES_ROOT",
        data_root / "external" / "hormuz_research" / "sources",
    )
    monkeypatch.setattr(hormuz_store, "APPENDIX_ROOT", data_root / "hormuz_appendix")
    monkeypatch.setattr(weekly_stocks, "WEEKLY_DIR", data_root / "_weekly_stocks")
    monkeypatch.setattr(
        stock_research, "STOCK_RESEARCH_ROOT", data_root / "stock_research"
    )
    monkeypatch.setattr(
        serena_analysis, "ANALYSIS_ROOT", data_root / "serena_analysis"
    )
    monkeypatch.setattr(
        serena_analysis, "TRAINING_ROOT", data_root / "serena_training"
    )
    monkeypatch.setattr(desk_store, "DESK_ROOT", data_root / "market_desk")
    monkeypatch.setattr(
        annotation_store, "ANNOTATIONS_ROOT", data_root / "report_annotations"
    )
    monkeypatch.setattr(market_brief, "BRIEFS_ROOT", data_root / "market_briefs")
    monkeypatch.setattr(news_brief, "BRIEFS_ROOT", data_root / "news_briefs")


def _is_test_fake_thread(thread: threading.Thread) -> bool:
    """Threads a test body starts itself (fake process pipes, drips) never
    touch data paths and may block until the process exits."""
    module = getattr(getattr(thread, "_target", None), "__module__", "") or ""
    return module == "conftest" or module.startswith(("test_", "tests."))


@pytest.fixture(autouse=True)
def _join_test_threads(_isolate_data_dir):
    """Let background work a test started finish while its data dirs are
    still redirected. Job workers (console hydrate, stock aggregates) run on
    daemon threads; one that outlives its test writes into the live ``data/``
    after the monkeypatches are undone. Console session dispatchers idle on
    their queue for minutes, so those are waited on until their turns are
    done instead of joined.
    """
    from server import console_session

    before = set(threading.enumerate())
    with console_session._DISPATCHERS_LOCK:
        known = set(console_session._DISPATCHERS)
    yield
    deadline = time.monotonic() + 5.0
    with console_session._DISPATCHERS_LOCK:
        dispatchers = [d for key, d in console_session._DISPATCHERS.items() if key not in known]
    for dispatcher in dispatchers:
        while dispatcher.pending_ids and time.monotonic() < deadline:
            time.sleep(0.01)
    idle_workers = {dispatcher.thread for dispatcher in dispatchers}
    for thread in threading.enumerate():
        if thread in before or thread in idle_workers or not thread.is_alive():
            continue
        if _is_test_fake_thread(thread):
            continue
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            break
        thread.join(remaining)


@pytest.fixture(autouse=True)
def _no_real_claude_cli(request, monkeypatch):
    """Fail closed on real Claude CLI spawns in unit tests.

    Every spawn path checks ``claude_runner.is_available()`` first, so forcing
    it False makes an unmocked path return its "claude not on PATH" error
    instead of silently launching a real (paid, minutes-long) subprocess.
    This actually happened: on machines with ``claude`` installed, the fast
    memo pipeline tests reached the real bilingual parallel wrapper, spawned
    per-section Claude processes, and only then fell back to the patched
    monolithic fake — one such test took 58s with claude on PATH vs 0.9s
    without. CI has no claude binary, so this also makes local runs match CI.
    Tests that need availability monkeypatch ``is_available`` themselves
    (their patch runs after this one and wins); e2e tests keep the real
    check.
    """
    if request.node.get_closest_marker("e2e"):
        return
    from server import claude_runner

    monkeypatch.setattr(claude_runner, "is_available", lambda: False)


@pytest.fixture(autouse=True)
def _no_real_gemini_key(request, monkeypatch):
    """Fail closed on real Gemini API calls, the sibling of the Claude guard.

    Importing ``server.main`` anywhere in the session ``load_dotenv()``s the
    project .env into ``os.environ`` for every test that follows, so a
    developer's real GEMINI_API_KEY leaks in and any unstubbed call would
    reach the live API and spend money — and, worse, tests would pass or
    fail depending on whether some earlier test happened to import the app.
    Clearing the keys makes ``gemini_runner.is_available()`` False by
    default; tests that want a key set one themselves (their patch runs
    after this one and wins).
    """
    if request.node.get_closest_marker("e2e"):
        return
    for name in ("GEMINI_API_KEY", "BSH_GEMINI_API_KEY", "GOOGLE_API_KEY"):
        monkeypatch.delenv(name, raising=False)


@pytest.fixture(autouse=True)
def _neutral_structure_flag(monkeypatch):
    """Tests must not inherit the developer's .env: importing server.main
    anywhere in the session load_dotenv()s the project .env into
    os.environ, and a BSH_MEMO_STRUCTURE_V2=1 there flips every
    pipeline test to the v2 structure mid-session (observed 2026-09-11
    as order-dependent failures). Tests that exercise v2 set it themselves.

    Since 2026-09-22 the pipeline switches default ON (server/memo_flags.py),
    so deleting the variable is no longer neutral. Every switch is pinned
    OFF here: the historical baseline these tests were written against.
    tests/test_memo_flags.py checks the real defaults on a clean env.
    """
    from server import memo_flags

    for name in memo_flags.DEFAULTS:
        monkeypatch.setenv(name, "0")


@pytest.fixture(autouse=True)
def _disable_auth(monkeypatch):
    """Let API tests through without credentials. ``require_api_token`` now
    fails closed, so we clear any shared token AND opt into anonymous dev
    access explicitly. Tests that exercise auth itself override these with
    their own ``monkeypatch`` in the test body. Env is restored on exit.
    """
    monkeypatch.delenv("BSH_RESEARCH_API_TOKEN", raising=False)
    monkeypatch.setenv("BSH_ALLOW_ANON_DEV", "1")
    # The /api/jobs/active TTL cache is keyed on time only; tests monkeypatch
    # data roots, so a cache entry from one test would otherwise leak into the
    # next. Clear it before each test.
    try:
        from server import api as _api

        _api._active_jobs_cache = None
    except Exception:
        pass


@pytest.fixture
def tmp_consoles(monkeypatch, tmp_path):
    """Redirect ``console_store.CONSOLES_ROOT`` to a tmp dir for a single
    test, plus the two source-file stores it can pull from."""
    from server import console_store, files_store, research_store

    consoles_root = tmp_path / "consoles"
    uploads_root = tmp_path / "uploads"
    research_root = tmp_path / "research"

    monkeypatch.setattr(console_store, "CONSOLES_ROOT", consoles_root)
    monkeypatch.setattr(files_store, "UPLOADS_ROOT", uploads_root)
    monkeypatch.setattr(research_store, "RESEARCH_ROOT", research_root)

    return {
        "consoles": consoles_root,
        "uploads": uploads_root,
        "research": research_root,
        "root": tmp_path,
    }


@pytest.fixture
def png_bytes() -> bytes:
    """Minimal valid 1x1 PNG."""
    # 8-byte signature + IHDR + IDAT + IEND. Hand-crafted minimal PNG.
    return bytes.fromhex(
        "89504e470d0a1a0a"            # signature
        "0000000d49484452"            # IHDR length(13) + tag
        "00000001000000010806000000"  # 1x1 RGBA, bit depth 8
        "1f15c4890000000a4944415478"  # CRC + IDAT length + tag + start
        "9c6300010000000500010d0a2db4"
        "0000000049454e44ae426082"    # IEND
    )


@pytest.fixture
def jpeg_bytes() -> bytes:
    """Tiny JPEG header (enough to pass our magic-byte sniff)."""
    # FF D8 FF E0 ... full minimal JPEG isn't needed — sniffer only reads
    # the first 3 bytes for JPEG.
    return b"\xff\xd8\xff\xe0\x00\x10JFIF\x00\x01\x01\x00\x00\x01\x00\x01\x00\x00" \
           + b"\xff\xd9"


@pytest.fixture
def webp_bytes() -> bytes:
    """RIFF/WEBP container header — magic-byte check needs bytes 0-3 == RIFF
    and 8-11 == WEBP."""
    return b"RIFF\x10\x00\x00\x00WEBPVP8 " + b"\x00" * 16


@pytest.fixture
def pdf_bytes() -> bytes:
    """Tiny PDF — just the magic header is enough for our sniffer."""
    return b"%PDF-1.4\n%mock\n1 0 obj <<>> endobj\n%%EOF\n"


@pytest.fixture
def doc_bytes() -> bytes:
    """OLE2 compound-binary header — same magic shared with .xls/.ppt; the
    sniffer relies on the .doc extension to disambiguate."""
    return b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1" + b"\x00" * 32


@pytest.fixture
def docx_bytes() -> bytes:
    """ZIP local-file header — same magic shared with .xlsx/.pptx; the
    sniffer relies on the .docx extension to disambiguate."""
    return b"PK\x03\x04" + b"\x00" * 32


@pytest.fixture
def real_docx_bytes() -> bytes:
    """An actual Word file, table and all. The sniffer is happy with the
    header fixtures above, but anything that reaches
    ``console_store.save_attachment`` now has its text extracted, so a
    document has to really be one."""
    import io

    from docx import Document

    doc = Document()
    doc.add_paragraph("Pre-money valuation of $42M.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Liquidation pref"
    table.rows[0].cells[1].text = "1x non-participating"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def empty_docx_bytes() -> bytes:
    """A real .docx with nothing in it — the "no text came back" case."""
    import io

    from docx import Document

    buf = io.BytesIO()
    Document().save(buf)
    return buf.getvalue()


@pytest.fixture
def real_pptx_bytes() -> bytes:
    import io

    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Series B"
    slide.placeholders[1].text = "ARR $4.2M"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.fixture
def real_xlsx_bytes() -> bytes:
    """A minimal workbook in the layout Excel writes: shared strings,
    a relationship id per sheet, ``t="s"`` cells pointing into the table.
    openpyxl isn't a dependency, so the fixture spells the parts out."""
    import io
    import zipfile

    def sheet(rows):
        body = "".join(
            "<row r=\"{}\">{}</row>".format(
                i,
                "".join(
                    '<c r="{}{}" t="s"><v>{}</v></c>'.format(chr(65 + j), i, v)
                    for j, v in enumerate(row)
                ),
            )
            for i, row in enumerate(rows, start=1)
        )
        return (
            '<?xml version="1.0"?><worksheet xmlns="http://schemas.openxmlformats.org'
            '/spreadsheetml/2006/main"><sheetData>' + body + "</sheetData></worksheet>"
        )

    strings = ["Holder", "Shares", "Ada Lovelace", "1,000,000"]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as bundle:
        bundle.writestr(
            "xl/workbook.xml",
            '<?xml version="1.0"?><workbook xmlns="http://schemas.openxmlformats.org'
            '/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org'
            '/officeDocument/2006/relationships"><sheets>'
            '<sheet name="Cap table" sheetId="1" r:id="rId1"/></sheets></workbook>',
        )
        bundle.writestr(
            "xl/_rels/workbook.xml.rels",
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats'
            '.org/package/2006/relationships"><Relationship Id="rId1" Target'
            '="worksheets/sheet1.xml"/></Relationships>',
        )
        bundle.writestr(
            "xl/sharedStrings.xml",
            '<?xml version="1.0"?><sst xmlns="http://schemas.openxmlformats.org'
            '/spreadsheetml/2006/main">'
            + "".join(f"<si><t>{s}</t></si>" for s in strings)
            + "</sst>",
        )
        bundle.writestr("xl/worksheets/sheet1.xml", sheet([(0, 1), (2, 3)]))
    return buf.getvalue()
