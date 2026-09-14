"""Founder-feedback round (2026-09-13): the executive summary must say
WHICH dimensions carry the case before explaining them, every risk must
say which aspect it concentrates on and how big the impact is, and the
prose must teach rather than assert. These tests pin the contract text,
the renderer's bold-lead bullets, and the market-sizing research focus."""

from __future__ import annotations

import pytest

from server import claude_runner, memo_analysis, memo_docx_renderer, memo_structure

PROFILES = {
    "late_v2": memo_structure.load_structure("late", 2),
    "late_compact": memo_structure.load_structure("late_compact"),
    "growth": memo_structure.load_structure("growth"),
    "early": memo_structure.load_structure("early"),
}


@pytest.mark.parametrize("name", sorted(PROFILES))
def test_exec_contract_opens_with_case_summary_and_pins_highlights(name):
    structure = PROFILES[name]
    spec = " ".join(
        structure.section_specs()[structure.section_for_role("exec").id].split()
    )
    assert "case-summary sentence" in spec
    assert '"component": "investment_highlights"' in spec
    assert "EXACTLY three items" in spec
    assert "headline VERBATIM" in spec
    assert '"component": "key_risks"' in spec
    assert "Impact: <pinned impact verbatim>" in spec
    assert "<Area label>" in spec


@pytest.mark.parametrize("name", sorted(PROFILES))
def test_market_contract_triangulates_external_estimates(name):
    structure = PROFILES[name]
    text = " ".join("\n".join(structure.section_specs().values()).split())
    assert (
        "never dropped" in text
        or "never fewer than three" in text
        or "nothing is dropped" in text
    )
    assert "the memo adopts" in text


def test_risk_card_v2_has_verdict_and_impact_rows():
    contract = claude_runner.MEMO_RISK_REGISTER_CONTRACT_V2
    for label in (
        "`Risk Type` / `风险类型`",
        "`Verdict` / `一句话结论`",
        "`Impact` / `影响有多大`",
        "`Why it matters` / `为什么重要`",
        "`What we watch` / `跟踪信号`",
        "`Mitigation` / `缓释措施`",
        "`Likelihood` / `可能性`",
        "`Risk Rating` / `风险评分`",
    ):
        assert label in contract
    assert "EXACTLY these eight" in contract
    assert "professor" in contract
    compact = claude_runner.MEMO_RISK_REGISTER_CONTRACT_COMPACT
    assert "Impact: <pinned impact VERBATIM>" in compact
    assert "<Area label>" in compact
    # v1 stays the frozen five-row card.
    assert "Verdict" not in claude_runner.MEMO_RISK_REGISTER_CONTRACT


def test_addendum_teaches_instead_of_asserting():
    addendum = claude_runner.MEMO_STRUCTURE_V2_ADDENDUM
    assert "## Teach, don't assert" in addendum
    assert "Define a term the first time it appears" in addendum
    assert "the middle step missing" in addendum
    # Both founder-review worked examples ride the cached context.
    assert "sevenfold in seven months" in addendum
    assert "the price, not the\n  growth, decides" in addendum.replace(
        "the price, not the growth, decides", "the price, not the\n  growth, decides"
    )
    assert "Risk 1 — Valuation & exit" in addendum
    # The teach block is v2-only: the v1 shared context stays byte-identical.
    assert "Teach, don't assert" not in claude_runner.HUMAN_EXEC_MEMO_VOICE_CONTRACT


def test_market_sizing_pass_keeps_every_external_estimate():
    spec = next(
        p for p in memo_analysis._FAST_MEMO_PASSES if p.pass_id == "market_sizing"
    )
    assert "NEVER discard an estimate" in spec.focus
    assert "Stratpace" in spec.focus
    assert "the URL" in spec.focus
    assert "which definition the memo adopts" in spec.focus


def test_spine_prompt_describes_highlights_area_and_impact(tmp_path, monkeypatch):
    captured: dict = {}

    def fake_run(**kwargs):
        captured["prompt"] = kwargs["prompt"]
        return None, "stub"

    monkeypatch.setattr(claude_runner, "_run_memo_local_json_artifact", fake_run)
    claude_runner.run_memo_fast_english_spine(
        run_dir=tmp_path,
        company_name="Acme",
        common_context="",
        add_dirs=[],
        progress=None,
        structure=PROFILES["late_v2"],
    )
    prompt = captured["prompt"]
    assert "`highlights`: EXACTLY three" in prompt
    assert "at least 60% of its max" in prompt
    assert "each risk ALSO carries `area`" in prompt
    assert "`impact`: what it costs the investment" in prompt
    # A MOIC just under the floor must not round onto it (live run
    # 2026-09-14: "returns 1.5x, the floor" vs "1.49x, under the floor").
    assert "within 0.05x of the return floor" in prompt
    assert 'never the words "the memo"' in prompt


def test_split_lead_sentence():
    split = memo_docx_renderer.split_lead_sentence
    assert split("Acme leads its category. It holds 40% of spend.") == (
        "Acme leads its category.",
        " It holds 40% of spend.",
    )
    # Periods inside numbers or abbreviations without a following space
    # do not end the lead.
    assert split("The price is $1.5T against $2.6T of value. Impact: none.") == (
        "The price is $1.5T against $2.6T of value.",
        " Impact: none.",
    )
    assert split("Only one sentence here.") == ("Only one sentence here.", "")
    assert split("市场足够大。证据一。") == ("市场足够大。", "证据一。")
    assert split("no terminal punctuation") == ("no terminal punctuation", "")


def test_split_lead_sentence_takes_a_markdown_bold_lead_as_written():
    """Live compact run 2026-09-14: writers wrapped the pinned headline
    in "**...**" and the asterisks printed literally in both languages."""
    split = memo_docx_renderer.split_lead_sentence
    assert split("**Acme leads its category.** It holds 40% of spend [S1].") == (
        "Acme leads its category.",
        " It holds 40% of spend [S1].",
    )
    assert split("**Acme leads its category.**It holds 40%.") == (
        "Acme leads its category.",
        " It holds 40%.",
    )
    assert split("**市场足够大。**证据一。") == ("市场足够大。", "证据一。")
    assert split("A **stray** marker ends here. Rest.") == (
        "A stray marker ends here.",
        " Rest.",
    )


def test_add_run_never_prints_markdown_asterisks():
    from docx import Document

    paragraph = Document().add_paragraph()
    memo_docx_renderer._add_run(paragraph, "A **bold** word, cited [S1].", locale="en")
    assert "**" not in paragraph.text
    plain = Document().add_paragraph()
    memo_docx_renderer._add_run(plain, "**Only markers**", locale="en")
    assert plain.text == "Only markers"


@pytest.mark.parametrize("name", sorted(PROFILES))
def test_exec_contract_forbids_markdown_and_memo_self_reference(name):
    """The lint bans "the memo" in body prose, so the contract must not
    suggest "the memo's own estimate"; and "rendered bold" invited
    markdown asterisks."""
    structure = PROFILES[name]
    spec = " ".join(
        structure.section_specs()[structure.section_for_role("exec").id].split()
    )
    assert "no asterisks or markdown" in spec
    assert "rendered bold" not in spec
    assert "the memo's own" not in spec
    addendum = claude_runner.MEMO_STRUCTURE_V2_ADDENDUM
    assert "the memo's own" not in " ".join(addendum.split())
    assert 'Never write "the memo"' in addendum


def test_highlight_and_risk_bullets_render_bold_lead(tmp_path):
    from docx import Document

    document = Document()
    block = {
        "type": "bullets",
        "component": "investment_highlights",
        "items": [
            {"en": "Acme leads its category. It holds 40% of spend.", "zh": ""}
        ],
    }
    memo_docx_renderer._add_block(document, block, "en")
    plain = {
        "type": "bullets",
        "items": [{"en": "A plain bullet. Not bold.", "zh": ""}],
    }
    memo_docx_renderer._add_block(document, plain, "en")
    paragraphs = document.paragraphs
    lead, rest = paragraphs[-2].runs[0], paragraphs[-2].runs[1]
    assert lead.text == "• Acme leads its category." and lead.bold is True
    assert rest.text == " It holds 40% of spend." and not rest.bold
    assert len(paragraphs[-1].runs) == 1 and not paragraphs[-1].runs[0].bold
