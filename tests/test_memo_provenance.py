"""Provenance (founder feedback 2026-09-13: "click to see where a number
came from"): sources carry URLs, the spine pins calculation notes, and
inline [S#]/[C#] citations render as links to bookmarked appendix rows.
Unknown citations fail validation; translations must keep the tokens."""

from __future__ import annotations

import copy

from docx import Document

from server import claude_runner, memo_docx_renderer, memo_pin_check, memo_quality_lint, memo_structure

V2 = memo_structure.load_structure("late", 2)


def test_pin_echo_matching_ignores_citation_tokens():
    """Live compact run 2026-09-14: the writer cited each scorecard
    why-line inside its final period ("...buyer is feasible [S15, S44].")
    and all nine exact-match pins "failed", costing a repair round."""
    haystack = "Exit certainty \n 2 / 3 \n Draft S-1 filed; no M&A buyer is feasible [S15, S44]."
    norm = memo_pin_check._norm(haystack)
    squashed = memo_pin_check._squash(haystack)
    assert memo_pin_check._contains(norm, squashed, "Draft S-1 filed; no M&A buyer is feasible.")
    # A pin that itself carries a token matches a differently cited echo.
    assert memo_pin_check._contains(norm, squashed, "no M&A buyer is feasible [C3].")
    assert not memo_pin_check._contains(norm, squashed, "an M&A buyer is feasible.")
    # Token digits never count as numbers in the echo.
    assert not memo_pin_check._contains(norm, squashed, "15")


def test_citation_ids_parse_single_and_grouped_tokens():
    assert memo_docx_renderer.citation_ids("Revenue is $65B [S3]. Base 1.5x [C2, S4].") == [
        "S3",
        "C2",
        "S4",
    ]
    assert memo_docx_renderer.citation_ids("[WV memo] [S] [internal]") == []
    assert memo_docx_renderer.citation_anchor("S3") == "src_S3"
    assert memo_docx_renderer.citation_anchor("C2") == "calc_C2"


def test_add_run_renders_citations_as_superscript_internal_links():
    document = Document()
    paragraph = document.add_paragraph()
    memo_docx_renderer._add_run(
        paragraph, "Run-rate is $65B [S3]. The base case returns 1.5x [C2, S4].", locale="en"
    )
    xml = paragraph._p.xml
    assert 'w:anchor="src_S3"' in xml
    assert 'w:anchor="calc_C2"' in xml
    assert 'w:anchor="src_S4"' in xml
    assert "<w:vertAlign" in xml and "superscript" in xml
    assert "[S3]" not in paragraph.text  # the bracket is consumed by the link
    plain = document.add_paragraph()
    run = memo_docx_renderer._add_run(plain, "No citations here [WV memo].", locale="en")
    assert run.text == "No citations here [WV memo]."
    assert "w:hyperlink" not in plain._p.xml


def _package() -> dict:
    return {
        "schema_version": 1,
        "structure": V2.meta(),
        "company": {"name": {"en": "Acme", "zh": "Acme"}},
        "sections": [
            {
                "id": "executive_summary",
                "blocks": [
                    {
                        "type": "paragraph",
                        "text": {
                            "en": "Run-rate is $65B [S1]. Base case 1.5x [C1].",
                            "zh": "年化收入为 $65B [S1]。基准情形 1.5x [C1]。",
                        },
                    }
                ],
            }
        ],
        "sources": [
            {
                "id": "S1",
                "title": {"en": "Bloomberg run-rate report", "zh": "彭博报道"},
                "class": "press reporting",
                "treatment": {"en": "Growth anchor only.", "zh": "仅作增长锚。"},
                "as_of": "2026-08-17",
                "url": "https://www.bloomberg.com/example",
            },
            {
                "id": "S2",
                "title": {"en": "Company deck", "zh": "公司材料"},
                "class": "company-reported",
                "treatment": {"en": "Company claims.", "zh": "公司口径。"},
                "as_of": "2026-07-01",
            },
        ],
        "calculations": [
            {
                "id": "C1",
                "label": {"en": "Base-case MOIC", "zh": "基准情形回报倍数"},
                "inputs": [
                    {"name": "2029 revenue", "value": "$200B", "ref": "S2"},
                    {"name": "exit multiple", "value": "13x", "ref": "assumption"},
                ],
                "formula": "13 × $200B = $2.6T; $2.6T ÷ $1.75T = 1.5x",
                "result": "1.5x",
                "meaning": {"en": "A 15% annual return.", "zh": "年化 15% 的回报。"},
            }
        ],
    }


def test_document_bookmarks_sources_and_calculations_and_links_urls():
    package = _package()
    document = memo_docx_renderer._build_document(package, "en")
    xml = document.element.xml
    assert 'w:name="src_S1"' in xml and 'w:name="src_S2"' in xml
    assert 'w:name="calc_C1"' in xml
    assert "Calculation notes" in xml
    assert "13 × $200B = $2.6T" in xml
    external = [
        rel
        for rel in document.part.rels.values()
        if rel.reltype.endswith("/hyperlink") and rel.is_external
    ]
    assert [rel.target_ref for rel in external] == ["https://www.bloomberg.com/example"]
    assert "Calculation notes" in memo_docx_renderer._toc_titles(package, "en")
    zh = memo_docx_renderer._build_document(package, "zh").element.xml
    assert "计算说明" in zh and 'w:name="calc_C1"' in zh


def test_validation_flags_unknown_citations_bad_urls_and_bad_calc_ids():
    good = _package()
    assert memo_docx_renderer._citation_errors(good) == []
    assert memo_docx_renderer._calculation_errors(good) == []
    errors: list[str] = []
    memo_docx_renderer._validate_source(good["sources"][0], "sources[0]", errors)
    assert errors == []
    dangling = copy.deepcopy(good)
    dangling["sections"][0]["blocks"][0]["text"]["en"] += " See [S9] and [C4]."
    errors = memo_docx_renderer._citation_errors(dangling)
    assert any("unknown_citation [S9]" in e for e in errors)
    assert any("unknown_citation [C4]" in e for e in errors)
    # The full validator chain surfaces the same finding.
    assert any(
        "unknown_citation [S9]" in e
        for e in memo_docx_renderer._package_validation_errors(dangling)
    )
    bad_url = copy.deepcopy(good["sources"][0])
    bad_url["url"] = "bloomberg.com/example"
    errors = []
    memo_docx_renderer._validate_source(bad_url, "sources[0]", errors)
    assert any("url must be an http(s) URL" in e for e in errors)
    bad_calc = copy.deepcopy(good)
    bad_calc["calculations"].append({"id": "X1", "formula": "", "result": "1x"})
    bad_calc["calculations"].append({"id": "C1", "formula": "1+1", "result": "2"})
    errors = memo_docx_renderer._calculation_errors(bad_calc)
    assert any("id must look like C1" in e for e in errors)
    assert any("formula is required" in e for e in errors)
    assert any("C1 is duplicated" in e for e in errors)


def test_package_calculations_become_bilingual_slots():
    notes = claude_runner._package_calculations(
        {
            "calculations": [
                {
                    "id": "C1",
                    "label": "Base-case MOIC",
                    "inputs": [{"name": "revenue", "value": "$200B", "ref": "S2"}],
                    "formula": "13 × $200B = $2.6T",
                    "result": "$2.6T",
                    "meaning": "Worth 1.5x the entry.",
                },
                "junk",
            ]
        }
    )
    assert notes == [
        {
            "id": "C1",
            "label": {"en": "Base-case MOIC", "zh": ""},
            "inputs": [{"name": "revenue", "value": "$200B", "ref": "S2"}],
            "formula": "13 × $200B = $2.6T",
            "result": "$2.6T",
            "meaning": {"en": "Worth 1.5x the entry.", "zh": ""},
        }
    ]
    assert claude_runner._package_calculations(None) == []


def test_translation_adoption_requires_matching_citations():
    source = {"en": "Revenue is $65B [S3] and 1.5x [C2].", "zh": ""}
    claude_runner._adopt_zh_translations(
        source, {"en": source["en"], "zh": "收入为 $65B [S3]，回报 1.5x [C1]。"}
    )
    assert source["zh"] == ""  # renumbered token -> rejected
    claude_runner._adopt_zh_translations(
        source, {"en": source["en"], "zh": "收入为 $65B [S3]，回报 1.5x [C2]。"}
    )
    assert source["zh"].endswith("[C2]。")


def test_lint_allows_citation_tokens_only_for_v2():
    def block(text: str):
        return memo_quality_lint._TextBlock(
            kind="paragraph",
            text=text,
            location="executive_summary/p1",
            section="executive_summary",
            allowed_trace_section=False,
            operating_table=False,
        )

    def leaks(text: str, structure) -> int:
        return sum(
            1
            for f in memo_quality_lint._lint_blocks([block(text)], structure)
            if f.code == "source_token_leak"
        )

    cited = "Run-rate is $65B [S3]. Base 1.5x [C2, S4]."
    assert leaks(cited, V2) == 0
    assert leaks(cited, memo_structure.LATE) == 1
    # Internal tokens stay banned everywhere.
    assert leaks("See [WV memo] and [S3].", V2) == 1


def test_spine_schema_and_gate_cover_calculations():
    schema = claude_runner.memo_fast_english_spine_schema(V2)
    facts = schema["properties"]["shared_facts"]
    assert "calculations" in facts["required"]
    assert facts["properties"]["calculations"]["items"]["properties"]["id"]["pattern"]
    from tests.test_memo_pins_v2 import _good_shared_facts

    good = _good_shared_facts()
    assert memo_pin_check.check_spine_pins_v2(good, V2) == []
    block = claude_runner._render_shared_facts_block(good, V2)
    assert "Calculation notes (cite the id in square brackets" in block
    assert "- C2 Fair value range: 25x × $24M = $600M" in block
    missing = copy.deepcopy(good)
    missing["calculations"] = [missing["calculations"][1]]  # fair value only
    problems = memo_pin_check.check_spine_pins_v2(missing, V2)
    assert any("base scenario MOIC 1.1x" in p for p in problems)
    dup = copy.deepcopy(good)
    dup["calculations"][1]["id"] = "C1"
    assert any(
        "duplicated" in p for p in memo_pin_check.check_spine_pins_v2(dup, V2)
    )


def test_contracts_describe_urls_and_citations():
    assert "`url` (optional)" in claude_runner.MEMO_PACKAGE_SOURCES_CONTRACT
    assert "[C2]" in claude_runner.HUMAN_EXEC_MEMO_VOICE_CONTRACT
    assert "## Citations (structure v2)" in claude_runner.MEMO_STRUCTURE_V2_ADDENDUM
    assert "keep every one exactly as written" in claude_runner._MEMO_BILINGUAL_STYLE
    pass_props = claude_runner.MEMO_FAST_PASS_SCHEMA["properties"]["supporting_evidence"][
        "items"
    ]["properties"]
    assert "url" in pass_props
