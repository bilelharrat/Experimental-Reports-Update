"""What an attached document looks like by the time the model sees it.

The bug these cover: an analyst attached a .docx and got "I wasn't able
to open that .docx directly — the Read tool reported it as binary".
"""

import io
import shutil

import pytest

from server import attachment_text, claude_runner


@pytest.fixture
def staged(tmp_path):
    """A workdir attachments dir, with a helper to put a file in it."""
    adir = tmp_path / "work" / "attachments"
    adir.mkdir(parents=True)

    def put(stored_name: str, text: str = "", display: str = "") -> str:
        (adir / stored_name).write_bytes(b"binary")
        if text:
            attachment_text.write_text(
                adir, stored_name, text, display_name=display or stored_name
            )
        return stored_name

    put.dir = adir  # type: ignore[attr-defined]
    return put


# ---- Extraction ---------------------------------------------------------


def test_word_tables_survive_as_rows():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Terms")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Pref"
    table.rows[0].cells[1].text = "1x"
    table.rows[1].cells[0].text = "Pool"
    table.rows[1].cells[1].text = "20%"
    buf = io.BytesIO()
    doc.save(buf)

    text = attachment_text.extract(
        buf.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="terms.docx",
    )
    assert "Terms" in text
    assert "Pref | 1x" in text
    assert "Pool | 20%" in text


def test_a_spreadsheet_comes_back_sheet_by_sheet(real_xlsx_bytes):
    text = attachment_text.extract(
        real_xlsx_bytes,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="cap.xlsx",
    )
    assert text.startswith("## Cap table")
    assert "Holder\tShares" in text
    assert "Ada Lovelace\t1,000,000" in text


def test_utf16_text_is_decoded_not_mangled():
    data = "Burn is ¥3.2M a month".encode("utf-16")
    assert attachment_text.extract(data, "text/plain", filename="n.txt") == (
        "Burn is ¥3.2M a month"
    )


def test_rtf_falls_back_to_stripping_control_words(monkeypatch):
    """CI has neither textutil nor LibreOffice; the prose still comes out."""
    monkeypatch.setattr(shutil, "which", lambda name: None)
    monkeypatch.setattr(
        "server.docx_pdf.soffice_path", lambda: None, raising=False
    )
    rtf = rb"{\rtf1\ansi\deff0 {\fonttbl{\f0 Arial;}}\f0\fs24 Pre-money $42M.\par}"
    assert "Pre-money $42M." in attachment_text.extract(
        rtf, "application/rtf", filename="terms.rtf"
    )


def test_a_broken_file_costs_nothing():
    """A file that resists extraction is worth an empty string, never a
    failed upload."""
    assert attachment_text.extract(b"PK\x03\x04garbage", "text/html") is not None
    assert (
        attachment_text.extract(
            b"PK\x03\x04garbage",
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document",
        )
        == ""
    )


# ---- What the prompt carries -------------------------------------------


def test_the_text_is_inlined_under_the_original_filename(staged):
    stored = staged("abc.docx", "Pre-money valuation of $42M.", "term-sheet.docx")

    block = attachment_text.prompt_block(staged.dir, [stored])

    assert "term-sheet.docx" in block
    assert "Pre-money valuation of $42M." in block
    # Nothing to open: the words are already here.
    assert "Read tool" not in block


def test_an_image_is_named_as_a_path_for_the_read_tool(staged):
    stored = staged("def.png")

    block = attachment_text.prompt_block(staged.dir, [stored])

    assert "`attachments/def.png`" in block
    assert "Read tool" in block


def test_a_long_document_is_capped_and_points_at_the_rest(staged):
    stored = staged("ghi.docx", "x" * 200_000, "long.docx")

    block = attachment_text.prompt_block(staged.dir, [stored])

    assert len(block) < 150_000
    assert "truncated" in block
    assert "`attachments/ghi.extract.txt`" in block


def test_only_a_file_nobody_could_read_pins_the_ask_to_claude(staged):
    doc = staged("jkl.docx", "Readable.", "memo.docx")
    scan = staged("mno.png")

    assert attachment_text.needs_native_read(staged.dir, []) is False
    assert attachment_text.needs_native_read(staged.dir, [doc]) is False
    assert attachment_text.needs_native_read(staged.dir, [doc, scan]) is True


def test_the_ask_command_carries_the_document_text(tmp_path, monkeypatch, staged):
    """End of the line: what `claude -p` is actually handed."""
    captured: dict = {}

    monkeypatch.setattr(claude_runner, "is_available", lambda: True)
    monkeypatch.setattr(claude_runner, "claude_path", lambda: "claude")
    def fake_spawn(cmd, cwd=None):
        captured["cmd"] = list(cmd)
        return type("H", (), {"proc": type("P", (), {"pid": 1})()})()

    monkeypatch.setattr(claude_runner, "_spawn_console", fake_spawn)
    monkeypatch.setattr(
        claude_runner, "_consume_stream",
        lambda handle, **kwargs: {"ok": True, "usage": {}, "cost_usd": 0.0},
    )
    skill = tmp_path / "skill.md"
    skill.write_text("persona", encoding="utf-8")
    stored = staged("pqr.docx", "Liquidation pref | 1x", "term-sheet.docx")

    claude_runner.run_console_ask(
        claude_session_id="sid-1",
        work_dir=staged.dir.parent,
        user_prompt="What is unusual here?",
        skill_path=skill,
        progress=type("P", (), {"emit": lambda *a, **k: None})(),
        attachments=[stored],
        tools="",
    )

    prompt = captured["cmd"][captured["cmd"].index("-p") + 1]
    assert "What is unusual here?" in prompt
    assert "term-sheet.docx" in prompt
    # The point of the whole exercise: with tools off, the text is still
    # in front of the model.
    assert "Liquidation pref | 1x" in prompt
