from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone

from docx import Document
from fastapi import HTTPException
import pytest

from server import (
    api,
    claude_runner,
    docx_pdf,
    job_progress,
    memo_analysis,
    memo_prep,
    storage,
)


@pytest.fixture
def memo_env(monkeypatch, tmp_path):
    data_root = tmp_path / "data"
    monkeypatch.setattr(storage, "DATA_DIR", data_root)
    monkeypatch.setattr(storage, "COMPANIES_FILE", data_root / "companies.yaml")
    monkeypatch.setattr(storage, "REPORTS_DIR", data_root / "reports")
    monkeypatch.setattr(storage, "THREADS_DIR", data_root / "threads")
    monkeypatch.setattr(memo_prep, "DATA_DIR", data_root)
    monkeypatch.setattr(memo_prep, "MEMOS_ROOT", data_root / "memos")
    monkeypatch.setattr(
        memo_prep,
        "SETTINGS_FILE",
        data_root / "settings" / "serena_background.md",
    )
    monkeypatch.setattr(memo_prep, "COMPANIES_FILE", data_root / "companies.yaml")
    monkeypatch.setenv("BSH_MEMO_GENERATE_INTERNAL", "1")
    monkeypatch.setenv("BSH_MEMO_FAST_PIPELINE", "0")
    return data_root


def _events(path):
    return [
        json.loads(line)
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


def _make_memo_report(data_root):
    run_id = "2026-05-22__191917"
    run_dir = (
        data_root
        / "memos"
        / "generalist-inc"
        / f"{run_id}__generalist-inc__memo-run"
    )
    memo_dir = run_dir / "memo"
    memo_dir.mkdir(parents=True)
    en_path = memo_dir / f"Generalist, Inc. - Investment Memo - {run_id}.docx"
    zh_path = memo_dir / f"Generalist, Inc. - 投资备忘录 - {run_id}.docx"
    internal_md = memo_dir / f"Generalist, Inc. - Internal Diligence Memo - {run_id}.md"
    internal_docx = memo_dir / f"Generalist, Inc. - Internal Diligence Memo - {run_id}.docx"
    _write_clean_memo_docx(en_path)
    _write_clean_memo_docx(zh_path)

    report = storage.create_report_record(
        company_id="generalist-inc",
        company_name="Generalist, Inc.",
        report_type=memo_prep.REPORT_TYPE,
        audience="Internal",
        language="en",
        kind="investment_memo_latestage",
        status="analyzing",
        progress=95,
        stage="Rendering PDF previews",
        run_id=run_id,
        run_dir=memo_prep._rel(run_dir),
        memo_files=[
            {"language": "en", "path": memo_prep._rel(en_path)},
            {"language": "zh", "path": memo_prep._rel(zh_path)},
        ],
        internal_memo_files=[
            {
                "kind": "internal_diligence_memo",
                "language": "en",
                "markdown_path": memo_prep._rel(internal_md),
                "path": memo_prep._rel(internal_docx),
            }
        ],
    )
    return report, run_dir


def _write_clean_memo_docx(path):
    document = Document()
    document.add_paragraph("I. Executive Summary")
    document.add_paragraph(
        "Generalist builds automation infrastructure. We recommend participating "
        "where deployment depth and valuation support are visible."
    )
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Treatment"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = (
        "Not disclosed; model uses customer-count proxy and valuation sensitivity."
    )
    document.add_paragraph("VI. Sources, Source Classes, and Fact Reference Index")
    document.add_paragraph("[S1] Company materials, company-reported.")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def _write_bad_memo_docx(path):
    document = Document()
    document.add_paragraph("I. Executive Summary")
    document.add_paragraph(
        "ZaiNar has no battery cost [WV SPV memo] and a hard IP wall "
        "(present-state)."
    )
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Last priced valuation"
    table.cell(1, 1).text = "~$1.0B post-money [companies.yaml]"
    document.save(path)


def _write_internal_memo_markdown(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        "\n".join([
            "# Internal Diligence Memo - Generalist, Inc.",
            "",
            "## Internal Recommendation",
            "Recommendation: Proceed.",
            "",
            "| Item | View |",
            "|---|---|",
            "| Suggested allocation | $5-10M with deployment-depth sensitivity |",
            "| Conviction | Medium |",
            "",
            "## Allocation Rationale",
            "The suggested allocation reflects scarce access, stage fit, and "
            "remaining deployment proof that can change commitment sizing.",
            "",
            "## Internal Diligence Priorities",
            "- Deployment depth shapes allocation size.",
            "- Valuation support shapes conviction.",
            "- SPV economics shape effective entry.",
            "",
            "## Risk Controls And Downside Sensitivities",
            "Customer proof that fails to support valuation weakens commitment sizing.",
            "",
            "Internal use only.",
        ]),
        encoding="utf-8",
    )


def _memo_package(body_en=None, body_zh=None):
    return {
        "schema_version": 1,
        "company": {
            "name": "Generalist, Inc.",
            "descriptor": {
                "en": "Industrial automation systems",
                "zh": "工业自动化系统",
            },
            "stage": "Late-stage / pre-IPO",
            "sector": "AI Robotics",
            "location": "San Francisco, CA",
            "round": "Series D context",
        },
        "run": {"run_id": "2026-05-22__191917", "as_of": "2026-05-22"},
        "sections": [
            {
                "id": "executive_summary",
                "blocks": [
                    {
                        "type": "paragraph",
                        "text": {
                            "en": body_en
                            or (
                                "Generalist builds automation infrastructure. "
                                "We recommend participating where deployment depth "
                                "and valuation support are visible."
                            ),
                            "zh": (
                                body_zh
                                or (
                                    "Generalist 构建自动化基础设施。BSH 只有在确认部署深度"
                                    "和估值支撑后才应继续推进。"
                                )
                            ),
                        },
                    },
                    {
                        "type": "table",
                        "component": "key_metrics_snapshot",
                        "title": {
                            "en": "Key Metrics Snapshot",
                            "zh": "关键指标快照",
                        },
                        "headers": [
                            {"en": "Metric", "zh": "指标"},
                            {"en": "Treatment", "zh": "处理方式"},
                        ],
                        "rows": [
                            [
                                {"en": "Revenue", "zh": "收入"},
                                {
                                    "en": (
                                        "Not disclosed; model uses customer-count "
                                        "proxy and valuation sensitivity."
                                    ),
                                    "zh": "未披露；模型使用客户数量代理和估值敏感因素。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "table",
                        "component": "deal_terms",
                        "title": {
                            "en": "Headline Terms",
                            "zh": "核心交易条款",
                        },
                        "headers": [
                            {"en": "Term", "zh": "条款"},
                            {"en": "Detail", "zh": "细节"},
                        ],
                        "rows": [
                            [
                                {"en": "Round", "zh": "轮次"},
                                {
                                    "en": "Series D context with valuation modeled by scenario range.",
                                    "zh": "Series D 背景，估值通过情景区间建模。",
                                },
                            ]
                        ],
                    },
                ],
            },
            {
                "id": "company_overview",
                "blocks": [
                    {
                        "type": "paragraph",
                        "text": {
                            "en": "The company sells automation systems for repeatable industrial workflows.",
                            "zh": "该公司销售面向可重复工业流程的自动化系统。",
                        },
                    },
                    {
                        "type": "table",
                        "component": "board",
                        "title": {"en": "Board of Directors", "zh": "董事会"},
                        "headers": [
                            {"en": "Member", "zh": "成员"},
                            {"en": "Strategic value", "zh": "战略价值"},
                        ],
                        "rows": [
                            [
                                {"en": "Lead investor director", "zh": "领投方董事"},
                                {
                                    "en": "Adds governance support for late-stage scaling.",
                                    "zh": "为后期扩张提供治理支持。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "table",
                        "component": "revenue",
                        "title": {"en": "Revenue Picture", "zh": "收入情况"},
                        "headers": [
                            {"en": "Component", "zh": "组成"},
                            {"en": "Treatment", "zh": "处理方式"},
                        ],
                        "rows": [
                            [
                                {"en": "Recognized revenue", "zh": "已确认收入"},
                                {
                                    "en": "Not disclosed; model uses deployment proxies.",
                                    "zh": "未披露；模型使用部署代理指标。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "table",
                        "component": "key_operating_metrics",
                        "title": {
                            "en": "Key Operating Metrics",
                            "zh": "关键运营指标",
                        },
                        "headers": [
                            {"en": "Metric", "zh": "指标"},
                            {"en": "Read", "zh": "解读"},
                        ],
                        "rows": [
                            [
                                {"en": "Gross margin", "zh": "毛利率"},
                                {
                                    "en": "Unavailable; margin path remains a valuation sensitivity.",
                                    "zh": "不可得；利润率路径仍是估值敏感因素。",
                                },
                            ]
                        ],
                    },
                ],
            },
            {
                "id": "investment_highlights",
                "blocks": [
                    {
                        "type": "bullets",
                        "items": [
                            {
                                "en": "Deployment proof creates a concrete diligence path.",
                                "zh": "部署验证提供了具体的尽调路径。",
                            },
                            {
                                "en": "Repeatable production usage can support the expansion case.",
                                "zh": "可重复的生产环境使用可支持扩张判断。",
                            }
                        ],
                    },
                    {
                        "type": "table",
                        "component": "competitive_analysis",
                        "title": {"en": "Competitive Analysis", "zh": "竞争分析"},
                        "headers": [
                            {"en": "Competitor", "zh": "竞争方"},
                            {"en": "Weakness vs. company", "zh": "相对弱点"},
                        ],
                        "rows": [
                            [
                                {"en": "Legacy vendors", "zh": "传统厂商"},
                                {
                                    "en": "Less integrated into targeted workflows.",
                                    "zh": "在目标流程中的集成较弱。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "table",
                        "component": "replacement_coexistence",
                        "title": {
                            "en": "Replacement vs. Coexistence",
                            "zh": "替代与共存",
                        },
                        "headers": [
                            {"en": "Workflow", "zh": "流程"},
                            {"en": "Read", "zh": "解读"},
                        ],
                        "rows": [
                            [
                                {"en": "Industrial workflow", "zh": "工业流程"},
                                {
                                    "en": "Coexists first, then replaces manual workflow layers.",
                                    "zh": "先共存，再替代人工流程层。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "table",
                        "component": "moat",
                        "title": {"en": "Moat", "zh": "护城河"},
                        "headers": [
                            {"en": "Component", "zh": "组成"},
                            {"en": "Durability", "zh": "持久性"},
                        ],
                        "rows": [
                            [
                                {"en": "Workflow lock-in", "zh": "流程锁定"},
                                {
                                    "en": "Strengthens as production integrations deepen.",
                                    "zh": "随着生产集成加深而增强。",
                                },
                            ]
                        ],
                    },
                ],
            },
            {
                "id": "investment_risk",
                "blocks": [
                    {
                        "type": "bullets",
                        "items": [
                            {
                                "en": "Hardware integration may slow gross-margin expansion.",
                                "zh": "硬件集成可能拖慢毛利率提升。",
                            },
                            {
                                "en": "Enterprise adoption may remain services-heavy.",
                                "zh": "企业采用可能继续偏服务交付。",
                            }
                        ],
                    },
                    {
                        "type": "table",
                        "component": "risk_register",
                        "title": {"en": "Risk Register", "zh": "风险清单"},
                        "headers": [
                            {"en": "Risk", "zh": "风险"},
                            {"en": "Severity", "zh": "严重性"},
                            {"en": "Likelihood", "zh": "可能性"},
                            {"en": "Mitigation", "zh": "缓释方式"},
                        ],
                        "rows": [
                            [
                                {"en": "Services-heavy adoption", "zh": "服务交付占比高"},
                                {"en": "Medium", "zh": "中"},
                                {"en": "Medium", "zh": "中"},
                                {
                                    "en": "Track repeatable deployment margin by cohort.",
                                    "zh": "按批次跟踪可重复部署利润率。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "callout",
                        "component": "disconfirming_evidence",
                        "tone": "warning",
                        "title": {
                            "en": "Bear-Case Evidence",
                            "zh": "熊市情景证据",
                        },
                        "items": [
                            {
                                "en": "Revenue disclosure remains limited relative to valuation.",
                                "zh": "相对于估值，收入披露仍有限。",
                            }
                        ],
                    },
                ],
            },
            {
                "id": "financial_forecast_valuation",
                "blocks": [
                    {
                        "type": "paragraph",
                        "text": {
                            "en": "The valuation case should use conservative scenario ranges.",
                            "zh": "估值判断应使用保守情景区间。",
                        },
                    },
                    {
                        "type": "table",
                        "component": "time_base_integrity",
                        "title": {
                            "en": "Time-Base Integrity Table",
                            "zh": "时间基准一致性表",
                        },
                        "headers": [
                            {"en": "Event", "zh": "事件"},
                            {"en": "Valuation", "zh": "估值"},
                        ],
                        "rows": [
                            [
                                {"en": "Current round", "zh": "当前轮次"},
                                {
                                    "en": "Modeled against latest disclosed valuation date.",
                                    "zh": "基于最近披露估值日期建模。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "table",
                        "component": "growth_bridge",
                        "title": {"en": "Growth Bridge Table", "zh": "增长桥接表"},
                        "headers": [
                            {"en": "Bridge item", "zh": "桥接项"},
                            {"en": "Model treatment", "zh": "模型处理"},
                        ],
                        "rows": [
                            [
                                {"en": "New deployments", "zh": "新增部署"},
                                {
                                    "en": "Credited only where repeatable usage is visible.",
                                    "zh": "仅在可重复使用可见时计入。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "table",
                        "component": "scenario_analysis",
                        "title": {"en": "Scenario Analysis", "zh": "情景分析"},
                        "headers": [
                            {"en": "Scenario", "zh": "情景"},
                            {"en": "Valuation outcome", "zh": "估值结果"},
                        ],
                        "rows": [
                            [
                                {"en": "Base", "zh": "基准"},
                                {
                                    "en": "Moderate step-up if deployments repeat with margin evidence.",
                                    "zh": "若部署可重复且利润率有证据，则温和上调。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "heading",
                        "level": 2,
                        "component": "investment_decision",
                        "text": {
                            "en": "Investment Decision / Closing View",
                            "zh": "投资决策 / 结论观点",
                        },
                    },
                    {
                        "type": "paragraph",
                        "text": {
                            "en": (
                                "We recommend participating where customer proof "
                                "and margin evidence support valuation."
                            ),
                            "zh": "若客户验证和利润率证据支撑估值，我们建议参与。",
                        },
                    },
                    {
                        "type": "table",
                        "component": "evidence_thresholds",
                        "title": {
                            "en": "Evidence Thresholds For Step-Up Support",
                            "zh": "支持估值上调的证据阈值",
                        },
                        "headers": [
                            {"en": "Evidence threshold", "zh": "证据阈值"},
                            {"en": "Valuation treatment", "zh": "估值处理"},
                        ],
                        "rows": [
                            [
                                {"en": "Repeatable production usage", "zh": "可重复生产使用"},
                                {
                                    "en": (
                                        "Supports a stronger step-up only when "
                                        "visible across multiple customers."
                                    ),
                                    "zh": "只有在多个客户中可见时，才支持更强的估值上调。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "paragraph",
                        "component": "disclosures",
                        "text": {
                            "en": (
                                "Disclosure: distributed only to the named recipients; "
                                "not an offer to sell securities. Any investment is "
                                "governed by definitive subscription documents and may "
                                "result in partial or total loss."
                            ),
                            "zh": (
                                "披露：仅供指定接收方使用，并非证券出售要约。任何投资均以"
                                "最终认购文件为准，并可能产生部分或全部损失。"
                            ),
                        },
                    },
                ],
            },
        ],
        "sources": [
            {
                "id": "S1",
                "title": "Company investor materials",
                "class": {"en": "Company material", "zh": "公司材料"},
                "treatment": {
                    "en": "Used for product, customer, and funding context.",
                    "zh": "用于产品、客户和融资背景。",
                },
                "as_of": "2026-05-22",
            }
        ],
    }


def _write_memo_package(run_dir, *, body_en=None, body_zh=None):
    path = run_dir / "logs" / "memo_package.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(
            _memo_package(body_en=body_en, body_zh=body_zh),
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return path


def _force_zh_to_english(value):
    if isinstance(value, dict):
        if "en" in value and "zh" in value:
            value["zh"] = value["en"]
        for child in value.values():
            _force_zh_to_english(child)
    elif isinstance(value, list):
        for item in value:
            _force_zh_to_english(item)


def test_memo_run_completes_when_optional_pdf_render_fails(
    memo_env, monkeypatch
):
    monkeypatch.setenv("BSH_MEMO_RENDER_PDF_PREVIEWS", "1")
    report, run_dir = _make_memo_report(memo_env)
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit(
        "job_init",
        kind="memo",
        title="Investment memo - Generalist, Inc.",
        report_id=report["id"],
        company_id="generalist-inc",
        run_id=report["run_id"],
    )

    def fake_run_investment_memo(**kwargs):
        _write_memo_package(run_dir)
        kwargs["progress"].emit(
            "claude_action",
            action="result",
            subtype="success",
            cost_usd=1.25,
            duration_ms=1234,
        )
        return {"ok": True, "cost_usd": 1.25, "duration_ms": 1234}

    monkeypatch.setattr(
        claude_runner, "run_investment_memo", fake_run_investment_memo
    )

    def fake_run_internal_diligence_memo(**kwargs):
        _write_internal_memo_markdown(kwargs["internal_markdown_path"])
        return {"ok": True, "cost_usd": 0.25, "duration_ms": 250}

    monkeypatch.setattr(
        claude_runner,
        "run_internal_diligence_memo",
        fake_run_internal_diligence_memo,
    )
    monkeypatch.setattr(
        docx_pdf,
        "convert_docx_to_pdf",
        lambda *_args, **_kwargs: (False, "Word conversion timed out"),
    )

    memo_analysis._run(report["id"])

    updated = storage.get_report(report["id"])
    assert updated["status"] == "complete"
    assert updated["progress"] == 100
    assert updated["stage"] == "Memo ready"
    assert not any("pdf_path" in f for f in updated["memo_files"])
    assert updated["internal_memo_files"][0]["path"].endswith(".docx")
    assert "pdf_path" not in updated["internal_memo_files"][0]
    assert (run_dir / "memo" / f"Generalist, Inc. - Internal Diligence Memo - {report['run_id']}.docx").exists()

    events = _events(memo_prep.stream_path(run_dir))
    assert events[-1]["type"] == "done"
    assert events[-1]["internal_memo_paths"]["md"].endswith(".md")
    assert any(
        e.get("type") == "claude_action"
        and e.get("tool") == "docx→pdf"
        and e.get("is_error")
        for e in events
    )


def test_memo_run_skips_pdf_and_internal_memo_by_default(memo_env, monkeypatch):
    monkeypatch.delenv("BSH_MEMO_RENDER_PDF_PREVIEWS", raising=False)
    monkeypatch.delenv("BSH_MEMO_GENERATE_INTERNAL", raising=False)
    report, run_dir = _make_memo_report(memo_env)
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit(
        "job_init",
        kind="memo",
        title="Investment memo — Generalist, Inc.",
        report_id=report["id"],
        company_id="generalist-inc",
        run_id=report["run_id"],
    )

    def fake_run_investment_memo(**kwargs):
        _write_memo_package(run_dir)
        kwargs["progress"].emit(
            "claude_action",
            action="result",
            subtype="success",
            cost_usd=1.25,
            duration_ms=1234,
        )
        return {"ok": True, "cost_usd": 1.25, "duration_ms": 1234}

    monkeypatch.setattr(
        claude_runner, "run_investment_memo", fake_run_investment_memo
    )
    monkeypatch.setattr(
        claude_runner,
        "run_internal_diligence_memo",
        lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("internal memo should be opt-in")
        ),
    )
    monkeypatch.setattr(
        docx_pdf,
        "convert_docx_to_pdf",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("PDF previews should be opt-in")
        ),
    )

    memo_analysis._run(report["id"])

    updated = storage.get_report(report["id"])
    assert updated["status"] == "complete"
    assert updated["stage"] == "Memo ready"
    assert not any("pdf_path" in f for f in updated["memo_files"])
    internal_docx = run_dir / "memo" / (
        f"Generalist, Inc. - Internal Diligence Memo - {report['run_id']}.docx"
    )
    assert not internal_docx.exists()

    events = _events(memo_prep.stream_path(run_dir))
    assert any(e.get("stage") == "pdf_previews_skipped" for e in events)
    assert any(e.get("stage") == "internal_memo_skipped" for e in events)
    assert events[-1]["type"] == "done"
    assert "internal_memo_paths" not in events[-1]


def test_memo_fast_pipeline_runs_parallel_passes_and_finalizes(
    memo_env, monkeypatch
):
    monkeypatch.setenv("BSH_MEMO_FAST_PIPELINE", "1")
    monkeypatch.delenv("BSH_MEMO_GENERATE_INTERNAL", raising=False)
    monkeypatch.delenv("BSH_MEMO_RENDER_PDF_PREVIEWS", raising=False)
    report, run_dir = _make_memo_report(memo_env)
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit(
        "job_init",
        kind="memo",
        title="Investment memo — Generalist, Inc.",
        report_id=report["id"],
        company_id="generalist-inc",
        run_id=report["run_id"],
    )
    called_passes = []

    def fake_analysis_pass(**kwargs):
        called_passes.append(kwargs["pass_id"])
        return {
            "summary": f"{kwargs['pass_label']} summary.",
            "key_findings": [
                {
                    "claim": "Commercial proof",
                    "finding": "Evidence supports a scoped diligence path.",
                    "evidence_class": "company-reported",
                    "implication": "Use as conditional support.",
                    "confidence": "medium",
                }
            ],
            "supporting_evidence": [],
            "disconfirming_evidence": [],
            "open_questions": [],
            "memo_uses": ["Use in the investment case."],
            "claude_cost_usd": 0.01,
            "claude_duration_ms": 100,
        }, None

    def fake_english_package(**kwargs):
        assert (run_dir / "analysis" / "fast").exists()
        return {
            "analysis_artifacts": {
                "claim_register_md": "# Claim Register\n\n- Commercial proof: supported.",
                "scenario_swim_lanes_md": "# Scenario Swim Lanes\n\n- Base: participate with contract-conversion sensitivity.",
                "downside_scenario_md": "# Downside Scenario\n\n- Deployment stalls.",
                "countercase_md": "# Countercase\n\n- Pass if valuation support fails.",
                "source_treatment_assumptions_md": (
                    "# Source Treatment And Assumptions\n\n- Revenue: not disclosed."
                ),
                "risk_sensitivities_md": (
                    "# Risk Sensitivities\n\n"
                    "1. Binding contract conversion supports valuation."
                ),
            },
            "memo_package": _memo_package(body_zh=""),
            "claude_cost_usd": 0.10,
            "claude_duration_ms": 500,
        }, None

    def fake_bilingual_package(**kwargs):
        assert kwargs["english_package_path"].name == "memo_package.en.json"
        return {
            "memo_package": _memo_package(),
            "claude_cost_usd": 0.05,
            "claude_duration_ms": 400,
        }, None

    monkeypatch.setattr(
        claude_runner,
        "run_memo_fast_analysis_pass",
        fake_analysis_pass,
    )
    monkeypatch.setattr(
        claude_runner,
        "run_memo_fast_english_package",
        fake_english_package,
    )
    monkeypatch.setattr(
        claude_runner,
        "run_memo_fast_bilingual_package",
        fake_bilingual_package,
    )
    monkeypatch.setattr(
        claude_runner,
        "run_investment_memo",
        lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("fast path should not call legacy memo runner")
        ),
    )

    memo_analysis._run(report["id"])

    updated = storage.get_report(report["id"])
    assert updated["status"] == "complete"
    assert updated["stage"] == "Memo ready"
    assert len(called_passes) == len(memo_analysis._FAST_MEMO_PASSES)
    assert (run_dir / "logs" / "memo_package.en.json").exists()
    assert (run_dir / "logs" / "memo_package.json").exists()
    assert (run_dir / "analysis" / "claim_register.md").exists()

    events = _events(memo_prep.stream_path(run_dir))
    assert any(e.get("stage") == "memo_fast_parallel_dispatch" for e in events)
    phase_events = [e for e in events if e.get("type") == "phase_timing"]
    phases = {(e.get("phase"), e.get("status")) for e in phase_events}
    assert ("memo_background_run", "started") in phases
    assert ("memo_background_run", "finished") in phases
    assert ("memo_fast_parallel_analysis", "finished") in phases
    assert ("memo_fast_english_package", "finished") in phases
    assert ("memo_fast_chinese_package", "finished") in phases
    assert ("memo_docx_render", "finished") in phases
    assert ("memo_chinese_parity_gate", "finished") in phases
    assert ("memo_quality_gate", "finished") in phases
    assert ("memo_pdf_previews", "skipped") in phases
    assert ("memo_internal_diligence", "skipped") in phases
    assert any(
        e.get("phase") == "memo_fast_parallel_analysis"
        and e.get("worker_count") == memo_analysis._memo_fast_max_workers()
        for e in phase_events
    )
    scanned = job_progress.scan_progress_state(memo_prep.stream_path(run_dir))
    assert any(
        e.get("phase") == "memo_docx_render"
        for e in scanned.get("phase_timings") or []
    )
    assert any(e.get("stage") == "pdf_previews_skipped" for e in events)
    assert events[-1]["type"] == "done"


def test_fast_synthesis_artifacts_normalize_legacy_private_labels(tmp_path):
    memo_analysis._write_fast_synthesis_artifacts(
        tmp_path,
        {
            "claim_register_md": "# Claim Register\n",
            "scenario_swim_lanes_md": "# Scenario Swim Lanes\n",
            "pre_mortem_md": "# Pre-Mortem\n\n- Deployment stalls.",
            "reverse_ic_md": "# Reverse IC\n\n- Valuation support fails.",
            "validation_log_md": "# Validation Log\n\n- Revenue is undisclosed.",
            "risk_sensitivities_md": "# Risk Sensitivities\n",
        },
    )

    analysis_dir = tmp_path / "analysis"
    assert not (analysis_dir / "pre_mortem.md").exists()
    assert not (analysis_dir / "reverse_ic.md").exists()
    assert not (analysis_dir / "validation_log.md").exists()
    assert (analysis_dir / "downside_scenario.md").read_text(
        encoding="utf-8"
    ).startswith("# Downside Scenario")
    assert (analysis_dir / "countercase.md").read_text(
        encoding="utf-8"
    ).startswith("# Countercase")
    assert (analysis_dir / "source_treatment_assumptions.md").read_text(
        encoding="utf-8"
    ).startswith("# Source Treatment And Assumptions")


def test_memo_fast_pipeline_retries_transient_english_package_failure(
    memo_env, monkeypatch
):
    monkeypatch.setenv("BSH_MEMO_FAST_PIPELINE", "1")
    monkeypatch.setenv("BSH_MEMO_FAST_ENGLISH_PACKAGE_RETRIES", "1")
    monkeypatch.setenv("BSH_MEMO_FAST_RETRY_BACKOFF_SEC", "0")
    monkeypatch.delenv("BSH_MEMO_GENERATE_INTERNAL", raising=False)
    monkeypatch.delenv("BSH_MEMO_RENDER_PDF_PREVIEWS", raising=False)
    report, run_dir = _make_memo_report(memo_env)
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit(
        "job_init",
        kind="memo",
        title="Investment memo — Generalist, Inc.",
        report_id=report["id"],
        company_id="generalist-inc",
        run_id=report["run_id"],
    )
    called_passes = []
    english_calls = []

    def fake_analysis_pass(**kwargs):
        called_passes.append(kwargs["pass_id"])
        return {
            "summary": f"{kwargs['pass_label']} summary.",
            "key_findings": [],
            "supporting_evidence": [],
            "disconfirming_evidence": [],
            "open_questions": [],
            "memo_uses": [],
            "claude_cost_usd": 0.01,
            "claude_duration_ms": 100,
        }, None

    def fake_english_package(**kwargs):
        english_calls.append(kwargs["run_id"])
        progress = kwargs["progress"]
        if len(english_calls) == 1:
            progress.emit(
                "claude_action",
                action="thinking",
                text=(
                    "API Error: The socket connection was closed unexpectedly. "
                    "For more information, pass verbose: true in the second "
                    "argument to fetch()"
                ),
            )
            progress.emit(
                "claude_action",
                action="result",
                subtype="success",
                cost_usd=0.71,
                duration_ms=254522,
            )
            return None, (
                "claude exited 1: API Error: The socket connection was closed "
                "unexpectedly"
            )
        progress.emit(
            "claude_action",
            action="result",
            subtype="success",
            cost_usd=0.10,
            duration_ms=500,
        )
        return {
            "analysis_artifacts": {
                "claim_register_md": "# Claim Register\n",
                "scenario_swim_lanes_md": "# Scenario Swim Lanes\n",
                "downside_scenario_md": "# Downside Scenario\n",
                "countercase_md": "# Countercase\n",
                "source_treatment_assumptions_md": "# Source Treatment And Assumptions\n",
                "risk_sensitivities_md": "# Risk Sensitivities\n",
            },
            "memo_package": _memo_package(body_zh=""),
            "claude_cost_usd": 0.10,
            "claude_duration_ms": 500,
        }, None

    monkeypatch.setattr(
        claude_runner,
        "run_memo_fast_analysis_pass",
        fake_analysis_pass,
    )
    monkeypatch.setattr(
        claude_runner,
        "run_memo_fast_english_package",
        fake_english_package,
    )
    monkeypatch.setattr(
        claude_runner,
        "run_memo_fast_bilingual_package",
        lambda **_kwargs: (
            {
                "memo_package": _memo_package(),
                "claude_cost_usd": 0.05,
                "claude_duration_ms": 400,
            },
            None,
        ),
    )

    memo_analysis._run(report["id"])

    updated = storage.get_report(report["id"])
    assert updated["status"] == "complete"
    assert len(called_passes) == len(memo_analysis._FAST_MEMO_PASSES)
    assert len(english_calls) == 2
    assert updated["claude_cost_usd"] >= 0.94

    events = _events(memo_prep.stream_path(run_dir))
    assert any(
        e.get("stage") == "memo_fast_english_package_retry_scheduled"
        for e in events
    )
    attempt_events = [
        e
        for e in events
        if e.get("type") == "phase_timing"
        and e.get("phase") == "memo_fast_english_package_attempt"
    ]
    assert any(e.get("status") == "failed" and e.get("transient") for e in attempt_events)
    assert any(e.get("status") == "finished" and e.get("attempt") == 2 for e in attempt_events)
    phase3_finished = [
        e
        for e in events
        if e.get("type") == "phase_timing"
        and e.get("phase") == "memo_fast_english_package"
        and e.get("status") == "finished"
    ]
    assert phase3_finished[-1]["attempts"] == 2
    assert phase3_finished[-1]["cost_usd"] >= 0.81
    assert not any(
        e.get("type") == "thread_failed"
        and e.get("thread") == claude_runner._MEMO_PHASE3_THREAD
        for e in events
    )


def test_memo_package_voice_cleanup_removes_quality_gate_terms(memo_env):
    report, run_dir = _make_memo_report(memo_env)
    bad_text = (
        "The relevant competitive surface is broader than legacy positioning "
        "vendors. We frame it as technology paradigms rather than a logo list, "
        "because the durable pricing-power question is whether incumbent "
        "paradigms absorb the function ZaiNar performs. We back infrastructure "
        "because this is why we want exposure. We still need a "
        "claim-scope review because we want exposure to network positioning. "
        "claim-scope and freedom-to-operate read before we underwrite the "
        "licensing fallback. The instrument has no preference, no voting, "
        "and no information rights at the LP level. This is the right way "
        "to view the underwriting: a pipeline-conversion case. Confirm before "
        "funding: final terms match the disclosed A2 economics. We still need "
        "a claim-scope read that the sponsor implies is covered by the IP "
        "portfolio. Memo language was closing imminent. Diligence Thresholds. "
        "Next Diligence Actions. The sponsor itself acknowledges that a "
        "meaningful portion is in MOU form and that pipeline figures are "
        "company-provided and unaudited. Sponsor explicitly discloses 18 to "
        "36 month carrier sales cycles. The competitor list embedded in the "
        "registry understates the threat surface. Revenue is not documented "
        "in source material. Closing Confirmations. We recommend proceeding "
        "with a participation in the Wisdom Ventures ZaiNar SPV subject to "
        "the closing confirmations below. Valuation Sensitivity Bars. Stop "
        "or Revisit Conditions. What Would Make Us Revisit or Decline. "
        "Immediate Confirmation Work. Closing bar: signed mix supports the case. "
        "Cross-check DoD signings on SAM.gov and USAspending.gov. Patent counsel "
        "claim-scope and freedom-to-operate read supports durable patent leverage. "
        "Source two non-investor technical references through the BSH and partner networks. "
        "BSH thesis fit relies on the late-stage financial-return exception. "
        "The disclosed founders are not Asian-immigrant per the BSH preference. "
        "Late-stage rules allow financial return to justify thesis exceptions if moat "
        "and multiple are compelling; the case therefore rests on the IP, channel, "
        "and contracted-traction case clearing on its own. Stop / Revisit Triggers. "
        "Series A2 slips materially beyond the May 2026 \"closing imminent\" framing, "
        "or reprices above approximately $3.53B pre-money, in which case the cap "
        "binds and the SAFE discount benefit erodes. The binding-contract share "
        "inside the $500M+ figure proves to be a small fraction of the headline, "
        "or the Kajima per-site economic is restated below approximately $5M ARR "
        "per site on a recurring basis. Series B pricing materially below the "
        "disclosed $15B target on a recapitalization or down-round path, shifting "
        "the SPV from a paper-mark outcome into a flat-to-modest carry for the "
        "holding period."
    )
    package_path = _write_memo_package(run_dir, body_en=bad_text)
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir), truncate=True)

    rewrite_count = memo_analysis._clean_memo_package_voice(package_path, stream)

    assert rewrite_count >= 1
    package_text = package_path.read_text(encoding="utf-8")
    assert "underwrite" not in package_text.lower()
    assert "underwriting" not in package_text.lower()
    assert "We frame" not in package_text
    assert "We back" not in package_text
    assert "we want exposure" not in package_text
    assert "BSH should" not in package_text
    assert "we recommend participating in network positioning" in package_text
    assert "BSH invests in infrastructure" in package_text
    assert "why the opportunity fits BSH's mandate" in package_text
    assert "information rights" not in package_text
    assert "We should confirm" not in package_text
    assert "Confirm before funding" not in package_text
    assert "We still need" not in package_text
    assert "sponsor implies" not in package_text
    assert "Memo language was" not in package_text
    assert "Diligence Thresholds" not in package_text
    assert "Next Diligence Actions" not in package_text
    assert "Closing Confirmations" not in package_text
    assert "Expected bar" not in package_text
    assert "subject to" not in package_text
    assert "Valuation Sensitivity Bars" not in package_text
    assert "Stop or Revisit Conditions" not in package_text
    assert "Stop / Revisit Triggers" not in package_text
    assert "What Would Make Us Revisit" not in package_text
    assert "Immediate Confirmation Work" not in package_text
    assert "Closing bar" not in package_text
    assert "Cross-check" not in package_text
    assert "Source two" not in package_text
    assert "sponsor itself acknowledges" not in package_text
    assert "Sponsor explicitly discloses" not in package_text
    assert "embedded in the registry" not in package_text
    assert "source material" not in package_text
    assert "late-stage financial-return exception" not in package_text
    assert "Asian-immigrant" not in package_text
    assert "BSH preference" not in package_text
    assert "thesis exceptions" not in package_text
    assert "paper-mark outcome" not in package_text
    assert "flat-to-modest carry" not in package_text
    assert "technical moat, channel access, and contracted traction" in package_text
    assert "Series A2 closing timing and pricing remain material" in package_text
    assert "Commercial quality depends on the binding-contract share" in package_text
    assert "flat-to-modest return profile" in package_text
    memo_paths_abs = memo_analysis._memo_paths_abs(report)
    memo_analysis.memo_docx_renderer.render_memos(
        package_path,
        out_en=memo_paths_abs["en"],
        out_zh=memo_paths_abs["zh"],
        manifest_path=run_dir / "logs" / "run_manifest.md",
        inventory_path=run_dir / "logs" / "file_inventory.md",
    )

    lint_result = memo_analysis.memo_quality_lint.lint_memo_docx(
        memo_paths_abs["en"]
    )
    assert not any(
        finding.code in {"sell_side_voice_violation", "meta_process_language"}
        for finding in lint_result.findings
    )
    events = _events(memo_prep.stream_path(run_dir))
    assert any(e.get("stage") == "memo_package_voice_cleanup" for e in events)


def test_memo_fast_pipeline_packet_mode_skips_parallel_passes(
    memo_env, monkeypatch, tmp_path
):
    monkeypatch.setenv("BSH_MEMO_FAST_PIPELINE", "1")
    monkeypatch.delenv("BSH_MEMO_GENERATE_INTERNAL", raising=False)
    report, run_dir = _make_memo_report(memo_env)
    packet_dir = tmp_path / "packet"
    packet_dir.mkdir()
    (packet_dir / "memo_packet.md").write_text("# Packet\n", encoding="utf-8")
    storage.update_report(
        report["id"],
        analysis_session_id="session-1",
        analysis_session_approved=True,
    )
    monkeypatch.setattr(
        memo_analysis.serena_analysis,
        "session_dir",
        lambda company_id, session_id: packet_dir,
    )
    monkeypatch.setattr(
        claude_runner,
        "run_memo_fast_analysis_pass",
        lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("approved packet mode should skip analysis fan-out")
        ),
    )

    def fake_english_package(**kwargs):
        assert kwargs["analysis_session_path"] == packet_dir
        return {
            "analysis_artifacts": {
                "claim_register_md": "# Claim Register\n",
                "scenario_swim_lanes_md": "# Scenario Swim Lanes\n",
                "downside_scenario_md": "# Downside Scenario\n",
                "countercase_md": "# Countercase\n",
                "source_treatment_assumptions_md": "# Source Treatment And Assumptions\n",
                "risk_sensitivities_md": "# Risk Sensitivities\n",
            },
            "memo_package": _memo_package(body_zh=""),
            "claude_cost_usd": 0.10,
            "claude_duration_ms": 500,
        }, None

    monkeypatch.setattr(
        claude_runner,
        "run_memo_fast_english_package",
        fake_english_package,
    )
    monkeypatch.setattr(
        claude_runner,
        "run_memo_fast_bilingual_package",
        lambda **_kwargs: (
            {
                "memo_package": _memo_package(),
                "claude_cost_usd": 0.05,
                "claude_duration_ms": 400,
            },
            None,
        ),
    )

    memo_analysis._run(report["id"])

    updated = storage.get_report(report["id"])
    assert updated["status"] == "complete"
    events = _events(memo_prep.stream_path(run_dir))
    assert any(e.get("stage") == "memo_fast_packet_mode" for e in events)
    assert not (run_dir / "analysis" / "fast").exists()


def test_memo_fast_pipeline_draft_packet_still_runs_parallel_passes(
    memo_env, monkeypatch, tmp_path
):
    monkeypatch.setenv("BSH_MEMO_FAST_PIPELINE", "1")
    monkeypatch.delenv("BSH_MEMO_GENERATE_INTERNAL", raising=False)
    report, run_dir = _make_memo_report(memo_env)
    packet_dir = tmp_path / "packet"
    packet_dir.mkdir()
    (packet_dir / "memo_packet.md").write_text("# Draft Packet\n", encoding="utf-8")
    storage.update_report(
        report["id"],
        analysis_session_id="session-1",
        analysis_session_approved=False,
    )
    monkeypatch.setattr(
        memo_analysis.serena_analysis,
        "session_dir",
        lambda company_id, session_id: packet_dir,
    )
    called_passes = []

    def fake_analysis_pass(**kwargs):
        called_passes.append(kwargs["pass_id"])
        return {
            "summary": f"{kwargs['pass_label']} summary.",
            "key_findings": [],
            "supporting_evidence": [],
            "disconfirming_evidence": [],
            "open_questions": [],
            "memo_uses": [],
            "claude_cost_usd": 0.01,
            "claude_duration_ms": 100,
        }, None

    def fake_english_package(**kwargs):
        assert kwargs["analysis_session_path"] is None
        return {
            "analysis_artifacts": {
                "claim_register_md": "# Claim Register\n",
                "scenario_swim_lanes_md": "# Scenario Swim Lanes\n",
                "downside_scenario_md": "# Downside Scenario\n",
                "countercase_md": "# Countercase\n",
                "source_treatment_assumptions_md": "# Source Treatment And Assumptions\n",
                "risk_sensitivities_md": "# Risk Sensitivities\n",
            },
            "memo_package": _memo_package(body_zh=""),
            "claude_cost_usd": 0.10,
            "claude_duration_ms": 500,
        }, None

    monkeypatch.setattr(
        claude_runner,
        "run_memo_fast_analysis_pass",
        fake_analysis_pass,
    )
    monkeypatch.setattr(
        claude_runner,
        "run_memo_fast_english_package",
        fake_english_package,
    )
    monkeypatch.setattr(
        claude_runner,
        "run_memo_fast_bilingual_package",
        lambda **_kwargs: (
            {
                "memo_package": _memo_package(),
                "claude_cost_usd": 0.05,
                "claude_duration_ms": 400,
            },
            None,
        ),
    )

    memo_analysis._run(report["id"])

    updated = storage.get_report(report["id"])
    assert updated["status"] == "complete"
    assert len(called_passes) == len(memo_analysis._FAST_MEMO_PASSES)
    events = _events(memo_prep.stream_path(run_dir))
    assert any(e.get("stage") == "memo_fast_parallel_dispatch" for e in events)
    assert not any(e.get("stage") == "memo_fast_packet_mode" for e in events)
    assert (run_dir / "analysis" / "fast").exists()


def test_memo_run_completes_with_warnings_when_chinese_parity_gate_finds_p0(
    memo_env, monkeypatch
):
    monkeypatch.setenv("BSH_MEMO_RENDER_PDF_PREVIEWS", "1")
    monkeypatch.delenv("BSH_MEMO_GENERATE_INTERNAL", raising=False)
    report, run_dir = _make_memo_report(memo_env)
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit(
        "job_init",
        kind="memo",
        title="Investment memo — Generalist, Inc.",
        report_id=report["id"],
        company_id="generalist-inc",
        run_id=report["run_id"],
    )

    def fake_run_investment_memo(**kwargs):
        package = _memo_package(
            body_zh=(
                "Generalist builds automation infrastructure. We recommend participating "
                "where deployment depth and valuation support are visible."
            )
        )
        for section in package["sections"]:
            if section["id"] == "company_overview":
                _force_zh_to_english(section)
        executive_table = package["sections"][0]["blocks"][1]
        executive_table["title"]["zh"] = "Key Metrics Snapshot"
        executive_table["headers"][0]["zh"] = "Metric"
        executive_table["headers"][1]["zh"] = "Treatment"
        executive_table["rows"][0][0]["zh"] = "Revenue"
        executive_table["rows"][0][1]["zh"] = (
            "Not disclosed; model uses customer-count proxy and valuation sensitivity."
        )
        path = run_dir / "logs" / "memo_package.json"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(
            json.dumps(package, ensure_ascii=False),
            encoding="utf-8",
        )
        kwargs["progress"].emit(
            "claude_action",
            action="result",
            subtype="success",
            cost_usd=1.25,
            duration_ms=1234,
        )
        return {"ok": True, "cost_usd": 1.25, "duration_ms": 1234}

    monkeypatch.setattr(
        claude_runner, "run_investment_memo", fake_run_investment_memo
    )
    def fake_convert_docx_to_pdf(docx_path, pdf_path):
        pdf_path.write_bytes(b"%PDF-1.4\n")
        return True, None

    monkeypatch.setattr(docx_pdf, "convert_docx_to_pdf", fake_convert_docx_to_pdf)

    memo_analysis._run(report["id"])

    updated = storage.get_report(report["id"])
    assert updated["status"] == "complete_with_warnings"
    assert updated["stage"] == "Memo ready (quality warnings)"
    assert updated["artifacts_available"] is True
    assert updated["memo_chinese_parity"]["p0_count"] >= 1
    assert updated["quality_warnings"]
    assert any("parity" in w.lower() for w in updated["quality_warnings"])
    assert all(f.get("pdf_path") for f in updated["memo_files"])
    assert (run_dir / "logs" / "memo_chinese_parity.md").exists()

    detail = api._report_detail(updated)
    assert detail["artifacts_available"] is True
    assert detail["memo_chinese_parity"]["p0_count"] >= 1
    assert detail["quality_warnings"]
    assert detail["download_urls"]["en"].endswith("language=en")
    assert detail["preview_urls"]["en"].endswith("language=en")

    events = _events(memo_prep.stream_path(run_dir))
    assert events[-1]["type"] == "done"
    assert events[-1]["quality_warnings"]
    warning_stages = [
        e for e in events
        if e.get("type") == "stage" and e.get("stage") == "chinese_parity_warning"
    ]
    assert warning_stages and warning_stages[0]["findings"]


def test_report_detail_advertises_internal_memo_urls(memo_env):
    report, run_dir = _make_memo_report(memo_env)
    internal_docx = run_dir / "memo" / (
        f"Generalist, Inc. - Internal Diligence Memo - {report['run_id']}.docx"
    )
    internal_pdf = internal_docx.with_suffix(".pdf")
    _write_clean_memo_docx(internal_docx)
    internal_pdf.write_bytes(b"%PDF-1.4\n")
    storage.update_report(
        report["id"],
        internal_memo_files=[
            {
                "kind": "internal_diligence_memo",
                "language": "en",
                "markdown_path": memo_prep._rel(internal_docx.with_suffix(".md")),
                "path": memo_prep._rel(internal_docx),
                "pdf_path": memo_prep._rel(internal_pdf),
            }
        ],
    )

    detail = api._report_detail(storage.get_report(report["id"]))

    assert detail["download_urls"]["internal"].endswith("artifact=internal")
    assert detail["preview_urls"]["internal"].endswith("artifact=internal")


def test_report_detail_advertises_partial_analysis_artifacts(memo_env):
    report, run_dir = _make_memo_report(memo_env)
    analysis_dir = run_dir / "analysis"
    analysis_dir.mkdir()
    (analysis_dir / "pressure_tests.md").write_text(
        "# Pressure tests\n\nGenerated before auth failure.\n",
        encoding="utf-8",
    )
    (analysis_dir / "custom_notes.md").write_text(
        "# Custom notes\n",
        encoding="utf-8",
    )
    storage.update_report(
        report["id"],
        status="failed_during_analysis",
        stage="Claude skill run failed",
        error="Failed to authenticate. API Error: 403 Request not allowed",
        failure_phase="analysis",
        failure_detail="Failed to authenticate. API Error: 403 Request not allowed",
        artifacts_available=False,
        memo_files=[],
        internal_memo_files=[],
    )

    detail = api._report_detail(storage.get_report(report["id"]))

    assert "download_urls" not in detail
    assert [a["filename"] for a in detail["analysis_artifacts"]] == [
        "pressure_tests.md",
        "custom_notes.md",
    ]
    assert detail["analysis_artifacts"][0]["label"] == "Arithmetic / pressure tests"
    assert detail["analysis_artifacts"][0]["download_url"].endswith(
        "artifact=analysis&file=pressure_tests.md"
    )
    assert detail["analysis_artifacts"][1]["label"] == "Custom Notes"
    assert api.ReportDetail(**detail).analysis_artifacts == detail["analysis_artifacts"]

    response = api.download_memo(
        report["id"],
        artifact="analysis",
        analysis_file="pressure_tests.md",
    )
    assert response.path.endswith("pressure_tests.md")
    with pytest.raises(HTTPException):
        api.download_memo(
            report["id"],
            artifact="analysis",
            analysis_file="../pressure_tests.md",
        )


def test_claude_transient_error_classifier_excludes_provider_limits():
    assert claude_runner.is_transient_claude_error(
        "claude exited 1: API Error: The socket connection was closed unexpectedly"
    )
    assert claude_runner.is_transient_claude_error("fetch failed")
    assert claude_runner.is_transient_claude_error(
        "memo English package stalled after 180s without output"
    )
    assert not claude_runner.is_transient_claude_error(
        "Failed to authenticate. API Error: 403 Request not allowed"
    )
    assert not claude_runner.is_transient_claude_error(
        "usage limit reached; try again later"
    )


def test_failed_memo_report_can_resume_from_analysis_artifacts(
    memo_env, monkeypatch
):
    report, run_dir = _make_memo_report(memo_env)
    analysis_dir = run_dir / "analysis"
    analysis_dir.mkdir()
    (analysis_dir / "pressure_tests.md").write_text(
        "# Pressure tests\n\nCompleted before provider failure.\n",
        encoding="utf-8",
    )
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit(
        "job_init",
        kind="memo",
        title="Investment memo — Generalist, Inc.",
        report_id=report["id"],
        company_id="generalist-inc",
        run_id=report["run_id"],
    )
    stream.emit("error", error="Failed to authenticate. API Error: 403 Request not allowed")
    storage.update_report(
        report["id"],
        status="failed_during_analysis",
        stage="Claude skill run failed",
        error="Failed to authenticate. API Error: 403 Request not allowed",
        failure_phase="analysis",
        failure_detail="Failed to authenticate. API Error: 403 Request not allowed",
        artifacts_available=False,
    )

    monkeypatch.setattr(
        claude_runner,
        "run_investment_memo",
        lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("resume must not rerun the full memo analysis")
        ),
    )

    def fake_resume_package(**kwargs):
        _write_memo_package(run_dir)
        kwargs["progress"].emit(
            "claude_action",
            action="result",
            subtype="success",
            cost_usd=0.5,
            duration_ms=500,
        )
        return {"ok": True, "cost_usd": 0.5, "duration_ms": 500, "resumed": True}

    monkeypatch.setattr(
        claude_runner,
        "run_resume_memo_package",
        fake_resume_package,
    )
    monkeypatch.setattr(
        claude_runner,
        "run_internal_diligence_memo",
        lambda **kwargs: (
            _write_internal_memo_markdown(kwargs["internal_markdown_path"])
            or {"ok": True, "cost_usd": 0.25, "duration_ms": 250}
        ),
    )
    monkeypatch.setattr(
        docx_pdf,
        "convert_docx_to_pdf",
        lambda docx_path, pdf_path: (pdf_path.write_bytes(b"%PDF-1.4\n") or True, None),
    )

    assert api._report_detail(storage.get_report(report["id"]))["resume_available"] is True

    memo_analysis._resume(report["id"])

    updated = storage.get_report(report["id"])
    assert updated["status"] == "complete"
    assert updated["stage"] == "Memo ready"
    assert updated["artifacts_available"] is True
    assert updated["failure_phase"] is None
    assert (run_dir / "logs" / "memo_package.json").exists()
    assert list((run_dir / "logs").glob("stream.before_resume.*.jsonl"))

    events = _events(memo_prep.stream_path(run_dir))
    assert events[0]["type"] == "job_init"
    assert events[0]["resumed"] is True
    assert not any(
        event.get("error") == "Failed to authenticate. API Error: 403 Request not allowed"
        for event in events
    )
    assert events[-1]["type"] == "done"
    assert events[-1]["recovered"] is True


def test_resume_retries_transient_resume_package_failure(
    memo_env, monkeypatch
):
    monkeypatch.delenv("BSH_MEMO_GENERATE_INTERNAL", raising=False)
    monkeypatch.setenv("BSH_MEMO_RESUME_PACKAGE_RETRIES", "1")
    monkeypatch.setenv("BSH_MEMO_FAST_RETRY_BACKOFF_SEC", "0")
    report, run_dir = _make_memo_report(memo_env)
    analysis_dir = run_dir / "analysis"
    analysis_dir.mkdir()
    (analysis_dir / "pressure_tests.md").write_text(
        "# Pressure tests\n\nCompleted before provider failure.\n",
        encoding="utf-8",
    )
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit(
        "job_init",
        kind="memo",
        title="Investment memo - Generalist, Inc.",
        report_id=report["id"],
        company_id="generalist-inc",
        run_id=report["run_id"],
    )
    storage.update_report(
        report["id"],
        status="failed_during_analysis",
        stage="Claude skill run failed",
        error="claude exited 1",
        failure_phase="analysis",
        failure_detail="claude exited 1",
        artifacts_available=False,
    )
    calls = []

    def fake_resume_package(**kwargs):
        calls.append(kwargs["run_id"])
        if len(calls) == 1:
            return {
                "ok": False,
                "error": (
                    "claude exited 1: API Error: The socket connection was "
                    "closed unexpectedly"
                ),
                "cost_usd": 0.71,
                "duration_ms": 254522,
            }
        _write_memo_package(run_dir)
        return {"ok": True, "cost_usd": 0.5, "duration_ms": 500, "resumed": True}

    monkeypatch.setattr(
        claude_runner,
        "run_resume_memo_package",
        fake_resume_package,
    )
    monkeypatch.setattr(
        claude_runner,
        "run_internal_diligence_memo",
        lambda **_kwargs: (_ for _ in ()).throw(
            AssertionError("internal memo should be opt-in")
        ),
    )

    memo_analysis._resume(report["id"])

    updated = storage.get_report(report["id"])
    assert updated["status"] == "complete"
    assert len(calls) == 2
    events = _events(memo_prep.stream_path(run_dir))
    assert any(e.get("stage") == "resume_package_retry_scheduled" for e in events)
    attempts = [
        e
        for e in events
        if e.get("type") == "phase_timing"
        and e.get("phase") == "memo_resume_package_attempt"
    ]
    assert any(e.get("status") == "failed" and e.get("transient") for e in attempts)
    assert any(e.get("status") == "finished" and e.get("attempt") == 2 for e in attempts)


def test_resume_regenerates_existing_invalid_memo_package(
    memo_env, monkeypatch
):
    report, run_dir = _make_memo_report(memo_env)
    analysis_dir = run_dir / "analysis"
    analysis_dir.mkdir()
    (analysis_dir / "pressure_tests.md").write_text(
        "# Pressure tests\n\nCompleted before renderer failure.\n",
        encoding="utf-8",
    )
    invalid_package = _memo_package()
    invalid_package["sections"][0]["blocks"].append(
        {
            "type": "table",
            "title": {"en": "Invalid summary", "zh": "无效摘要"},
            "headers": [
                {"en": "Metric", "zh": "指标"},
                {"en": "Value", "zh": "数值"},
            ],
            "rows": [
                [
                    {"en": "Total", "zh": "合计"},
                    {"en": "", "zh": ""},
                ]
            ],
        }
    )
    package_path = run_dir / "logs" / "memo_package.json"
    package_path.parent.mkdir(parents=True, exist_ok=True)
    package_path.write_text(
        json.dumps(invalid_package, ensure_ascii=False),
        encoding="utf-8",
    )
    storage.update_report(
        report["id"],
        status="failed_during_analysis",
        stage="Renderer contract failed",
        failure_phase="renderer_contract",
        failure_detail="Invalid memo package",
    )

    def fake_resume_package(**kwargs):
        assert not package_path.exists()
        assert list((run_dir / "logs").glob("memo_package.invalid.*.json"))
        _write_memo_package(run_dir)
        kwargs["progress"].emit(
            "claude_action",
            action="result",
            subtype="success",
            cost_usd=0.5,
            duration_ms=500,
        )
        return {"ok": True, "cost_usd": 0.5, "duration_ms": 500, "resumed": True}

    monkeypatch.setattr(
        claude_runner,
        "run_resume_memo_package",
        fake_resume_package,
    )
    monkeypatch.setattr(
        claude_runner,
        "run_internal_diligence_memo",
        lambda **kwargs: (
            _write_internal_memo_markdown(kwargs["internal_markdown_path"])
            or {"ok": True, "cost_usd": 0.25, "duration_ms": 250}
        ),
    )
    monkeypatch.setattr(
        docx_pdf,
        "convert_docx_to_pdf",
        lambda docx_path, pdf_path: (pdf_path.write_bytes(b"%PDF-1.4\n") or True, None),
    )

    memo_analysis._resume(report["id"])

    updated = storage.get_report(report["id"])
    assert updated["status"] == "complete"
    events = _events(memo_prep.stream_path(run_dir))
    assert any(event.get("stage") == "resume_package_invalid" for event in events)
    assert events[-1]["type"] == "done"


def test_resume_regenerates_quality_failed_memo_package(
    memo_env, monkeypatch
):
    report, run_dir = _make_memo_report(memo_env)
    analysis_dir = run_dir / "analysis"
    analysis_dir.mkdir()
    (analysis_dir / "pressure_tests.md").write_text(
        "# Pressure tests\n\nCompleted before quality gate failure.\n",
        encoding="utf-8",
    )
    package_path = _write_memo_package(
        run_dir,
        body_en=(
            "Confirm before funding: Series A2 closes at or above $3.0B "
            "pre-money with a named institutional lead."
        ),
    )
    quality_lint_path = run_dir / "logs" / "memo_quality_lint.md"
    quality_lint_path.write_text(
        "\n".join([
            "# Memo Quality Lint",
            "",
            "- status: failed",
            "- p0_count: 1",
            "",
            "| Severity | Code | Location | Snippet | Suggestion |",
            "|---|---|---|---|---|",
            "| P0 | sell_side_voice_violation | paragraph 1 | Confirm before funding | Rewrite buyer-side diligence wording. |",
        ]),
        encoding="utf-8",
    )
    storage.update_report(
        report["id"],
        status="failed_quality_gate",
        stage="Memo failed quality gate",
        failure_phase="quality_gate",
        failure_detail="Generated memo failed the DOCX quality gate",
        artifacts_available=True,
        memo_quality_lint={"p0_count": 1, "finding_count": 1},
    )

    stored = storage.get_report(report["id"])
    summary = api._report_summary(stored)
    detail = api._report_detail(stored)
    assert summary["resume_available"] is True
    assert detail["resume_available"] is True
    storage.update_report(
        report["id"],
        status="analyzing",
        stage="Resume queued",
        progress=98,
        resume_from_status="failed_quality_gate",
        resume_from_failure_phase="quality_gate",
        resume_from_failure_detail="Generated memo failed the DOCX quality gate",
        failure_phase=None,
        failure_detail=None,
    )

    def fake_resume_package(**kwargs):
        assert not package_path.exists()
        archives = list((run_dir / "logs").glob("memo_package.quality_failed.*.json"))
        assert archives
        assert kwargs["quality_lint_path"] == quality_lint_path
        assert kwargs["prior_package_path"] == archives[0]
        _write_memo_package(run_dir)
        kwargs["progress"].emit(
            "claude_action",
            action="result",
            subtype="success",
            cost_usd=0.5,
            duration_ms=500,
        )
        return {"ok": True, "cost_usd": 0.5, "duration_ms": 500, "resumed": True}

    monkeypatch.setattr(
        claude_runner,
        "run_resume_memo_package",
        fake_resume_package,
    )
    monkeypatch.setattr(
        claude_runner,
        "run_internal_diligence_memo",
        lambda **kwargs: (
            _write_internal_memo_markdown(kwargs["internal_markdown_path"])
            or {"ok": True, "cost_usd": 0.25, "duration_ms": 250}
        ),
    )
    monkeypatch.setattr(
        docx_pdf,
        "convert_docx_to_pdf",
        lambda docx_path, pdf_path: (pdf_path.write_bytes(b"%PDF-1.4\n") or True, None),
    )

    memo_analysis._resume(report["id"])

    updated = storage.get_report(report["id"])
    assert updated["status"] == "complete"
    assert updated["stage"] == "Memo ready"
    assert updated["failure_phase"] is None
    events = _events(memo_prep.stream_path(run_dir))
    assert any(
        event.get("stage") == "resume_package_quality_failed"
        for event in events
    )
    assert not any(
        event.get("stage") == "resume_package_reuse"
        for event in events
    )
    assert events[-1]["type"] == "done"


def test_resume_memo_endpoint_queues_failed_report(memo_env, monkeypatch):
    report, run_dir = _make_memo_report(memo_env)
    analysis_dir = run_dir / "analysis"
    analysis_dir.mkdir()
    (analysis_dir / "pressure_tests.md").write_text(
        "# Pressure tests\n",
        encoding="utf-8",
    )
    storage.update_report(
        report["id"],
        status="failed_during_analysis",
        stage="Claude skill run failed",
        failure_phase="analysis",
        failure_detail="Failed to authenticate. API Error: 403 Request not allowed",
    )
    called = []
    monkeypatch.setattr(
        memo_analysis,
        "start_resume",
        lambda report_id: called.append(report_id),
    )

    response = api.resume_memo_report(_admin_request(), report["id"])

    assert called == [report["id"]]
    assert response.status == "analyzing"
    assert response.stage == "Resume queued"
    queued = storage.get_report(report["id"])
    assert queued["resume_from_status"] == "failed_during_analysis"
    assert queued["resume_from_failure_phase"] == "analysis"
    assert (
        queued["resume_from_failure_detail"]
        == "Failed to authenticate. API Error: 403 Request not allowed"
    )


def test_resume_memo_endpoint_preserves_original_resume_source(
    memo_env, monkeypatch
):
    report, run_dir = _make_memo_report(memo_env)
    analysis_dir = run_dir / "analysis"
    analysis_dir.mkdir()
    (analysis_dir / "pressure_tests.md").write_text(
        "# Pressure tests\n",
        encoding="utf-8",
    )
    storage.update_report(
        report["id"],
        status="failed_during_analysis",
        stage="Memo resume failed",
        failure_phase="resume",
        failure_detail="API Error: socket closed",
        resume_from_status="failed_quality_gate",
        resume_from_failure_phase="quality_gate",
        resume_from_failure_detail="Generated memo failed quality gate",
    )
    called = []
    monkeypatch.setattr(
        memo_analysis,
        "start_resume",
        lambda report_id: called.append(report_id),
    )

    response = api.resume_memo_report(_admin_request(), report["id"])

    assert called == [report["id"]]
    assert response.status == "analyzing"
    queued = storage.get_report(report["id"])
    assert queued["resume_from_status"] == "failed_quality_gate"
    assert queued["resume_from_failure_phase"] == "quality_gate"
    assert queued["resume_from_failure_detail"] == "Generated memo failed quality gate"


def test_resume_continues_from_quality_failed_archive_after_transport_error(
    memo_env, monkeypatch
):
    report, run_dir = _make_memo_report(memo_env)
    analysis_dir = run_dir / "analysis"
    analysis_dir.mkdir()
    (analysis_dir / "pressure_tests.md").write_text(
        "# Pressure tests\n\nCompleted before transport failure.\n",
        encoding="utf-8",
    )
    original_package = _write_memo_package(
        run_dir,
        body_en="Confirm before funding: buyer-side checkpoint language.",
    )
    archived_package = original_package.with_name(
        "memo_package.quality_failed.20260623T114731Z.json"
    )
    original_package.replace(archived_package)
    quality_lint_path = run_dir / "logs" / "memo_quality_lint.md"
    quality_lint_path.write_text(
        "# Memo Quality Lint\n\nP0 sell_side_voice_violation\n",
        encoding="utf-8",
    )
    storage.update_report(
        report["id"],
        status="failed_during_analysis",
        stage="Memo resume failed",
        failure_phase="resume",
        failure_detail="API Error: socket closed",
        artifacts_available=True,
    )

    def fake_resume_package(**kwargs):
        assert kwargs["quality_lint_path"] == quality_lint_path
        assert kwargs["prior_package_path"] == archived_package
        _write_memo_package(run_dir)
        kwargs["progress"].emit(
            "claude_action",
            action="result",
            subtype="success",
            cost_usd=0.5,
            duration_ms=500,
        )
        return {"ok": True, "cost_usd": 0.5, "duration_ms": 500, "resumed": True}

    monkeypatch.setattr(
        claude_runner,
        "run_resume_memo_package",
        fake_resume_package,
    )
    monkeypatch.setattr(
        claude_runner,
        "run_internal_diligence_memo",
        lambda **kwargs: (
            _write_internal_memo_markdown(kwargs["internal_markdown_path"])
            or {"ok": True, "cost_usd": 0.25, "duration_ms": 250}
        ),
    )
    monkeypatch.setattr(
        docx_pdf,
        "convert_docx_to_pdf",
        lambda docx_path, pdf_path: (pdf_path.write_bytes(b"%PDF-1.4\n") or True, None),
    )

    memo_analysis._resume(report["id"])

    updated = storage.get_report(report["id"])
    assert updated["status"] == "complete"
    events = _events(memo_prep.stream_path(run_dir))
    assert any(
        event.get("stage") == "resume_package_quality_failed_continue"
        for event in events
    )
    assert not any(
        event.get("stage") == "resume_package_reuse"
        for event in events
    )


def test_memo_run_completes_with_warnings_when_docx_quality_gate_finds_p0(
    memo_env, monkeypatch
):
    monkeypatch.setenv("BSH_MEMO_RENDER_PDF_PREVIEWS", "1")
    monkeypatch.delenv("BSH_MEMO_GENERATE_INTERNAL", raising=False)
    report, run_dir = _make_memo_report(memo_env)
    analysis_dir = run_dir / "analysis"
    analysis_dir.mkdir(exist_ok=True)
    (analysis_dir / "pressure_tests.md").write_text("# Notes\n", encoding="utf-8")
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit(
        "job_init",
        kind="memo",
        title="Investment memo — Generalist, Inc.",
        report_id=report["id"],
        company_id="generalist-inc",
        run_id=report["run_id"],
    )

    def fake_run_investment_memo(**kwargs):
        _write_memo_package(
            run_dir,
            body_en=(
                "ZaiNar has no battery cost [WV SPV memo] and a hard IP wall "
                "(present-state)."
            ),
        )
        kwargs["progress"].emit(
            "claude_action",
            action="result",
            subtype="success",
            cost_usd=1.25,
            duration_ms=1234,
        )
        return {"ok": True, "cost_usd": 1.25, "duration_ms": 1234}

    monkeypatch.setattr(
        claude_runner, "run_investment_memo", fake_run_investment_memo
    )
    def fake_convert_docx_to_pdf(docx_path, pdf_path):
        pdf_path.write_bytes(b"%PDF-1.4\n")
        return True, None

    monkeypatch.setattr(docx_pdf, "convert_docx_to_pdf", fake_convert_docx_to_pdf)

    memo_analysis._run(report["id"])

    updated = storage.get_report(report["id"])
    assert updated["status"] == "complete_with_warnings"
    assert updated["stage"] == "Memo ready (quality warnings)"
    assert updated["artifacts_available"] is True
    assert updated["memo_quality_lint"]["p0_count"] >= 1
    assert updated["quality_warnings"]
    assert all(f.get("pdf_path") for f in updated["memo_files"])
    assert (run_dir / "logs" / "memo_quality_lint.md").exists()
    assert updated["renderer_contract"]["errors"] == []

    summary = api._report_summary(updated)
    assert summary["artifacts_available"] is True
    assert summary["memo_quality_lint"]["p0_count"] >= 1
    assert summary["download_urls"]["en"].endswith("language=en")
    # Complete-with-warnings keeps Resume available for regeneration.
    assert summary["resume_available"] is True
    detail = api._report_detail(updated)
    assert detail["artifacts_available"] is True
    assert detail["memo_quality_lint"]["p0_count"] >= 1
    assert detail["quality_warnings"]
    assert detail["download_urls"]["en"].endswith("language=en")
    assert detail["preview_urls"]["en"].endswith("language=en")

    events = _events(memo_prep.stream_path(run_dir))
    assert events[-1]["type"] == "done"
    assert events[-1]["quality_warnings"]
    warning_stages = [
        e for e in events
        if e.get("type") == "stage" and e.get("stage") == "quality_gate_warning"
    ]
    assert warning_stages and warning_stages[0]["findings"]


def test_memo_run_fails_when_memo_package_missing(memo_env, monkeypatch):
    report, run_dir = _make_memo_report(memo_env)
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit(
        "job_init",
        kind="memo",
        title="Investment memo — Generalist, Inc.",
        report_id=report["id"],
        company_id="generalist-inc",
        run_id=report["run_id"],
    )

    def fake_run_investment_memo(**kwargs):
        kwargs["progress"].emit(
            "claude_action",
            action="result",
            subtype="success",
            cost_usd=1.25,
            duration_ms=1234,
        )
        return {"ok": True, "cost_usd": 1.25, "duration_ms": 1234}

    monkeypatch.setattr(
        claude_runner, "run_investment_memo", fake_run_investment_memo
    )
    monkeypatch.setattr(
        docx_pdf,
        "convert_docx_to_pdf",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("PDF rendering should not run after package failure")
        ),
    )

    memo_analysis._run(report["id"])

    updated = storage.get_report(report["id"])
    assert updated["status"] == "failed_during_analysis"
    assert updated["stage"] == "Renderer contract failed"
    assert updated["failure_phase"] == "renderer_contract"
    assert "memo_package.json" in updated["failure_detail"]
    assert updated["renderer_contract"]["errors"]

    events = _events(memo_prep.stream_path(run_dir))
    assert events[-1]["type"] == "error"
    assert events[-1]["phase"] == "renderer_contract"
    assert "memo_package.json" in events[-1]["error"]
    assert "memo_package missing" in events[-1]["contract_errors"]
    assert any(
        file.get("label") == "memo_package" and file.get("exists") is False
        for file in events[-1]["expected_files"]
    )

    detail = api._report_detail(updated)
    assert detail["failure_detail"] == updated["failure_detail"]
    assert detail["renderer_contract"]["errors"]


def test_memo_run_blocks_generated_renderer_scripts(memo_env, monkeypatch):
    report, run_dir = _make_memo_report(memo_env)
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit(
        "job_init",
        kind="memo",
        title="Investment memo — Generalist, Inc.",
        report_id=report["id"],
        company_id="generalist-inc",
        run_id=report["run_id"],
    )

    def fake_run_investment_memo(**kwargs):
        (run_dir / "build_memo.py").write_text(
            "# generated renderer should be blocked\n",
            encoding="utf-8",
        )
        kwargs["progress"].emit(
            "claude_action",
            action="result",
            subtype="success",
            cost_usd=1.25,
            duration_ms=1234,
        )
        return {"ok": True, "cost_usd": 1.25, "duration_ms": 1234}

    monkeypatch.setattr(
        claude_runner, "run_investment_memo", fake_run_investment_memo
    )
    monkeypatch.setattr(
        docx_pdf,
        "convert_docx_to_pdf",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("PDF rendering should not run after renderer violation")
        ),
    )

    memo_analysis._run(report["id"])

    updated = storage.get_report(report["id"])
    assert updated["status"] == "failed_during_analysis"
    assert updated["stage"] == "Generated renderer script blocked"

    events = _events(memo_prep.stream_path(run_dir))
    assert events[-1]["type"] == "error"
    assert events[-1]["phase"] == "renderer_contract"
    assert events[-1]["generated_renderer_scripts"]


def test_recover_stale_memo_report_emits_missing_done(memo_env):
    report, run_dir = _make_memo_report(memo_env)
    _write_memo_package(run_dir)
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit(
        "job_init",
        kind="memo",
        title="Investment memo — Generalist, Inc.",
        report_id=report["id"],
        company_id="generalist-inc",
        run_id=report["run_id"],
    )
    stream.emit("thread_started", thread="Source treatment and assumptions")
    stream.emit(
        "claude_action",
        action="result",
        subtype="success",
        cost_usd=9.02,
        duration_ms=1845017,
    )

    recovered = memo_analysis.recover_stale_reports()

    assert recovered == 1
    updated = storage.get_report(report["id"])
    assert updated["status"] == "complete"
    assert updated["progress"] == 100
    assert updated["claude_cost_usd"] == 9.02

    events = _events(memo_prep.stream_path(run_dir))
    assert any(
        event.get("type") == "thread_finished"
        and event.get("thread") == "Source treatment and assumptions"
        for event in events
    )
    assert events[-1]["type"] == "done"
    assert events[-1]["recovered"] is True

    state = api._scan_progress_state(memo_prep.stream_path(run_dir))
    assert state["terminated"] is True


def test_recover_stale_memo_report_repairs_failed_record_when_stream_done(
    memo_env,
):
    report, run_dir = _make_memo_report(memo_env)
    package_path = _write_memo_package(run_dir)
    memo_paths_abs = memo_analysis._memo_paths_abs(report)
    memo_analysis.memo_docx_renderer.render_memos(
        package_path,
        out_en=memo_paths_abs["en"],
        out_zh=memo_paths_abs["zh"],
        manifest_path=run_dir / "logs" / "run_manifest.md",
        inventory_path=run_dir / "logs" / "file_inventory.md",
    )
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir), truncate=True)
    stream.emit(
        "job_init",
        kind="memo",
        title="Resume investment memo - Generalist, Inc.",
        report_id=report["id"],
        company_id="generalist-inc",
        run_id=report["run_id"],
        resumed=True,
    )
    stream.emit(
        "done",
        report_id=report["id"],
        memo_paths={k: str(v) for k, v in memo_paths_abs.items()},
        cost_usd=2.25,
        duration_ms=111000,
        recovered=True,
    )
    storage.update_report(
        report["id"],
        status="failed_during_analysis",
        stage="Memo resume failed",
        error="API Error: The socket connection was closed unexpectedly.",
        failure_phase="resume",
        failure_detail="API Error: The socket connection was closed unexpectedly.",
        artifacts_available=False,
    )

    recovered = memo_analysis.recover_stale_reports()

    assert recovered == 1
    updated = storage.get_report(report["id"])
    assert updated["status"] == "complete"
    assert updated["stage"] == "Memo ready"
    assert updated["error"] is None
    assert updated["failure_phase"] is None
    assert updated["artifacts_available"] is True
    assert updated["claude_cost_usd"] == 2.25
    assert updated["claude_duration_ms"] == 111000
    assert updated["memo_quality_lint"]["p0_count"] == 0
    assert updated["memo_chinese_parity"]["p0_count"] == 0
    assert all(
        item["exists"]
        for item in updated["renderer_contract"]["expected_files"]
    )

    events = _events(memo_prep.stream_path(run_dir))
    assert sum(1 for event in events if event.get("type") == "done") == 1


def test_recover_stale_memo_report_blocks_generated_renderer_scripts(memo_env):
    report, run_dir = _make_memo_report(memo_env)
    (run_dir / "build_memo.py").write_text(
        "# generated renderer should be blocked during recovery\n",
        encoding="utf-8",
    )
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit(
        "job_init",
        kind="memo",
        title="Investment memo — Generalist, Inc.",
        report_id=report["id"],
        company_id="generalist-inc",
        run_id=report["run_id"],
    )
    stream.emit(
        "claude_action",
        action="result",
        subtype="success",
        cost_usd=9.02,
        duration_ms=1845017,
    )

    recovered = memo_analysis.recover_stale_reports()

    assert recovered == 0
    updated = storage.get_report(report["id"])
    assert updated["status"] == "failed_during_analysis"
    assert updated["stage"] == "Generated renderer script blocked"

    events = _events(memo_prep.stream_path(run_dir))
    assert events[-1]["type"] == "error"
    assert events[-1]["phase"] == "renderer_contract"
    assert events[-1]["recovered"] is True
    assert events[-1]["generated_renderer_scripts"]


def test_active_memo_job_registers_subtask_completion(memo_env):
    report, run_dir = _make_memo_report(memo_env)
    stream_path = memo_prep.stream_path(run_dir)
    stream_path.parent.mkdir(parents=True, exist_ok=True)
    base_ts = datetime.now(timezone.utc) - timedelta(seconds=10)

    def ts(offset_s: int) -> str:
        return (base_ts + timedelta(seconds=offset_s)).isoformat()

    stream_path.write_text(
        "\n".join(
            json.dumps(event)
            for event in [
                {
                    "type": "job_init",
                    "ts": ts(0),
                    "kind": "memo",
                    "title": "Investment memo — Generalist, Inc.",
                    "report_id": report["id"],
                    "company_id": "generalist-inc",
                    "run_id": report["run_id"],
                },
                {
                    "type": "thread_started",
                    "ts": ts(1),
                    "thread": "Pressure tests",
                },
                {
                    "type": "claude_action",
                    "ts": ts(2),
                    "action": "tool_use",
                    "tool": "Write",
                    "thread": "Pressure tests",
                    "preview": "analysis/pressure_tests.md",
                },
                {
                    "type": "thread_finished",
                    "ts": ts(5),
                    "thread": "Pressure tests",
                },
                {
                    "type": "thread_started",
                    "ts": ts(6),
                    "thread": "Source treatment and assumptions",
                },
            ]
        )
        + "\n",
        encoding="utf-8",
    )

    state = api._scan_progress_state(stream_path)

    assert state["terminated"] is False
    assert state["thread_count"] == 2
    assert state["thread_done_count"] == 1
    assert state["thread_failed_count"] == 0
    assert state["open_thread_count"] == 1
    assert {
        (thread["name"], thread["status"], thread["event_count"])
        for thread in state["threads"]
    } == {
        ("Pressure tests", "done", 3),
        ("Source treatment and assumptions", "running", 1),
    }
    pressure = next(
        thread for thread in state["threads"]
        if thread["name"] == "Pressure tests"
    )
    assert pressure["elapsed_ms"] == 4000
    assert pressure["last_event_at"] == ts(5)
    assert pressure["latest_action"]["action"] == "tool_use"

    memo_job = next(
        job for job in api._memo_kind_records()
        if job.get("report_id") == report["id"]
    )
    assert memo_job["thread_count"] == 2
    assert memo_job["thread_done_count"] == 1
    assert memo_job["open_thread_count"] == 1


def test_active_memo_job_includes_not_started_phase_rows(memo_env):
    report, run_dir = _make_memo_report(memo_env)
    stream_path = memo_prep.stream_path(run_dir)
    stream_path.parent.mkdir(parents=True, exist_ok=True)
    base_ts = datetime.now(timezone.utc) - timedelta(seconds=10)

    def ts(offset_s: int) -> str:
        return (base_ts + timedelta(seconds=offset_s)).isoformat()

    stream_path.write_text(
        "\n".join(
            json.dumps(event)
            for event in [
                {
                    "type": "job_init",
                    "ts": ts(0),
                    "kind": "memo",
                    "title": "Investment memo — Generalist, Inc.",
                    "report_id": report["id"],
                    "company_id": "generalist-inc",
                    "run_id": report["run_id"],
                },
                {
                    "type": "thread_planned",
                    "ts": ts(1),
                    "thread": "Phase 1 - Intake and setup",
                    "title": "Phase 1 - Intake and setup",
                    "phase_index": 1,
                    "estimate_ms": 150000,
                    "description": "Setup usually takes about 2m 30s.",
                },
                {
                    "type": "thread_planned",
                    "ts": ts(2),
                    "thread": "Phase 2 - Parallel analysis passes",
                    "title": "Phase 2 - Parallel analysis passes",
                    "phase_index": 2,
                },
                {
                    "type": "thread_started",
                    "ts": ts(3),
                    "thread": "Phase 1 - Intake and setup",
                },
            ]
        )
        + "\n",
        encoding="utf-8",
    )

    state = api._scan_progress_state(stream_path)

    assert state["thread_count"] == 2
    assert state["open_thread_count"] == 1
    phase1 = next(
        thread for thread in state["threads"]
        if thread["name"] == "Phase 1 - Intake and setup"
    )
    phase2 = next(
        thread for thread in state["threads"]
        if thread["name"] == "Phase 2 - Parallel analysis passes"
    )
    assert phase1["status"] == "running"
    assert phase1["estimate_ms"] == 150000
    assert phase1["event_count"] == 1
    assert phase2["status"] == "not_started"
    assert phase2["event_count"] == 0
    assert phase2["started_at"] is None


# ---- orphan demotion (restart-killed runs must become resume-eligible) ----

def _age_stream(run_dir, seconds_past_threshold=60):
    import os
    import time as _time

    path = memo_prep.stream_path(run_dir)
    old = _time.time() - memo_analysis.ORPHAN_IDLE_THRESHOLD_SEC - seconds_past_threshold
    os.utime(path, (old, old))


def test_recover_demotes_orphaned_analyzing_report(memo_env):
    report, run_dir = _make_memo_report(memo_env)
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit("stage", stage="analyzing", message="working")
    _age_stream(run_dir)

    recovered = memo_analysis.recover_stale_reports()

    assert recovered == 1
    updated = storage.get_report(report["id"])
    assert updated["status"] == "failed_during_analysis"
    assert updated["error"] == "orphaned by server restart"
    assert updated["failure_phase"] == "orphaned"
    events = _events(memo_prep.stream_path(run_dir))
    assert events[-1]["type"] == "error"
    assert events[-1]["phase"] == "orphaned"


def test_recover_leaves_fresh_analyzing_report_alone(memo_env):
    report, run_dir = _make_memo_report(memo_env)
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit("stage", stage="analyzing", message="working")

    memo_analysis.recover_stale_reports()

    assert storage.get_report(report["id"])["status"] == "analyzing"


def test_recover_leaves_run_with_live_worker_alone(memo_env, monkeypatch):
    report, run_dir = _make_memo_report(memo_env)
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit("stage", stage="analyzing", message="working")
    _age_stream(run_dir)
    monkeypatch.setattr(memo_analysis, "_memo_worker_alive", lambda _rid: True)

    memo_analysis.recover_stale_reports()

    assert storage.get_report(report["id"])["status"] == "analyzing"


def test_atexit_hook_fails_active_runs(memo_env):
    report, run_dir = _make_memo_report(memo_env)
    memo_analysis._register_active_run(report["id"])
    try:
        memo_analysis._fail_active_runs_at_exit()
    finally:
        memo_analysis._unregister_active_run(report["id"])

    updated = storage.get_report(report["id"])
    assert updated["status"] == "failed_during_analysis"
    assert updated["error"] == "interrupted by server shutdown"
    events = _events(memo_prep.stream_path(run_dir))
    assert events[-1]["type"] == "error"
    assert events[-1]["phase"] == "shutdown"


def _admin_request():
    """Stub Request for calling gated handlers as plain functions:
    resolves to the anon-dev admin role in _caller_role."""
    from types import SimpleNamespace

    return SimpleNamespace(
        state=SimpleNamespace(auth_kind="anon_dev", session_email=None),
        cookies={},
        headers={},
    )


# ---- Phase 4.6: restart-interrupted runs auto-resume at startup ----

def _make_resumable_failed_report(memo_env, failure_phase):
    report, run_dir = _make_memo_report(memo_env)
    analysis_dir = run_dir / "analysis"
    analysis_dir.mkdir(exist_ok=True)
    (analysis_dir / "pressure_tests.md").write_text("# Notes\n", encoding="utf-8")
    storage.update_report(
        report["id"],
        status="failed_during_analysis",
        stage="Analysis interrupted",
        failure_phase=failure_phase,
        failure_detail="interrupted by server shutdown",
    )
    return report


def test_auto_resume_picks_up_shutdown_interrupted_runs(memo_env, monkeypatch):
    report = _make_resumable_failed_report(memo_env, "shutdown")
    called = []
    monkeypatch.setattr(
        memo_analysis, "start_resume", lambda report_id: called.append(report_id)
    )

    resumed = api.resume_interrupted_memo_runs()

    assert resumed == 1
    assert called == [report["id"]]
    queued = storage.get_report(report["id"])
    assert queued["status"] == "analyzing"
    assert queued["resume_from_failure_phase"] == "shutdown"


def test_auto_resume_leaves_orphaned_and_failed_runs_parked(memo_env, monkeypatch):
    # Orphan-sweep demotions may be weeks old — a human decides those.
    _make_resumable_failed_report(memo_env, "orphaned")
    called = []
    monkeypatch.setattr(
        memo_analysis, "start_resume", lambda report_id: called.append(report_id)
    )

    assert api.resume_interrupted_memo_runs() == 0
    assert called == []


def test_auto_resume_respects_kill_switch(memo_env, monkeypatch):
    _make_resumable_failed_report(memo_env, "shutdown")
    monkeypatch.setenv("BSH_MEMO_AUTO_RESUME", "0")
    called = []
    monkeypatch.setattr(
        memo_analysis, "start_resume", lambda report_id: called.append(report_id)
    )

    assert api.resume_interrupted_memo_runs() == 0
    assert called == []


def test_chained_resume_regenerates_quality_failed_package(memo_env, monkeypatch):
    """Renderer failure → resume → quality failure → resume must REGENERATE
    the package, not reuse the one the quality gate just rejected. The
    endpoint's resume_from_* fields keep the ORIGINAL failure (provenance),
    so the worker relies on resume_last_* for the current one."""
    report, run_dir = _make_memo_report(memo_env)
    analysis_dir = run_dir / "analysis"
    analysis_dir.mkdir(exist_ok=True)
    (analysis_dir / "pressure_tests.md").write_text("# Notes\n", encoding="utf-8")
    (run_dir / "logs").mkdir(exist_ok=True)
    (run_dir / "logs" / "memo_package.json").write_text("{}", encoding="utf-8")
    (run_dir / "logs" / "memo_quality_lint.md").write_text("# P0\n", encoding="utf-8")
    # State as the resume endpoint leaves it on the SECOND resume of a run
    # whose FIRST failure was the renderer contract:
    storage.update_report(
        report["id"],
        status="analyzing",
        resume_from_status="failed_during_analysis",
        resume_from_failure_phase="renderer_contract",
        resume_last_status="failed_quality_gate",
        resume_last_failure_phase="quality_gate",
        failure_phase=None,
    )
    regen_calls = []

    def fake_resume_package(**kwargs):
        regen_calls.append(kwargs)
        return {"ok": False, "error": "stop test here"}

    monkeypatch.setattr(
        claude_runner, "run_resume_memo_package", fake_resume_package
    )
    monkeypatch.setattr(
        claude_runner, "is_transient_claude_error", lambda _m: False
    )

    memo_analysis._resume(report["id"])

    # The worker must have archived the rejected package and called the
    # regeneration pass — not taken the package-reuse shortcut.
    assert len(regen_calls) >= 1
    assert not (run_dir / "logs" / "memo_package.json").exists()
    assert regen_calls[0].get("quality_lint_path") is not None


def test_resume_endpoint_records_current_failure_in_resume_last(memo_env, monkeypatch):
    report, run_dir = _make_memo_report(memo_env)
    analysis_dir = run_dir / "analysis"
    analysis_dir.mkdir(exist_ok=True)
    (analysis_dir / "pressure_tests.md").write_text("# Notes\n", encoding="utf-8")
    storage.update_report(
        report["id"],
        status="failed_quality_gate",
        failure_phase="quality_gate",
        # Stale provenance from an earlier, different failure:
        resume_from_status="failed_during_analysis",
        resume_from_failure_phase="renderer_contract",
    )
    monkeypatch.setattr(
        memo_analysis, "start_resume", lambda report_id: None
    )

    api.resume_memo_report(_admin_request(), report["id"])

    queued = storage.get_report(report["id"])
    # Original provenance preserved...
    assert queued["resume_from_status"] == "failed_during_analysis"
    assert queued["resume_from_failure_phase"] == "renderer_contract"
    # ...current failure captured for the worker.
    assert queued["resume_last_status"] == "failed_quality_gate"
    assert queued["resume_last_failure_phase"] == "quality_gate"


# ---- English-package validation at generation time (phase 3) ----

def test_english_package_validation_accepts_blank_zh():
    from server import memo_docx_renderer

    package = _memo_package(body_zh="")
    assert memo_docx_renderer.english_package_validation_errors(package) == []


def test_english_package_validation_flags_analysis_pass_source_vocabulary():
    """Both fresh runs on record emitted sources shaped like fast-pass
    evidence (label/source_class/detail) and died at the render gate — the
    generation-time validator must catch exactly that."""
    from server import memo_docx_renderer

    package = _memo_package(body_zh="")
    package["sources"] = [
        {
            "id": "S1",
            "label": {"en": "Company registry", "zh": ""},
            "source_class": "company-reported",
            "detail": {"en": "Registry and launch disclosures.", "zh": ""},
            "as_of": "2026-03-10",
        }
    ]
    errors = memo_docx_renderer.english_package_validation_errors(package)
    joined = "; ".join(errors)
    assert "sources[0].title" in joined
    assert "sources[0].class" in joined
    assert "sources[0].treatment" in joined


def test_phase3_retries_on_validation_failure_with_feedback(memo_env, monkeypatch):
    """An invalid English package must trigger a same-phase retry with the
    validation errors fed back, not a failure 20 minutes later at render."""
    monkeypatch.setenv("BSH_MEMO_FAST_PIPELINE", "1")
    monkeypatch.setenv("BSH_MEMO_FAST_ENGLISH_PACKAGE_RETRIES", "1")
    monkeypatch.delenv("BSH_MEMO_GENERATE_INTERNAL", raising=False)
    monkeypatch.delenv("BSH_MEMO_RENDER_PDF_PREVIEWS", raising=False)
    report, run_dir = _make_memo_report(memo_env)
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit("job_init", kind="memo", report_id=report["id"])

    def fake_analysis_pass(**kwargs):
        return {
            "summary": "s",
            "key_findings": [],
            "supporting_evidence": [],
            "disconfirming_evidence": [],
            "open_questions": [],
            "memo_uses": [],
        }, None

    english_calls = []

    def fake_english_package(**kwargs):
        english_calls.append(kwargs.get("validation_feedback"))
        package = _memo_package(body_zh="")
        if len(english_calls) == 1:
            # First attempt: the wrong source vocabulary.
            package["sources"] = [
                {
                    "id": "S1",
                    "label": {"en": "Registry", "zh": ""},
                    "source_class": "company-reported",
                    "detail": {"en": "d", "zh": ""},
                    "as_of": "2026-03-10",
                }
            ]
        return {
            "analysis_artifacts": {},
            "memo_package": package,
        }, None

    def fake_bilingual(**kwargs):
        return {"memo_package": _memo_package()}, None

    monkeypatch.setattr(
        claude_runner, "run_memo_fast_analysis_pass", fake_analysis_pass
    )
    monkeypatch.setattr(
        claude_runner, "run_memo_fast_english_package", fake_english_package
    )
    monkeypatch.setattr(
        claude_runner,
        "run_memo_fast_bilingual_package_parallel",
        fake_bilingual,
    )

    result = memo_analysis._run_fast_memo_pipeline(
        report_id=report["id"],
        report=storage.get_report(report["id"]),
        run_dir=run_dir,
        stream=stream,
        company_name="Generalist, Inc.",
        company_slug="generalist-inc",
        run_id=str(report["run_id"]),
        memo_paths_abs=memo_analysis._memo_paths_abs(report),
        analysis_session_path=None,
        lessons_path=None,
    )

    assert len(english_calls) == 2
    # First attempt had no feedback; the retry carried the exact errors.
    assert english_calls[0] is None
    assert "sources[0].title" in english_calls[1]
    assert result.get("ok") is True


# ---- Resume-path validation of the freshly regenerated package ----

def test_resume_retries_with_feedback_when_regenerated_package_invalid(
    memo_env, monkeypatch
):
    """A resume regeneration that writes a contract-violating package must
    be validated immediately and retried with the errors fed back — not
    handed to the renderer to fail the whole run (the ZaiNar 2026-06-24
    failure mode: plain-string prose table cells)."""
    monkeypatch.setenv("BSH_MEMO_RESUME_PACKAGE_RETRIES", "1")
    report, run_dir = _make_memo_report(memo_env)
    analysis_dir = run_dir / "analysis"
    analysis_dir.mkdir(exist_ok=True)
    (analysis_dir / "pressure_tests.md").write_text("# Notes\n", encoding="utf-8")
    (run_dir / "logs").mkdir(exist_ok=True)
    storage.update_report(
        report["id"],
        status="analyzing",
        resume_from_status="failed_during_analysis",
        resume_from_failure_phase="renderer_contract",
        failure_phase=None,
    )
    package_path = run_dir / "logs" / "memo_package.json"
    feedback_seen = []

    def fake_resume_package(**kwargs):
        feedback_seen.append(kwargs.get("validation_feedback"))
        package = _memo_package()
        package["sections"][0]["blocks"].append(
            {
                "type": "table",
                "title": {"en": "Patents", "zh": "专利"},
                "headers": [
                    {"en": "Metric", "zh": "指标"},
                    {"en": "Value", "zh": "数值"},
                ],
                # Plain-string prose cell: the renderer must reject this.
                "rows": [[{"en": "Patents", "zh": "专利"}, "120+ filed / 90+ issued"]],
            }
        )
        package_path.write_text(
            json.dumps(package, ensure_ascii=False), encoding="utf-8"
        )
        return {"ok": True}

    monkeypatch.setattr(
        claude_runner, "run_resume_memo_package", fake_resume_package
    )
    monkeypatch.setattr(
        claude_runner, "is_transient_claude_error", lambda _m: False
    )

    memo_analysis._resume(report["id"])

    assert len(feedback_seen) == 2
    assert feedback_seen[0] is None
    assert "must be bilingual" in feedback_seen[1]
    # Invalid attempts were archived, never handed to the renderer.
    assert not package_path.exists()
    assert list((run_dir / "logs").glob("memo_package.invalid.*.json"))
    final = storage.get_report(report["id"])
    assert final["status"] == "failed_during_analysis"


def test_fast_pipeline_repairs_blank_zh_after_bilingual_merge(
    memo_env, monkeypatch
):
    """If per-section Chinese fill leaves a blank zh, the pipeline must run
    one monolithic repair pass and proceed — not fail at render time."""
    monkeypatch.setenv("BSH_MEMO_FAST_PIPELINE", "1")
    monkeypatch.delenv("BSH_MEMO_GENERATE_INTERNAL", raising=False)
    report, run_dir = _make_memo_report(memo_env)
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit("job_init", kind="memo", report_id=report["id"])

    def fake_analysis_pass(**kwargs):
        return {
            "summary": "s",
            "key_findings": [],
            "supporting_evidence": [],
            "disconfirming_evidence": [],
            "open_questions": [],
            "memo_uses": [],
        }, None

    def fake_english_package(**kwargs):
        return {
            "analysis_artifacts": {},
            "memo_package": _memo_package(body_zh=""),
        }, None

    def fake_parallel_bilingual(**kwargs):
        package = _memo_package()
        # One translation missed by the per-section fan-out.
        package["sections"][0]["blocks"][0]["text"]["zh"] = ""
        return {"memo_package": package}, None

    repair_calls = []

    def fake_monolithic_bilingual(**kwargs):
        repair_calls.append(kwargs)
        return {"memo_package": _memo_package()}, None

    monkeypatch.setattr(
        claude_runner, "run_memo_fast_analysis_pass", fake_analysis_pass
    )
    monkeypatch.setattr(
        claude_runner, "run_memo_fast_english_package", fake_english_package
    )
    monkeypatch.setattr(
        claude_runner,
        "run_memo_fast_bilingual_package_parallel",
        fake_parallel_bilingual,
    )
    monkeypatch.setattr(
        claude_runner,
        "run_memo_fast_bilingual_package",
        fake_monolithic_bilingual,
    )

    result = memo_analysis._run_fast_memo_pipeline(
        report_id=report["id"],
        report=storage.get_report(report["id"]),
        run_dir=run_dir,
        stream=stream,
        company_name="Generalist, Inc.",
        company_slug="generalist-inc",
        run_id=str(report["run_id"]),
        memo_paths_abs=memo_analysis._memo_paths_abs(report),
        analysis_session_path=None,
        lessons_path=None,
    )

    assert len(repair_calls) == 1
    assert result.get("ok") is True
    final_package = json.loads(
        (run_dir / "logs" / "memo_package.json").read_text(encoding="utf-8")
    )
    assert final_package["sections"][0]["blocks"][0]["text"]["zh"]


# ---- Pre-render quality gate feeds the generation retry loops ----

_SCAFFOLD_BODY_EN = (
    "Generalist builds automation infrastructure, but commercial "
    "repeatability is still unproven across verticals."
)


def test_prerender_quality_error_flags_scaffold_phrase():
    package = _memo_package(body_en=_SCAFFOLD_BODY_EN)
    error = memo_analysis._memo_package_prerender_quality_error(package)
    assert error is not None
    assert "scaffold_label" in error


def test_prerender_quality_error_passes_clean_package():
    assert (
        memo_analysis._memo_package_prerender_quality_error(_memo_package())
        is None
    )


def test_phase3_retries_on_quality_gate_finding(memo_env, monkeypatch):
    """Banned memo vocabulary must be caught at English-package time and fed
    back into the same retry loop — not discovered post-render, where each
    miss costs a whole regeneration round (the ZaiNar 'still unproven'
    failure mode)."""
    monkeypatch.setenv("BSH_MEMO_FAST_PIPELINE", "1")
    monkeypatch.setenv("BSH_MEMO_FAST_ENGLISH_PACKAGE_RETRIES", "1")
    monkeypatch.delenv("BSH_MEMO_GENERATE_INTERNAL", raising=False)
    monkeypatch.delenv("BSH_MEMO_RENDER_PDF_PREVIEWS", raising=False)
    report, run_dir = _make_memo_report(memo_env)
    stream = job_progress.ProgressLog(memo_prep.stream_path(run_dir))
    stream.emit("job_init", kind="memo", report_id=report["id"])

    def fake_analysis_pass(**kwargs):
        return {
            "summary": "s",
            "key_findings": [],
            "supporting_evidence": [],
            "disconfirming_evidence": [],
            "open_questions": [],
            "memo_uses": [],
        }, None

    english_calls = []

    def fake_english_package(**kwargs):
        english_calls.append(kwargs.get("validation_feedback"))
        body = _SCAFFOLD_BODY_EN if len(english_calls) == 1 else None
        return {
            "analysis_artifacts": {},
            "memo_package": _memo_package(body_en=body, body_zh=""),
        }, None

    def fake_bilingual(**kwargs):
        return {"memo_package": _memo_package()}, None

    monkeypatch.setattr(
        claude_runner, "run_memo_fast_analysis_pass", fake_analysis_pass
    )
    monkeypatch.setattr(
        claude_runner, "run_memo_fast_english_package", fake_english_package
    )
    monkeypatch.setattr(
        claude_runner,
        "run_memo_fast_bilingual_package_parallel",
        fake_bilingual,
    )

    result = memo_analysis._run_fast_memo_pipeline(
        report_id=report["id"],
        report=storage.get_report(report["id"]),
        run_dir=run_dir,
        stream=stream,
        company_name="Generalist, Inc.",
        company_slug="generalist-inc",
        run_id=str(report["run_id"]),
        memo_paths_abs=memo_analysis._memo_paths_abs(report),
        analysis_session_path=None,
        lessons_path=None,
    )

    assert len(english_calls) == 2
    assert english_calls[0] is None
    assert "scaffold_label" in english_calls[1]
    assert result.get("ok") is True


def test_resume_retries_then_delivers_with_warnings_on_quality_findings(
    memo_env, monkeypatch
):
    """Quality findings in a regenerated package get ONE feedback retry;
    if the final attempt still has findings, the memo is delivered as
    complete_with_warnings instead of blocking the run."""
    monkeypatch.setenv("BSH_MEMO_RESUME_PACKAGE_RETRIES", "1")
    monkeypatch.delenv("BSH_MEMO_GENERATE_INTERNAL", raising=False)
    report, run_dir = _make_memo_report(memo_env)
    analysis_dir = run_dir / "analysis"
    analysis_dir.mkdir(exist_ok=True)
    (analysis_dir / "pressure_tests.md").write_text("# Notes\n", encoding="utf-8")
    (run_dir / "logs").mkdir(exist_ok=True)
    storage.update_report(
        report["id"],
        status="analyzing",
        resume_from_status="failed_during_analysis",
        resume_from_failure_phase="renderer_contract",
        failure_phase=None,
    )
    package_path = run_dir / "logs" / "memo_package.json"
    feedback_seen = []

    def fake_resume_package(**kwargs):
        feedback_seen.append(kwargs.get("validation_feedback"))
        package = _memo_package(body_en=_SCAFFOLD_BODY_EN)
        package_path.write_text(
            json.dumps(package, ensure_ascii=False), encoding="utf-8"
        )
        return {"ok": True}

    monkeypatch.setattr(
        claude_runner, "run_resume_memo_package", fake_resume_package
    )
    monkeypatch.setattr(
        claude_runner, "is_transient_claude_error", lambda _m: False
    )

    memo_analysis._resume(report["id"])

    assert len(feedback_seen) == 2
    assert feedback_seen[0] is None
    assert "scaffold_label" in feedback_seen[1]
    # Attempt 1 was archived for the feedback retry; the final attempt's
    # package was kept and delivered with warnings.
    assert package_path.exists()
    assert list((run_dir / "logs").glob("memo_package.quality_failed.*.json"))
    final = storage.get_report(report["id"])
    assert final["status"] == "complete_with_warnings"
    assert final["quality_warnings"]
    events = _events(memo_prep.stream_path(run_dir))
    assert events[-1]["type"] == "done"


def test_resume_of_complete_with_warnings_regenerates_package(
    memo_env, monkeypatch
):
    """Resuming a complete-with-warnings report must take the quality
    regeneration path (archive the delivered package, rewrite from analysis
    artifacts) — the same loop that used to serve failed_quality_gate."""
    report, run_dir = _make_memo_report(memo_env)
    analysis_dir = run_dir / "analysis"
    analysis_dir.mkdir(exist_ok=True)
    (analysis_dir / "pressure_tests.md").write_text("# Notes\n", encoding="utf-8")
    (run_dir / "logs").mkdir(exist_ok=True)
    package_path = run_dir / "logs" / "memo_package.json"
    package_path.write_text(
        json.dumps(_memo_package(), ensure_ascii=False), encoding="utf-8"
    )
    (run_dir / "logs" / "memo_quality_lint.md").write_text(
        "# P0\n", encoding="utf-8"
    )
    storage.update_report(
        report["id"],
        status="analyzing",
        resume_last_status="complete_with_warnings",
        quality_warnings=["Memo quality gate found 1 P0 finding."],
        failure_phase=None,
    )
    regen_calls = []

    def fake_resume_package(**kwargs):
        regen_calls.append(kwargs)
        return {"ok": False, "error": "stop test here"}

    monkeypatch.setattr(
        claude_runner, "run_resume_memo_package", fake_resume_package
    )
    monkeypatch.setattr(
        claude_runner, "is_transient_claude_error", lambda _m: False
    )

    memo_analysis._resume(report["id"])

    # The delivered-with-warnings package was archived and regeneration ran
    # instead of the package-reuse shortcut.
    assert len(regen_calls) >= 1
    assert not package_path.exists()
    assert list((run_dir / "logs").glob("memo_package.quality_failed.*.json"))


def test_resume_available_for_complete_with_warnings(memo_env):
    report, run_dir = _make_memo_report(memo_env)
    analysis_dir = run_dir / "analysis"
    analysis_dir.mkdir(exist_ok=True)
    (analysis_dir / "pressure_tests.md").write_text("# Notes\n", encoding="utf-8")
    storage.update_report(report["id"], status="complete_with_warnings")

    assert api._report_resume_available(storage.get_report(report["id"])) is True
    # Plain complete stays non-resumable.
    storage.update_report(report["id"], status="complete")
    assert api._report_resume_available(storage.get_report(report["id"])) is False
