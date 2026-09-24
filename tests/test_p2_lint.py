"""Quality lint round 2: estimate discipline in v2 memos (I9,
``figure_without_anchor``) and cross-section restatement (A3,
``claim_restated`` / ``figure_restated_across_sections``). All P1/P2 —
none may block."""
from __future__ import annotations

from docx import Document

from server import memo_quality_lint, memo_structure

V2 = memo_structure.load_structure("late", 2)
V1 = memo_structure.LATE


def _docx(path, items):
    document = Document()
    for item in items:
        if isinstance(item, str):
            document.add_paragraph(item)
            continue
        table = document.add_table(rows=len(item), cols=len(item[0]))
        for row_index, row in enumerate(item):
            for col_index, value in enumerate(row):
                table.cell(row_index, col_index).text = value
    document.save(path)
    return path


def _by_code(result, code):
    return [f for f in result.findings if f.code == code]


def test_figure_without_anchor_fires_only_in_v2_prose(tmp_path):
    path = _docx(
        tmp_path / "memo.docx",
        [
            "I. Executive Summary",
            "The $10M primary cash infusion provides only 8 to 10 months of runway, forcing a raise by early 2027.",
            "Revenue reached $4.02M in the trailing twelve months S7, on a 40x multiple C3.",
            "Dividing the $1,000M mark by the $450M+ aggregate gives 2.2x [C3] against the round.",
            "Gross margin is not disclosed; we assume 75% for the software track pending a document.",
            "ARR: Not disclosed — no document on file; the $24M figure is a placeholder.",
            [["Metric", "Value"], ["Runway", "8 to 10 months"], ["Burn", "$1.2M a month"]],
            "IV. Investment Risk",
            "Risk 1: a 30% down-round is the base case if conversion stalls before the next raise.",
        ],
    )
    v2 = memo_quality_lint.lint_memo_docx(path, V2)
    anchors = _by_code(v2, "figure_without_anchor")
    assert {f.severity for f in anchors} == {"P1"}
    snippets = " | ".join(f.snippet for f in anchors)
    assert "8 to 10 months of runway" in snippets and "30% down-round" in snippets
    # Cited sentences (bare ids or bracketed), "not disclosed" sentences and
    # table cells are not flagged.
    assert "trailing twelve months" not in snippets
    assert "2.2x" not in snippets
    assert "assume 75%" not in snippets and "placeholder" not in snippets
    assert "$1.2M a month" not in snippets
    assert len(anchors) == 2
    assert not v2.has_blocking_findings

    v1 = memo_quality_lint.lint_memo_docx(path, V1)
    assert _by_code(v1, "figure_without_anchor") == []


def test_claim_restated_in_three_sections_is_p1(tmp_path):
    sentence = "The patent estate is the strongest asset in the file and standardization is what makes it worth less."
    path = _docx(
        tmp_path / "memo.docx",
        [
            "I. Executive Summary", sentence + " Something else follows.",
            "II. Company Overview", "A different point entirely about the team.",
            "III. Investment Highlights", sentence + " More follows here.",
            "V. Financial Forecast & Valuation", "A different point about the exit.",
            "IV. Investment Risk", "Risk 1 is the one below. " + sentence,
        ],
    )
    result = memo_quality_lint.lint_memo_docx(path, V1)
    claims = _by_code(result, "claim_restated")
    assert len(claims) == 1 and claims[0].severity == "P1"
    assert "3 sections" in claims[0].snippet
    assert not result.has_blocking_findings
    # Twice is allowed.
    twice = _docx(tmp_path / "twice.docx", ["I. Executive Summary", sentence, "IV. Investment Risk", sentence])
    assert _by_code(memo_quality_lint.lint_memo_docx(twice, V1), "claim_restated") == []


def test_figure_restated_across_four_sections_is_p2(tmp_path):
    path = _docx(
        tmp_path / "memo.docx",
        [
            "I. Executive Summary", "The mark implies 2.2x the $450M+ aggregate of contracts.",
            "II. Company Overview", "Headcount is 52 employees; the ratio to bookings is 2.2x today.",
            "III. Investment Highlights", "ARR of $24M is a placeholder, but 2.2x coverage of the raise holds.",
            "IV. Investment Risk", "A 2.2x multiple on unsigned MOUs is the first risk.",
            "V. Financial Forecast & Valuation", "Revenue of $24M would put the round at 2.2x contracts.",
            [["Case", "MOIC"], ["Base", "2.2x"], ["Bull", "3.1x"]],
        ],
    )
    result = memo_quality_lint.lint_memo_docx(path, V1)
    figures = _by_code(result, "figure_restated_across_sections")
    assert len(figures) == 1 and figures[0].severity == "P2"
    assert figures[0].snippet.startswith("2.2x is stated in 5 sections")
    # $24M with its metric (arr / revenue) sits in two sections only.
    assert not result.has_blocking_findings
    # A pinned figure is exempt from the figure check.
    pinned = memo_quality_lint.lint_memo_docx(path, V1, pinned_values=["2.2x"])
    assert _by_code(pinned, "figure_restated_across_sections") == []


def test_round_figures_need_their_metric_to_count(tmp_path):
    path = _docx(
        tmp_path / "memo.docx",
        [
            "I. Executive Summary", "Revenue of $100M is the target.",
            "II. Company Overview", "The company raised $100M in total.",
            "III. Investment Highlights", "A $100M valuation step is plausible.",
            "IV. Investment Risk", "A $100M shortfall in bookings would matter.",
            "V. Financial Forecast & Valuation", "Exit at $100M revenue in 2029.",
        ],
    )
    result = memo_quality_lint.lint_memo_docx(path, V1)
    assert _by_code(result, "figure_restated_across_sections") == []
