"""Renderer round 2 (I22): canonical source-class labels the renderer owns,
the fixed Chinese for the prompts' placeholders, one heading where a
sub-heading and its table title say the same thing, and the optional
"Terms used" glossary block."""
from __future__ import annotations

import copy

from docx import Document

import test_memo_docx_renderer as base
from server import memo_docx_renderer as r


def _headings(path):
    document = Document(path)
    return [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]


def _cells(path):
    document = Document(path)
    return [cell.text for table in document.tables for row in table.rows for cell in row.cells]


# ---- (a) source classes -------------------------------------------------------


def test_every_prompt_prescribed_class_has_one_chinese_label():
    prescribed = {
        "unverified registry value (no document on file)": "未经核实的登记值（无文件佐证）",
        "secondary press report, original article not retrieved": "二手媒体报道（未获取原文）",
        "government program page": "政府项目页面",
        "company-reported press release": "公司新闻稿",
        "independent trade press": "独立行业媒体",
        "third-party company profile": "第三方公司资料",
        "competitor company materials": "竞争对手公司材料",
        "third-party private-market data": "第三方私募市场数据",
        "third-party market research": "第三方市场研究",
        "public patent-office record": "公开专利记录",
        "crowdsourced org-chart aggregator": "众包组织架构汇总站",
        "investor/intermediary": "投资方/中介资料",
        "company-reported": "公司自述",
        "third-party market data": "第三方市场数据",
        "public filings": "公开申报文件",
    }
    for en, zh in prescribed.items():
        assert r.canonical_source_class_zh(en) == zh, en
        assert r._source_class_label({"class": en}, "zh") == zh
        # Case, hyphen and underscore variants are the same class.
        assert r.canonical_source_class_zh(en.upper().replace("-", "_")) == zh


def test_canonical_zh_overrides_the_models_own_translation():
    source = {"class": {"en": "unverified registry value (no document on file)", "zh": "BSH 内部资料"}}
    assert r._source_class_label(source, "zh") == "未经核实的登记值（无文件佐证）"
    # The English half is printed as written.
    assert r._source_class_label(source, "en") == "unverified registry value (no document on file)"
    unknown = {"class": {"en": "a class nobody prescribed", "zh": "模型自己的译法"}}
    assert r._source_class_label(unknown, "zh") == "模型自己的译法"


def test_a_qualified_class_matches_by_its_known_prefix():
    assert r.canonical_source_class_zh("third-party market data (aggregator, methodology not disclosed)") == "第三方市场数据"
    assert r.canonical_source_class_zh("company-reported, distributed via tier-1 wire") == "公司自述"
    assert r._source_class_label({"class": "third-party market data (aggregator)"}, "en") == (
        "Third-party market data (aggregator)"
    )
    assert r.canonical_source_class_zh("something entirely new") is None
    assert r._source_class_label({"class": "something entirely new"}, "zh") == "其他来源"


# ---- (b) placeholders ----------------------------------------------------------


def test_placeholders_get_their_fixed_chinese():
    value = {
        "en": "Allocation: [TO BE DETERMINED BY IC]; close: [date to set]",
        "zh": "配置：[TO BE DETERMINED BY IC]；交割：[Date To Set]",
    }
    assert r._loc(value, "zh") == "配置：【待投委会确定】；交割：【日期待定】"
    assert r._loc(value, "en") == value["en"]
    assert r.apply_zh_placeholders("【待定】") == "【待定】"


def test_owner_roles_alone_in_a_cell_are_translated():
    row = [{"en": "Owner", "zh": "Owner"}, {"en": "Deal lead / Legal", "zh": "Deal lead / Legal"}, {"en": "Finance", "zh": "Finance"}]
    assert r._row_values(row, "zh") == ["Owner", "交易负责人/法务", "财务"]
    assert r._row_values(row, "en") == ["Owner", "Deal lead / Legal", "Finance"]
    # Not alone in the cell: left to the translation.
    assert r.owner_role_zh("the deal lead signs off") is None
    # ZaiNar 2026-09-23 (A3 IC memo): "partner" printed in English in the
    # Chinese owner column.
    assert r.owner_role_zh("Partner") == "合伙人"
    assert r.owner_role_zh("partner / legal") == "合伙人/法务"


# ---- (c) heading / table-title dedupe ------------------------------------------


def test_merged_heading_titles_keeps_the_longer_form():
    blocks = [
        {"type": "heading", "text": {"en": "Revenue Picture", "zh": "收入情况"}},
        {"type": "table", "title": {"en": "Revenue picture — what is disclosed and what it will carry", "zh": "收入情况——已披露内容"}, "rows": [[{"en": "a", "zh": "甲"}]]},
        {"type": "heading", "text": {"en": "Key Operating Metrics", "zh": "核心经营指标"}},
        {"type": "table", "title": {"en": "Key operating metrics", "zh": "核心经营指标"}, "rows": [[{"en": "a", "zh": "甲"}]]},
        {"type": "heading", "text": {"en": "Scenarios", "zh": "情景"}},
        {"type": "table", "title": {"en": "Outcome scenarios against the mark", "zh": "情景结果"}, "rows": [[{"en": "a", "zh": "甲"}]]},
    ]
    merged = r.merged_heading_titles(blocks)
    assert merged[0][0]["en"].startswith("Revenue picture — what") and merged[1] == (None, True)
    assert merged[2] == (None, False) and merged[3] == (None, True)
    assert 4 not in merged and 5 not in merged  # one-word heading, a different title


def test_rendered_memo_prints_each_heading_once(tmp_path):
    package = base._package()
    section = package["sections"][0]
    section["blocks"] = [
        {"type": "heading", "level": 2, "text": {"en": "Executive Summary", "zh": "执行摘要"}},
        *section["blocks"],
        {"type": "heading", "level": 2, "text": {"en": "Revenue Picture", "zh": "收入情况"}},
        {
            "type": "table",
            "title": {"en": "Revenue picture — what is disclosed", "zh": "收入情况——已披露内容"},
            "headers": [{"en": "Metric", "zh": "指标"}, {"en": "Value", "zh": "数值"}],
            "rows": [[{"en": "ARR", "zh": "ARR"}, {"en": "Not disclosed", "zh": "未披露"}]],
        },
        {"type": "heading", "level": 2, "text": {"en": "Board and Leadership", "zh": "董事会与管理层"}},
        {"type": "paragraph", "text": {"en": "Two founders sit on the board.", "zh": "两位创始人进入董事会。"}},
        {
            "type": "table",
            "title": {"en": "Board of Directors", "zh": "董事会与管理层"},
            "headers": [{"en": "Name", "zh": "姓名"}, {"en": "Role", "zh": "职务"}],
            "rows": [[{"en": "A. Founder", "zh": "A. Founder"}, {"en": "CEO", "zh": "CEO"}]],
        },
    ]
    out_en, out_zh = tmp_path / "memo" / "en.docx", tmp_path / "memo" / "zh.docx"
    r.render_memos(package, out_en=out_en, out_zh=out_zh)
    en, zh = _headings(out_en), _headings(out_zh)
    # Only the executive summary's headings (the base package's company
    # overview has a "Revenue Picture" of its own).
    en = en[: en.index("II. Company Overview")]
    zh = zh[: next(i for i, h in enumerate(zh) if h.startswith("二"))]
    # The section title is printed by the renderer; the block restating it is not.
    assert sum(1 for h in en if r._title_key(h).endswith("executive summary")) == 1
    assert [h for h in en if "revenue picture" in h.lower()] == ["Revenue picture — what is disclosed"]
    assert sum(1 for h in zh if "收入情况" in h) == 1
    # The Chinese translation collapsed two English titles into one: once.
    assert sum(1 for h in zh if h == "董事会与管理层") == 1
    assert sum(1 for h in en if h in {"Board and Leadership", "Board of Directors"}) == 2
    for i, heading in enumerate(zh[1:], start=1):
        assert r._title_key(heading) != r._title_key(zh[i - 1]), heading


# ---- (d) glossary -----------------------------------------------------------------


def _glossary():
    return [
        {"term": {"en": "CSI", "zh": "信道状态信息（CSI）"}, "definition": {"en": "Channel state information.", "zh": "无线信道的状态信息。"}},
        {"term": {"en": "TDoA", "zh": "到达时间差（TDoA）"}, "definition": {"en": "Time difference of arrival.", "zh": "信号到达时间差。"}},
    ]


def test_package_glossary_is_optional_validated_and_rendered_after_the_summary(tmp_path):
    package = base._package()
    assert r.validate_package(package) == []
    package["glossary"] = _glossary()
    assert r.validate_package(package) == []
    out_en, out_zh = tmp_path / "memo" / "en.docx", tmp_path / "memo" / "zh.docx"
    r.render_memos(package, out_en=out_en, out_zh=out_zh)
    en = _headings(out_en)
    terms_at = en.index("Terms used")
    first_sections = [i for i, h in enumerate(en) if h.startswith(("I. ", "II. "))]
    assert first_sections[0] < terms_at < first_sections[1]
    assert "术语说明" in _headings(out_zh)
    cells = _cells(out_zh)
    assert "信道状态信息（CSI）" in cells and "信号到达时间差。" in cells

    bad = copy.deepcopy(package)
    bad["glossary"] = [{"term": {"en": "X", "zh": ""}}]
    errors = r._package_validation_errors(bad)
    assert any("glossary.items[0].term.zh" in e for e in errors)
    assert any("glossary.items[0].definition" in e for e in errors)


def test_glossary_block_inside_a_section(tmp_path):
    package = base._package()
    package["sections"][0]["blocks"].append(
        {"type": "glossary", "component": "glossary", "items": _glossary()}
    )
    assert r.validate_package(package) == []
    out_en, out_zh = tmp_path / "memo" / "en.docx", tmp_path / "memo" / "zh.docx"
    r.render_memos(package, out_en=out_en, out_zh=out_zh)
    assert "Terms used" in _headings(out_en)
    assert "Time difference of arrival." in _cells(out_en)
    empty = copy.deepcopy(package)
    empty["sections"][0]["blocks"][-1]["items"] = []
    assert any("glossary" in e and "non-empty" in e for e in r._package_validation_errors(empty))


def test_fallback_classes_from_a_live_gemini_package():
    """ZaiNar 2026-09-23 (Gemini, v2): plain classes the map does not know.
    "patent office registry" rendered as BSH 内部资料 — our own material —
    for a public blog post; "tier-1 tech journalism" fell through to 其他来源
    because \\bjournal\\b never matches "journalism"."""
    expected = {
        "patent office registry": "公开申报与法律文件",
        "tier-1 tech journalism": "媒体报道",
        "industry trade press": "媒体报道",
        "commercial partnership announcement": "公司披露",
        "public market data": "研究与市场数据",
        # Genuinely ours still reads as ours.
        "BSH internal notes": "BSH 内部资料",
        "diligence call summary": "BSH 内部资料",
    }
    for raw, zh in expected.items():
        assert r._source_class_label({"class": raw}, "zh") == zh, raw


def test_template_label_lines_render_in_chinese():
    """ZaiNar 2026-09-23 (Claude v2): the translator kept the founder's
    template lines verbatim, so the Chinese recommendation box read
    "Proposed amount: 【待投委会确定】" and "Strategy: [... — IC to select]"."""
    assert r.apply_zh_placeholders("Proposed amount: [TO BE DETERMINED BY IC]") == "拟投金额：【待投委会确定】"
    assert r.apply_zh_placeholders("Allocation: [TO BE DETERMINED BY IC]") == "额度分配：【待投委会确定】"
    assert r.apply_zh_placeholders("Strategy: [重仓 / 跟投 / 卡位 — IC to select]") == (
        "投资策略：【重仓 / 跟投 / 卡位 — 投委会选定】"
    )
    # A translated line keeps its own label; only the bracket is normalised.
    assert r.apply_zh_placeholders("投资策略：[重仓 / 跟投 / 卡位 — 投委会选定]") == (
        "投资策略：【重仓 / 跟投 / 卡位 — 投委会选定】"
    )
    # A label followed by real content is not a template line.
    assert r.apply_zh_placeholders("Allocation: 5% of the fund [TO BE DETERMINED BY IC]") == (
        "Allocation: 5% of the fund 【待投委会确定】"
    )
    # ZaiNar 2026-09-23 (Gemini): the translator's own variants of both.
    assert r.apply_zh_placeholders("拟投资金额：[由投委会决定 / TO BE DETERMINED BY IC]") == (
        "拟投资金额：【待投委会确定】"
    )
    assert r.apply_zh_placeholders("投资策略：[重仓 / 跟投 / 卡位 — 由投委会选择]") == (
        "投资策略：【重仓 / 跟投 / 卡位 — 投委会选定】"
    )
    # Brackets that only mention the IC are left alone.
    assert r.apply_zh_placeholders("[see the IC notes / appendix]") == "[see the IC notes / appendix]"


def test_research_reports_are_research_not_media_and_titles_are_quoted_as_published():
    for cls in ("analyst report", "industry report", "third-party research report", "academic journal article"):
        assert r._source_class_label({"class": cls}, "zh") == "研究与市场数据"
    assert r._source_class_label({"class": "news report"}, "zh") == "媒体报道"
    title = "36氪：某公司完成1亿美元C轮融资"
    assert r._loc(title, "zh", rewrite=False) == title
    assert r._loc(title, "zh") != title  # body text still gets the money form
