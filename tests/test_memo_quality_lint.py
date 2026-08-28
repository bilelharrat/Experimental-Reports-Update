from __future__ import annotations

from docx import Document

from server import memo_quality_lint


def _save_docx(path, paragraphs=(), tables=()):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    for rows in tables:
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for row_index, row in enumerate(rows):
            for col_index, value in enumerate(row):
                table.cell(row_index, col_index).text = value
    document.save(path)


def test_linter_catches_source_scaffold_fuzzy_em_dash_and_disclosure_gaps(tmp_path):
    path = tmp_path / "bad-memo.docx"
    _save_docx(
        path,
        paragraphs=[
            "I. Executive Summary",
            "ZaiNar has no battery cost [WV SPV memo].",
            "CRITICAL REALITY CHECK (for BSH)",
            "A real technical asset behind a hard IP wall (present-state).",
            "A soft instrument into an unclosed round.",
            "Revenue not disclosed. ARR not disclosed.",
            "Series A2, ~$3.0B pre-money — closing imminent.",
            "What Is Not Yet Underwritten",
            "Conditional Yes at the minimum ticket.",
        ],
        tables=[
            [
                ["Metric", "Value"],
                ["Last priced valuation", "~$1.0B post-money [companies.yaml]"],
            ],
        ],
    )

    result = memo_quality_lint.lint_memo_docx(path)
    codes = {finding.code for finding in result.findings}

    assert result.has_blocking_findings is True
    assert {
        "source_token_leak",
        "operating_table_source_token",
        "internal_artifact_leak",
        "scaffold_label",
        "banned_fuzzy_phrase",
        "em_dash_bridge",
        "disclosure_gap_without_treatment",
        "sell_side_voice_violation",
    }.issubset(codes)


def test_linter_allows_source_ids_in_fact_index_and_source_class_in_tables(tmp_path):
    path = tmp_path / "clean-memo.docx"
    document = Document()
    document.add_paragraph("I. Executive Summary")
    document.add_paragraph(
        "ZaiNar is company-reported to have named commercial momentum. "
        "Revenue is not disclosed; model treatment uses a customer-count "
        "proxy and valuation sensitivity."
    )
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Treatment"
    table.cell(1, 0).text = "Commercial signals"
    table.cell(1, 1).text = "company-reported; no source token here"
    table.cell(2, 0).text = "Recognized revenue"
    table.cell(2, 1).text = (
        "Not disclosed; model does not infer ARR from MOU headlines."
    )
    document.add_paragraph("VI. Sources, Source Classes, and Fact Reference Index")
    document.add_paragraph(
        "[S1] WV SPV memo, company-originated, used for transaction terms."
    )
    document.add_paragraph(
        "[S2] companies.yaml, internal registry, used only for fact index."
    )
    document.save(path)

    result = memo_quality_lint.lint_memo_docx(path)

    assert result.p0_findings == []


def test_linter_does_not_flag_wv_deal_party_as_artifact(tmp_path):
    # "WV" is Wisdom Ventures, a real deal counterparty (WV-ZaiNar SPV) — not an
    # internal file artifact. It must not trip internal_artifact_leak.
    path = tmp_path / "wv-memo.docx"
    _save_docx(
        path,
        paragraphs=[
            "I. Executive Summary",
            "BSH co-invests through the WV-ZaiNar SPV on a SAFE, bearing "
            "20% non-WV carry.",
        ],
    )

    result = memo_quality_lint.lint_memo_docx(path)

    assert not any(f.code == "internal_artifact_leak" for f in result.findings)


def test_linter_treats_disclosure_gap_via_sibling_note_cell(tmp_path):
    # A "Not disclosed" value paired with treatment / characterization in the
    # row's Note column is adequately treated and must not be flagged.
    path = tmp_path / "row-treated.docx"
    _save_docx(
        path,
        paragraphs=["II. Company Overview"],
        tables=[
            [
                ["Metric", "Value", "Note"],
                ["Gross margin / burn / NRR", "Not disclosed", "Valuation sensitivity"],
                ["Revenue at date", "Not disclosed", "Undefined (no denominator)"],
            ],
        ],
    )

    result = memo_quality_lint.lint_memo_docx(path)

    assert not any(
        f.code == "disclosure_gap_without_treatment" for f in result.findings
    )


def test_linter_accepts_natural_treatment_verb_forms(tmp_path):
    # Real cells from the 2026-08-21 Tenstorrent run: attempt 1 treated its
    # gaps with "our downside case assumes ..." and "we value ... take a
    # discount", which the exact-word term list missed, costing a 12-minute
    # full-package retry. Stemmed terms must accept these.
    path = tmp_path / "verb-forms.docx"
    _save_docx(
        path,
        paragraphs=["II. Company Overview"],
        tables=[
            [
                ["Topic", "Treatment"],
                [
                    "Preference stack",
                    "Detailed seniority is not disclosed, and our downside "
                    "case assumes the preference absorbs most of a weak "
                    "outcome.",
                ],
                [
                    "Team premium",
                    "Retention terms are not disclosed. We value the team as "
                    "an execution input and take a discount rather than a "
                    "premium when pricing it.",
                ],
            ],
        ],
    )

    result = memo_quality_lint.lint_memo_docx(path)

    assert not any(
        f.code == "disclosure_gap_without_treatment" for f in result.findings
    )


def test_linter_accepts_hyphenated_back_verbs(tmp_path):
    # "we back-solve" / "we back-test" are modeling verbs, not the banned
    # sell-side "we back <asset>" phrasing. The 2026-08-21 ZaiNar run showed
    # the voice rewrite mangling "we back-solve" into "BSH invests in-solve";
    # this lint pattern shared the same word-boundary hole.
    path = tmp_path / "back-solve.docx"
    _save_docx(
        path,
        paragraphs=[
            "II. Company Overview",
            "We back-solve a prior-year revenue base near $8,600,000 and "
            "back-test the resulting multiple.",
        ],
    )

    result = memo_quality_lint.lint_memo_docx(path)

    assert not any(
        f.code == "sell_side_voice_violation" for f in result.findings
    )


def test_linter_still_flags_plain_we_back(tmp_path):
    path = tmp_path / "we-back.docx"
    _save_docx(
        path,
        paragraphs=[
            "II. Company Overview",
            "We back the company because the channel is durable.",
        ],
    )

    result = memo_quality_lint.lint_memo_docx(path)

    assert any(f.code == "sell_side_voice_violation" for f in result.findings)


def test_linter_still_flags_untreated_disclosure_gap(tmp_path):
    # A bare "Not disclosed" cell with no treatment anywhere in its row must
    # still be flagged — the row-aware fix must not neuter the gate.
    path = tmp_path / "row-untreated.docx"
    _save_docx(
        path,
        paragraphs=["II. Company Overview"],
        tables=[
            [
                ["Metric", "Value"],
                ["Active users", "Not disclosed"],
            ],
        ],
    )

    result = memo_quality_lint.lint_memo_docx(path)

    assert any(
        f.code == "disclosure_gap_without_treatment" for f in result.findings
    )


def test_linter_blocks_buyer_side_language_in_final_body(tmp_path):
    path = tmp_path / "buyer-side-language.docx"
    _save_docx(
        path,
        paragraphs=[
            "I. Executive Summary",
            "Recommendation: Conditional Yes at the minimum ticket.",
            "We back infrastructure because this is why we want exposure.",
            "Top 3 Gating Questions (for BSH)",
            "What Is Not Yet Underwritten",
            "BSH target allocation: $5-10M.",
            "The final section says what BSH should do and the position size.",
            "Decision Posture: proceed if confirmed.",
            "Wisdom Ventures should be able to share the executed term sheet.",
            "Federal contract identifiers are realistically obtainable.",
            "These items should be closeable in the diligence window.",
            "Confirm before funding: the signed-vs-MOU split matches the model.",
            "Confirm the A2 closes on disclosed terms.",
            "Confirm final terms before signing subscription documents.",
            "We still need a claim-scope read before giving credit to patents.",
            "Diligence Thresholds",
            "Next Diligence Actions",
            "Require data room access, a named institutional lead, MFN, "
            "down-round protection, information rights, and voting rights.",
        ],
    )

    result = memo_quality_lint.lint_memo_docx(path)

    assert result.has_blocking_findings is True
    assert any(f.code == "sell_side_voice_violation" for f in result.findings)


def test_linter_blocks_third_person_recommendation_voice(tmp_path):
    path = tmp_path / "third-person-recommendation.docx"
    _save_docx(
        path,
        paragraphs=[
            "I. Executive Summary",
            "The opportunity offered to investors is a sponsor-backed SPV interest.",
            "The base case credits commercial pull before revenue is disclosed.",
            "The recommendation is Proceed if confirmed: participate in the SPV.",
            "The right posture is proceed if confirmed only if contracts bind.",
        ],
    )

    result = memo_quality_lint.lint_memo_docx(path)
    snippets = " ".join(f.snippet for f in result.findings)

    assert result.has_blocking_findings is True
    assert any(f.code == "sell_side_voice_violation" for f in result.findings)
    assert "The recommendation is" in snippets
    assert "The opportunity offered to investors is" in snippets


def test_linter_blocks_meta_process_language(tmp_path):
    path = tmp_path / "meta-process-language.docx"
    _save_docx(
        path,
        paragraphs=[
            "I. Executive Summary",
            "The memo frames ZaiNar as a scarce technical asset.",
            "This memo treats ZaiNar as a scarce technical asset.",
            "Our memo recommends participating through the SPV.",
            "The analysis suggests revenue should be treated as unproven.",
            "Our analysis points to conditional participation.",
            "This document outlines the key risks.",
            "This section covers the investment risk.",
            "The framework points toward conditional participation.",
            "Memo language was closing imminent as of May 2026.",
            "The sponsor implies a licensing fallback under standardization.",
            "The sponsor itself flags standardization risk.",
            "The sponsor acknowledges the figure is MOU-heavy.",
            "Sponsor explicitly discloses 18 to 36 month carrier sales cycles.",
            "Pipeline figures inside the memo are company-provided.",
            "The competitor list embedded in the registry understates threat.",
            "Revenue is not documented in source material.",
            "We outline the investment case below.",
            "We discuss the downside case later.",
        ],
    )

    result = memo_quality_lint.lint_memo_docx(path)
    snippets = " ".join(f.snippet for f in result.findings)

    assert result.has_blocking_findings is True
    assert any(f.code == "meta_process_language" for f in result.findings)
    assert "The memo" in snippets
    assert "The analysis" in snippets
    assert "This memo" in snippets
    assert "Our memo" in snippets
    assert "Our analysis" in snippets
    assert "This document" in snippets
    assert "Memo language was" in snippets
    assert "The sponsor implies" in snippets
    assert "The sponsor acknowledges" in snippets
    assert "inside the memo" in snippets
    assert "embedded in the registry" in snippets
    assert "We outline" in snippets
    assert "We discuss" in snippets


def test_linter_blocks_meta_process_language_in_source_index(tmp_path):
    path = tmp_path / "source-index-meta-process.docx"
    _save_docx(
        path,
        paragraphs=[
            "I. Executive Summary",
            "We recommend participating where valuation support is visible.",
            "VI. Sources, Source Classes, and Fact Reference Index",
            "[S1] Company materials, used for this memo.",
        ],
    )

    result = memo_quality_lint.lint_memo_docx(path)
    snippets = " ".join(f.snippet for f in result.findings)

    assert result.has_blocking_findings is True
    assert any(f.code == "meta_process_language" for f in result.findings)
    assert "this memo" in snippets


def test_linter_allows_neutral_meta_references_in_source_index(tmp_path):
    """Section VI describes the memo's own sourcing by definition — the
    2026-08-28 benchmark runs both burned a surgical-repair round on these
    exact neutral-article phrasings. Demonstrative forms stay banned."""
    path = tmp_path / "source-index-neutral-meta.docx"
    _save_docx(
        path,
        paragraphs=[
            "I. Executive Summary",
            "We recommend participating where valuation support is visible.",
            "VI. Sources, Source Classes, and Fact Reference Index",
            (
                "[S1] Q2 FY2027 report: anchor for reported revenue; every "
                "headline figure in the memo carries this quarter end."
            ),
            (
                "[S2] Company registry entry: the analyzed entity is resolved "
                "to NVIDIA Corporation and the registry legal, headquarters "
                "and founding fields are not used."
            ),
            "[S3] Company materials, used for this memo.",
        ],
    )

    result = memo_quality_lint.lint_memo_docx(path)
    meta = [f for f in result.findings if f.code == "meta_process_language"]
    snippets = " ".join(f.snippet for f in meta)

    # Neutral-article references in the index are treatment language.
    assert "the memo" not in snippets
    assert "the registry" not in snippets
    # The demonstrative form is still process leakage, even in Section VI.
    assert "this memo" in snippets


def test_linter_allows_mandatory_disclosure_language(tmp_path):
    """The renderer REQUIRES the legal disclosures sentence; the meta gate
    must not flag its 'This document …' phrasing (Run B/C false positive)."""
    path = tmp_path / "disclosure-language.docx"
    _save_docx(
        path,
        paragraphs=[
            "V. Financial Forecast & Valuation",
            (
                "This document is a confidential summary prepared for "
                "existing and prospective limited partners and is not an "
                "offer to sell securities; any investment is made only "
                "through definitive subscription documents available to "
                "accredited investors and may result in partial or total "
                "loss of capital."
            ),
            # Control: 'This document' WITHOUT disclosure markers stays
            # banned in body sections.
            "This document outlines the key risks.",
        ],
    )

    result = memo_quality_lint.lint_memo_docx(path)
    meta = [f for f in result.findings if f.code == "meta_process_language"]

    # Exactly one finding: the control paragraph. Two would mean the
    # disclosure sentence was flagged as well.
    assert len(meta) == 1
    assert "outlines the key risks" in meta[0].snippet


def test_linter_blocks_internal_questionnaire_language(tmp_path):
    path = tmp_path / "gating-questions.docx"
    _save_docx(
        path,
        paragraphs=[
            "I. Executive Summary",
            "Top 3 Gating Questions",
            "Top 3 Decision Questions (for BSH)",
            "Open Questions",
            "Does binding deployment evidence support the current allocation?",
            "Can carrier conversion produce upside without assuming all MOUs become revenue?",
            "The final contract split must be met before BSH funds.",
        ],
    )

    result = memo_quality_lint.lint_memo_docx(path)

    assert result.has_blocking_findings is True
    assert any(f.code == "sell_side_voice_violation" for f in result.findings)


def test_linter_blocks_copied_packet_process_labels(tmp_path):
    path = tmp_path / "packet-labels.docx"
    _save_docx(
        path,
        paragraphs=[
            "I. Executive Summary",
            "Prompt: Explain why the claim matters.",
            "Design prompt: Use a ladder diagram.",
            "Reviewer prompts:",
            "Confidence: medium",
            "Source traces:",
            "Readiness Reviews And Waivers",
            "Memo Uses",
            "Pre-Mortem",
            "Reverse IC",
            "Validation & Assumptions Log",
        ],
    )

    result = memo_quality_lint.lint_memo_docx(path)

    assert result.has_blocking_findings is True
    assert any(f.code == "packet_process_label" for f in result.findings)


def test_linter_blocks_confirmation_sections_and_conditions(tmp_path):
    path = tmp_path / "closing-confirmations.docx"
    _save_docx(
        path,
        paragraphs=[
            "I. Executive Summary",
            "Closing Confirmations",
            "We proceed once final A2 terms, SAFE conversion mechanics, and the binding-contract split are confirmed.",
            "We would revisit if the A2 slips beyond the expected closing window or prices below the current mark.",
            "The recommendation is Proceed if confirmed.",
            "Investment Conditions",
            "Expected Bars",
            "Valuation Sensitivity Bars",
            "Stop or Revisit Conditions",
            "Stop / Revisit Triggers",
            "What Would Make Us Revisit or Decline",
            "Immediate Confirmation Work",
            "BSH thesis fit relies on the late-stage financial-return exception.",
            "The disclosed founders are not Asian-immigrant per the BSH preference.",
            "Series B pricing shifts the SPV from a paper-mark outcome into flat-to-modest carry.",
            "Closing bar: contract split supports the case.",
            "Cross-check DoD signings on SAM.gov and USAspending.gov.",
            "Source two non-investor technical references through the network.",
            "Patent counsel claim-scope and freedom-to-operate read supports durable patent leverage.",
        ],
    )

    result = memo_quality_lint.lint_memo_docx(path)

    assert result.has_blocking_findings is True
    assert any(f.code == "sell_side_voice_violation" for f in result.findings)


def test_linter_blocks_imperative_confirm_items_and_heading(tmp_path):
    path = tmp_path / "confirm-items.docx"
    _save_docx(
        path,
        paragraphs=[
            "I. Executive Summary",
            "Closing Confirmations",
            "Confirm the A2 lead investor identity and final pre-money.",
            "Expected bar: A2 lead investor identity and final pre-money match the disclosed Series A2 economics.",
        ],
    )

    result = memo_quality_lint.lint_memo_docx(path)
    snippets = " ".join(f.snippet for f in result.findings)

    assert result.has_blocking_findings is True
    assert "Confirm the A2 lead" in snippets
    assert "Closing Confirmations" in snippets
