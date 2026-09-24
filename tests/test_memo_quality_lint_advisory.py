"""Non-blocking (P1/P2) lint findings added 2026-09-22 — invented sizing,
vehicle assertions without deal terms, a commit below the firm's hurdle,
duplicated table rows, a figure restated within one section, and the
"rather than" density the Chinese memo mirrors as 而非 — plus the negated
stock slogans, which join the existing P0 slogan ban.

None of the advisory findings may block: every test that expects one also
checks that the gate still passes.
"""

from __future__ import annotations

from docx import Document

from server import memo_quality_lint


def _docx(path, items):
    """items: strings become paragraphs, lists of rows become tables — in
    order, so a table lands in the section of the heading before it."""
    document = Document()
    for item in items:
        if isinstance(item, str):
            document.add_paragraph(item)
            continue
        table = document.add_table(rows=len(item), cols=len(item[0]))
        for row_index, row in enumerate(item):
            for col_index, value in enumerate(row):
                table.cell(row_index, col_index).text = value
    document.save(path)
    return path


def _codes(result, severity=None):
    return {
        f.code
        for f in result.findings
        if severity is None or f.severity == severity
    }


# ---- sizing ------------------------------------------------------------------


def test_an_amount_after_bsh_commits_without_a_sizing_input_is_p1(tmp_path):
    path = _docx(
        tmp_path / "m.docx",
        [
            "I. Executive Summary",
            "Recommendation: BSH commits $15M to Tarnwell at the Series B terms.",
        ],
    )
    result = memo_quality_lint.lint_memo_docx(path)
    assert "sizing_without_input" in _codes(result, "P1")
    assert result.has_blocking_findings is False

    supplied = memo_quality_lint.lint_memo_docx(path, sizing_supplied=True)
    assert "sizing_without_input" not in _codes(supplied)


def test_the_amount_free_recommendation_is_clean(tmp_path):
    path = _docx(
        tmp_path / "m.docx",
        [
            "I. Executive Summary",
            "Recommendation: BSH commits to Tarnwell at the Series B terms.",
        ],
    )
    assert "sizing_without_input" not in _codes(memo_quality_lint.lint_memo_docx(path))


# ---- deal terms --------------------------------------------------------------


def test_vehicle_assertions_without_deal_terms_are_p1_only_when_known(tmp_path):
    path = _docx(
        tmp_path / "m.docx",
        [
            "I. Executive Summary",
            "BSH participates through an SPV holding a SAFE at a $1.2B cap.",
            "No vehicle or terms on file (pipeline stage: Sourced). The SPV "
            "question is open.",
        ],
    )
    no_terms = memo_quality_lint.lint_memo_docx(path, deal_terms_on_file=False)
    findings = [f for f in no_terms.findings if f.code == "deal_vehicle_without_terms"]
    assert len(findings) == 1  # the "No vehicle or terms on file" line is exempt
    assert findings[0].severity == "P1"
    assert no_terms.has_blocking_findings is False
    # Unknown (None) and terms on file (True) never flag.
    assert "deal_vehicle_without_terms" not in _codes(memo_quality_lint.lint_memo_docx(path))
    assert "deal_vehicle_without_terms" not in _codes(
        memo_quality_lint.lint_memo_docx(path, deal_terms_on_file=True)
    )


# ---- hurdle ------------------------------------------------------------------


def _scenario_memo(tmp_path, base_moic: str, call: str):
    return _docx(
        tmp_path / "m.docx",
        [
            "I. Executive Summary",
            call,
            "V. Financial Forecast & Valuation",
            [
                ["Scenario", "Exit value", "Gross MOIC", "IRR"],
                ["Bear", "$0.9B", "0.8x", "-6%"],
                ["Base", "$2.6B", base_moic, "15%"],
                ["Bull", "$5.0B", "3.1x", "40%"],
            ],
        ],
    )


def test_a_commit_below_the_hurdle_is_p1(tmp_path):
    path = _scenario_memo(
        tmp_path, "1.5x-2.4x", "Recommendation: BSH commits to Tarnwell at the Series B terms."
    )
    result = memo_quality_lint.lint_memo_docx(path, hurdle_moic=2.0)
    finding = next(f for f in result.findings if f.code == "recommendation_below_hurdle")
    assert finding.severity == "P1"
    assert "1.5x" in finding.suggestion and "2x" in finding.suggestion
    assert result.has_blocking_findings is False


def test_no_hurdle_finding_without_a_policy_above_it_or_on_a_watch(tmp_path):
    commit = "Recommendation: BSH commits to Tarnwell at the Series B terms."
    below = _scenario_memo(tmp_path, "1.5x", commit)
    assert "recommendation_below_hurdle" not in _codes(memo_quality_lint.lint_memo_docx(below))
    above = _scenario_memo(tmp_path, "2.4x", commit)
    assert "recommendation_below_hurdle" not in _codes(
        memo_quality_lint.lint_memo_docx(above, hurdle_moic=2.0)
    )
    watch = _scenario_memo(tmp_path, "1.5x", "Recommendation: watch Tarnwell — the trigger is a term sheet.")
    assert "recommendation_below_hurdle" not in _codes(
        memo_quality_lint.lint_memo_docx(watch, hurdle_moic=2.0)
    )


# ---- repetition --------------------------------------------------------------


def test_a_row_repeated_in_a_second_table_is_p1(tmp_path):
    path = _docx(
        tmp_path / "m.docx",
        [
            "II. Company Overview",
            [["Metric", "Value"], ["Post-money valuation", "$1.2B"], ["ARR", "$30M"]],
            [["Term", "Value"], ["Post-money valuation", "$1.2B"], ["Instrument", "SAFE"]],
        ],
    )
    result = memo_quality_lint.lint_memo_docx(path)
    dupes = [f for f in result.findings if f.code == "duplicate_table_row"]
    assert len(dupes) == 1 and dupes[0].severity == "P1"
    assert "table 2" in dupes[0].location
    assert result.has_blocking_findings is False


def test_risk_card_rows_and_figure_free_rows_are_not_duplicates(tmp_path):
    card = [["Risk Type", "Valuation & exit"], ["Risk Rating", "9/10: one assumption"]]
    path = _docx(tmp_path / "m.docx", ["IV. Investment Risk", card, card])
    assert "duplicate_table_row" not in _codes(memo_quality_lint.lint_memo_docx(path))


def test_a_figure_restated_more_than_twice_in_one_section_is_p1(tmp_path):
    items = [
        "V. Financial Forecast & Valuation",
        "The round prices the company at $1.2B.",
        "At $1.2B the entry multiple is 40x.",
        "The base case exits above $1.2B only if growth holds.",
    ]
    path = _docx(tmp_path / "m.docx", items)
    result = memo_quality_lint.lint_memo_docx(path)
    repeated = [f for f in result.findings if f.code == "number_repeated_in_section"]
    assert len(repeated) == 1 and repeated[0].severity == "P1"
    assert "$1.2b appears 3 times" in repeated[0].snippet
    assert result.has_blocking_findings is False
    # A pinned figure the echo gate requires is exempt.
    pinned = memo_quality_lint.lint_memo_docx(path, pinned_values=["$1.2B post-money"])
    assert "number_repeated_in_section" not in _codes(pinned)


def test_a_figure_repeated_down_a_table_column_is_not_restatement(tmp_path):
    # The v2 scenario table carries the same dilution on every row; that is
    # layout, and a clean memo must not pick up a P1 for it.
    scenarios = [
        ["Case", "Exit value", "Dilution", "MOIC"],
        ["Bear", "$0.8B", "15%", "0.6x"],
        ["Base", "$2.4B", "15%", "1.8x"],
        ["Bull", "$4.0B", "15%", "3.0x"],
    ]
    path = _docx(tmp_path / "m.docx", ["V. Financial Forecast & Valuation", scenarios])
    assert "number_repeated_in_section" not in _codes(memo_quality_lint.lint_memo_docx(path))


def test_the_same_figure_in_different_sections_is_not_repetition(tmp_path):
    path = _docx(
        tmp_path / "m.docx",
        [
            "I. Executive Summary",
            "The round prices the company at $1.2B.",
            "II. Company Overview",
            "The March round set $1.2B.",
            "V. Financial Forecast & Valuation",
            "At $1.2B the entry multiple is 40x.",
        ],
    )
    assert "number_repeated_in_section" not in _codes(memo_quality_lint.lint_memo_docx(path))


# ---- contrast density --------------------------------------------------------


def test_rather_than_density_is_a_p2_note(tmp_path):
    path = _docx(
        tmp_path / "m.docx",
        [
            "I. Executive Summary",
            "The pipeline is demand rather than revenue. The price rests on growth, "
            "not margin. Customers buy uptime rather than hardware. The moat is "
            "distribution, not IP.",
        ],
    )
    result = memo_quality_lint.lint_memo_docx(path)
    finding = next(f for f in result.findings if f.code == "rather_than_density")
    assert finding.severity == "P2"
    assert "而非" in finding.suggestion
    assert result.has_blocking_findings is False


def test_idioms_and_sparse_contrasts_are_not_flagged(tmp_path):
    words = " ".join(["Revenue grew in every quarter of the year."] * 40)
    path = _docx(
        tmp_path / "m.docx",
        [
            "I. Executive Summary",
            words,
            "The price is high, not yet absurd. It is not only growth, not least margin.",
            "Customers buy uptime rather than hardware.",
        ],
    )
    assert "rather_than_density" not in _codes(memo_quality_lint.lint_memo_docx(path))


# ---- negated slogans (P0, like the positive ones) -----------------------------


def test_negated_stock_slogans_are_banned_like_the_positive_ones(tmp_path):
    for bad in (
        "We do not recommend participating at this price.",
        "We are not being offered an allocation.",
        "We are not participating through the SPV.",
    ):
        path = _docx(tmp_path / "m.docx", ["I. Executive Summary", bad])
        result = memo_quality_lint.lint_memo_docx(path)
        assert "sell_side_voice_violation" in _codes(result, "P0"), bad
        assert memo_quality_lint.banned_body_voice_phrase(bad), bad
