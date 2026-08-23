from __future__ import annotations

import copy
import json

from docx import Document
import pytest

from server import (
    internal_memo_renderer,
    memo_chinese_parity,
    memo_docx_renderer,
    memo_quality_lint,
)


def _package() -> dict:
    return {
        "schema_version": 1,
        "company": {
            "name": "Generalist, Inc.",
            "descriptor": {
                "en": "Embodied AI systems for industrial automation",
                "zh": "面向工业自动化的具身智能系统",
            },
            "stage": "Late-stage / pre-IPO",
            "sector": "AI Robotics",
            "location": "San Francisco, CA",
            "round": "Series D context",
        },
        "run": {
            "run_id": "2026-06-22__120000",
            "as_of": "2026-06-22",
        },
        "sections": [
            {
                "id": "executive_summary",
                "blocks": [
                    {
                        "type": "heading",
                        "level": 2,
                        "text": {
                            "en": "Investment Opportunity",
                            "zh": "投资机会",
                        },
                    },
                    {
                        "type": "paragraph",
                        "text": {
                            "en": (
                                "We recommend evaluating Generalist through a "
                                "deployment-depth investment case because production proof is "
                                "visible while revenue disclosure remains limited."
                            ),
                            "zh": (
                                "我们建议以部署深度为核心情景承销 Generalist，"
                                "因为生产环境证据已经出现，但收入披露仍然有限。"
                            ),
                        },
                    },
                    {
                        "type": "table",
                        "component": "key_metrics_snapshot",
                        "title": {
                            "en": "Key Metrics Snapshot",
                            "zh": "关键指标快照",
                        },
                        "headers": [
                            {"en": "Metric", "zh": "指标"},
                            {"en": "Treatment", "zh": "处理方式"},
                        ],
                        "rows": [
                            [
                                {"en": "Revenue", "zh": "收入"},
                                {
                                    "en": "Use a deployment-depth proxy until audited ARR is available.",
                                    "zh": "在审计 ARR 可得前，使用部署深度作为代理指标。",
                                },
                            ],
                            [
                                {"en": "Customer depth", "zh": "客户深度"},
                                {
                                    "en": "Model a step-up only where repeatable production usage is visible.",
                                    "zh": "只有在可重复生产环境使用可见时，才在模型中计入估值上调。",
                                },
                            ],
                        ],
                    },
                    {
                        "type": "table",
                        "component": "deal_terms",
                        "title": {
                            "en": "Headline Terms",
                            "zh": "核心交易条款",
                        },
                        "headers": [
                            {"en": "Term", "zh": "条款"},
                            {"en": "Detail", "zh": "细节"},
                        ],
                        "rows": [
                            [
                                {"en": "Instrument", "zh": "工具"},
                                {
                                    "en": "Preferred equity with final economics modeled by scenario range.",
                                    "zh": "优先股，最终经济性通过情景区间建模。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "callout",
                        "tone": "warning",
                        "title": {
                            "en": "Valuation Sensitivity",
                            "zh": "估值敏感性",
                        },
                        "items": [
                            {
                                "en": "Production deployment depth and margin path drive sizing.",
                                "zh": "生产部署深度和利润率路径决定配置规模。",
                            },
                            {
                                "en": "Separate contracted pilots from recognized revenue.",
                                "zh": "区分已签试点与已确认收入。",
                            },
                        ],
                    },
                ],
            },
            {
                "id": "company_overview",
                "blocks": [
                    {
                        "type": "paragraph",
                        "text": {
                            "en": (
                                "Generalist sells automation systems for "
                                "industrial workflows where repeatability and "
                                "integration depth drive adoption."
                            ),
                            "zh": (
                                "Generalist 销售面向工业流程的自动化系统，"
                                "可重复性和集成深度决定采用速度。"
                            ),
                        },
                    },
                    {
                        "type": "table",
                        "component": "board",
                        "title": {
                            "en": "Board of Directors",
                            "zh": "董事会",
                        },
                        "headers": [
                            {"en": "Member", "zh": "成员"},
                            {"en": "Strategic value", "zh": "战略价值"},
                        ],
                        "rows": [
                            [
                                {"en": "Lead investor director", "zh": "领投方董事"},
                                {
                                    "en": "Adds governance depth for late-stage scaling.",
                                    "zh": "为后期扩张增加治理深度。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "table",
                        "component": "revenue",
                        "title": {
                            "en": "Revenue Picture",
                            "zh": "收入情况",
                        },
                        "headers": [
                            {"en": "Component", "zh": "组成"},
                            {"en": "Treatment", "zh": "处理方式"},
                        ],
                        "rows": [
                            [
                                {"en": "Recognized revenue", "zh": "已确认收入"},
                                {
                                    "en": "Not disclosed; scenario model uses deployment proxies.",
                                    "zh": "未披露；情景模型使用部署代理指标。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "table",
                        "component": "key_operating_metrics",
                        "title": {
                            "en": "Key Operating Metrics",
                            "zh": "关键运营指标",
                        },
                        "headers": [
                            {"en": "Metric", "zh": "指标"},
                            {"en": "Read", "zh": "解读"},
                        ],
                        "rows": [
                            [
                                {"en": "Gross margin", "zh": "毛利率"},
                                {
                                    "en": "Unavailable; margin path remains a valuation sensitivity.",
                                    "zh": "不可得；利润率路径仍是估值敏感因素。",
                                },
                            ]
                        ],
                    },
                ],
            },
            {
                "id": "investment_highlights",
                "blocks": [
                    {
                        "type": "bullets",
                        "items": [
                            {
                                "en": "Customer proof supports the BSH infrastructure-scarcity thesis.",
                                "zh": "客户验证支持 BSH 的基础设施稀缺性判断。",
                            },
                            {
                                "en": "The category can support expansion if deployments repeat.",
                                "zh": "如果部署可重复，该品类具备扩张空间。",
                            },
                        ],
                    },
                    {
                        "type": "table",
                        "component": "competitive_analysis",
                        "title": {
                            "en": "Competitive Analysis",
                            "zh": "竞争分析",
                        },
                        "headers": [
                            {"en": "Competitor", "zh": "竞争方"},
                            {"en": "Weakness vs. company", "zh": "相对弱点"},
                        ],
                        "rows": [
                            [
                                {"en": "Legacy automation vendors", "zh": "传统自动化厂商"},
                                {
                                    "en": "Integration depth is weaker in targeted workflows.",
                                    "zh": "在目标流程中的集成深度较弱。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "table",
                        "component": "replacement_coexistence",
                        "title": {
                            "en": "Replacement vs. Coexistence",
                            "zh": "替代与共存",
                        },
                        "headers": [
                            {"en": "Workflow", "zh": "流程"},
                            {"en": "Read", "zh": "解读"},
                        ],
                        "rows": [
                            [
                                {"en": "Industrial automation", "zh": "工业自动化"},
                                {
                                    "en": "Mostly coexists before replacing manual workflow layers.",
                                    "zh": "主要先共存，再替代人工流程层。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "table",
                        "component": "moat",
                        "title": {
                            "en": "Moat",
                            "zh": "护城河",
                        },
                        "headers": [
                            {"en": "Component", "zh": "组成"},
                            {"en": "Durability", "zh": "持久性"},
                        ],
                        "rows": [
                            [
                                {"en": "Workflow lock-in", "zh": "流程锁定"},
                                {
                                    "en": "Strengthens as production integrations deepen.",
                                    "zh": "随着生产集成加深而增强。",
                                },
                            ]
                        ],
                    },
                ],
            },
            {
                "id": "investment_risk",
                "blocks": [
                    {
                        "type": "paragraph",
                        "text": {
                            "en": (
                                "Risk concentrates in commercial proof, not in whether "
                                "the technology works in production."
                            ),
                            "zh": "风险集中在商业验证，而非技术能否在生产环境运行。",
                        },
                    },
                    {
                        "type": "heading",
                        "level": 3,
                        "text": {
                            "en": "Risk 1: Deployments may stay services-heavy and cap margins",
                            "zh": "风险 1：部署可能持续偏服务交付，压制利润率",
                        },
                    },
                    {
                        "type": "table",
                        "component": "risk_register",
                        "layout": "key_value",
                        "headers": [],
                        "rows": [
                            [
                                {"en": "Risk Type", "zh": "风险类型"},
                                {"en": "Commercial", "zh": "商业"},
                            ],
                            [
                                {"en": "Why it matters", "zh": "为什么重要"},
                                {
                                    "en": (
                                        "Each deployment still needs on-site engineering, so "
                                        "gross margin sits below the software multiple the "
                                        "round implies."
                                    ),
                                    "zh": (
                                        "每次部署仍需要现场工程支持，因此毛利率低于本轮估值"
                                        "所隐含的软件业务倍数。"
                                    ),
                                },
                            ],
                            [
                                {"en": "What we watch", "zh": "跟踪信号"},
                                {
                                    "en": (
                                        "Repeatable deployment margin by customer cohort, "
                                        "reported quarterly."
                                    ),
                                    "zh": "按客户批次的可重复部署利润率，按季度跟踪。",
                                },
                            ],
                            [
                                {"en": "Likelihood", "zh": "可能性"},
                                {
                                    "en": (
                                        "Medium: every deployment to date has needed "
                                        "on-site work."
                                    ),
                                    "zh": "中：迄今每次部署都需要现场支持。",
                                },
                            ],
                            [
                                {"en": "Risk Rating", "zh": "风险评分"},
                                {
                                    "en": "7/10: margin path drives the exit multiple.",
                                    "zh": "7/10：利润率路径决定退出倍数。",
                                },
                            ],
                        ],
                    },
                    {
                        "type": "heading",
                        "level": 3,
                        "text": {
                            "en": "Risk 2: Hardware integration may slow gross-margin expansion",
                            "zh": "风险 2：硬件集成可能拖慢毛利率提升",
                        },
                    },
                    {
                        "type": "table",
                        "component": "risk_register",
                        "layout": "key_value",
                        "headers": [],
                        "rows": [
                            [
                                {"en": "Risk Type", "zh": "风险类型"},
                                {"en": "Technology", "zh": "技术"},
                            ],
                            [
                                {"en": "Why it matters", "zh": "为什么重要"},
                                {
                                    "en": (
                                        "Custom hardware per site adds cost the software "
                                        "roadmap cannot amortize quickly."
                                    ),
                                    "zh": "每个站点的定制硬件带来软件路线图难以快速摊薄的成本。",
                                },
                            ],
                            [
                                {"en": "What we watch", "zh": "跟踪信号"},
                                {
                                    "en": "Bill-of-materials cost per deployment over 2026.",
                                    "zh": "2026 年内每次部署的物料成本。",
                                },
                            ],
                            [
                                {"en": "Likelihood", "zh": "可能性"},
                                {
                                    "en": (
                                        "Low: the reference design has held across "
                                        "recent sites."
                                    ),
                                    "zh": "低：参考设计在近期站点保持稳定。",
                                },
                            ],
                            [
                                {"en": "Risk Rating", "zh": "风险评分"},
                                {
                                    "en": "5/10: meaningful but monitorable.",
                                    "zh": "5/10：影响明确但可跟踪。",
                                },
                            ],
                        ],
                    },
                    {
                        "type": "heading",
                        "level": 3,
                        "text": {
                            "en": "Risk 3: Enterprise buyers may pause pilots in a downturn",
                            "zh": "风险 3：企业客户可能在下行周期暂停试点",
                        },
                    },
                    {
                        "type": "table",
                        "component": "risk_register",
                        "layout": "key_value",
                        "headers": [],
                        "rows": [
                            [
                                {"en": "Risk Type", "zh": "风险类型"},
                                {"en": "Market", "zh": "市场"},
                            ],
                            [
                                {"en": "Why it matters", "zh": "为什么重要"},
                                {
                                    "en": (
                                        "Pilot budgets are discretionary, so a capex pullback "
                                        "delays the conversion the base case counts on."
                                    ),
                                    "zh": (
                                        "试点预算属于可自由支配支出，资本开支收缩会推迟基准"
                                        "情景所依赖的转化。"
                                    ),
                                },
                            ],
                            [
                                {"en": "What we watch", "zh": "跟踪信号"},
                                {
                                    "en": "Pilot-to-production conversion rate each quarter.",
                                    "zh": "每季度试点转生产的转化率。",
                                },
                            ],
                            [
                                {"en": "Likelihood", "zh": "可能性"},
                                {
                                    "en": (
                                        "Medium: pilot budgets tighten first in a "
                                        "downturn."
                                    ),
                                    "zh": "中：下行周期试点预算最先收缩。",
                                },
                            ],
                            [
                                {"en": "Risk Rating", "zh": "风险评分"},
                                {
                                    "en": "4/10: real but limited effect on the case.",
                                    "zh": "4/10：真实存在但对投资判断影响有限。",
                                },
                            ],
                        ],
                    },
                    {
                        "type": "callout",
                        "component": "disconfirming_evidence",
                        "tone": "warning",
                        "title": {
                            "en": "Bear-Case Evidence",
                            "zh": "熊市情景证据",
                        },
                        "items": [
                            {
                                "en": "Revenue disclosure remains limited relative to valuation.",
                                "zh": "相对于估值，收入披露仍然有限。",
                            }
                        ],
                    },
                ],
            },
            {
                "id": "financial_forecast_valuation",
                "blocks": [
                    {
                        "type": "paragraph",
                        "text": {
                            "en": (
                                "The valuation case should use scenario ranges "
                                "until audited revenue and margin data are available."
                            ),
                            "zh": (
                                "在审计收入和利润率数据可得前，估值判断应使用情景区间。"
                            ),
                        },
                    },
                    {
                        "type": "table",
                        "component": "time_base_integrity",
                        "title": {
                            "en": "Time-Base Integrity Table",
                            "zh": "时间基准一致性表",
                        },
                        "headers": [
                            {"en": "Event", "zh": "事件"},
                            {"en": "Valuation", "zh": "估值"},
                        ],
                        "rows": [
                            [
                                {"en": "Current round", "zh": "当前轮次"},
                                {
                                    "en": "Modeled against latest disclosed valuation date.",
                                    "zh": "基于最近披露估值日期建模。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "table",
                        "component": "growth_bridge",
                        "title": {
                            "en": "Growth Bridge Table",
                            "zh": "增长桥接表",
                        },
                        "headers": [
                            {"en": "Bridge item", "zh": "桥接项"},
                            {"en": "Model treatment", "zh": "模型处理"},
                        ],
                        "rows": [
                            [
                                {"en": "New deployments", "zh": "新增部署"},
                                {
                                    "en": "Credited only where repeatable usage is visible.",
                                    "zh": "仅在可重复使用可见时计入。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "table",
                        "component": "scenario_analysis",
                        "title": {
                            "en": "Scenario Analysis",
                            "zh": "情景分析",
                        },
                        "headers": [
                            {"en": "Scenario", "zh": "情景"},
                            {"en": "Valuation outcome", "zh": "估值结果"},
                        ],
                        "rows": [
                            [
                                {"en": "Base", "zh": "基准"},
                                {
                                    "en": "Moderate step-up if deployments repeat with margin evidence.",
                                    "zh": "若部署可重复且利润率有证据，则温和上调。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "heading",
                        "level": 2,
                        "component": "investment_decision",
                        "text": {
                            "en": "Investment Decision / Closing View",
                            "zh": "投资决策 / 结论观点",
                        },
                    },
                    {
                        "type": "paragraph",
                        "text": {
                            "en": (
                                "We recommend participating only where valuation support "
                                "tracks repeatable customer deployment and margin evidence."
                            ),
                            "zh": (
                                "我们建议仅在估值支撑与可重复客户部署和利润率证据相匹配时参与。"
                            ),
                        },
                    },
                    {
                        "type": "table",
                        "component": "evidence_thresholds",
                        "title": {
                            "en": "Evidence Thresholds For Step-Up Support",
                            "zh": "支持估值上调的证据阈值",
                        },
                        "headers": [
                            {"en": "Evidence threshold", "zh": "证据阈值"},
                            {"en": "Valuation treatment", "zh": "估值处理"},
                        ],
                        "rows": [
                            [
                                {"en": "Repeatable production usage", "zh": "可重复生产使用"},
                                {
                                    "en": "Supports a stronger step-up only when visible across multiple customers.",
                                    "zh": "只有在多个客户中可见时，才支持更强的估值上调。",
                                },
                            ]
                        ],
                    },
                    {
                        "type": "paragraph",
                        "component": "disclosures",
                        "text": {
                            "en": (
                                "Disclosure: distributed only to the named recipients; "
                                "not an offer to sell securities. Any investment is governed "
                                "by definitive subscription documents and may result in "
                                "partial or total loss."
                            ),
                            "zh": (
                                "披露：本备忘录仅供指定接收方使用，并非证券出售要约。"
                                "任何投资均以最终认购文件为准，并可能产生部分或全部损失。"
                            ),
                        },
                    },
                ],
            },
        ],
        "sources": [
            {
                "id": "S1",
                "title": "Company investor materials",
                "class": {"en": "Company material", "zh": "公司材料"},
                "treatment": {
                    "en": "Used for product, customer, and funding context.",
                    "zh": "用于产品、客户和融资背景。",
                },
                "as_of": "2026-06-22",
            }
        ],
    }


def _all_text(path):
    document = Document(path)
    text = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text.append(cell.text)
    return "\n".join(text)


def test_parameterized_renderer_writes_bilingual_docx_and_logs(tmp_path):
    package_path = tmp_path / "logs" / "memo_package.json"
    package_path.parent.mkdir(parents=True)
    package_path.write_text(json.dumps(_package()), encoding="utf-8")
    out_en = tmp_path / "memo" / "Generalist, Inc. - Investment Memo - 2026-06-22__120000.docx"
    out_zh = tmp_path / "memo" / "Generalist, Inc. - 投资备忘录 - 2026-06-22__120000.docx"

    result = memo_docx_renderer.render_memos(
        package_path,
        out_en=out_en,
        out_zh=out_zh,
        manifest_path=tmp_path / "logs" / "run_manifest.md",
        inventory_path=tmp_path / "logs" / "file_inventory.md",
    )

    assert result["ok"] is True
    assert out_en.exists()
    assert out_zh.exists()
    assert "Generalist, Inc." in _all_text(out_en)
    assert "I. EXECUTIVE SUMMARY" in _all_text(out_en)
    assert "BSH target allocation" not in _all_text(out_en)
    assert "target allocation" not in _all_text(out_en)
    assert "BSH ticket size" not in _all_text(out_en)
    assert "关键指标快照" in _all_text(out_zh)
    assert "server.memo_docx_renderer" in (
        tmp_path / "logs" / "run_manifest.md"
    ).read_text(encoding="utf-8")
    assert (tmp_path / "logs" / "validation.txt").exists()
    assert (tmp_path / "logs" / "validation_cn.txt").exists()
    validation_text = (tmp_path / "logs" / "validation.txt").read_text(
        encoding="utf-8"
    )
    assert "## Content Coverage" in validation_text
    assert "- risk_register: present" in validation_text
    assert "- growth_bridge: present" in validation_text
    assert "- evidence_thresholds: present" in validation_text
    assert "- disclosures: present" in validation_text

    lint_result = memo_quality_lint.lint_memo_docx(out_en)
    assert lint_result.has_blocking_findings is False


def test_internal_memo_markdown_renderer_writes_docx(tmp_path):
    md_path = tmp_path / "internal.md"
    out_path = tmp_path / "internal.docx"
    md_path.write_text(
        "\n".join([
            "# Internal Diligence Memo - Generalist",
            "",
            "## Internal Recommendation",
            "Recommendation: Proceed with a staged allocation.",
            "",
            "| Item | View |",
            "|---|---|",
            "| Suggested allocation | $5-10M with valuation sensitivity |",
            "| Conviction | Medium |",
            "",
            "## Internal Diligence Priorities",
            "- Deployment depth shapes allocation size.",
            "- SPV economics shape effective entry.",
            "",
            "Internal use only.",
        ]),
        encoding="utf-8",
    )

    result = internal_memo_renderer.render_internal_memo(md_path, out_path)

    assert result["ok"] is True
    assert out_path.exists()
    text = _all_text(out_path)
    assert "Internal Diligence Memo" in text
    assert "Suggested allocation" in text


def test_generated_renderer_script_detector_flags_old_pattern(tmp_path):
    (tmp_path / "build_memo.py").write_text("# old generated renderer\n", encoding="utf-8")
    (tmp_path / "logs").mkdir()
    (tmp_path / "logs" / "build_memos.js").write_text("// old renderer\n", encoding="utf-8")

    matches = memo_docx_renderer.find_generated_renderer_scripts(tmp_path)

    assert {path.name for path in matches} == {"build_memo.py", "build_memos.js"}


def test_renderer_rejects_missing_required_section():
    package = copy.deepcopy(_package())
    package["sections"] = [
        section
        for section in package["sections"]
        if section["id"] != "investment_highlights"
    ]

    with pytest.raises(memo_docx_renderer.MemoRenderError, match="investment_highlights"):
        memo_docx_renderer.validate_package(package)


def test_renderer_rejects_missing_chinese_translation():
    package = copy.deepcopy(_package())
    package["sections"][0]["blocks"][0]["text"].pop("zh")

    with pytest.raises(memo_docx_renderer.MemoRenderError, match=r"\.zh is required"):
        memo_docx_renderer.validate_package(package)


def test_renderer_accepts_language_neutral_plain_cells():
    package = copy.deepcopy(_package())
    package["sections"][0]["blocks"][2]["rows"].append(
        [
            {"en": "Effective entry valuation", "zh": "有效入场估值"},
            "~$2.55B",
        ]
    )
    package["sections"][0]["blocks"][2]["rows"].append(
        [
            {"en": "Lead investor", "zh": "领投方"},
            "Koch Disruptive Technologies",
        ]
    )
    package["sections"][0]["blocks"][2]["rows"].append(
        [{"en": "Data vintage", "zh": "数据时点"}, "2026-05"]
    )

    memo_docx_renderer.validate_package(package)


def test_renderer_rejects_plain_cells_containing_english_prose():
    package = copy.deepcopy(_package())
    package["sections"][0]["blocks"][2]["rows"].append(
        [
            {"en": "Patents", "zh": "专利"},
            "120+ filed / 90+ issued",
        ]
    )

    with pytest.raises(
        memo_docx_renderer.MemoRenderError, match="must be bilingual"
    ):
        memo_docx_renderer.validate_package(package)


def test_is_language_neutral_text():
    neutral = [
        "~$2.55B",
        "2026-05",
        "+180%",
        "~106x-125x",
        "$500M+",
        "Steve Jurvetson",
        "NextNav (NN)",
        "HERE",
    ]
    prose = [
        "120+ filed / 90+ issued",
        "$10M/site (LOI expected)",
        "~$36M (+~$16M expected)",
        "Not disclosed",
        "该轮次尚未定价",
        "",
    ]
    for value in neutral:
        assert memo_docx_renderer._is_language_neutral_text(value), value
    for value in prose:
        assert not memo_docx_renderer._is_language_neutral_text(value), value


def test_renderer_renders_plain_scalar_cells_in_both_locales(tmp_path):
    package = copy.deepcopy(_package())
    package["sections"][0]["blocks"][2]["rows"].append(
        [{"en": "Valuation", "zh": "估值"}, "~$2.55B"]
    )

    result = memo_docx_renderer.render_memos(
        package,
        out_en=tmp_path / "memo" / "memo_en.docx",
        out_zh=tmp_path / "memo" / "memo_zh.docx",
    )

    assert result["ok"]
    for locale in ("en", "zh"):
        document = Document(result["outputs"][locale])
        cell_texts = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        assert "~$2.55B" in cell_texts


def test_renderer_rejects_unsupported_block_type():
    package = copy.deepcopy(_package())
    package["sections"][0]["blocks"].append({
        "type": "chart",
        "title": {"en": "Unsupported", "zh": "不支持"},
    })

    with pytest.raises(memo_docx_renderer.MemoRenderError, match="unsupported"):
        memo_docx_renderer.validate_package(package)


def test_renderer_rejects_missing_sources():
    package = copy.deepcopy(_package())
    package["sources"] = []

    with pytest.raises(memo_docx_renderer.MemoRenderError, match="sources"):
        memo_docx_renderer.validate_package(package)


def test_renderer_rejects_unknown_section_id_without_visible_title():
    package = copy.deepcopy(_package())
    package["sections"].insert(
        2,
        {
            "id": "commercial_traction",
            "blocks": [
                {
                    "type": "paragraph",
                    "text": {
                        "en": "Commercial traction is summarized here.",
                        "zh": "此处概述商业进展。",
                    },
                }
            ],
        },
    )

    with pytest.raises(memo_docx_renderer.MemoRenderError, match="must provide"):
        memo_docx_renderer.validate_package(package)


def test_renderer_rejects_package_missing_required_content_component():
    package = copy.deepcopy(_package())
    for section in package["sections"]:
        if section["id"] == "investment_highlights":
            section["blocks"] = [
                block
                for block in section["blocks"]
                if block.get("component") != "moat"
            ]

    with pytest.raises(memo_docx_renderer.MemoRenderError, match="moat"):
        memo_docx_renderer.validate_package(package)


def test_renderer_rejects_heading_only_required_section():
    package = copy.deepcopy(_package())
    package["sections"][1]["blocks"] = [
        {
            "type": "heading",
            "level": 2,
            "text": {"en": "Overview", "zh": "概览"},
        }
    ]

    with pytest.raises(memo_docx_renderer.MemoRenderError, match="substantive"):
        memo_docx_renderer.validate_package(package)


def test_renderer_rejects_spacer_only_required_section():
    package = copy.deepcopy(_package())
    package["sections"][1]["blocks"] = [{"type": "spacer"}]

    with pytest.raises(memo_docx_renderer.MemoRenderError, match="substantive"):
        memo_docx_renderer.validate_package(package)


def test_renderer_rejects_table_with_headers_but_no_rows():
    package = copy.deepcopy(_package())
    package["sections"][1]["blocks"] = [
        {
            "type": "table",
            "title": {"en": "Overview Table", "zh": "概览表"},
            "headers": [
                {"en": "Metric", "zh": "指标"},
                {"en": "Treatment", "zh": "处理方式"},
            ],
            "rows": [],
        }
    ]

    with pytest.raises(memo_docx_renderer.MemoRenderError, match="substantive"):
        memo_docx_renderer.validate_package(package)


def test_renderer_accepts_substantive_required_sections():
    memo_docx_renderer.validate_package(_package())


def test_is_numbered_section_heading():
    assert memo_docx_renderer._is_numbered_section_heading(
        "VI. Sources, Source Classes, and Fact Reference Index"
    )
    assert memo_docx_renderer._is_numbered_section_heading("六、资料、来源分类与事实索引")
    assert memo_docx_renderer._is_numbered_section_heading("I. Executive Summary")
    # Legitimate sub-headings are unnumbered and must be preserved.
    assert not memo_docx_renderer._is_numbered_section_heading("Investment Opportunity")
    assert not memo_docx_renderer._is_numbered_section_heading("风险清单")
    assert not memo_docx_renderer._is_numbered_section_heading("")


def test_renderer_drops_numbered_section_heading_restatement(tmp_path):
    # Reproduces the zainar parity failure: the sources section carried a
    # heading block restating its own numbered title. The renderer already
    # emits the section title, so the restatement produced a doubled heading —
    # counted asymmetrically across EN/ZH (en=7, zh=6) and tripping the gate.
    package = copy.deepcopy(_package())
    package["sections"].append(
        {
            "id": "sources",
            "blocks": [
                {
                    "type": "heading",
                    "level": 2,
                    "text": {
                        "en": "VI. Sources, Source Classes, and Fact Reference Index",
                        "zh": "六、资料、来源分类与事实索引",
                    },
                },
                {
                    "type": "paragraph",
                    "text": {
                        "en": "Source index follows.",
                        "zh": "以下为来源索引。",
                    },
                },
            ],
        }
    )
    package_path = tmp_path / "logs" / "memo_package.json"
    package_path.parent.mkdir(parents=True)
    package_path.write_text(json.dumps(package, ensure_ascii=False), encoding="utf-8")
    out_en = tmp_path / "memo" / "Generalist, Inc. - Investment Memo.docx"
    out_zh = tmp_path / "memo" / "Generalist, Inc. - 投资备忘录.docx"

    memo_docx_renderer.render_memos(package_path, out_en=out_en, out_zh=out_zh)

    # Count *section headings* (paragraphs), not TOC table entries. Before the
    # fix the restated heading rendered as a second "sources" section heading in
    # English (en=2) while Chinese counted one (zh=1) — an asymmetry that failed
    # the gate. Both must now resolve to exactly one.
    en_shape = memo_chinese_parity._extract_docx_shape(out_en, "en")
    zh_shape = memo_chinese_parity._extract_docx_shape(out_zh, "zh")
    assert en_shape.section_ids.count("sources") == 1, en_shape.section_ids
    assert zh_shape.section_ids.count("sources") == 1, zh_shape.section_ids
    assert en_shape.section_count == zh_shape.section_count

    result = memo_chinese_parity.lint_chinese_memo_pair(out_en, out_zh)
    assert result.has_blocking_findings is False
    assert not any(
        finding.code == "section_heading_count_mismatch" for finding in result.findings
    )


def test_chinese_parity_passes_for_renderer_output(tmp_path):
    package_path = tmp_path / "logs" / "memo_package.json"
    package_path.parent.mkdir(parents=True)
    package_path.write_text(
        json.dumps(_package(), ensure_ascii=False),
        encoding="utf-8",
    )
    out_en = tmp_path / "memo" / "Generalist, Inc. - Investment Memo.docx"
    out_zh = tmp_path / "memo" / "Generalist, Inc. - 投资备忘录.docx"

    memo_docx_renderer.render_memos(package_path, out_en=out_en, out_zh=out_zh)

    result = memo_chinese_parity.lint_chinese_memo_pair(out_en, out_zh)

    assert result.has_blocking_findings is False


def test_chinese_parity_fails_when_zh_section_missing(tmp_path):
    out_en = tmp_path / "en.docx"
    out_zh = tmp_path / "zh.docx"
    _write_parity_docx(out_en, locale="en")
    _write_parity_docx(out_zh, locale="zh", omit_section="investment_risk")

    result = memo_chinese_parity.lint_chinese_memo_pair(out_en, out_zh)

    assert result.has_blocking_findings is True
    assert any(finding.code == "zh_core_section_missing" for finding in result.findings)


def test_chinese_parity_fails_when_zh_has_no_cjk_body(tmp_path):
    out_en = tmp_path / "en.docx"
    out_zh = tmp_path / "zh.docx"
    _write_parity_docx(out_en, locale="en")
    _write_parity_docx(
        out_zh,
        locale="zh",
        english_only_section="executive_summary",
    )

    result = memo_chinese_parity.lint_chinese_memo_pair(out_en, out_zh)

    assert result.has_blocking_findings is True
    assert any(
        finding.code == "zh_core_section_no_cjk_body"
        for finding in result.findings
    )


def test_chinese_parity_fails_when_table_counts_diverge(tmp_path):
    out_en = tmp_path / "en.docx"
    out_zh = tmp_path / "zh.docx"
    _write_parity_docx(out_en, locale="en")
    _write_parity_docx(out_zh, locale="zh", extra_table=True)

    result = memo_chinese_parity.lint_chinese_memo_pair(out_en, out_zh)

    assert result.has_blocking_findings is True
    assert any(finding.code == "table_count_mismatch" for finding in result.findings)


def _write_parity_docx(
    path,
    *,
    locale: str,
    omit_section: str | None = None,
    english_only_section: str | None = None,
    extra_table: bool = False,
):
    document = Document()
    titles = {
        "en": [
            ("executive_summary", "I. Executive Summary"),
            ("company_overview", "II. Company Overview"),
            ("investment_highlights", "III. Investment Highlights"),
            ("investment_risk", "IV. Investment Risk"),
            (
                "financial_forecast_valuation",
                "V. Financial Forecast & Valuation",
            ),
            (
                "sources",
                "VI. Sources, Source Classes, and Fact Reference Index",
            ),
        ],
        "zh": [
            ("executive_summary", "I. 执行摘要"),
            ("company_overview", "II. 公司概览"),
            ("investment_highlights", "III. 投资亮点"),
            ("investment_risk", "IV. 投资风险"),
            ("financial_forecast_valuation", "V. 财务预测与估值"),
            ("sources", "VI. 来源、来源类别与事实索引"),
        ],
    }
    bodies = {
        "en": "The memo preserves structure, tables, risks, valuation, and sensitivity treatment.",
        "zh": "本备忘录保留相同结构、表格、风险、估值判断和敏感性处理。",
    }
    for section_id, title in titles[locale]:
        if section_id == omit_section:
            continue
        document.add_paragraph(title)
        body_locale = "en" if section_id == english_only_section else locale
        document.add_paragraph(bodies[body_locale])
        if section_id == "executive_summary":
            table_locale = "en" if section_id == english_only_section else locale
            headers = {
                "en": ("Metric", "Treatment"),
                "zh": ("指标", "处理方式"),
            }
            values = {
                "en": ("Revenue", "Use scenario ranges for valuation support."),
                "zh": ("收入", "使用情景区间支持估值判断。"),
            }
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = headers[table_locale][0]
            table.cell(0, 1).text = headers[table_locale][1]
            table.cell(1, 0).text = values[table_locale][0]
            table.cell(1, 1).text = values[table_locale][1]
    if extra_table:
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "额外表格"
    document.save(path)


# ---- Per-risk card format gate (generation-time only) ----


def _risk_section(package: dict) -> dict:
    return next(
        section
        for section in package["sections"]
        if section["id"] == "investment_risk"
    )


def test_valid_risk_cards_pass_generation_gate():
    assert memo_docx_renderer.english_package_validation_errors(_package()) == []


def test_legacy_risk_register_fails_generation_gate_but_still_renders():
    """Old single-table registers must be rejected at generation time (so the
    retry loop reformats them) while render-time validation keeps accepting
    them — packages from runs that predate the card format still re-render."""
    package = copy.deepcopy(_package())
    section = _risk_section(package)
    section["blocks"] = [
        {
            "type": "paragraph",
            "text": {
                "en": "Risk concentrates in commercial proof.",
                "zh": "风险集中在商业验证。",
            },
        },
        {
            "type": "table",
            "component": "risk_register",
            "title": {"en": "Risk Register", "zh": "风险清单"},
            "headers": [
                {"en": "Risk", "zh": "风险"},
                {"en": "Severity", "zh": "严重性"},
                {"en": "Likelihood", "zh": "可能性"},
            ],
            "rows": [
                [
                    {"en": "Services-heavy deployment", "zh": "服务交付占比较高"},
                    {"en": "Medium", "zh": "中"},
                    {"en": "Medium", "zh": "中"},
                ]
            ],
        },
        {
            "type": "callout",
            "component": "disconfirming_evidence",
            "tone": "warning",
            "title": {"en": "Bear-Case Evidence", "zh": "熊市情景证据"},
            "items": [
                {
                    "en": "Revenue disclosure remains limited relative to valuation.",
                    "zh": "相对于估值，收入披露仍然有限。",
                }
            ],
        },
    ]
    errors = memo_docx_renderer.english_package_validation_errors(package)
    assert any("per-risk cards" in error for error in errors)
    # Render-time validation is unchanged.
    memo_docx_renderer.validate_package(package)


def test_risk_card_missing_layout_is_flagged():
    package = copy.deepcopy(_package())
    card = _risk_section(package)["blocks"][2]
    del card["layout"]
    errors = memo_docx_renderer.english_package_validation_errors(package)
    assert any('"layout": "key_value"' in error for error in errors)


def test_risk_card_rating_format_is_enforced():
    package = copy.deepcopy(_package())
    card = _risk_section(package)["blocks"][2]
    card["rows"][4][1] = {"en": "High", "zh": "高"}
    errors = memo_docx_renderer.english_package_validation_errors(package)
    assert any("'N/10: short reason'" in error for error in errors)


def test_risk_card_likelihood_format_is_enforced():
    package = copy.deepcopy(_package())
    card = _risk_section(package)["blocks"][2]
    card["rows"][3][1] = {"en": "7/10: quite likely.", "zh": "7/10：可能性较高。"}
    errors = memo_docx_renderer.english_package_validation_errors(package)
    assert any("'High|Medium|Low: short reason'" in error for error in errors)


def test_risk_card_missing_likelihood_row_is_flagged():
    package = copy.deepcopy(_package())
    card = _risk_section(package)["blocks"][2]
    del card["rows"][3]
    errors = memo_docx_renderer.english_package_validation_errors(package)
    assert any("exactly five two-cell rows" in error for error in errors)


def test_risk_cards_must_be_ordered_by_rating():
    package = copy.deepcopy(_package())
    section = _risk_section(package)
    blocks = section["blocks"]
    # Move the 4/10 card (blocks[5:7]) ahead of the 7/10 card (blocks[1:3]).
    section["blocks"] = (
        [blocks[0]] + blocks[5:7] + blocks[1:3] + blocks[3:5] + blocks[7:]
    )
    errors = memo_docx_renderer.english_package_validation_errors(package)
    assert any("ordered by Risk Rating" in error for error in errors)


def test_key_value_risk_cards_render_label_column(tmp_path):
    package_path = tmp_path / "logs" / "memo_package.json"
    package_path.parent.mkdir(parents=True)
    package_path.write_text(json.dumps(_package()), encoding="utf-8")
    out_en = tmp_path / "memo" / "en.docx"
    out_zh = tmp_path / "memo" / "zh.docx"

    result = memo_docx_renderer.render_memos(
        package_path,
        out_en=out_en,
        out_zh=out_zh,
        manifest_path=tmp_path / "logs" / "run_manifest.md",
        inventory_path=tmp_path / "logs" / "file_inventory.md",
    )
    assert result["ok"] is True

    def _card_tables(path):
        labels = [
            "Risk Type",
            "Why it matters",
            "What we watch",
            "Likelihood",
            "Risk Rating",
        ]
        zh_labels = ["风险类型", "为什么重要", "跟踪信号", "可能性", "风险评分"]
        found = []
        for table in Document(path).tables:
            first_column = [row.cells[0].text.strip() for row in table.rows]
            if first_column in (labels, zh_labels):
                found.append(table)
        return found

    en_cards = _card_tables(out_en)
    assert len(en_cards) == 3
    assert all(len(card.columns) == 2 for card in en_cards)
    en_text = _all_text(out_en)
    assert "Risk 1: Deployments may stay services-heavy and cap margins" in en_text
    assert "7/10: margin path drives the exit multiple." in en_text
    zh_cards = _card_tables(out_zh)
    assert len(zh_cards) == 3
    zh_text = _all_text(out_zh)
    assert "风险 1：部署可能持续偏服务交付，压制利润率" in zh_text
    assert "7/10：利润率路径决定退出倍数。" in zh_text


# ---- Deterministic structural auto-repair (repair_package_structure) ----


def test_repair_derives_missing_callout_title_from_body():
    """The Axiom-run killer: a callout without a title must be repaired by
    deriving a short title from the callout's own content, not fail the run."""
    package = copy.deepcopy(_package())
    package["sections"][0]["blocks"].append(
        {
            "type": "callout",
            "body": {
                "en": (
                    "Valuation support depends on scenario ranges. The "
                    "step-up carries no commercial de-risking."
                ),
                "zh": "估值支撑取决于情景区间。",
            },
            "items": [],
        }
    )
    assert memo_docx_renderer.english_package_validation_errors(package)

    repaired, repairs = memo_docx_renderer.repair_package_structure(package)
    assert any("derived callout title" in r for r in repairs)
    title = repaired["sections"][0]["blocks"][-1]["title"]
    assert title["en"].startswith("Valuation support depends on scenario ranges")
    # The original input is never mutated.
    assert "title" not in package["sections"][0]["blocks"][-1]
    assert not memo_docx_renderer.english_package_validation_errors(repaired)


def test_repair_renames_analysis_pass_source_vocabulary():
    package = copy.deepcopy(_package())
    package["sources"] = [
        {
            "label": {"en": "Company registry", "zh": ""},
            "source_class": "company-reported",
            "detail": {"en": "Registry and launch disclosures.", "zh": ""},
            "as_of": "2026-03-10",
        }
    ]
    repaired, repairs = memo_docx_renderer.repair_package_structure(package)
    source = repaired["sources"][0]
    assert source["title"] == {"en": "Company registry", "zh": ""}
    assert source["class"] == "company-reported"
    assert source["treatment"] == {
        "en": "Registry and launch disclosures.",
        "zh": "",
    }
    assert source["id"] == "S1"
    assert repairs
    assert not memo_docx_renderer.english_package_validation_errors(repaired)


def test_repair_wraps_plain_english_strings_as_bilingual():
    package = copy.deepcopy(_package())
    package["sections"][0]["blocks"].append(
        {
            "type": "bullets",
            "items": ["120+ patents filed and 90+ issued across the portfolio."],
        }
    )
    repaired, repairs = memo_docx_renderer.repair_package_structure(package)
    item = repaired["sections"][0]["blocks"][-1]["items"][0]
    assert item == {
        "en": "120+ patents filed and 90+ issued across the portfolio.",
        "zh": "",
    }
    assert repairs
    assert not memo_docx_renderer.english_package_validation_errors(repaired)


def test_repair_normalizes_block_type_synonyms():
    package = copy.deepcopy(_package())
    package["sections"][0]["blocks"].append(
        {
            "type": "bullet_list",
            "items": [{"en": "A substantive bullet point.", "zh": "要点。"}],
        }
    )
    repaired, repairs = memo_docx_renderer.repair_package_structure(package)
    assert repaired["sections"][0]["blocks"][-1]["type"] == "bullets"
    assert any("normalized" in r for r in repairs)


def test_repair_is_noop_on_valid_package():
    package = copy.deepcopy(_package())
    repaired, repairs = memo_docx_renderer.repair_package_structure(package)
    assert repairs == []
    assert repaired == package
