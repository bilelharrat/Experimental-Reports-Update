"""The Chinese IC decision memo renders the placeholders and owner roles
in Chinese (live 2026-09-23 they printed as English inside the Chinese)."""
from __future__ import annotations

from server import internal_memo_renderer


def test_placeholders_and_owner_roles_are_localized():
    text = "\n".join(
        [
            "出资金额：[TO BE DETERMINED BY IC]。",
            "| 事项 | 重要性 | 负责人 | 截止日期 |",
            "|---|---|---|---|",
            "| 专利审查 | 决定护城河 | legal | [date to set] |",
            "| 渠道 | 分销 | deal lead | 2027-03-31 |",
            "| Arm 条款 | 版税 | [owner to assign] | [date to set] |",
        ]
    )
    out = internal_memo_renderer.localize_zh_placeholders(text)
    assert "TO BE DETERMINED" not in out and "date to set" not in out and "owner to assign" not in out
    assert "【待投委会确定】" in out and "【日期待定】" in out and "【负责人待定】" in out
    assert "| 法务 |" in out and "| 交易负责人 |" in out
    # the separator row and ordinary cells are untouched
    assert "|---|---|---|---|" in out and "专利审查" in out


def test_english_ic_memo_is_left_alone(tmp_path):
    md = tmp_path / "ic.md"
    body = "# The call\n" + ("Pass. " * 60) + "\nAmount: [TO BE DETERMINED BY IC].\n"
    md.write_text(body, encoding="utf-8")
    internal_memo_renderer.render_internal_memo(md, tmp_path / "ic.docx", locale="en")
    import docx

    text = "\n".join(p.text for p in docx.Document(tmp_path / "ic.docx").paragraphs)
    assert "[TO BE DETERMINED BY IC]" in text


def test_english_header_rows_of_the_two_ic_tables_are_localized():
    text = "\n".join(
        [
            "| Item | Why it matters | Owner | Due |",
            "|---|---|---|---|",
            "| 专利 | 护城河 | legal | [date to set] |",
            "",
            "| Figure | Value | Source |",
            "|---|---|---|",
            "| Value | $1.0B | [S1] |",
        ]
    )
    out = internal_memo_renderer.localize_zh_placeholders(text).splitlines()
    assert out[0] == "| 事项 | 为什么重要 | 负责人 | 截止时间 |"
    assert out[4] == "| 数字 | 取值 | 来源 |"
    # A data row whose first cell reads "Value" is not a header row.
    assert out[6] == "| Value | $1.0B | [S1] |"


def test_the_ic_memo_opts_in_to_the_lp_money_form_and_the_buffett_memo_does_not():
    line = "最高可接受估值约为 3.93 亿美元，相当于 [TO BE DETERMINED BY IC]。"
    assert internal_memo_renderer.localize_zh_placeholders(line, money_form=True) == (
        "最高可接受估值约为 $393M，相当于 【待投委会确定】。"
    )
    # The default (the Buffett memo's path) keeps its own Chinese money form.
    assert internal_memo_renderer.localize_zh_placeholders(line) == "最高可接受估值约为 3.93 亿美元，相当于 【待投委会确定】。"
