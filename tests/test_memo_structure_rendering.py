"""Renderer-owned document structure (R27, R5, R13, R21, R25, R33, G2, G8).

Real Word headings with stable bookmarks and a linked contents list;
renderer-owned numbering (I. II. / 一、二、) with extra sections after the
core ones and ONE sources section, numbered as the profile says; page one
= cover facts + contents (no decision summary); honest cover dates; review
stamp, footer and file metadata; fitted tables; the risk summary; the
sources section with tiers and stacked dates; the one-language render; and
the legacy ``strict_sources=False`` re-render mode.
"""
from __future__ import annotations

import copy
import json

import pytest
from docx import Document
from docx.oxml.ns import qn

import test_memo_docx_renderer as base
from server import (
    memo_chinese_parity,
    memo_docx_renderer,
    memo_quality_lint,
    memo_structure,
)


def _render(tmp_path, package, **kwargs):
    out_en = tmp_path / "memo" / "en.docx"
    out_zh = tmp_path / "memo" / "zh.docx"
    result = memo_docx_renderer.render_memos(
        package, out_en=out_en, out_zh=out_zh, **kwargs
    )
    return result, Document(out_en), Document(out_zh)


def _headings(doc, level=1):
    return [p.text for p in doc.paragraphs if p.style.name == f"Heading {level}"]


def _bookmarks(doc, prefix="bsh_"):
    return [
        b.get(qn("w:name"))
        for b in doc.element.body.iter(qn("w:bookmarkStart"))
        if b.get(qn("w:name")).startswith(prefix)
    ]


def _text(doc):
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _with_extras(package):
    """A package with the extra sections late v1 runs add: a numbered model
    decision section, a sources-like section and disclosures — plus one
    extra placed BETWEEN two core sections."""
    package = copy.deepcopy(package)
    sections = package["sections"]
    sections.insert(
        2,
        {
            "id": "market_notes",
            "title": {"en": "Market Notes", "zh": "市场补充"},
            "blocks": [
                {
                    "type": "paragraph",
                    "text": {
                        "en": "Industrial automation budgets are moving to software-defined cells.",
                        "zh": "工业自动化预算正在转向软件定义的产线单元。",
                    },
                }
            ],
        },
    )
    sections += [
        {
            "id": "investment_decision",
            "title": {"en": "VI. Investment Decision", "zh": "六、投资决策"},
            "blocks": [
                {
                    "type": "paragraph",
                    "text": {
                        "en": "We recommend a measured allocation. The price is supported by deployments.",
                        "zh": "我们建议适度配置。价格有部署证据支撑。",
                    },
                }
            ],
        },
        {
            "id": "source_index",
            "title": {"en": "Source Classes and Fact Index", "zh": "资料类别与事实索引"},
            "blocks": [
                {
                    "type": "heading",
                    "level": 2,
                    "text": {"en": "Sources", "zh": "来源"},
                },
                {
                    "type": "paragraph",
                    "text": {
                        "en": "Company materials are weighed as claims, press as corroboration.",
                        "zh": "公司材料按公司口径采信，媒体报道作为佐证。",
                    },
                },
            ],
        },
        {
            "id": "disclosures",
            "title": {"en": "Disclosures", "zh": "披露声明"},
            "blocks": [
                {
                    "type": "paragraph",
                    "text": {
                        "en": "Not an offer to sell securities; any investment may result in partial or total loss.",
                        "zh": "本文件并非证券出售要约；任何投资均可能产生部分或全部损失。",
                    },
                }
            ],
        },
    ]
    return package


# ---- structure: headings, numbering, bookmarks, contents --------------------


def test_extra_sections_are_numbered_after_the_core_and_the_gates_still_find_them(tmp_path):
    package = _with_extras(base._package())
    _result, en, zh = _render(tmp_path, package)
    assert _headings(en) == [
        "I. Executive Summary",
        "II. Company Overview",
        "III. Investment Highlights",
        "IV. Investment Risk",
        "V. Financial Forecast & Valuation",
        "VI. Market Notes",
        "VII. Investment Decision",
        "VIII. Disclosures",
        # the profile numbers its sources back matter, after the extras
        "IX. Sources, Source Classes, and Fact Reference Index",
    ]
    assert _headings(zh) == [
        "一、执行摘要",
        "二、公司概览",
        "三、投资亮点",
        "四、投资风险",
        "五、财务预测与估值",
        "六、市场补充",
        "七、投资决策",
        "八、披露声明",
        "九、来源、来源类别与事实索引",
    ]
    # the sources-like section folded into the ONE sources section (its restated
    # "Sources" heading dropped), its prose kept
    assert "Company materials are weighed as claims" in _text(en)
    assert "Source Classes and Fact Index" not in "\n".join(_headings(en, 1) + _headings(en, 2))

    en_shape = memo_chinese_parity._extract_docx_shape(tmp_path / "memo" / "en.docx", "en")
    zh_shape = memo_chinese_parity._extract_docx_shape(tmp_path / "memo" / "zh.docx", "zh")
    for shape in (en_shape, zh_shape):
        assert shape.section_ids.count("sources") == 1, shape.section_ids
        assert "investment_risk" in shape.section_ids
    parity = memo_chinese_parity.lint_chinese_memo_pair(
        tmp_path / "memo" / "en.docx", tmp_path / "memo" / "zh.docx"
    )
    assert not parity.p0_findings, [f.to_dict() for f in parity.p0_findings]

    blocks = memo_quality_lint._extract_docx_blocks(tmp_path / "memo" / "en.docx")
    sections = {block.section for block in blocks}
    assert memo_structure.LATE.numbered_lint_key("risk") in sections
    assert any(block.allowed_trace_section for block in blocks if "sources" in block.section)
    lint = memo_quality_lint.lint_memo_docx(tmp_path / "memo" / "en.docx")
    assert not lint.p0_findings, [f.to_dict() for f in lint.p0_findings]


def test_headings_are_word_headings_with_stable_bookmarks_and_linked_contents(tmp_path):
    package = _with_extras(base._package())
    _result, en, zh = _render(tmp_path, package)
    heading1 = en.styles["Heading 1"].element
    assert heading1.find(qn("w:pPr")).find(qn("w:outlineLvl")).get(qn("w:val")) == "0"
    fonts = heading1.find(qn("w:rPr")).find(qn("w:rFonts"))
    assert fonts.get(qn("w:ascii")) == "Arial"
    assert fonts.get(qn("w:asciiTheme")) is None
    assert not any(h.isupper() for h in _headings(en))  # no .upper()
    sections = _bookmarks(en, "bsh_sec_")
    assert sections == [f"bsh_sec_{n}" for n in range(1, 10)]
    assert _bookmarks(en) == _bookmarks(zh)
    for paragraph in en.paragraphs:
        if paragraph.style.name == "Heading 1":
            names = [b.get(qn("w:name")) for b in paragraph._p.iter(qn("w:bookmarkStart"))]
            assert names and names[0].startswith("bsh_sec_")
            # the bookmark follows pPr, which stays the paragraph's first child
            assert paragraph._p[0].tag == qn("w:pPr")
    anchors = [
        h.get(qn("w:anchor"))
        for h in en.element.body.iter(qn("w:hyperlink"))
        if (h.get(qn("w:anchor")) or "").startswith("bsh_sec_")
    ]
    assert anchors == sections  # one contents link per level-1 heading
    # no fake page-number column and no Word TOC field
    assert "TOC" not in en.element.body.xml
    assert _bookmarks(en, "bsh_sub_")  # level-2 headings carry bsh_sub_<n>_<m>
    assert memo_docx_renderer.section_bookmark(3) == "bsh_sec_3"
    assert memo_docx_renderer.subsection_bookmark(3, 2) == "bsh_sub_3_2"


def test_toc_titles_are_the_bare_titles_in_document_order():
    package = _with_extras(base._package())
    titles = memo_docx_renderer._toc_titles(package, "en")
    assert titles[:6] == [
        "Executive Summary",
        "Company Overview",
        "Investment Highlights",
        "Investment Risk",
        "Financial Forecast & Valuation",
        "Market Notes",
    ]
    assert titles[-1] == "Sources, Source Classes, and Fact Reference Index"


# ---- page one: cover facts and contents --------------------------------------


def _body_order(doc):
    """(kind, element) for every body child, in order."""
    out = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            out.append(("page_break" if child.findall(".//" + qn("w:br")) and any(
                br.get(qn("w:type")) == "page" for br in child.iter(qn("w:br"))
            ) else "p", child))
        elif child.tag == qn("w:tbl"):
            out.append(("tbl", child))
    return out


def test_page_one_is_the_cover_facts_and_the_contents(tmp_path):
    package = _with_extras(base._package())
    _result, en, zh = _render(tmp_path, package)
    assert not hasattr(memo_docx_renderer, "decision_summary")
    for doc, masthead, contents_title in (
        (en, "BERKELEY SUMMIT HOUSE  ·  Confidential Investment Memo", "Contents"),
        (zh, "BERKELEY SUMMIT HOUSE  ·  机密投资备忘录", "目录"),
    ):
        order = _body_order(doc)
        first_break = next(i for i, (kind, _el) in enumerate(order) if kind == "page_break")
        page_one = order[:first_break]
        texts = ["".join(t.text or "" for t in el.iter(qn("w:t"))) for kind, el in page_one if kind == "p"]
        assert texts[0] == masthead
        assert texts[1] == "Generalist, Inc."
        assert any(text.startswith(("Written ", "撰写于 ")) for text in texts)
        assert contents_title in texts
        tables = [el for kind, el in page_one if kind == "tbl"]
        # the cover facts table, then the contents list — no decision summary
        assert len(tables) == 2
        anchors = [h.get(qn("w:anchor")) for h in tables[1].iter(qn("w:hyperlink"))]
        assert anchors == [f"bsh_sec_{n}" for n in range(1, 10)]
        # two columns: the section numeral and the linked title, no page numbers
        grid = tables[1].find(qn("w:tblGrid"))
        assert len(list(grid)) == 2
        all_text = _text(doc)
        for word in ("Decision summary", "决策摘要", "Top risks", "主要风险", "What moves the view"):
            assert word not in all_text
        # the first section opens page two
        heading = order[first_break + 1][1]
        assert heading.find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val")) == "Heading1"
    contents = [
        c.text for c in en.tables[1].columns[0].cells
    ]
    assert contents[:2] == ["I.", "II."] and contents[-1] == "IX."
    assert [c.text for c in zh.tables[1].columns[0].cells][-1] == "九、"


def test_run_checks_line_opens_the_sources_section_not_page_one(tmp_path):
    package = base._package()
    package["run"]["checks"] = {
        "fact_check": {"verified": 30, "found_elsewhere": 10, "derived": 2, "not_traced": 8, "thin_corpus": False},
        "gates": {"quality": "passed", "chinese_parity": "warnings", "fact_check": "passed"},
    }
    package["decision"] = {"evidence_cutoff": "2026-01-01", "recommendation": {"en": "Buy", "zh": "买入"}}
    _result, en, zh = _render(tmp_path, package)
    for doc, sources_heading, line in (
        (
            en,
            "VI. Sources, Source Classes, and Fact Reference Index",
            "50 figures: 30 verified · 10 found elsewhere · 8 not traced · 2 derived"
            " · Checks: quality passed · Chinese parity warnings · fact check passed",
        ),
        (
            zh,
            "六、来源、来源类别与事实索引",
            "共 50 个数字：30 个已核实 · 10 个在其他来源找到 · 8 个未能溯源 · 2 个为推算"
            " · 检查：质量检查 通过 · 中文一致性 有警告 · 事实核查 通过",
        ),
    ):
        paragraphs = [p for p in doc.paragraphs if p.text.strip()]
        index = next(i for i, p in enumerate(paragraphs) if p.text == sources_heading)
        assert paragraphs[index + 1].text == line
        assert paragraphs[index + 1].style.name == memo_structure.DERIVED_PARAGRAPH_STYLE
        assert [p.text for p in paragraphs].count(line) == 1
    # the pipeline no longer stamps a decision block; a stray one is ignored
    en_text = _text(en)
    assert "Evidence through 2026-01-01" not in en_text
    assert "Buy" not in [p.text for p in en.paragraphs]


def test_cover_dates_never_print_a_run_id(tmp_path):
    package = base._package()
    package["run"] = {"run_id": "2026-08-31__224828"}
    package["sources"][0]["as_of"] = "2026-06-13"
    _result, en, zh = _render(tmp_path, package)
    en_text, zh_text = _text(en), _text(zh)
    assert "2026-08-31__224828" not in en_text + zh_text
    assert "Written 2026-08-31 · Evidence through 2026-06-13" in en_text
    assert "撰写于 2026年8月31日 · 证据截至 2026年6月13日" in zh_text
    assert en.core_properties.created.date().isoformat() == "2026-08-31"


def test_the_stated_evidence_ceiling_wins_unless_it_repeats_the_run_date():
    package = base._package()
    package["run"] = {"run_id": "x", "as_of": "2026-06-22", "evidence_ceiling": "2026-06-13"}
    assert memo_docx_renderer._evidence_through(package, "2026-06-22") == "2026-06-13"
    package["run"]["evidence_cutoff"] = "2026-06-22"  # the run date: derive instead
    package["sources"][0]["as_of"] = "2026-05-02"
    assert memo_docx_renderer._evidence_through(package, "2026-06-22") == "2026-05-02"


@pytest.mark.parametrize(
    "company_update, expected_en, expected_zh",
    [
        ({"round": "Series D at $2.5B"}, "Series D at $2.5B", "Series D at $2.5B"),
        # the package SAYS there is none: private company -> "No round on offer"
        ({"round": ""}, "No round on offer", "暂无在售轮次"),
        ({"round": None}, "No round on offer", "暂无在售轮次"),
        # absent = unknown (a legacy package): no row rather than a guess
        ({"round": "__absent__"}, None, None),
        # public companies never print the row
        ({"round": "", "ticker": "ACME"}, None, None),
    ],
)
def test_round_row(company_update, expected_en, expected_zh):
    package = base._package()
    company = package["company"]
    if company_update.get("round") == "__absent__":
        company.pop("round")
    else:
        company.update(company_update)
    rows_en = dict(memo_docx_renderer._cover_fact_rows(package, "en"))
    rows_zh = dict(memo_docx_renderer._cover_fact_rows(package, "zh"))
    assert rows_en.get("Round") == expected_en
    assert rows_zh.get("轮次") == expected_zh


def test_location_falls_back_to_hq():
    package = base._package()
    package["company"].pop("location")
    package["company"]["hq"] = {"en": "Belmont, California", "zh": "加利福尼亚州贝尔蒙特"}
    assert dict(memo_docx_renderer._cover_fact_rows(package, "zh"))["地点"] == "加利福尼亚州贝尔蒙特"


# ---- header stamp, footer, file metadata ------------------------------------


@pytest.mark.parametrize(
    "review, en_stamp, zh_stamp, status",
    [
        (None, "DRAFT — AI-generated, not reviewed", "草稿 — AI 生成，未经审阅", "Draft"),
        ({"state": "in_review"}, "DRAFT — AI-generated, not reviewed", "草稿 — AI 生成，未经审阅", "Draft"),
        (
            {"state": "approved", "reviewer": "J. Chen", "reviewed_at": "2026-09-20T10:00:00Z"},
            "Reviewed by J. Chen, 2026-09-20",
            "已审阅：J. Chen，2026年9月20日",
            "Reviewed",
        ),
        ({"state": "withdrawn"}, "WITHDRAWN — do not rely on this memo", "已撤回 — 请勿依据本备忘录", "Withdrawn"),
    ],
)
def test_review_stamp_sits_in_the_headers_only(tmp_path, review, en_stamp, zh_stamp, status):
    package = base._package()
    if review:
        package["run"]["review"] = review
    _result, en, zh = _render(tmp_path, package)
    for doc, stamp in ((en, en_stamp), (zh, zh_stamp)):
        section = doc.sections[0]
        assert section.first_page_header.paragraphs[0].text == stamp
        assert section.header.paragraphs[0].text.startswith(stamp)
        assert stamp not in _text(doc)  # never a body paragraph
    assert en.core_properties.content_status == status


def test_footer_carries_page_x_of_y_fields_and_the_confidentiality_line(tmp_path):
    _result, en, zh = _render(tmp_path, base._package())
    for doc, words, line in (
        (en, ("Page ", " of "), "BSH Confidential · 2026-06-22 · EN"),
        (zh, ("第 ", " 页，共 "), "BSH 机密 · 2026年6月22日 · 中文"),
    ):
        for footer in (doc.sections[0].footer, doc.sections[0].first_page_footer):
            paragraph = footer.paragraphs[0]
            xml = paragraph._p.xml
            instructions = [i.text.strip() for i in paragraph._p.iter(qn("w:instrText"))]
            assert instructions == ["PAGE", "NUMPAGES"]
            kinds = [f.get(qn("w:fldCharType")) for f in paragraph._p.iter(qn("w:fldChar"))]
            assert kinds == ["begin", "separate", "end", "begin", "separate", "end"]
            assert all(word in paragraph.text for word in words)
            assert paragraph.text.endswith(line)
            # the literal words share the run with the field characters, so a
            # viewer that drops field runs prints no stray "Page  of "
            field_run = paragraph.runs[0]._r
            assert field_run.find(qn("w:fldChar")) is not None and words[0] in xml


def test_file_metadata_names_the_memo_and_the_renderer(tmp_path):
    _result, en, zh = _render(tmp_path, base._package())
    for doc, title, language in (
        (en, "Generalist, Inc. — Investment Memo", "en-US"),
        (zh, "Generalist, Inc. — 投资备忘录", "zh-CN"),
    ):
        props = doc.core_properties
        assert props.title == title
        assert props.author == props.last_modified_by == "Berkeley Summit House"
        assert props.category == "Confidential"
        assert props.created.date().isoformat() == "2026-06-22"
        assert props.revision == 1
        assert props.identifier == props.version == memo_docx_renderer.RENDERER_VERSION
        assert props.language == language
        assert "python-docx" not in (props.comments or "")


def test_bilingual_company_name_names_the_chinese_file(tmp_path):
    package = base._package()
    package["company"]["name"] = {"en": "TSMC", "zh": "台积电（TSMC）"}
    _result, en, zh = _render(tmp_path, package)
    assert zh.core_properties.title == "台积电（TSMC） — 投资备忘录"
    assert "台积电（TSMC） | BSH 机密投资备忘录" in zh.sections[0].header.paragraphs[0].text
    assert en.core_properties.title == "TSMC — Investment Memo"


# ---- Chinese document hygiene -------------------------------------------------


def test_chinese_document_is_tagged_zh_cn_and_english_is_left_alone(tmp_path):
    _result, en, zh = _render(tmp_path, base._package())

    def langs(doc):
        rpr = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
        theme = doc.settings.element.find(qn("w:themeFontLang"))
        return rpr.find(qn("w:lang")).get(qn("w:eastAsia")), theme.get(qn("w:eastAsia"))

    assert langs(zh) == ("zh-CN", "zh-CN")
    assert langs(en) == ("en-US", "ja-JP")  # python-docx defaults, untouched


def test_silent_english_fallbacks_are_reported_in_validation_cn(tmp_path):
    # A blank zh the validator does not police (the cover descriptor) goes
    # all the way through render_memos into validation_cn.txt.
    package = base._package()
    package["company"]["descriptor"]["zh"] = ""
    memo_docx_renderer.render_memos(
        package,
        out_en=tmp_path / "memo" / "en.docx",
        out_zh=tmp_path / "memo" / "zh.docx",
    )
    report = (tmp_path / "logs" / "validation_cn.txt").read_text(encoding="utf-8")
    assert "- zh_fallback_count: 1" in report
    assert "- P1 zh_blank_translation · front_matter: “Embodied AI systems" in report
    assert report.rstrip().endswith("- status: passed")

    # A missing zh key (which validation would reject) prints English; the
    # builder records it against the section it sits in.
    package = base._package()
    block = package["sections"][1]["blocks"][0]
    block["text"] = {"en": block["text"]["en"]}
    sink: list = []
    token = memo_docx_renderer._ZH_FALLBACKS.set(sink)
    try:
        memo_docx_renderer._build_document(package, "zh")
    finally:
        memo_docx_renderer._ZH_FALLBACKS.reset(token)
    assert [(section, kind) for section, kind, _en in sink] == [
        ("company_overview", "english_printed")
    ]
    # the English render collects nothing
    memo_docx_renderer._build_document(package, "en")
    assert memo_docx_renderer._ZH_FALLBACKS.get() is None


def test_a_clean_render_reports_zero_fallbacks(tmp_path):
    memo_docx_renderer.render_memos(
        base._package(),
        out_en=tmp_path / "memo" / "en.docx",
        out_zh=tmp_path / "memo" / "zh.docx",
    )
    report = (tmp_path / "logs" / "validation_cn.txt").read_text(encoding="utf-8")
    assert "- zh_fallback_count: 0" in report
    assert "## Chinese fallbacks" not in (tmp_path / "logs" / "validation.txt").read_text(encoding="utf-8")


# ---- blocks: markdown bold, callouts, bullets, risk cards -------------------


def test_markdown_bold_renders_as_bold_runs_and_never_as_asterisks():
    document = Document()
    memo_docx_renderer._add_block(
        document,
        {"type": "paragraph", "text": {"en": "**Decision.** BSH passes at $2T and a **hard** cap.", "zh": ""}},
        "en",
    )
    memo_docx_renderer._add_block(
        document,
        {"type": "bullets", "component": "investment_decision", "items": [{"en": "**Core bet.** Coding stays.", "zh": ""}]},
        "en",
    )
    paragraph, bullet = document.paragraphs[-2], document.paragraphs[-1]
    assert "**" not in paragraph.text + bullet.text
    runs = [(run.text, bool(run.bold)) for run in paragraph.runs]
    assert runs[0] == ("Decision.", True)
    assert ("hard", True) in runs
    assert all(not bold for text, bold in runs if text not in {"Decision.", "hard"})
    assert (bullet.runs[0].text, bool(bullet.runs[0].bold)) == ("Core bet.", True)
    assert bullet._p.pPr.numPr is not None  # a real bullet


def test_callout_title_opens_the_cell_without_a_blank_line():
    document = Document()
    memo_docx_renderer._add_block(
        document,
        {
            "type": "callout",
            "title": {"en": "Valuation Sensitivity", "zh": "估值敏感性"},
            "items": [{"en": "Margin path drives sizing.", "zh": "利润率路径决定规模。"}],
        },
        "en",
    )
    cell = document.tables[-1].cell(0, 0)
    assert cell.paragraphs[0].text == "Valuation Sensitivity"
    assert cell.paragraphs[1]._p.pPr.numPr is not None


def test_risk_cards_drop_the_repeated_title_and_bold_only_the_token(tmp_path):
    package = base._package()
    risk = next(s for s in package["sections"] if s["id"] == "investment_risk")
    risk["blocks"][2]["title"] = {"en": "Risk 1", "zh": "风险 1"}
    _result, en, zh = _render(tmp_path, package)
    for doc in (en, zh):
        assert not any(p.text.strip() in {"Risk 1", "风险 1"} for p in doc.paragraphs)
    likelihood = next(
        row.cells[1]
        for table in en.tables
        for row in table.rows
        if row.cells[0].text == "Likelihood"
    )
    runs = likelihood.paragraphs[0].runs
    assert (runs[0].text, bool(runs[0].bold)) == ("Medium:", True)
    assert not runs[1].bold


def test_risk_summary_table_leads_the_cards_in_both_languages(tmp_path):
    _result, en, zh = _render(tmp_path, base._package())
    assert "Risk summary" in _text(en) and "风险一览" in _text(zh)
    summary = next(t for t in en.tables if t.rows[0].cells[0].text == "#" and len(t.columns) == 5)
    assert [c.text for c in summary.rows[0].cells] == ["#", "Risk", "Area", "Likelihood", "Rating"]
    assert [c.text for c in summary.rows[1].cells] == [
        "1",
        "Deployments may stay services-heavy and cap margins",
        "Commercial",
        "Medium",
        "7/10",
    ]
    assert len(en.tables) == len(zh.tables)


def test_risk_summary_handles_the_eight_row_v2_cards():
    labels = [label for _p, label in memo_docx_renderer._RISK_CARD_ROW_LABELS_V2]
    rows = [[label, {"en": f"High: {label} value", "zh": f"高：{label}"}] for label in labels]
    rows[-1][1] = {"en": "9/10: the price", "zh": "9/10：价格"}
    section = {
        "id": "investment_risk",
        "blocks": [
            {"type": "heading", "level": 3, "text": {"en": "Risk 1: The price is the risk", "zh": "风险 1：价格即风险"}},
            {"type": "table", "layout": "key_value", "headers": [], "rows": rows},
        ],
    }
    card = memo_docx_renderer._risk_cards(section)[0]
    assert memo_docx_renderer._rating_token(memo_docx_renderer._card_row(card, "risk rating")) == "9/10"
    assert memo_docx_renderer._likelihood_token(memo_docx_renderer._card_row(card, "likelihood"), "zh") == "高"
    assert memo_docx_renderer._card_title(card) == {"en": "The price is the risk", "zh": "价格即风险"}


def test_no_risk_summary_for_a_bullets_format_profile():
    from dataclasses import replace

    package = base._package()
    risk = next(s for s in package["sections"] if s["id"] == "investment_risk")
    bullets = replace(memo_structure.LATE, risk_format="bullets")

    def summary_titles(structure):
        document = Document()
        memo_docx_renderer._ensure_derived_styles(document)
        context = memo_docx_renderer._RenderContext(
            package=package, structure=structure, locale="en"
        )
        memo_docx_renderer._add_section_blocks(document, risk, context)
        return [p.text for p in document.paragraphs if p.text == "Risk summary"]

    assert summary_titles(bullets) == []
    assert summary_titles(memo_structure.LATE) == ["Risk summary"]


# ---- tables ---------------------------------------------------------------------


def test_tables_have_fitted_fixed_widths_header_rows_and_right_aligned_numbers(tmp_path):
    package = base._package()
    snapshot = package["sections"][0]["blocks"][2]
    snapshot["headers"] = [
        {"en": "Metric", "zh": "指标"},
        {"en": "Figure", "zh": "数值"},
        {"en": "Note", "zh": "说明"},
    ]
    snapshot["rows"] = [
        [{"en": "ARR", "zh": "ARR"}, "~$24M", {"en": "Company-reported; medium confidence in the growth rate.", "zh": "公司自述；增速置信度中等。"}],
        [{"en": "Valuation", "zh": "估值"}, "$1.0B+", {"en": "Dated February mark, not a current price.", "zh": "2 月的估值标记，并非当前价格。"}],
    ]
    _result, en, _zh = _render(tmp_path, package)
    section = en.sections[0]
    text_width = int((section.page_width - section.left_margin - section.right_margin) / 635)
    for table in en.tables:
        tbl = table._tbl
        grid = [int(col.get(qn("w:w"))) for col in tbl.find(qn("w:tblGrid"))]
        assert tbl.tblPr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"
        assert int(tbl.tblPr.find(qn("w:tblW")).get(qn("w:w"))) == sum(grid)
        assert abs(sum(grid) - text_width) <= 2
        for row in table.rows:
            for col, cell in enumerate(row.cells):
                assert int(cell._tc.tcPr.find(qn("w:tcW")).get(qn("w:w"))) == grid[col]
    metrics = next(t for t in en.tables if [c.text for c in t.rows[0].cells] == ["Metric", "Figure", "Note"])
    assert metrics.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    grid = [int(col.get(qn("w:w"))) for col in metrics._tbl.find(qn("w:tblGrid"))]
    assert grid[2] > grid[0] and grid[2] > grid[1]  # the prose column is the wide one
    figure_cell = metrics.rows[1].cells[1]
    assert figure_cell.paragraphs[0].alignment == 2  # WD_ALIGN_PARAGRAPH.RIGHT
    assert metrics.rows[1]._tr.trPr.find(qn("w:cantSplit")) is not None


def test_width_planner_keeps_numbers_whole_and_sums_to_the_page():
    widths = memo_docx_renderer._plan_column_widths(
        ["Scenario", "EV", "Narrative"],
        [["Bear", "$320M–$480M", "Conversion stays under 10% and growth decelerates to about 60% as reference customers stall."]],
        9746,
    )
    assert sum(widths) == 9746
    assert widths[1] >= len("$320M–$480M") * memo_docx_renderer._CHAR_TWIPS
    assert widths[2] == max(widths)


# ---- sources appendix -------------------------------------------------------------


def test_sources_appendix_has_ids_links_domains_tiers_dates_and_labels(tmp_path):
    package = base._package()
    package["company"]["website"] = "https://www.generalist.ai"
    package["sources"] = [
        {
            "id": "S1",
            "title": {"en": "CNBC — Generalist raises Series D", "zh": "CNBC —— Generalist 完成 D 轮融资"},
            "class": "business_press",
            "treatment": {"en": "Round size and lead.", "zh": "融资规模与领投方。"},
            "as_of": "2026-06-22",
            "url": "https://www.cnbc.com/2026/06/17/generalist-series-d.html",
            "data_period": "2026-05",
            "retrieved_at": "2026-06-21",
        },
        {
            "id": "S2",
            "title": "Generalist newsroom",
            "class": "third_party_market_data",
            "treatment": {"en": "Deployment count.", "zh": "部署数量。"},
            "published_at": "2026-06",
            "url": "https://news.generalist.ai/deployments",
        },
        {
            "id": "S3",
            "title": "Some blog",
            "class": "weird_new_class",
            "treatment": {"en": "Context only.", "zh": "仅作背景。"},
            "as_of": "undated",
            "url": "https://someblog.example/post",
        },
    ]
    _result, en, zh = _render(tmp_path, package)
    table = next(t for t in en.tables if t.rows[0].cells[0].text == "ID")
    assert [c.text for c in table.rows[0].cells] == ["ID", "Source", "Class", "Tier", "Treatment", "Dates"]
    s1 = [c.text for c in table.rows[1].cells]
    assert s1[0] == "S1"
    assert s1[1] == "CNBC — Generalist raises Series D\ncnbc.com"
    assert s1[2] == "Business press"
    assert s1[3] == "B"
    # the slug date beats the run-day as_of; period and retrieval stack in one cell
    assert s1[5] == "17 Jun 2026 · covers May 2026 · retrieved 21 Jun 2026"
    s2 = [c.text for c in table.rows[2].cells]
    assert s2[2] == "Third-party market data" and s2[3] == "A-" and s2[5] == "Jun 2026"
    s3 = [c.text for c in table.rows[3].cells]
    assert s3[2] == "Weird New Class" and s3[3] == "C" and s3[5] == "Undated"
    assert 'w:name="src_S1"' in table._tbl.xml
    zh_table = next(t for t in zh.tables if t.rows[0].cells[0].text == "编号")
    assert [c.text for c in zh_table.rows[0].cells] == ["编号", "来源", "类别", "等级", "处理方式", "日期"]
    assert zh_table.rows[1].cells[5].text == "2026年6月17日发布 · 覆盖2026年5月 · 2026年6月21日检索"
    assert zh_table.rows[2].cells[2].text == "第三方市场数据"
    assert zh_table.rows[3].cells[2].text == "其他来源"  # never a raw slug
    assert "Tier: A = regulator" in _text(en) and "来源等级：A = " in _text(zh)


def test_a_key_metric_resting_only_on_tier_c_is_marked(tmp_path):
    package = base._package()
    package["sources"].append(
        {
            "id": "S2",
            "title": "Aggregator write-up",
            "class": "aggregator",
            "treatment": {"en": "Low weight.", "zh": "低权重。"},
            "as_of": "2026-05",
            "url": "https://buildmvpfast.com/blog/generalist",
        }
    )
    snapshot = package["sections"][0]["blocks"][2]
    snapshot["rows"][0][1] = {
        "en": "Gross margin is about 40% [S2].",
        "zh": "毛利率约为 40% [S2]。",
    }
    _result, en, zh = _render(tmp_path, package)
    for doc, header, footnote in (
        (en, ["Metric", "Treatment"], "† Single third-party estimate, not corroborated by a filing or the company."),
        (zh, ["指标", "处理方式"], "† 单一第三方估算，未经监管申报文件或公司披露佐证。"),
    ):
        snapshot_table = next(t for t in doc.tables if [c.text for c in t.rows[0].cells] == header)
        assert snapshot_table.rows[1].cells[0].text.endswith(" †")
        assert not snapshot_table.rows[2].cells[0].text.endswith("†")
        # the footnote sits under the in-body table, once (page one no
        # longer restates the metrics)
        assert [p.text for p in doc.paragraphs].count(footnote) == 1
    # the footnote is renderer text: neither gate reads it
    lint = memo_quality_lint.lint_memo_docx(tmp_path / "memo" / "en.docx")
    assert not lint.p0_findings
    parity = memo_chinese_parity.lint_chinese_memo_pair(tmp_path / "memo" / "en.docx", tmp_path / "memo" / "zh.docx")
    assert not parity.p0_findings


# ---- legacy re-render and date leniency -------------------------------------------


def test_strict_sources_false_re_renders_a_package_that_predates_the_url_rule(tmp_path, monkeypatch):
    monkeypatch.delenv("BSH_MEMO_SOURCE_URL_REQUIRED", raising=False)
    package = base._package()
    package["sources"].append(
        {
            "id": "S2",
            "title": "Trade press coverage",
            "class": "press reporting",
            "treatment": {"en": "Corroboration.", "zh": "佐证。"},
            "as_of": "2026-05-01",
        }
    )
    with pytest.raises(memo_docx_renderer.MemoRenderError, match=r"sources\[1\]\.url is required"):
        memo_docx_renderer.render_memos(
            package, out_en=tmp_path / "a" / "memo" / "en.docx", out_zh=tmp_path / "a" / "memo" / "zh.docx"
        )
    result = memo_docx_renderer.render_memos(
        package,
        out_en=tmp_path / "b" / "memo" / "en.docx",
        out_zh=tmp_path / "b" / "memo" / "zh.docx",
        strict_sources=False,
    )
    assert result["ok"] and result["renderer_version"] == memo_docx_renderer.RENDERER_VERSION
    assert len(result["warnings"]) == 1 and "sources[1].url is required" in result["warnings"][0]
    report = (tmp_path / "b" / "logs" / "validation.txt").read_text(encoding="utf-8")
    assert "Source warnings" in report
    # a path goes through load_package with the same leniency
    path = tmp_path / "pkg.json"
    path.write_text(json.dumps(package), encoding="utf-8")
    assert memo_docx_renderer.render_memos(
        path,
        out_en=tmp_path / "c" / "memo" / "en.docx",
        out_zh=tmp_path / "c" / "memo" / "zh.docx",
        strict_sources=False,
    )["ok"]
    # everything else stays fail-closed in legacy mode
    broken = copy.deepcopy(package)
    broken["sources"][1]["url"] = "not-a-url"
    with pytest.raises(memo_docx_renderer.MemoRenderError, match="http"):
        memo_docx_renderer.validate_package(broken, strict_sources=False)


def test_published_at_or_data_period_alone_dates_a_source():
    package = base._package()
    source = package["sources"][0]
    source.pop("as_of")
    source["published_at"] = "2026-05"
    memo_docx_renderer.validate_package(package)
    source.pop("published_at")
    source["data_period"] = "undated"
    memo_docx_renderer.validate_package(package)
    source.pop("data_period")
    with pytest.raises(memo_docx_renderer.MemoRenderError, match=r"as_of is required"):
        memo_docx_renderer.validate_package(package)


# ---- the gates read derived blocks and real bullets as before --------------------


def test_gates_skip_renderer_derived_blocks_but_count_their_tables(tmp_path):
    from server import memo_docx_renderer as r

    def build(derived: bool):
        document = Document()
        r._ensure_derived_styles(document)
        paragraph = r._derived_paragraph(document) if derived else document.add_paragraph()
        paragraph.add_run("The investment view is that BSH commits.")
        table = document.add_table(rows=1, cols=1)
        if derived:
            r._mark_derived_table(table)
        table.cell(0, 0).text = "Not disclosed"
        path = tmp_path / f"derived_{derived}.docx"
        document.save(path)
        return path

    plain = memo_quality_lint.lint_memo_docx(build(False))
    assert any(f.code == "sell_side_voice_violation" for f in plain.findings)
    derived_path = build(True)
    assert not memo_quality_lint.lint_memo_docx(derived_path).findings
    shape = memo_chinese_parity._extract_docx_shape(derived_path, "en")
    assert shape.table_count == 1 and not shape.blocks and shape.paragraph_count == 0


def test_gates_read_a_real_bullet_with_its_bullet_mark(tmp_path):
    document = Document()
    memo_docx_renderer._add_bullet(document, "Confirm the ARR before funding.", locale="en")
    path = tmp_path / "bullet.docx"
    document.save(path)
    blocks = memo_quality_lint._extract_docx_blocks(path)
    assert blocks[0].text == "• Confirm the ARR before funding."
    shape = memo_chinese_parity._extract_docx_shape(path, "en")
    assert shape.blocks[0].text == "• Confirm the ARR before funding."


# ---- template choice per run (active_structure version) ----------------------------


def test_active_structure_version_v2_runs_the_v2_chain_with_the_flag_off(monkeypatch):
    monkeypatch.setenv("BSH_MEMO_STRUCTURE_V2", "0")
    assert memo_structure.active_structure("late") is memo_structure.LATE
    chosen = memo_structure.active_structure("late", version="v2")
    assert (chosen.stage, chosen.version) == ("late", 2)
    compact = memo_structure.active_structure("growth", mode="compact", version="v2")
    assert compact.stage == "late_compact" and compact.declared_stage == "growth"
    monkeypatch.setenv("BSH_MEMO_STRUCTURE_V2", "1")
    flagged = memo_structure.active_structure("growth", mode="compact")
    assert flagged == compact


def test_active_structure_version_v1_pins_the_standard_memo_with_the_flag_on(monkeypatch):
    monkeypatch.setenv("BSH_MEMO_STRUCTURE_V2", "1")
    assert memo_structure.active_structure("late", version="v1") is memo_structure.LATE
    assert memo_structure.active_structure("growth", mode="compact", version="V1") is memo_structure.LATE


def test_active_structure_without_a_version_follows_the_flag(monkeypatch):
    monkeypatch.setenv("BSH_MEMO_STRUCTURE_V2", "1")
    assert memo_structure.active_structure("late").version == 2
    assert memo_structure.active_structure("late", version=None).version == 2
    monkeypatch.setenv("BSH_MEMO_STRUCTURE_V2", "0")
    assert memo_structure.active_structure("late", version=None) is memo_structure.LATE
    assert memo_structure.active_structure("late", version="") is memo_structure.LATE


# ---- one language at a time (English-first delivery, Retry Chinese) ----------------


def _blank_zh(value):
    """The package as the pipeline holds it before the Chinese exists."""
    if isinstance(value, dict):
        out = {key: _blank_zh(item) for key, item in value.items()}
        if "en" in out and "zh" in out:
            out["zh"] = ""
        return out
    if isinstance(value, list):
        return [_blank_zh(item) for item in value]
    return value


def test_render_memo_locale_delivers_english_before_the_chinese_exists(tmp_path):
    package = base._package()
    english_only = _blank_zh(package)
    # the pair render refuses a package with no Chinese
    with pytest.raises(memo_docx_renderer.MemoRenderError, match=r"\.zh is required"):
        memo_docx_renderer.validate_package(english_only)
    run_dir = tmp_path / "run"
    out_en = run_dir / "memo" / "memo_en.docx"
    out_zh = run_dir / "memo" / "memo_zh.docx"
    result = memo_docx_renderer.render_memo_locale(english_only, "en", out_en)
    assert result["ok"] and result["locale"] == "en"
    assert result["output"] == str(out_en) and result["outputs"] == {"en": str(out_en)}
    assert result["validation"] == {"en": str(run_dir / "logs" / "validation.txt")}
    assert result["renderer_version"] == memo_docx_renderer.RENDERER_VERSION
    assert "zh_fallback_count" not in result
    assert out_en.exists() and not out_zh.exists()
    assert not (run_dir / "logs" / "validation_cn.txt").exists()
    assert "Generalist, Inc." in _text(Document(out_en))
    inventory = (run_dir / "logs" / "file_inventory.md").read_text(encoding="utf-8")
    assert f"- memo_en: `{out_en}`" in inventory and "memo_zh" not in inventory
    manifest = (run_dir / "logs" / "run_manifest.md").read_text(encoding="utf-8")
    assert manifest.count("## Analysis Finalization") == 1 and f"- memo_en: `{out_en}`" in manifest
    # a Chinese render of the English-only package fails like the pair would
    with pytest.raises(memo_docx_renderer.MemoRenderError, match=r"\.zh is required"):
        memo_docx_renderer.render_memo_locale(english_only, "zh", out_zh)
    assert not out_zh.exists()

    # "Retry Chinese": the finished package renders ONLY the Chinese file
    before = out_en.read_bytes()
    zh_result = memo_docx_renderer.render_memo_locale(package, "zh", out_zh)
    assert zh_result["ok"] and zh_result["outputs"] == {"zh": str(out_zh)}
    assert zh_result["zh_fallback_count"] == 0
    assert out_en.read_bytes() == before  # the English file is never touched
    report = (run_dir / "logs" / "validation_cn.txt").read_text(encoding="utf-8")
    assert "- zh_fallback_count: 0" in report
    inventory = (run_dir / "logs" / "file_inventory.md").read_text(encoding="utf-8").splitlines()
    assert [line for line in inventory if line.startswith(("- memo_", "- validation_"))] == [
        f"- memo_en: `{out_en}`",
        f"- memo_zh: `{out_zh}`",
        f"- validation_en: `{run_dir / 'logs' / 'validation.txt'}`",
        f"- validation_zh: `{run_dir / 'logs' / 'validation_cn.txt'}`",
    ]
    manifest = (run_dir / "logs" / "run_manifest.md").read_text(encoding="utf-8")
    assert manifest.count("## Analysis Finalization") == 2


def test_render_memo_locale_matches_the_pair_render(tmp_path):
    package = base._package()
    pair = memo_docx_renderer.render_memos(
        package,
        out_en=tmp_path / "pair" / "memo" / "en.docx",
        out_zh=tmp_path / "pair" / "memo" / "zh.docx",
    )
    for locale in ("en", "zh"):
        single = memo_docx_renderer.render_memo_locale(
            package, locale, tmp_path / "single" / "memo" / f"{locale}.docx"
        )
        assert _text(Document(single["output"])) == _text(Document(pair["outputs"][locale]))
        name = memo_docx_renderer._VALIDATION_FILE_NAMES[locale]
        single_report = (tmp_path / "single" / "logs" / name).read_text(encoding="utf-8")
        pair_report = (tmp_path / "pair" / "logs" / name).read_text(encoding="utf-8")
        assert single_report.replace(str(tmp_path / "single"), "") == pair_report.replace(
            str(tmp_path / "pair"), ""
        )
    with pytest.raises(ValueError):
        memo_docx_renderer.render_memo_locale(package, "fr", tmp_path / "x.docx")


def test_the_late_stage_renderer_declares_the_review_stamp():
    from server import memo_prep, report_rerender

    assert memo_docx_renderer.SUPPORTS_REVIEW_STAMP is True
    assert report_rerender.renderer_capabilities(memo_prep.LATESTAGE_KIND)["review"] is True


# ---- validator triage (R11 FIX 1) ---------------------------------------------------


def _risk_blocks(package):
    return next(s for s in package["sections"] if s["id"] == "investment_risk")["blocks"]


def test_a_blank_corner_header_is_a_normal_table():
    package = base._package()
    table = package["sections"][0]["blocks"][2]
    table["headers"] = ["", {"en": "2025", "zh": "2025"}, "2026"]
    table["rows"] = [[{"en": "Revenue", "zh": "收入"}, "$50M", "$86M"]]
    assert memo_docx_renderer.english_package_validation_errors(package) == []
    memo_docx_renderer.validate_package(package)
    table["headers"] = [{"en": "", "zh": ""}, "2025", "2026"]
    memo_docx_renderer.validate_package(package)
    # only the corner: a blank header anywhere else is still an error
    table["headers"] = ["Metric", "", "2026"]
    with pytest.raises(memo_docx_renderer.MemoRenderError, match=r"headers\[1\] is required"):
        memo_docx_renderer.validate_package(package)
    table["headers"] = ["", ""]
    with pytest.raises(memo_docx_renderer.MemoRenderError, match=r"headers\[0\] is required"):
        memo_docx_renderer.validate_package(package)


def test_risk_card_wording_checks_split_from_the_structural_ones():
    package = base._package()
    blocks = _risk_blocks(package)
    card = blocks[2]
    card["rows"][1][1] = {"en": "The risk is important.", "zh": "该风险很重要。"}  # no economic consequence
    card["rows"][2][1] = {"en": "Confirm the next financing round.", "zh": "确认下一轮融资。"}  # a command
    everything = memo_docx_renderer.english_package_validation_errors(package)
    assert any("economic consequence" in e for e in everything)
    assert any("must be a signal" in e for e in everything)
    structural = memo_docx_renderer.english_package_validation_errors(package, editorial_risk_checks=False)
    assert not any("economic consequence" in e or "must be a signal" in e for e in structural)
    findings = memo_docx_renderer.risk_card_quality_findings(package)
    assert any("economic consequence" in f for f in findings)
    assert any("must be a signal" in f for f in findings)
    assert sorted(structural + findings) == sorted(everything)


def test_a_missing_card_row_is_a_quality_finding_not_a_structural_error():
    package = base._package()
    del _risk_blocks(package)[2]["rows"][3]  # drop Likelihood
    assert any(
        "exactly five two-cell rows" in e
        for e in memo_docx_renderer.english_package_validation_errors(package)
    )
    assert not any(
        "two-cell rows" in e
        for e in memo_docx_renderer.english_package_validation_errors(package, editorial_risk_checks=False)
    )
    assert any("exactly five two-cell rows" in f for f in memo_docx_renderer.risk_card_quality_findings(package))


def test_structural_risk_card_errors_stay_blocking():
    package = base._package()
    card = _risk_blocks(package)[2]
    del card["layout"]
    card["rows"][4][1] = {"en": "High", "zh": "高"}
    structural = memo_docx_renderer.english_package_validation_errors(package, editorial_risk_checks=False)
    assert any('"layout": "key_value"' in e for e in structural)
    assert any("'N/10: short reason'" in e for e in structural)
    assert memo_docx_renderer.risk_card_quality_findings(package) == []
    assert memo_docx_renderer.risk_card_quality_findings("not a package") == []


# ---- private inventory, the BSH-diligence message, unseen links (A4 / R4 / G1) ----


def _private_source(source_class="BSH primary diligence", title="Series B board deck", **extra):
    source = {
        "id": "S9",
        "title": title,
        "class": source_class,
        "treatment": {"en": "Company claims.", "zh": "公司口径。"},
        "as_of": "2026-05",
    }
    source.update(extra)
    return source


INVENTORY = [
    {"id": "doc-1", "kind": "research_document", "title": "Series B board deck", "ref": "research/acme/1__deck.pdf"},
    {"id": "doc-2", "kind": "research_analysis", "title": "Customer reference call notes", "ref": "research/acme/2.md"},
]


def test_a_url_less_private_source_must_name_an_inventory_item(monkeypatch):
    monkeypatch.delenv("BSH_MEMO_SOURCE_URL_REQUIRED", raising=False)
    errors: list[str] = []
    memo_docx_renderer._validate_source(
        _private_source(), "sources[0]", errors, private_material=True, private_inventory=INVENTORY
    )
    assert errors == []
    errors = []
    memo_docx_renderer._validate_source(
        _private_source(title="Management KPI pack", private_ref="doc-2"),
        "sources[0]", errors, private_inventory=INVENTORY,
    )
    assert errors == []  # private_ref names the item
    errors = []
    memo_docx_renderer._validate_source(
        _private_source(title="Board minutes March"), "sources[0]", errors,
        private_material=True, private_inventory=INVENTORY,
    )
    assert len(errors) == 1
    assert "names none of the firm's private material" in errors[0]
    assert "“Series B board deck”" in errors[0] and "private_ref" in errors[0]
    # lenient re-render: a warning, not an error
    errors, warnings = [], []
    memo_docx_renderer._validate_source(
        _private_source(title="Board minutes March"), "sources[0]", errors,
        private_inventory=INVENTORY, strict_sources=False, warnings=warnings,
    )
    assert errors == [] and len(warnings) == 1
    # without an inventory, today's boolean rule stands
    errors = []
    memo_docx_renderer._validate_source(
        _private_source(title="Board minutes March"), "sources[0]", errors, private_material=True
    )
    assert errors == []


def test_bsh_diligence_with_nothing_on_file_says_drop_it(monkeypatch):
    monkeypatch.delenv("BSH_MEMO_SOURCE_URL_REQUIRED", raising=False)
    for kwargs in ({"private_material": False}, {"private_inventory": []}):
        errors: list[str] = []
        memo_docx_renderer._validate_source(_private_source(), "sources[0]", errors, **kwargs)
        assert len(errors) == 1, kwargs
        assert "no BSH document backs this; drop the source and remove or restate as unverified" in errors[0]
        assert "carry that page" not in errors[0].lower()
    # a merely "internal" class keeps the carry-the-URL message (RadixArk)
    errors = []
    memo_docx_renderer._validate_source(
        _private_source(source_class="internal document", title="Comparables tracker"),
        "sources[0]", errors, private_material=False,
    )
    assert "holds nothing private" in errors[0] and "Carry that page's URL" in errors[0]


def test_the_contract_registry_class_needs_no_url(monkeypatch):
    # The sources contract tells the writer to label a registry value with
    # no document this way; the validator must not then fail the run on it
    # (ZaiNar 2026-09-23__015923, title naming the BSH registry).
    monkeypatch.delenv("BSH_MEMO_SOURCE_URL_REQUIRED", raising=False)
    registry = _private_source(
        source_class="unverified registry value (no document on file)",
        title="BSH company registry entry for ZaiNar, Inc.",
    )
    for kwargs in ({"private_material": False}, {"private_inventory": []}, {"private_material": True}, {}):
        errors: list[str] = []
        memo_docx_renderer._validate_source(registry, "sources[19]", errors, **kwargs)
        assert errors == [], kwargs
    # it is the exact label, not any class that mentions a registry
    errors = []
    memo_docx_renderer._validate_source(
        _private_source(source_class="BSH registry diligence", title="BSH company registry entry"),
        "sources[0]", errors, private_material=False,
    )
    assert "no BSH document backs this" in errors[0]


def test_the_inventory_reaches_the_package_validator(monkeypatch):
    monkeypatch.delenv("BSH_MEMO_SOURCE_URL_REQUIRED", raising=False)
    package = base._package()
    package["sources"].append(_private_source(title="Board minutes March"))
    # the base source ("Company investor materials") names an item too
    package["run"]["private_inventory"] = INVENTORY + [
        {"id": "doc-3", "kind": "research_document", "title": "Company investor materials", "ref": "research/acme/3.pdf"}
    ]
    with pytest.raises(memo_docx_renderer.MemoRenderError, match="names none of the firm's private material"):
        memo_docx_renderer.validate_package(package)
    package["sources"][-1]["title"] = "Series B board deck"
    memo_docx_renderer.validate_package(package)
    package["run"]["private_inventory"] = "not a list"  # malformed -> ignored
    memo_docx_renderer.validate_package(package)


def test_an_unseen_url_prints_unlinked_while_the_source_still_carries_it(tmp_path):
    package = base._package()
    package["sources"] = [
        {
            "id": "S1",
            "title": {"en": "Reuters — Acme raises", "zh": "路透社——Acme 融资"},
            "class": "business_press",
            "treatment": {"en": "Round size.", "zh": "融资规模。"},
            "as_of": "2026-06-01",
            "url": "https://www.reuters.com/acme-raises",
        },
        {
            "id": "S2",
            "title": {"en": "Bloomberg — Acme margins", "zh": "彭博——Acme 利润率"},
            "class": "business_press",
            "treatment": {"en": "Margins.", "zh": "利润率。"},
            "as_of": "2026-06-02",
            "url": "https://www.bloomberg.com/acme-margins-v2",
        },
    ]
    package["run"]["source_url_status"] = {
        "checked_at": "2026-06-22T10:00:00Z",
        # S1's URL was never seen; S2's entry is stale (a repair changed it)
        "unseen": {"S1": "https://www.reuters.com/acme-raises", "S2": "https://www.bloomberg.com/acme-margins"},
    }
    _result, en, zh = _render(tmp_path, package)
    for doc, note in ((en, "link not verified"), (zh, "链接未经核实")):
        table = next(t for t in doc.tables if t.rows[0].cells[0].text in {"ID", "编号"})
        s1, s2 = table.rows[1].cells[1], table.rows[2].cells[1]
        assert "w:hyperlink" not in s1._tc.xml
        assert s1.paragraphs[1].text == f"reuters.com · {note}"
        assert "w:hyperlink" in s2._tc.xml and note not in s2.text
    external = [rel.target_ref for rel in en.part.rels.values() if rel.reltype.endswith("/hyperlink") and rel.is_external]
    assert external == ["https://www.bloomberg.com/acme-margins-v2"]


# ---- back matter follows the structure profile ----------------------------------


def test_back_matter_follows_the_profiles_numbered_flag_and_titles():
    from dataclasses import replace

    package = _with_extras(base._package())
    package["sections"].append(
        {
            "id": "validation_log",
            "title": {"en": "Validation", "zh": "验证"},
            "blocks": [{"type": "paragraph", "text": {"en": "Figures were checked against filings.", "zh": "数字已与申报文件核对。"}}],
        }
    )
    package["calculations"] = [
        {"id": "C1", "label": {"en": "Entry multiple", "zh": "入场倍数"}, "formula": "$2.4B ÷ $86M = 27.9x", "result": "27.9x"}
    ]
    plan = memo_docx_renderer._document_plan(package, memo_structure.LATE)
    assert [(e.kind, e.title["en"], e.label["en"]) for e in plan[-3:]] == [
        ("sources", "IX. Sources, Source Classes, and Fact Reference Index", "IX."),
        ("back_matter", "Appendix: Source Treatment And Assumptions", ""),
        ("calculations", "Calculation notes", ""),
    ]
    assert [e.title["zh"] for e in plan[-3:]] == ["九、来源、来源类别与事实索引", "附录：来源处理与假设", "计算说明"]
    assert [e.number for e in plan] == list(range(1, len(plan) + 1))
    # a profile whose sources entry is not numbered prints its title as written
    unnumbered = replace(
        memo_structure.LATE,
        pseudo_sections=tuple(
            replace(ps, numbered=False) if ps.id == "sources" else ps
            for ps in memo_structure.LATE.pseudo_sections
        ),
    )
    sources = next(e for e in memo_docx_renderer._document_plan(package, unnumbered) if e.kind == "sources")
    assert sources.title == {"en": "Sources, Source Classes, and Fact Reference Index", "zh": "来源、来源类别与事实索引"}
    assert sources.label == {"en": "", "zh": ""}


def test_the_parity_gate_finds_sources_numbered_after_extra_sections(tmp_path):
    package = _with_extras(base._package())
    _result, en, zh = _render(tmp_path, package)
    en_shape = memo_chinese_parity._extract_docx_shape(tmp_path / "memo" / "en.docx", "en", memo_structure.LATE.parity_patterns())
    zh_shape = memo_chinese_parity._extract_docx_shape(tmp_path / "memo" / "zh.docx", "zh", memo_structure.LATE.parity_patterns())
    # "IX. Sources, ..." / "九、来源、..." — the profile's own pattern names "VI"
    assert en_shape.section_ids[-1] == zh_shape.section_ids[-1] == "sources"
    source_rows = [b for b in zh_shape.blocks if b.section_id == "sources" and b.kind == "table_cell"]
    assert source_rows, "the Chinese sources table is attributed to the sources section"


def test_calculation_inputs_read_in_the_memo_language():
    calc = {
        "id": "C1",
        "inputs": [
            {"name": {"en": "post-money valuation", "zh": "投后估值"}, "value": "$2.4B", "ref": "S1"},
            {"name": "exit multiple", "value": "10x", "ref": "assumption"},
            {"name": "dilution", "value": "15%", "ref": "[S2, C3]"},
            {"name": "peer median", "value": "11x", "ref": "Bloomberg screen"},
        ],
    }
    assert memo_docx_renderer._calculation_inputs_text(calc, "en") == (
        "post-money valuation = $2.4B [S1]; exit multiple = 10x (assumption); "
        "dilution = 15% [S2, C3]; peer median = 11x (Bloomberg screen)"
    )
    assert memo_docx_renderer._calculation_inputs_text(calc, "zh") == (
        "投后估值 = $2.4B [S1]; exit multiple = 10x（假设）; "
        "dilution = 15% [S2, C3]; peer median = 11x（Bloomberg screen）"
    )


def test_the_deal_terms_input_source_reads_in_the_memo_language():
    # memo_returns' "BSH proceeds by case" note cites the deal record.
    calc = {"id": "C9", "inputs": [{"name": {"en": "BSH check", "zh": "BSH 出资额"}, "value": "$5M", "ref": "deal terms"}]}
    assert memo_docx_renderer._calculation_inputs_text(calc, "en") == "BSH check = $5M (deal terms)"
    assert memo_docx_renderer._calculation_inputs_text(calc, "zh") == "BSH 出资额 = $5M（交易条款）"
